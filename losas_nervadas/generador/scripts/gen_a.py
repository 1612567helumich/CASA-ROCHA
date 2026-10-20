"""Parte A del generador: DATOS, APOYOS, CARGAS, TABLAS."""
import json
from openpyxl.styles import Alignment
from openpyxl.comments import Comment
from xlstyle import *

import os
W = os.path.dirname(os.path.dirname(os.path.abspath(__file__))) + '/'
D = json.load(open(W + 'work/data.json'))
MOD = json.load(open(W + 'work/proyecto_modelo.json'))
N = {}            # nombres -> referencias absolutas
EDGE = {}         # (key, lado) -> fila en APOYOS
PROPROW = {}      # castillo propuesto -> celda de aceptación
SIDES = ['inferior', 'superior', 'izquierdo', 'derecho']
E2K = '001_EST_CROCHA_R3.2.e2k'
DXFA = '001_PLANO_CROCHA_R3.0.dxf'
DXFT = 'AREAS_TRIBUTARIAS_R3.2.dxf'
NAIM24 = 'Naim — CASA HABITACIÓN, MEMORIA DE CÁLCULO (sep-2024)'
NAIM21 = 'Naim — 21-03 MEMORIA DE CÁLCULO (jul-2021)'
NAIMPL = 'Naim — CASA HABITACIÓN, PLANOS EST-01 a EST-06'
NTCC = 'NTC-Concreto 2023 (CDMX)'
NTCCA = 'NTC-Criterios y Acciones 2023 (CDMX)'
NTCM = 'NTC-Mampostería 2023 (CDMX)'


def key(t):
    return '%s %s' % (t['name'], t['level'])


TABS = sorted(D['tableros'], key=lambda t: (t['level'], t['name'] if t['name'] != 'F2' else 'F00002'))
NPS = [t for t in D['tramos'] if t['tipo'] == 'LIBRE']
DMS = [t for t in D['tramos'] if t['tipo'] == 'MURO']
# --- lo que baja por los muros del primer nivel a las nervaduras principales del N+3.00 ---
#   WAZ: dalas DM del N+6.00 (azotea sobre muro) que quedan encima de cada NP del N+3.00 y fracción de su largo encima
#   transferencias «vía muro»: reacciones de las NP del N+6.00 cuyo extremo cae sobre un muro que está sobre una NP del N+3.00
WAZ = {}
for _t in NPS:
    if _t['level'] != 'N+3.00' or _t['muro_arriba'] <= 0:
        continue
    for _d in DMS:
        if _d['level'] != 'N+6.00' or _d['ori'] != _t['ori'] or abs(_d['c'] - _t['c']) > 0.05:
            continue
        _ov = min(_t['s1'], _d['s1']) - max(_t['s0'], _d['s0'])
        if _ov > 0.01:
            WAZ.setdefault(_t['np'], []).append((_d['np'], round(_ov / _d['L'], 4)))
for _x in D['transfer']:
    _x.setdefault('via', 'directo')
for _u in NPS:
    if _u['level'] != 'N+6.00':
        continue
    for _k in (0, 1):
        _px, _py = (_u['x0'], _u['y0']) if _k == 0 else (_u['x1'], _u['y1'])
        for _t in NPS:
            if _t['level'] != 'N+3.00' or _t['muro_arriba'] <= 0:
                continue
            _c, _sv = (_py, _px) if _t['ori'] == 'H' else (_px, _py)
            if abs(_c - _t['c']) < 0.05 and _t['s0'] - 0.01 <= _sv <= _t['s1'] + 0.01:
                D['transfer'].append(dict(de=_u['np'], extremo=_k, a=_t['np'], pos=round(min(max(_sv - _t['s0'], 0), _t['L']), 3), via='muro'))


# ---------- volados (D10): borde raíz, tableros de respaldo, orillas laterales y muro raíz ----------
def _seg(t, side):
    """Segmento de un borde del rectángulo del tablero: ('H' o 'V', coordenada fija, de, a)."""
    if side in ('inferior', 'superior'):
        return 'H', (t['y0'] if side == 'inferior' else t['y1']), min(t['x0'], t['x1']), max(t['x0'], t['x1'])
    return 'V', (t['x0'] if side == 'izquierdo' else t['x1']), min(t['y0'], t['y1']), max(t['y0'], t['y1'])


def _ov(a0, a1, b0, b1):
    return max(0.0, min(a1, b1) - max(a0, b0))


_TBK = {key(t): t for t in TABS}
BACK = {}      # tablero de respaldo -> dict(vol=clave del volado, side=su borde junto al muro raíz, ov=largo común en m)
VOLLAT = {}    # NP de una orilla lateral de volado (corre en la dirección del volado) -> (clave del volado, lado)
VOLROOT = {}   # clave del volado -> [(DM del muro raíz, largo común en m)]
for _t in TABS:
    if _t['modo'] != 'VOLADO':
        continue
    _v = _t['volado']; _rs = _v['raiz']; _o, _c, _a0, _a1 = _seg(_t, _rs)
    for _n in [x.strip() for x in (_t['bordes'][_rs].get('vecinos') or '').split(',') if x.strip()]:
        _nb = _TBK.get(_n + ' ' + _t['level'])
        if not _nb or _nb['modo'] != 'PERIMETRAL':
            continue
        for _q in SIDES:
            if _t['name'] in [x.strip() for x in (_nb['bordes'][_q].get('vecinos') or '').split(',')]:
                _o2, _c2, _b0, _b1 = _seg(_nb, _q)
                BACK[key(_nb)] = dict(vol=key(_t), side=_q, ov=round(_ov(_a0, _a1, _b0, _b1), 3), Ledge=round(_b1 - _b0, 3))
                break
    for _s in SIDES:
        if _s in (_rs, _v['borde_libre']):
            continue
        _e = _t['bordes'][_s]
        if _e['apoyo'].startswith('NERV') and _e.get('nps'):
            VOLLAT[_e['nps'][0]] = (key(_t), _s)
    VOLROOT[key(_t)] = [(_d['np'], round(_ov(_a0, _a1, _d['s0'], _d['s1']), 3)) for _d in DMS
                        if _d['level'] == _t['level'] and _d['ori'] == _o and abs(_d['c'] - _c) < 0.01 and _ov(_a0, _a1, _d['s0'], _d['s1']) > 0.01]
TBK, segm, ovl = _TBK, _seg, _ov   # para gen_b (import * no trae nombres con guion bajo)
# claro de respaldo del volado: el MÁS LARGO de sus tableros de respaldo (con 0.9·PP todo el respaldo queda en negativo)
VOLB = {}
for _k, _b in BACK.items():
    _nb = _TBK[_k]
    VOLB[_b['vol']] = max(VOLB.get(_b['vol'], 0), _nb['ay'] if _b['side'] in ('inferior', 'superior') else _nb['ax'])
for _t in TABS:
    if _t['modo'] == 'VOLADO':
        _t['volado']['Lb'] = VOLB[key(_t)]
# el área que el DXF da a la orilla libre del volado (nervadura de borde, sin apoyo propio) baja por el volado al muro raíz
_FREE = {}
for _t in TABS:
    if _t['modo'] == 'VOLADO':
        _e = _t['bordes'][_t['volado']['borde_libre']]
        for _np in _e.get('nps') or []:
            _FREE[(_np, _t['name'], _t['level'])] = key(_t)
_tr2 = []
for _x in D['trib']:
    _kv = _FREE.get((_x['tramo'], _x['tablero'], _x['level']))
    if _kv and VOLROOT.get(_kv):
        _tot = sum(o for _, o in VOLROOT[_kv])
        for _dm, _o in VOLROOT[_kv]:
            _tr2.append(dict(tramo=_dm, tablero=_x['tablero'], level=_x['level'], area=round(_x['area'] * _o / _tot, 4),
                             nota='Orilla libre del volado (el DXF la da a %s, que no tiene apoyo propio): baja por el volado al muro raíz, repartida por largo.' % _x['tramo']))
    else:
        _tr2.append(_x)
D['trib'] = _tr2
# ---------- layouts fijos de las hojas de tabla ----------
T0 = 9                     # primera fila de datos en TABLEROS
TROW = {key(t): T0 + i for i, t in enumerate(TABS)}
T1 = T0 + len(TABS) - 1
P0 = 9                     # PRINCIPALES
PROW = {t['np']: P0 + i for i, t in enumerate(NPS)}
P1 = P0 + len(NPS) - 1
R0 = 6                     # TRIB
TRIBROWS = [r for r in D['trib']]
R1 = R0 + len(TRIBROWS) - 1
X0 = 6                     # TRANSFER
X1 = X0 + len(D['transfer']) - 1
# TABLAS
TB_M = 5                   # fila con los valores de m (D..J)
CASOS = ['INTERIOR', 'BORDE-CORTO', 'BORDE-LARGO', 'ESQUINA', 'EXTREMO-LARGO', 'EXTREMO-CORTO', 'AISLADO']
MOMS = ['NIC', 'NIL', 'NDC', 'NDL', 'PC', 'PL']
TB_C0 = 6
TB_C1 = TB_C0 + len(CASOS) * len(MOMS) - 1
BARS = [(2, 0.64, 0.32, 0.248), (2.5, 0.79, 0.49, 0.384), (3, 0.95, 0.71, 0.559), (4, 1.27, 1.27, 0.994), (5, 1.59, 1.98, 1.552),
        (6, 1.91, 2.85, 2.235), (8, 2.54, 5.07, 3.973)]
TB_V0 = TB_C1 + 5
TB_V1 = TB_V0 + len(BARS) - 1
MALLAS = [('6x6-10/10', 3.43), ('6x6-8/8', 4.11), ('6x6-6/6', 4.88), ('6x6-4/4', 5.72)]
TB_W0 = TB_V1 + 5
TB_W1 = TB_W0 + len(MALLAS) - 1
TB_L0 = TB_W1 + 5          # listas
LISTS = [('Caso de tablero', ['AUTO'] + ['INTERIOR', 'BORDE-CORTO', 'BORDE-LARGO', 'ESQUINA', 'EXTREMO-LARGO', 'EXTREMO-CORTO', 'AISLADO']),
         ('Modo', ['PERIMETRAL', 'VOLADO', 'HUECO']), ('Destino', ['HABITACIÓN', 'AZOTEA']), ('Extremo de nervadura', ['FIRME', 'NERVADURA', 'PUNTA', 'LIBRE']), ('Losa', ['NERVADA', 'MACIZA'])]
for _j, (_nm, _vals) in enumerate(LISTS):
    N['LIST_%d' % _j] = "TABLAS!$%s$%d:$%s$%d" % (chr(66 + _j), TB_L0 + 1, chr(66 + _j), TB_L0 + len(_vals))
BARCOL = "TABLAS!$B$%d:$B$%d" % (TB_V0, TB_V1)
DBCOL = "TABLAS!$C$%d:$C$%d" % (TB_V0, TB_V1)
ABCOL = "TABLAS!$D$%d:$D$%d" % (TB_V0, TB_V1)
KGCOL = "TABLAS!$E$%d:$E$%d" % (TB_V0, TB_V1)
LDCOL = "TABLAS!$F$%d:$F$%d" % (TB_V0, TB_V1)
LDHCOL = "TABLAS!$G$%d:$G$%d" % (TB_V0, TB_V1)
LTCOL = "TABLAS!$H$%d:$H$%d" % (TB_V0, TB_V1)
LDOCOL = "TABLAS!$I$%d:$I$%d" % (TB_V0, TB_V1)
NBARS = len(BARS)
N['BARCOL'] = BARCOL; N['DBCOL'] = DBCOL; N['ABCOL'] = ABCOL; N['KGCOL'] = KGCOL; N['LDCOL'] = LDCOL; N['LDHCOL'] = LDHCOL; N['LTCOL'] = LTCOL; N['LDOCOL'] = LDOCOL
N['MROW'] = "TABLAS!$D$%d:$J$%d" % (TB_M, TB_M)
N['CKEY'] = "TABLAS!$B$%d:$B$%d" % (TB_C0, TB_C1)
N['CVAL'] = "TABLAS!$D$%d:$J$%d" % (TB_C0, TB_C1)
N['TKEYS'] = "TABLEROS!$B$%d:$B$%d" % (T0, T1)
N['NPKEYS'] = "PRINCIPALES!$B$%d:$B$%d" % (P0, P1)


def ftx(*parts):
    """Texto vivo para una celda: los literales se parten en trozos de ≤ 200 caracteres (Excel no acepta literales de más de 255); las tuplas son expresiones."""
    out = []
    for p_ in parts:
        if isinstance(p_, tuple):
            out.append(p_[0])
        else:
            out += ['"%s"' % p_[i:i + 200].replace('"', '""') for i in range(0, len(p_), 200)]
    return '=' + '&'.join(out)


def db(var):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (DBCOL, var, BARCOL)


def ab(var):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (ABCOL, var, BARCOL)


def dbz(var):
    """Diámetro que admite 0 = «sin estribo»."""
    return 'IF(%s=0,0,%s)' % (var, db(var))


def ld(var):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (LDCOL, var, BARCOL)


def ldo(var):
    """L_d de «otros casos» de la Tabla 14.4.2.4 (separación libre < 2d_b o recubrimiento < d_b)."""
    return 'INDEX(%s,MATCH(%s,%s,0))' % (LDOCOL, var, BARCOL)


# =====================================================================
def sheet_datos(wb):
    ws = wb.create_sheet('DATOS'); S = Narr(ws, N)
    S.title('LOSAS NERVADAS CASA ROCHA — DATOS DE ENTRADA')
    S.subtitle('Casa habitación de dos niveles · muros de block de 15 cm confinados con castillos K-15x15 y dalas D-15x20 · losa reticular LR-25 en entrepiso (N+3.00) y azotea (N+6.00), salvo F80032 y F80010 de azotea, que el modelo trae como losa maciza de 20 cm («Slab1») y así se diseñan (decisión D11). '
               'Se diseñan las nervaduras secundarias de cada tablero, las dos losas macizas y las nervaduras principales (dalas que quedan sin muro abajo y trabajan como trabe).')
    S.gap()
    S.notice('CÓMO SE USA: sólo se tocan las celdas AMARILLAS con número AZUL. Todo lo demás está calculado y se rompe si lo escribes encima. '
             'La columna «CÓMO SE CALCULA» dice de dónde sale el número y la columna «REFERENCIA» dice qué documento, norma o archivo lo respalda.')
    S.notice('ALCANCE: (1) nervaduras secundarias de los tableros de losa reticular con el método de coeficientes de la NTC-Concreto para losas apoyadas en su perímetro; '
             '(2) dos tableros en volado (F80012 y F80019), nervados con nervaduras de 20 cm en la dirección del volado (decisión D10); (3) dos losas macizas de 20 cm en azotea (F80032 y F80010, decisión D11) con el mismo método de coeficientes; '
             '(4) 34 nervaduras principales sin muro abajo diseñadas como trabe, incluida la revisión de muros sobre vigas de la NTC-Mampostería; (5) carga de la losa sobre los muros. '
             'NO incluye la escalera, el pergolado de vigas VP-15x40, la revisión de los muros por sismo ni la cimentación.', 54)
    S.notice('NORMAS: NTC-Concreto 2023, NTC-Criterios y Acciones 2023 y NTC-Mampostería 2023 (CDMX). Cada fórmula se cotejó contra el texto oficial del compendio NTC 2023 '
             'y se cita con su numeral y la página impresa de cada norma (p. ej. «6.7.8.1.4, p. 187»). Los criterios que la norma deja al calculista se marcan como «criterio de esta hoja». '
             'ACI 318-19 ya no se usa para nada que la NTC 2023 resuelva.', 54)
    S.gap()
    S.section('1. GEOMETRÍA DE LA LOSA NERVADA LR-25 Y RECUBRIMIENTOS')
    S.secnote('Casetón de poliestireno 60×60×19, nervadura de 10 cm y capa de 6 cm (decisión D3): el peralte sigue en 25 cm (LR-25). La nota 2 del plano de cortes y el E2K traen 60×60×20 con capa de 5 cm; '
              'con capa de 5 cm ninguna malla cumple a la vez diafragma, punzonamiento y acero mínimo junto a la carga (NERVADURA sección 11). Los modificadores del modelo ya no coinciden: se actualizan en ETABS (sección 4 y decisión D5).', 42)
    S.header()
    S.row('a_c — lado del casetón (separación libre entre nervaduras)', 60, 'cm', 'CAPTURA. Debe ser ≤ 80 cm y ≤ claro/6 para que la losa cuente como encasetonada (se revisa en TABLEROS).', DXFA + ' · NOTAS DE CORTES, nota 2 · ' + NTCC + ' 6.7.8.1.4, p. 187', 'input', '0', 'ac')
    S.row('h_c — altura del casetón', 19, 'cm', 'CAPTURA. Decisión D3: casetón de 19 cm para que, con la capa de 6 cm, la losa siga en h = 25 cm. El plano dice 20 cm: corregir la nota 2 del plano de cortes.', DXFA + ' · NOTAS DE CORTES, nota 2 (dice 20)', 'input', '0', 'hc')
    S.row('b_w — ancho de la nervadura secundaria', 10, 'cm', 'CAPTURA. Mínimo 10 cm en cualquier punto del peralte.', DXFA + ' · nota 2 · ' + NTCC + ' 6.7.8.1.2, p. 187', 'input', '0', 'bw')
    S.row('b_w de las nervaduras de los volados F80012 y F80019 (las que corren en la dirección del volado)', 20, 'cm',
          'CAPTURA. Decisión D10: el entrepiso es nervado y los volados también; con 10 cm el M⁻ del empotramiento pide una No. 8 que no cabe. Casetón de (s − b_w) × a_c × h_c = 50×60×19 (los 50 cm entre las nervaduras del volado): '
          's sigue en 70 cm y la retícula queda alineada con la del resto del entrepiso. Las nervaduras paralelas al muro siguen de 10 cm. Si con 20 cm no alcanzara, se pasa a 25 cm (casetón 45×60).',
          'Decisión D10 · ' + NTCC + ' 6.7.8.1.2 y 6.7.8.1.3, p. 187', 'input', '0', 'bw_vol')
    S.row('t — espesor de la capa de compresión', 6, 'cm', 'CAPTURA. Mínimo el mayor de a_c/13 y 5 cm; 6 cm si algún tablero tiene claro ≥ 6 m (diafragma). Decisión D3: 6 cm con malla 6x6-4/4, que cumple diafragma (6.8.6.4), punzonamiento y acero mínimo junto a la carga (6.7.6.1.2).',
          DXFA + ' · nota 2 (dice 5) · ' + NTCC + ' 6.7.8.3.1, p. 188 y 6.8.3.1.2, p. 189', 'input', '0', 't')
    S.row('h — peralte total de la losa', '=({hc}+{t})', 'cm', 'CALCULADA = h_c + t. Debe dar 25: es la LR-25 del modelo. Máximo 3.5·b_w.', E2K + ' · SHELLPROP "LR-25" SLABTHICKNESS 0.25 · ' + NTCC + ' 6.7.8.1.3, p. 187', 'key', '0', 'h')
    S.row('s — separación entre ejes de nervaduras', '=({ac}+{bw})', 'cm', 'CALCULADA = a_c + b_w. Es el ancho tributario de cada nervadura.', 'Geometría', 'key', '0', 's')
    S.row('a_c del volado — lado del casetón entre las nervaduras de 20 cm (s − b_w del volado)', '={s}-{bw_vol}', 'cm', 'CALCULADA = s − b_w del volado (s sigue en 70 cm). Casetón de los volados y de las hileras de respaldo que siguen con nervaduras de 20 cm (TABLEROS, «RESPALDO»): (s − b_w) × a_c × h_c = 50×60×19.', 'Decisión D10 · Geometría', 'calc', '0', 'ac_vol')
    S.row('b_p — ancho de la nervadura principal / dala', 15, 'cm', 'CAPTURA. Es el ancho de las dalas D-15x20 del modelo. En PRINCIPALES cada tramo lleva su propia sección propuesta.', E2K + ' · FRAMESECTION "002_D-15x20" B 0.15', 'input', '0', 'bp')
    S.row('b_borde — ancho de la DALA DE BORDE armada en los bordes discontinuos', 40, 'cm',
          'CAPTURA. Decisión D9: en todo borde discontinuo apoyado (sobre muro o sobre nervadura principal) va una DALA DE BORDE ARMADA de b_borde de ancho total en el peralte de la losa, en lugar de la primera franja junto al apoyo: '
          'sobre muro se apoya en el muro de 15 cm; sobre nervadura principal es la propia principal ensanchada a b_borde en el peralte de la losa, con estribos que abarcan ese ancho. Recibe las nervaduras en su paño interior. '
          'Armado mínimo propuesto (lo revisa y firma el Responsable): 4 No. 4 corridas en las esquinas (cubren las 2 No. 4 + 2 No. 3 de la dala D-15x20, NTC-M 6.2.2.2) y estribos cerrados No. 3 @ 20 cm que abarcan todo el ancho. '
          'Un macizo SIN armar no sirve: sería parte de la losa, el apoyo sería el paño del muro y no habría dónde anclar (NERVADURA sección 10). En las losas macizas de azotea (D11) es una dala de b_borde × h que baja h − h_maciza bajo la maciza. '
          'Debe ser ≥ 1.25·L_dh + r_np para anclar la barra inferior a 1.25 fy y el bastón superior con gancho de 180°. TABLEROS acomoda los casetones con este ancho en esos bordes.',
          'Decisión D9 · ' + NTCC + ' 6.7.8.1.7, p. 188 · 6.7.7.4.1.1, p. 179 · 14.4.3.1, p. 355 · ' + NTCM + ' 6.2.2.2, p. 62', 'input', '0', 'b_borde')
    S.row('h de las losas macizas de azotea F80032 y F80010', 20, 'cm',
          'CAPTURA. Decisión D11: el modelo las trae como «Slab1», losa maciza de 20 cm; se respeta. Su lecho inferior queda 5 cm más alto que el de la LR-25 vecina (desnivel en planos).',
          E2K + ' · SHELLPROP "Slab1" (maciza 0.20 m) · Decisión D11', 'input', '0', 'h_mac')
    S.row('¿Se anotan en planos las precauciones de 3.9.2? (casetón de medida fija que garantiza b_w y silletas que fijan el peralte del lecho superior)', 'SÍ', '—',
          'SELECTOR. b_w = 10 cm < 200 mm y, donde el lecho superior queda con d ≤ 200 mm, la norma pide calcular con 20 mm menos, salvo que se tomen precauciones que garanticen la dimensión y se anoten en planos. '
          'Si dices NO, la hoja usa b_w − 2 cm en cortante y en flexión negativa, y d − 2 cm en el lecho superior de cada tablero cuyo d sea ≤ 20 cm.',
          NTCC + ' 3.9.1 d) y e) y 3.9.2, p. 74 · 6.3.1.2 b), p. 117 (las nervaduras son vigas)', 'input', '@', 'ctl_bw')
    dv_list(ws, N['ctl_bw'].split('!')[1].replace('$', ''), '"SÍ,NO"')
    S.row('b_w de cálculo', '=IF({ctl_bw}="SÍ",{bw},{bw}-2)', 'cm', 'CALCULADA.', NTCC + ' 3.9.1 e), p. 74', 'key', '0.0', 'bwc')
    S.row('r — recubrimiento libre en losa y nervaduras secundarias (interior, A1)', 2.0, 'cm',
          'CAPTURA. El mayor de: 15 mm (colocación en losas), 0.75 × 25 mm = 19 mm (corrosión A1 con f\'c = 250) y 20 mm (losas y viguetas no expuestas). Manda 20 mm.',
          NTCC + ' 4.13.1.1 y 4.13.2.2 b), p. 83 · Tabla 4.13.3.1 y 4.13.3.2, p. 84 · Tabla 4.14.16.a, p. 85 · ' + NAIM24 + ' pág. 13', 'input', '0.0', 'rec')
    S.row('r_ext — recubrimiento de caras expuestas a la intemperie (barras ≤ No. 5)', 4.0, 'cm',
          'CAPTURA. Se usa en el lecho inferior de los volados (cara exterior), en su lecho superior si queda a la intemperie (selector de abajo) y en el lecho superior de azotea si la pones como expuesta. Con No. 6 o mayor son 5 cm.',
          NTCC + ' Tabla 4.14.16.a, p. 85 · Tabla 4.3.1, p. 76', 'input', '0.0', 'rec_ext')
    S.row('r_np — recubrimiento libre de las nervaduras principales que trabajan como trabe', 4.0, 'cm',
          'CAPTURA. Para «vigas» no expuestas la tabla pide 40 mm al estribo (las nervaduras ordinarias son «viguetas» y llevan 20 mm). '
          'También lo llevan las nervaduras secundarias de los tableros que no cumplen 6.7.8.1.4: esa losa se diseña como losa apoyada en vigas (6.7.8.1.9) y sus nervaduras ya no son «viguetas».',
          NTCC + ' Tabla 4.14.16.a, p. 85 (vigas: refuerzo principal y estribos, 40 mm) · 6.7.8.1.9 y comentario de 6.7.8, p. 188', 'input', '0.0', 'rec_np')
    S.row('¿La cara superior de la azotea queda expuesta directo a la intemperie?', 'NO', '—',
          'SELECTOR. Con relleno, entortado e impermeabilizante encima la hoja la toma como protegida (criterio del Responsable: la norma no lo dice expresamente). Si dices SÍ, el lecho superior de azotea usa r_ext.',
          NTCC + ' 4.3.2 y Tabla 4.3.1, p. 76 · Tabla 4.14.16.a, p. 85', 'input', '@', 'exp_az')
    dv_list(ws, N['exp_az'].split('!')[1].replace('$', ''), '"SÍ,NO"')
    S.row('¿La cara superior de los volados F80012 y F80019 queda expuesta a la intemperie?', 'NO', '—',
          'SELECTOR. Decisión D12: los dos volados del N+3.00 son terrazas con la cara superior PROTEGIDA con impermeabilizante y piso, igual que la azotea; el plano debe decir «terraza con impermeabilizante y piso». '
          'Con NO la exposición ya no manda arriba: la malla va al centro de la capa de 6 cm (≈ 2.4 cm por cara) y el lecho superior de las nervaduras del volado conserva el recubrimiento de viga r_np = 4 cm, porque esas nervaduras no cumplen 6.7.8.1.4 '
          '(losa sobre vigas, 6.7.8.1.9, decisión D7): del lado seguro. Abajo sigue r_ext (cara exterior). Si dices SÍ (cara a la intemperie, exposición A2) el lecho superior y la malla llevan r_ext y la malla no cabe con 2 cm abajo (NERVADURA sección 11).',
          NTCC + ' Tabla 4.3.1 c), p. 76 · 4.3.2, p. 76 · Tabla 4.14.16.a, p. 85 · 6.7.8.1.9, p. 188 · Decisión D12', 'input', '@', 'exp_vol')
    dv_list(ws, N['exp_vol'].split('!')[1].replace('$', ''), '"SÍ,NO"')
    S.row('No. de estribo en la nervadura secundaria (0 = sin estribo)', 0, '—',
          'CAPTURA. Las nervaduras de 10 cm no llevan estribo: la fuerza cortante la toma el concreto con la ec. 5.5.3.2.1 y donde no alcanza se maciza junto al apoyo. Si algún día se ponen, deben ser de varilla o alambre CORRUGADO (no alambrón liso).',
          NTCC + ' 5.5.3.8.1, p. 97 · Tabla 6.3.5.4.1 c), p. 123 · 2.4.2.1.1, p. 43 y Tabla 2.4.2.2.4.a, p. 46', 'input', 'General', 'est_sec')
    S.row('No. de varilla preferida — lecho INFERIOR de nervaduras secundarias', 4, '—', 'CAPTURA. Mínimo No. 4: toda nervadura lleva al menos 1 No. 4 abajo y 1 No. 4 arriba, corridas. Si con el máximo de varillas no alcanza, la hoja sube sola al siguiente número.',
          NTCC + ' 6.7.6.1.3, p. 175', 'input', 'General', 'var_inf')
    S.row('No. de varilla preferida — lecho SUPERIOR de nervaduras secundarias', 4, '—', 'CAPTURA. Misma regla: una varilla del lecho superior es corrida y las demás son bastones.', NTCC + ' 6.7.6.1.3, p. 175', 'input', 'General', 'var_sup')
    S.row('Máximo de varillas por lecho en una nervadura secundaria', 2, 'pzas', 'CAPTURA. Vale también para las nervaduras de 20 cm de los volados (D10). La hoja revisa además que quepan con la separación libre mínima.', 'Geometría · ' + NTCC + ' 14.2.1, p. 347', 'input', '0', 'nmax_sec')
    S.row('Mínimo de varillas por lecho y por nervadura', 1, 'pzas',
          'CAPTURA. 1 No. 4 corrida. Si As colocada < 4/3 de la requerida, la hoja pide 2 varillas y As mín de viga.',
          NTCC + ' 6.7.6.1.3, p. 175 · 6.3.5.1.1 y 6.3.5.1.3, p. 122', 'input', '0', 'nmin_sec')
    S.row('Tamaño máximo del agregado grueso', 1.27, 'cm', 'CAPTURA. 1/2". No más de 2/3 del recubrimiento ni de la separación libre entre barras: con r = 2 cm el máximo es 1.33 cm, así que el de 3/4" NO sirve en esta losa.',
          NTCC + ' 2.2.5.2.1, p. 32', 'input', '0.00', 'tmag')
    S.row('¿El agregado cabe? (≤ 2/3 de r y de la separación libre)', '=IF({tmag}<=2/3*MIN({rec},MAX(2.5,1.5*{tmag})),"CUMPLE","NO CUMPLE — usa agregado más chico")', '—', 'CALCULADA.', NTCC + ' 2.2.5.2.1, p. 32', 'calc', '@')
    S.row('Separación libre mínima entre varillas', '=MAX(2.5,1.5*{tmag})', 'cm', 'CALCULADA = el mayor de 25 mm, d_b y 1.5·TMA (d_b se revisa barra por barra).', NTCC + ' 14.2.1, p. 347', 'calc', '0.00', 'slib')
    S.row('Ancho libre para varillas dentro de la nervadura secundaria', '={bw}-2*{rec}-2*' + dbz('{est_sec}'), 'cm', 'CALCULADA = b_w − 2r − 2·estribo con r = 2 cm (tableros que cumplen 6.7.8.1.4). Informativa: TABLEROS revisa renglón por renglón con su r y su b_w; con r = 4 cm (tableros que no cumplen 6.7.8.1.4) quedan 2 cm y cabe una sola varilla de d_b ≤ 2 cm; en las nervaduras de 20 cm de los volados (D10) quedan 12 cm.', 'Geometría', 'calc', '0.00', 'av_sec')
    S.row('Malla electrosoldada en la capa de compresión', '6x6-4/4', '—',
          'SELECTOR. La lista está en TABLAS. Debe cubrir: contracción y temperatura (0.0018·t), cortante de diafragma (0.0025·t), flexión de la capa con 1 000 kg entre nervaduras, acero mínimo alrededor de la carga concentrada (6.7.6.1.2) y caber con 2 cm de recubrimiento por cara (NERVADURA sección 11). '
          'La malla va AL CENTRO de la capa, sobre silletas de ≈ 2.4 cm: (t − 2Ø)/2 por cara; así la revisa la sección 11 (d = t/2) y así se dibuja. Donde la cara superior quedara expuesta (selectores de arriba) bajaría a r_ext de la cara superior y se revisa aparte; con la decisión D12 no hay zonas expuestas. '
          'Decisión D3: capa de 6 cm con 6x6-4/4, que cumple todas las revisiones; con capa de 5 cm ninguna malla las cumplía. El plano EST-03 dice 6x6-10/10: corregirlo.',
          NTCC + ' 13.8.2.1.1, p. 345 · 6.7.8.1.8, p. 188 · 6.8.6.4, p. 196 · 6.7.6.1.2, p. 175 · Tabla 6.7.4.2.2.3, p. 171 · Tabla 4.14.16.a, p. 85 · ' + NAIMPL + ' · EST-03 (dice 6x6-10/10)', 'input', '@', 'malla')
    dv_list(ws, N['malla'].split('!')[1].replace('$', ''), '=TABLAS!$B$%d:$B$%d' % (TB_W0, TB_W1))
    S.row('Ø del alambre de la malla', '=INDEX(TABLAS!$C$%d:$C$%d,MATCH({malla},TABLAS!$B$%d:$B$%d,0))/10' % (TB_W0, TB_W1, TB_W0, TB_W1), 'cm',
          'CALCULADA. La malla va al centro de la capa y las varillas superiores de las nervaduras pasan DEBAJO de ella: su d se mide con el mayor de r del tablero y (t − 2Ø)/2 hasta la cara superior de la malla, menos 2·Ø.', 'TABLAS sección C', 'calc', '0.000', 'fim')
    S.gap()
    S.section('2. MATERIALES')
    S.header()
    S.row("f'c — resistencia del concreto", 250, 'kg/cm²', "CAPTURA. Concreto Clase 1A (250 ≤ f'c < 400). Es también el mínimo para estructuras de ductilidad baja y para exposición A1, A2 y B1.",
          E2K + " · MATERIAL \"Concreto f'c=250kg/cm2. Clase 1A\" · " + NTCC + ' 2.2.2 a), p. 30 · Tabla 2.2.6.3, p. 33 · 4.4.1.1 y 4.5.2 a), p. 78', 'input', '0', 'fc')
    S.row('fy — fluencia de la varilla corrugada', 4200, 'kg/cm²', 'CAPTURA. Grado 42, NMX-B-506. Es también el máximo para estribos.', E2K + ' · MATERIAL "Fy=4200" · ' + NTCC + ' Tabla 2.4.2.1.2.a, p. 43 y Tabla 2.4.2.2.4.a, p. 46', 'input', '0', 'fy')
    S.row('fy de la malla electrosoldada', 5000, 'kg/cm²', 'CAPTURA. Alambre de malla NMX-B-290; en contracción, temperatura y flexión de losas se permite hasta 7 000. En 2023 la cuantía de temperatura ya no depende de fy.',
          NTCC + ' 2.4.2.1.6 y 2.4.2.1.7, p. 44 · Tabla 2.4.2.2.4.a renglón «Otras», p. 46 · plano EST Casa Yépez (Naim), nota 1.2', 'input', '0', 'fym')
    S.row('fy de los estribos de nervaduras principales', 4200, 'kg/cm²', 'CAPTURA. Varilla corrugada; para cortante el máximo es 4 200.', NTCC + ' 5.5.2.4.3, p. 94 · Tabla 2.4.2.2.4.a, p. 46', 'input', '0', 'fyt')
    S.row('γc — peso volumétrico del concreto reforzado', 2.44, 't/m³', 'CAPTURA. El mismo del modelo (valor máximo probable).', E2K + ' · WEIGHTPERVOLUME 2440 · ' + NTCCA + ' 5.1.1.2, p. 9', 'input', '0.00', 'gc')
    S.row('γ del casetón de poliestireno', 0.015, 't/m³', 'CAPTURA. Poliestireno expandido. El modelo no lo cuenta; aquí sí.', DXFA + ' · simbología del corte «CASETÓN DE POLIESTIRENO»', 'input', '0.000', 'gpol')
    S.row('Ec — módulo de elasticidad', '=14000*SQRT({fc})', 'kg/cm²', "CALCULADA = 14 000·√f'c: Clase 1A con agregado grueso CALIZO (con basalto sería 11 000·√f'c). Especificarlo así en planos.",
          E2K + ' · E 2.2135944E+09 kgf/m² · ' + NTCC + ' Tabla 2.2.1, p. 31 y 2.2.8.4, p. 34', 'calc', '#,##0', 'Ec')
    S.row('Es — módulo de elasticidad del acero', 2000000, 'kg/cm²', 'CAPTURA.', NTCC + ' 2.4.2.2.2, p. 45', 'input', '#,##0', 'Es')
    S.row("f̄f — módulo de rotura medio (para el momento de agrietamiento)", '=2*SQRT({fc})', 'kg/cm²', "CALCULADA = 2·√f'c (Clase 1A). Se usa sólo en las flechas.", NTCC + ' Tabla 2.2.1, p. 31 · 13.4.3.3, p. 338', 'calc', '0.0', 'ffr')
    S.row("f''c — esfuerzo del bloque equivalente", '=0.85*{fc}', 'kg/cm²', "CALCULADA = 0.85·f'c (sin límite de f'c en 2023).", NTCC + ' 3.6.1 e), p. 68 · ' + NAIM24 + ' (su MR sale con 0.85 f\'c)', 'calc', '0.0', 'fpc')
    S.row('β1 — profundidad del bloque', '=IF({fc}<=300,0.85,MAX(0.65,1.05-{fc}/1400))', '—', "CALCULADA. 0.85 hasta f'c = 300 kg/cm².", NTCC + ' ec. 3.6.1, p. 68', 'calc', '0.00', 'beta1')
    S.row("f'm — resistencia de diseño a compresión del block", 50, 'kg/cm²', "CAPTURA. Block de concreto f'p = 100 kg/cm² con mortero tipo I. La casa es estructura Tipo I (tiene muros desplantados sobre vigas): no se pueden usar los valores por omisión del Cap. 12.",
          E2K + " · MATERIAL \"f'm=50kg/cm2 (Mortero Tipo I, f'p=100kg/cm2)\" · " + NTCM + ' Tabla 2.8.1.4.1, p. 24 · 1.6.1 c), p. 15 · 1.6.2, p. 16', 'input', '0', 'fm')
    S.row('Em — módulo de elasticidad de la mampostería', '=800*{fm}', 'kg/cm²', "CALCULADA = 800·f'm (block de concreto, cargas de corta duración). El E2K trae 40 000.", NTCM + ' ec. 2.8.5.4.a, p. 27 · ' + E2K + ' E 4.0E+08 kgf/m²', 'calc', '#,##0', 'Em')
    S.gap()
    S.section('3. FACTORES DE CARGA Y DE RESISTENCIA')
    S.header()
    S.row('FC_CM — factor de carga muerta', 1.3, '—', 'CAPTURA. Grupo B, cargas permanentes + variables (Grupo A sería 1.5).', NTCCA + ' 3.4.1 a), p. 7 (combinación 2.3.1 a), p. 5) · ' + NAIM24 + ' pág. 8 · ' + E2K + ' COMBO "Comb2.2"', 'input', '0.00', 'FCcm')
    S.row('FC_CV — factor de carga viva', 1.5, '—', 'CAPTURA. Grupo B (Grupo A sería 1.7).', NTCCA + ' 3.4.1 a), p. 7', 'input', '0.00', 'FCcv')
    S.row('FC_acc — factor de carga en combinaciones con acciones accidentales', 1.1, '—', 'CAPTURA. Se usa en la combinación de granizo en azotea: 1.1·(CM + W_g), sin carga viva.', NTCCA + ' 3.4.1 b), p. 7 · 7.3.1.3, p. 17', 'input', '0.00', 'FCacc')
    S.row('F_R — flexión', 0.9, '—', 'CAPTURA. Sólo vale si la sección es controlada por tensión (ε_t ≥ 0.005). Losas y vigas sin presfuerzo tienen que serlo, así que la hoja limita el acero para garantizarlo (p_max, sección 4).',
          NTCC + ' Tabla 3.8.2.1 a) y Tabla 3.8.2.2, p. 70 · 3.8.2.4, p. 71 · 6.3.3.3.1, p. 121 · 6.7.3.3.1, p. 169', 'input', '0.00', 'FRf')
    S.row('F_R — cortante', 0.75, '—', 'CAPTURA.', NTCC + ' Tabla 3.8.2.1 b), p. 70', 'input', '0.00', 'FRv')
    S.row('F_R — mampostería en compresión', 0.6, '—', 'CAPTURA. Para muros sobre vigas (esfuerzo en los extremos del muro).', NTCM + ' 4.1.6.2, p. 39 · 4.1.18.2.2, p. 44', 'input', '0.00', 'FRm')
    S.row('Fracción de la cuantía balanceada que se permite', 0.90, '—', 'CAPTURA. Tope de vigas de ductilidad baja; en la práctica manda antes ε_t ≥ 0.005.', NTCC + ' 6.3.5.2.1, p. 122', 'input', '0.00', 'pbf')
    S.row('Incremento de V_cR en losas encasetonadas', 1.1, '—', 'CAPTURA. Sólo se aplica en los tableros que cumplen las limitaciones geométricas de losa encasetonada (claro ≥ 6·a_c en las dos direcciones); en los demás vale 1.0.',
          NTCC + ' 6.7.8.1.5, p. 187 · 6.7.8.1.9, p. 188', 'input', '0.00', 'VcRinc')
    S.row('Incremento de cortante con bordes continuos y discontinuos', 1.15, '—', 'CAPTURA. Parte de la fórmula de cortante de losas perimetralmente apoyadas de la NTC 2017, que la 2023 ya no trae; se conserva como «método reconocido».',
          NTCC + ' 3.4.1.1, p. 64 (métodos reconocidos) · NTC-Concreto 2017', 'input', '0.00', 'Vinc')
    S.row('Relación M⁻/M⁺ para nervaduras principales apoyadas', 0.667, '—', 'CAPTURA. Equivale a wL²/12 contra wL²/8: arriba se arma para empotramiento, abajo para apoyo simple (envolvente conservadora).', 'Criterio de esta hoja — no viene de norma', 'input', '0.000', 'Mrat')
    S.gap()
    S.section('4. PROPIEDADES DE LA SECCIÓN NERVADA — Y COMPROBACIÓN CONTRA EL MODELO')
    S.secnote('Con la sección de la decisión D3 (casetón de 19 cm y capa de 6 cm) estos cocientes YA NO coinciden con los del E2K, que trae la de 20 + 5: el modelo de ETABS se actualiza con los modificadores de esta sección (decisión D5). '
              'Los valores del modelo se dejan para comparar.', 30)
    S.header()
    S.row('Volumen de concreto por m² de losa', '=({h}/100-({ac}/100)^2*({hc}/100)/({s}/100)^2)', 'm³/m²', 'CALCULADA = h − a_c²·h_c / s² (lo macizo menos los huecos de casetón).', 'Geometría', 'calc', '0.0000', 'Vm2')
    S.row('Peso propio geométrico de la losa nervada', '=({Vm2}*{gc}+({ac}/100)^2*({hc}/100)/({s}/100)^2*{gpol})*1000', 'kg/m²', 'CALCULADA = V·γc + casetón·γ_pol. A esto se le suma el incremento de la sección 5.', 'Geometría', 'key', '0.0', 'PP')
    S.row('Peso propio que trae el modelo', '=1503.75*0.25*0.6688', 'kg/m²', 'CALCULADA = 1503.75 kg/m³ × 0.25 m × WMOD 0.6688. Es el de la sección vieja (20 + 5) sin casetón; con la capa de 6 cm la losa pesa más: actualizar el WMOD del modelo (fila «WMOD para ETABS», decisión D5).', E2K + ' · MATERIAL "LR-25 equiv" 1503.75 kg/m³ y SHELLPROP "LR-25" WMOD 0.6688', 'calc', '0.0', 'PPe2k')
    S.row('A_T — área de la sección T de una nervadura', '={s}*{t}+{bw}*{hc}', 'cm²', 'CALCULADA = s·t + b_w·h_c.', 'Geometría', 'calc', '0.0', 'AT')
    S.row('ȳ — centroide medido desde arriba', '=({s}*{t}*{t}/2+{bw}*{hc}*({t}+{hc}/2))/{AT}', 'cm', 'CALCULADA.', 'Geometría', 'calc', '0.00', 'ybar')
    S.row('I_T — inercia bruta de una nervadura con su patín', '={s}*{t}^3/12+{s}*{t}*({ybar}-{t}/2)^2+{bw}*{hc}^3/12+{bw}*{hc}*({t}+{hc}/2-{ybar})^2', 'cm⁴', 'CALCULADA (Steiner).', 'Geometría', 'calc', '#,##0', 'IT')
    S.row('I_T / I de losa maciza del mismo peralte', '={IT}/({s}*{h}^3/12)', '—', 'CALCULADA. Es el modificador de inercia (M11, M22, M12) que va en ETABS con la sección de la decisión D3 (D5); el del modelo (fila de abajo) es el de la sección vieja.', 'Geometría · Decisión D5', 'key', '0.0000', 'Irat')
    S.row('Modificador de inercia en el modelo', 0.29914984, '—', 'DATO del modelo (sección vieja 20 + 5). Ya no coincide con la fila de arriba porque la sección cambió (D3): en ETABS va el valor de arriba (D5).', E2K + ' · SHELLPROP "LR-25" M11MOD 0.29914984', 'calc', '0.0000', 'Irat_e2k')
    S.row('h_eq — espesor de losa maciza con la misma inercia', '=(12*{IT}/{s})^(1/3)', 'cm', 'CALCULADA = ∛(12·I_T / s). La norma no da espesor mínimo especial para encasetonadas: la hoja compara h_eq (no los 25 cm) contra el espesor mínimo, del lado seguro.', 'Criterio de esta hoja · ' + NTCC + ' 6.7.3.1, p. 167–168', 'key', '0.00', 'heq')
    S.row('d inferior típico (varilla preferida)', '={h}-{rec}-' + dbz('{est_sec}') + '-' + db('{var_inf}') + '/2', 'cm', 'CALCULADA = h − r − estribo − d_b/2. En TABLEROS se recalcula con la varilla real, la dirección (la del claro largo va encima de la del corto) y el recubrimiento de cada tablero.', 'Geometría', 'key', '0.00', 'd_inf')
    S.row('d superior típico (varilla preferida)', '={h}-MAX({rec},({t}-2*{fim})/2)-2*{fim}-' + dbz('{est_sec}') + '-' + db('{var_sup}') + '/2', 'cm', 'CALCULADA = h − máx(r, (t − 2Ø)/2) − 2·Ø malla − estribo − d_b/2: la varilla va bajo la malla, que va al centro de la capa. Con r = 2 cm manda la malla centrada; con r_np = 4 cm (tableros que no cumplen 6.7.8.1.4) se descuenta además el espesor de la malla, del lado seguro. La reducción de 20 mm de 3.9.1 d) (d ≤ 200 mm sin la precaución anotada) se revisa renglón por renglón en TABLEROS.', 'Geometría · ' + NTCC + ' 3.9.1 d) y 3.9.2, p. 74', 'key', '0.00', 'd_sup')
    S.row('p_b — cuantía balanceada', "={fpc}/{fy}*6000*{beta1}/({fy}+6000)", '—', "CALCULADA = (f''c/fy)·6000β1/(fy + 6000).", NTCC + ' ec. 5.2.1.3.1, p. 88', 'calc', '0.0000', 'pb')
    S.row('p para ε_t = 0.005 (sección controlada por tensión)', "=0.375*{beta1}*{fpc}/{fy}", '—', "CALCULADA: c ≤ 0.375·d ⇒ a ≤ 0.375·β1·d ⇒ p ≤ 0.375·β1·f''c/fy, sobre el ancho que trabaja a compresión.", NTCC + ' Tabla 3.8.2.2, p. 70 · 3.8.2.4, p. 71 (ε_ty = 0.002) · 6.3.3.3.1, p. 121', 'calc', '0.0000', 'pten')
    S.row('p_max — cuantía máxima', '=MIN({pbf}*{pb},{pten})', '—', 'CALCULADA = el menor de 0.90·p_b y el de ε_t = 0.005 (manda el segundo).', NTCC + ' 6.3.5.2.1, p. 122 · 6.3.3.3.1, p. 121 · 6.7.3.3.1, p. 169', 'key', '0.0000', 'pmax')
    S.row('p_min — cuantía mínima por flexión (vigas y nervaduras)', '=MAX(0.8*SQRT({fc}),14)/{fy}', '—', "CALCULADA = el mayor de 0.80·√f'c/fy y 14/fy, sobre b_w·d. Con f'c = 250 manda 14/fy = 0.00333. Se dispensa si la As colocada es ≥ 4/3 de la requerida.",
          NTCC + ' 6.3.5.1.2 y 6.3.5.1.3, p. 122', 'calc', '0.00000', 'pmin')
    S.row('As mín de losa por nervadura = 0.0018·A_g', '=0.0018*({s}*{t}+{bw}*{hc})', 'cm²', 'CALCULADA. La cubre de sobra la No. 4 corrida (1.27 cm²).', NTCC + ' 6.7.6.1.1, p. 175', 'calc', '0.00', 'Asl')
    S.row('WMOD para ETABS — modificador de peso de la LR-25 con la sección nueva', '={Vm2}*{gc}*1000/(1503.75*{h}/100)', '—',
          'CALCULADA = V·γc / (1503.75 kg/m³ × h), con el material «LR-25 equiv» que ya trae el modelo (sin casetón, como el modelo). El modelo usa el mismo valor en los modificadores de membrana F11, F22 y F12: seguir ese criterio. Decisión D5.',
          E2K + ' · SHELLPROP "LR-25" (WMOD 0.6688 con la sección vieja) · Decisión D5', 'key', '0.0000', 'wmod')
    S.row('Carga muerta extra de la dala de borde (D9) por metro de borde discontinuo', '=({b_borde}-{bp})/100*({h}/100-{Vm2})*{gc}', 't/m',
          'CALCULADA = (b_borde − b_p)·(h − V)·γc: el concreto de la dala de borde que no tendría la losa nervada (servicio). Baja directo al muro (súmala en su revisión, MUROS) o a la nervadura principal de ese borde (PRINCIPALES ya la suma en «w peso extra»); '
          'en los momentos del tablero se desprecia porque está sobre el apoyo. El macizo de orilla que deja el acomodo de casetones (TABLEROS) puede ser algo más ancho; como en los demás bordes, su peso se toma con el promedio de la losa.',
          'Decisión D9 · Geometría', 'calc', '0.000', 'Wborde')
    S.gap()
    S.section('4b. SECCIÓN DE LOS VOLADOS (D10) Y DE LAS LOSAS MACIZAS DE AZOTEA (D11)')
    S.secnote('Volados F80012 y F80019: misma LR-25 (h_c = 19, t = 6), con las nervaduras de la dirección del volado de b_w = 20 cm y casetón 50×60×19; las paralelas al muro siguen de 10 cm. '
              'Las losas macizas son de h = 20 cm. Los modificadores de abajo son los que van en ETABS (D5).', 30)
    S.header()
    S.row('Volados — volumen de concreto por m²', '={h}/100-{ac_vol}/100*{ac}/100*{hc}/100/({s}/100)^2', 'm³/m²', 'CALCULADA = h − (s − b_w)·a_c·h_c / s².', 'Geometría · Decisión D10', 'calc', '0.0000', 'Vm2v')
    S.row('Volados — peso propio geométrico', '=({Vm2v}*{gc}+{ac_vol}/100*{ac}/100*{hc}/100/({s}/100)^2*{gpol})*1000', 'kg/m²', 'CALCULADA = V·γc + casetón·γ_pol. TABLEROS lo calcula renglón por renglón con la b_w de cada tablero.', 'Geometría · Decisión D10', 'key', '0.0', 'PPv')
    S.row('Volados — A_T de una nervadura con su patín (b_w del volado)', '={s}*{t}+{bw_vol}*{hc}', 'cm²', 'CALCULADA = s·t + b_w·h_c.', 'Geometría', 'calc', '0.0', 'ATv')
    S.row('Volados — ȳ medido desde arriba', '=({s}*{t}*{t}/2+{bw_vol}*{hc}*({t}+{hc}/2))/{ATv}', 'cm', 'CALCULADA.', 'Geometría', 'calc', '0.00', 'ybarv')
    S.row('Volados — I_T de una nervadura con su patín', '={s}*{t}^3/12+{s}*{t}*({ybarv}-{t}/2)^2+{bw_vol}*{hc}^3/12+{bw_vol}*{hc}*({t}+{hc}/2-{ybarv})^2', 'cm⁴', 'CALCULADA (Steiner).', 'Geometría', 'calc', '#,##0', 'ITv')
    S.row('Volados — modificador de inercia en la dirección del volado (M11 o M22)', '={ITv}/({s}*{h}^3/12)', '—', 'CALCULADA. Con el eje local 1 de la losa en X (el que pone ETABS por omisión en losas horizontales): F80012 (volado en X) lleva éste en M11 y el de la LR-25 (sección 4) en M22; F80019 (volado en Y) lleva el de la LR-25 en M11 y éste en M22. En M12 (torsión) va el de la LR-25 en los dos, que es el menor (lado conservador).', 'Geometría · Decisión D5', 'key', '0.0000', 'Iratv')
    S.row('Volados — WMOD para ETABS', '={Vm2v}*{gc}*1000/(1503.75*{h}/100)', '—', 'CALCULADA con el material «LR-25 equiv» del modelo (sin casetón). Crear en ETABS una losa «LR-25V» para los dos volados.', E2K + ' · Decisión D5', 'key', '0.0000', 'wmodv')
    S.row('Losas macizas — peso propio geométrico', '={h_mac}/100*{gc}*1000', 'kg/m²', 'CALCULADA = h·γc. En ETABS «Slab1» sin modificadores, con el concreto f\'c = 250 (hoy trae 4000 psi).', E2K + ' · SHELLPROP "Slab1" · Decisión D11', 'key', '0.0', 'PPm')
    S.gap()
    S.section('5. CARGAS VIVAS UNITARIAS E INCREMENTOS DE CARGA MUERTA')
    S.secnote('Valores de la tabla de cargas vivas de la NTC. Son los mismos que trae el modelo en cada tablero y los que usa Naim. En TABLEROS cada tablero dice si es HABITACIÓN o AZOTEA y de ahí toma W_m (diseño) y W (flechas diferidas).')
    S.header()
    S.row('Habitación — W_m (máxima, para diseño por gravedad)', 190, 'kg/m²', 'CAPTURA.', NTCCA + ' Tabla 6.1.2.2 inciso a), p. 11 · 6.1.2.1 a), p. 10 · ' + NAIM24 + ' pág. 5 · ' + E2K + ' AREALOAD "Carga Viva Máxima" 190', 'input', '0', 'cv_hab')
    S.row('Habitación — W_a (instantánea)', 100, 'kg/m²', 'CAPTURA. Informativa: no entra en el diseño por gravedad.', NTCCA + ' Tabla 6.1.2.2 a), p. 11 · 6.1.2.1 b), p. 10', 'input', '0', 'cva_hab')
    S.row('Habitación — W (media, para flechas diferidas)', 80, 'kg/m²', 'CAPTURA.', NTCCA + ' Tabla 6.1.2.2 a), p. 11 · 6.1.2.1 c), p. 10', 'input', '0', 'cvw_hab')
    S.row('Azotea con pendiente ≤ 5% — W_m', 100, 'kg/m²', 'CAPTURA.', NTCCA + ' Tabla 6.1.2.2 inciso h), p. 11 · ' + NAIM24 + ' pág. 6 · ' + E2K + ' AREALOAD 100', 'input', '0', 'cv_az')
    S.row('Azotea con pendiente ≤ 5% — W_a', 70, 'kg/m²', 'CAPTURA. Informativa.', NTCCA + ' Tabla 6.1.2.2 h), p. 11', 'input', '0', 'cva_az')
    S.row('Azotea con pendiente ≤ 5% — W (media)', 15, 'kg/m²', 'CAPTURA.', NTCCA + ' Tabla 6.1.2.2 h), p. 11', 'input', '0', 'cvw_az')
    S.row('Incremento del peso propio por losa colada en el lugar', 20, 'kg/m²', 'CAPTURA. Se suma al peso propio geométrico en todos los tableros.', NTCCA + ' 5.1.2.1, p. 9', 'input', '0', 'inc_losa')
    S.row('Incremento por capa de mortero sobre la losa (pegapiso, firme o entortado)', 20, 'kg/m²',
          'CAPTURA. Decisión D4: se cuentan los 20 + 20 kg/m² además de la sobrecarga del modelo (160 a 315 kg/m²), porque no se puede confirmar que ésta ya traiga la «carga adicional 0.04 t/m²» de Naim: lado seguro. No se aplica donde la carga muerta es favorable.',
          NTCCA + ' 5.1.2.1 y 5.1.2.2, p. 9 · ' + NAIM24 + ' pág. 5 (carga adicional 0.04 t/m²)', 'input', '0', 'inc_mort')
    S.row('W_g — granizo en azotea (Grupo B, pendiente < 5°)', '=900*0.125*1', 'kg/m²', 'CALCULADA = 900·H_s·C_s con H_s = 0.125 m y C_s = 1.0. Va con FC = 1.1 y sin carga viva.',
          NTCCA + ' ec. 7.3.2.1, p. 17 · Tablas 7.3.2.2 y 7.3.2.3, p. 18', 'calc', '0.0', 'Wg')
    S.gap()
    S.section('6. MUROS DE PLANTA ALTA QUE QUEDAN SOBRE LA LOSA')
    S.secnote('Los muros del primer nivel que no tienen muro abajo descargan en las nervaduras principales, con su propio peso y con la azotea que cargan. Su peso por metro sale de aquí.')
    S.header()
    S.row('γm — peso volumétrico del block', 1.9, 't/m³', 'CAPTURA. El mismo del modelo.', E2K + " · MATERIAL \"f'm=50kg/cm2...\" WEIGHTPERVOLUME 1900", 'input', '0.00', 'gm')
    S.row('t_m — espesor del muro', 0.15, 'm', 'CAPTURA. Mínimo 120 mm en estructuras Tipo I.', E2K + ' · SHELLPROP "Muro 15 cm" · ' + NTCM + ' 6.1.3, p. 60', 'input', '0.00', 'tm')
    S.row('Altura de entrepiso', 3.0, 'm', 'CAPTURA.', E2K + ' · STORY "N+6.00" HEIGHT 3', 'input', '0.00', 'hent')
    S.row('Altura del muro sobre la losa', '={hent}-{h}/100', 'm', 'CALCULADA = entrepiso − h. H/t = 18 ≤ 25.', 'Geometría · ' + NTCM + ' 6.1.3, p. 60', 'calc', '0.00', 'hm')
    S.row('Aplanados (las dos caras)', 0.06, 't/m²', 'CAPTURA. 2 cm por cara con mortero de 1.5 t/m³.', 'Práctica de ingeniería', 'input', '0.00', 'apl')
    S.row('w_muro — peso del muro por metro lineal (servicio)', '=({gm}*{tm}+{apl})*{hm}', 't/m', 'CALCULADA = (γm·t + aplanados)·altura. Los muros divisorios se cuantifican aparte de la carga viva.', 'Geometría · ' + NTCCA + ' 6.1.1.2 y 6.1.1.3, p. 10', 'key', '0.000', 'wmuro')
    S.gap()
    S.section('7. NERVADURAS PRINCIPALES — VALORES POR OMISIÓN')
    S.secnote('Cada tramo de PRINCIPALES arranca con estos valores; ahí mismo se pueden cambiar renglón por renglón.')
    S.header()
    S.row('No. de varilla longitudinal preferida', 4, '—', 'CAPTURA. Mínimo 2 No. 4 corridas en cada lecho donde haya tensión. Si no cabe, la hoja sube al siguiente número.', NTCC + ' 6.3.5.1.1, p. 122', 'input', 'General', 'var_np')
    S.row('No. de varilla del estribo', 3, '—', 'CAPTURA. Estribo cerrado mínimo No. 3, con dobleces a 135° y extensión de 6d_b ≥ 8 cm.', NTCC + ' 14.7.3.3 a), p. 375 · 14.3.4 a) y Tabla 14.3.2, p. 350', 'input', 'General', 'est_np')
    S.row('Mínimo de varillas por lecho', 2, 'pzas', 'CAPTURA.', NTCC + ' 6.3.5.1.1, p. 122 · 6.3.7.7.1, p. 134', 'input', '0', 'nmin_np')
    S.row('Separación máxima de estribos (tope de obra)', 20, 'cm', 'CAPTURA. Criterio propio. La norma pide d/2 ≤ 60 cm (o d/4 ≤ 30 cm si V_s > 1.1·√f\'c·b_w·d); con d > 40 cm manda este tope de 20 cm y con d ≤ 40 cm manda d/2.', 'Criterio de esta hoja · ' + NTCC + ' Tabla 6.3.7.6.2.2, p. 132', 'input', '0', 'smax_np')
    S.gap()
    S.section('8. LOSAS MACIZAS DE AZOTEA (F80032 Y F80010) — DECISIÓN D11')
    S.secnote('Se diseñan con el mismo método de coeficientes de la Tabla 3.4.2.1 (que es para losas macizas), por franja de 1 m, con r = 2 cm (lecho superior con r_ext si la azotea se marca expuesta en la sección 1). '
              'En TABLEROS son los renglones con «MACIZA» en la columna «Losa»; su armado sale como «No. x @ y cm».', 30)
    S.header()
    S.row('No. de varilla preferida en las losas macizas', 3, '—', 'CAPTURA. Si con ella la separación sale menor que la mínima de abajo, la hoja sube sola al siguiente número.', NTCC + ' 6.7.6.1.1, p. 175', 'input', 'General', 'var_mac')
    S.row('Separación mínima práctica en las losas macizas', 10, 'cm', 'CAPTURA. Criterio de obra; la libre mínima de 14.2.1 (2.5 cm) queda muy por debajo.', 'Criterio de esta hoja · ' + NTCC + ' 14.2.1, p. 347', 'input', '0', 'smin_mac')
    S.row('Separación máxima en las losas macizas (tope absoluto)', 45, 'cm',
          'CAPTURA. TABLEROS usa el menor de 2h y este valor en todas las secciones (lado seguro: es el tope de las secciones críticas de losas en dos direcciones y queda abajo de 5h y 45 cm del acero de temperatura y del diafragma).',
          'Criterio de esta hoja · ' + NTCC + ' 6.8.7.3, p. 196 (5t y 45 cm) · 13.8.2.1.1, p. 345', 'input', '0', 'smax_mac')
    S.row('As mín por metro = 0.0018·b·h (cada dirección, cerca de la cara en tensión)', '=0.0018*100*{h_mac}', 'cm²/m', 'CALCULADA. Contracción y temperatura y acero mínimo de losa; TABLEROS la aplica en los cuatro lechos.', NTCC + ' 6.7.6.1.1, p. 175 · 13.8.2.1.1, p. 345', 'calc', '0.00', 'Asm_mac')
    S.row('Desnivel del lecho inferior respecto a la LR-25 vecina', '={h}-{h_mac}', 'cm', 'CALCULADA = h − h_maciza. La cara superior queda al mismo nivel; abajo, la maciza queda más alta: anotarlo en planos (cimbra y plafón).', 'Decisión D11', 'key', '0')
    S.row('Dala de borde en los bordes discontinuos de las macizas (D9)', '=TEXT({b_borde},"0")&" × "&TEXT({h},"0")&" cm (baja "&TEXT({h}-{h_mac},"0")&" cm bajo la maciza)"', '—',
          'CALCULADA. Igual que en la LR-25: en todo borde discontinuo la maciza remata en una dala de borde armada de b_borde × h, al ras del lecho inferior de la LR-25 vecina. Su paño interior es el apoyo de la losa: ahí se anclan el bastón superior y la varilla inferior (NERVADURA sección 10). '
          'Sin ella, en el muro de 15 cm sólo quedan b_p − r y la No. 3 no se desarrolla. Si el borde está sobre una nervadura principal (NP-33), esa principal es la dala de borde: se ensancha a b_borde en el peralte de la losa.',
          'Decisión D9 · Decisión D11 · ' + NTCC + ' 6.7.7.4.1.1, p. 179 · 14.4.3.1, p. 355', 'calc', '@')
    S.gap()
    S.notice('DECISIONES TOMADAS (D1 a D12, RESUMEN sección 2; propuesta del calculista auxiliar para que la revise y firme el Responsable): castillos C1…C8 a cimentación (D1; C7 es nuevo); secciones propuestas de las principales, y en NP-15 y NP-26 (H/L < 0.6) la trabe toma todo el muro sin efecto de arco (D2); '
             'capa de 6 cm con casetón de 19 cm y malla 6x6-4/4 al centro de la capa (D3); se cuentan los 20 + 20 kg/m² de 5.1.2.1 (D4); tableros chicos como losa sobre vigas (D7); distribución de 3.4.3.1 donde no se cumple 3.4.2.1 c) (D8); dala de borde armada de 40 cm en los bordes discontinuos (D9); '
             'volados nervados con nervaduras de 20 cm que siguen dentro de sus tableros de respaldo, con vigas de orilla en los costados libres de F80012 (D10); F80032 y F80010 de azotea macizas de 20 cm (D11); cara superior de los volados protegida: terraza con impermeabilizante y piso (D12). '
             'LO QUE QUEDA DE TU LADO: (1) actualizar y volver a correr el modelo de ETABS (D5) y comparar; (2) pasar a planos las notas de D6 (RESUMEN sección 7); (3) coordinar con arquitectura el castillo nuevo C7 (llegada de la escalera, sin castillo arriba) '
             'y revisar NP-07 y NP-30 con el diseño de la escalera, que está fuera del alcance; (4) revisar y firmar D1 a D12.', 90)
    add_cf(ws, 'C1:F%d' % S.r)
    ws.sheet_view.zoomScale = 85
    return ws


# =====================================================================
def sheet_apoyos(wb):
    ws = wb.create_sheet('APOYOS')
    S = Narr(ws, N, {'A': 1.78, 'B': 30, 'C': 12, 'D': 9, 'E': 10, 'F': 30, 'G': 15, 'H': 30, 'I': 40}, last='I')
    S.title('APOYOS — DÓNDE SE APOYA CADA TABLERO, CUÁLES SON LAS NERVADURAS PRINCIPALES Y DÓNDE HAY VACÍOS')
    S.subtitle('Se cruzó el modelo (muros, castillos y dalas por nivel) con el plano arquitectónico (capa 004_Vacios y escalera) y con el DXF de áreas tributarias. '
               'Los muros que cargan la losa N+3.00 son los de planta baja; los que cargan la azotea N+6.00 son los del primer nivel.', 30)
    S.gap()
    S.section('A. LO QUE SE ENCONTRÓ')
    for txt in [
        '1) AZOTEA (N+6.00): casi todos los bordes caen sobre muros del primer nivel. Sólo 4 tramos cortos quedan sin muro (bordes de los tragaluces, la ventana del área de tendido y el borde del vacío de la escalera) y todos llegan a muros en sus dos extremos.',
        '2) ENTREPISO (N+3.00): en planta baja sólo hay muros en los ejes A y J y en dos núcleos (baño junto a la cochera y servicios atrás). La cochera mide 8.70 m libres y la sala tiene el vacío de doble altura: muchas dalas del modelo quedan sin muro abajo y en realidad son NERVADURAS PRINCIPALES que trabajan como trabe.',
        '3) En el modelo tal cual, varias de esas principales rematan en nudos que no tienen nada abajo: las esquinas de la cochera (x = 4.327 y = 1.919; x = 5.923 y = 2.921), las esquinas del vacío (x = 4.327 con y = 7.871 y 14.021), la esquina de la escalera (F-8) y la esquina C-10. Así quedan en volado de 4 a 8 m: con 25 cm de peralte NO funcionan.',
        '4) En esos mismos nudos el primer nivel SÍ tiene castillo, pero no baja a planta baja (31 castillos sin continuidad, lista F). Decisión D1: se bajan 7 de ellos a cimentación y se agrega uno nuevo, C7, en la llegada de la escalera (lista B, ACEPTADOS); C7 no tiene castillo arriba: coordinarlo con arquitectura. Con eso todas las principales quedan apoyadas en sus dos extremos. Toda bajada que no tenga muro abajo en planta baja es columna de 25×25 como mínimo (NTC-Concreto 2023 (CDMX) 6.4.2.1.1, p. 138).',
        '5) F80012 (área de tendido, fuera del eje A) y F80019 (terraza entre los ejes 14 y 15) son VOLADOS de 2.00 m. Sus orillas libres llevan nervadura de borde, no trabe. La de F80019 recibe el barandal de 333.7 kg/m que trae el modelo. Decisión D10: siguen nervados (LR-25) con las nervaduras de la dirección del volado de 20 cm (casetón 50×60×19), que siguen dentro de sus tableros de respaldo hasta donde la nervadura de 10 cm ya resiste (TABLEROS, «RESPALDO»). Los costados de F80012 (y = 14.021 y 17.446) corren en la dirección del volado y llevan también barandal: NP-09 y NP-14 se diseñan como vigas de orilla en volado (PRINCIPALES). F80019 tiene muros en sus dos costados (DM-43 y DM-53).',
        '6) VACÍOS: doble altura de la sala (4.17 × 6.00 m), cubo de la escalera curva, vacío sobre la escalera en azotea y tres tragaluces de 0.30 m en azotea. F80028 y F80029 son prácticamente hueco (quedan 0.60 y 0.53 m² de losa): se tratan como marco de hueco, no como tablero.',
        '7) En azotea, F80032 y F80010 están en el modelo como «Slab1» (losa maciza de 20 cm, concreto 4000 psi). Decisión D11: se respeta el modelo y se diseñan como LOSAS MACIZAS de 20 cm perimetralmente apoyadas (TABLEROS, columna «Losa»), con f\'c = 250 Clase 1A como el resto de la obra: en ETABS cambiar su concreto de 4000 psi a f\'c = 250.',
        '8) F2 (azotea, sobre la terraza del frente) no tiene cargas en el modelo. Aquí se le pone la misma sobrecarga que al tablero vecino F80023 (270 kg/m²); en ETABS hay que ponérsela (decisión D5).',
    ]:
        S.bullet(txt)
    S.gap()
    cf0 = S.r
    S.section('B. CASTILLOS QUE BAJAN A CIMENTACIÓN — ACEPTADOS (DECISIÓN D1)')
    S.secnote('Decisión D1: los ocho se aceptan y bajan a cimentación; donde una bajada queda sin muro abajo en planta baja es columna de 25×25 como mínimo (NTC-Concreto 2023 (CDMX) 6.4.2.1.1, p. 138). Si pones NO, los tramos que dependen de ese castillo pasan a volado en PRINCIPALES y la hoja te dice si así aguantan. '
              'OJO: si rechazas un castillo que parte un tramo en dos (C2, C4), junta a mano los dos renglones de PRINCIPALES en uno solo con el claro completo.', 42)
    r = S.r
    for k, t in enumerate(['Castillo', '¿Se acepta?', 'x (m)', 'y (m)', 'Ejes', 'Tramos que dependen', 'Por qué', '']):
        hdr_cell(ws, r, 2 + k, t)
    ws.merge_cells('H%d:I%d' % (r, r)); S.r += 1
    dep = {}
    for t in NPS:
        for k in (0, 1):
            p = t['ext'][k].get('prop')
            if p:
                dep.setdefault(p, []).append(t['np'])
    for c in D['castillos_prop']:
        r = S.r
        cell(ws, r, 2, '=IF(C%d="SÍ","%s — ACEPTADO (D1)","%s — RECHAZADO: ver PRINCIPALES")' % (r, c['id'], c['id']), 'calc', al='left', bold=True)
        cell(ws, r, 3, 'SÍ', 'input')
        dv_list(ws, 'C%d' % r, '"SÍ,NO"')
        cell(ws, r, 4, c['x'], 'calc', '0.000'); cell(ws, r, 5, c['y'], 'calc', '0.000')
        cell(ws, r, 6, c['ref'], 'calc', al='left')
        cell(ws, r, 7, ', '.join(dep.get(c['id'], [])), 'calc', al='left', wrap=True)
        ws.merge_cells('H%d:I%d' % (r, r)); cell(ws, r, 8, c['razon'], 'calc', al='left', sz=8, wrap=True); ws.cell(r, 9).border = BOX
        ws.row_dimensions[r].height = 36
        PROPROW[c['id']] = 'APOYOS!$C$%d' % r
        S.r += 1
    S.gap()
    S.section('C. BORDES DE CADA TABLERO — APOYO Y CONTINUIDAD')
    S.secnote('«Apoyo» dice qué hay bajo el borde: MURO, una NERVADURA PRINCIPAL (ver PRINCIPALES) o nada (LIBRE). «Continuidad» dice si del otro lado sigue losa: '
              'con eso TABLEROS escoge el caso de la tabla de coeficientes. La continuidad es AMARILLA: la puedes cambiar si el criterio automático no te convence.', 42)
    r = S.r
    for k, t in enumerate(['Tablero', 'Lado', 'Largo (m)', '% sobre muro', 'Apoyo', 'Continuidad', 'Del otro lado', 'Vacío o hueco junto']):
        hdr_cell(ws, r, 2 + k, t)
    S.r += 1
    for t in TABS:
        for side in SIDES:
            e = t['bordes'][side]; r = S.r
            cell(ws, r, 2, key(t), 'calc', al='left', bold=(side == 'inferior'))
            cell(ws, r, 3, side, 'calc')
            cell(ws, r, 4, e['L'], 'calc', '0.000')
            cell(ws, r, 5, e['frac_muro'], 'calc', '0%')
            cell(ws, r, 6, e['apoyo'], 'calc', al='left', wrap=True)
            cell(ws, r, 7, e['cont'], 'input')
            dv_list(ws, 'G%d' % r, '"CONTINUO,DISCONTINUO"')
            cell(ws, r, 8, e['vecinos'], 'calc', al='left', sz=8, wrap=True)
            cell(ws, r, 9, e['vacio'], 'calc', al='left', sz=8, wrap=True)
            if e.get('ovr'):
                ws.cell(r, 6).comment = Comment('Fijado a mano: el tablero es irregular y se usa su rectángulo representativo (ver TABLEROS).', 'Claude')
            EDGE[(key(t), side)] = r
            S.r += 1
    S.gap()
    S.section('D. VACÍOS Y HUECOS')
    r = S.r
    for k, t in enumerate(['Vacío / hueco', 'Nivel', 'Ancho X (m)', 'Largo Y (m)', 'x de – a (m)', 'y de – a (m)', 'Área (m²)', 'Fuente']):
        hdr_cell(ws, r, 2 + k, t)
    S.r += 1
    for lvl, lst in D['vacios'].items():
        for v in lst:
            xs = [p[0] for p in v['coords']]; ys = [p[1] for p in v['coords']]; r = S.r
            cell(ws, r, 2, v['nombre'], 'calc', al='left', wrap=True); cell(ws, r, 3, lvl)
            cell(ws, r, 4, max(xs) - min(xs), 'calc', '0.00'); cell(ws, r, 5, max(ys) - min(ys), 'calc', '0.00')
            cell(ws, r, 6, '%.3f – %.3f' % (min(xs), max(xs))); cell(ws, r, 7, '%.3f – %.3f' % (min(ys), max(ys)))
            cell(ws, r, 8, v['area'], 'calc', '0.00'); cell(ws, r, 9, v['fuente'], 'ref', al='left', wrap=True)
            ws.row_dimensions[r].height = 30; S.r += 1
    S.gap()
    S.section('E. CÓMO QUEDA SI NO SE BAJA NINGÚN CASTILLO (TAL CUAL EL E2K)')
    S.secnote('Tramos de nervadura principal del N+3.00 que en el modelo actual no tienen apoyo en uno o los dos extremos. Esto es lo que motiva la lista B.')
    for t in MOD['tramos']:
        if t['tipo'] != 'LIBRE' or t['level'] != 'N+3.00' or t.get('cond') not in ('VOLADO', 'SIN APOYO'):
            continue
        if (t['ori'] == 'H' and t['c'] in (14.021, 17.446) and t['s1'] <= 0.01) or (t['ori'] == 'V' and t['c'] == -2.0) or t['L'] < 0.6:
            continue
        eje = ('y = %.3f, de x = %.3f a %.3f' % (t['c'], t['s0'], t['s1'])) if t['ori'] == 'H' else ('x = %.3f, de y = %.3f a %.3f' % (t['c'], t['s0'], t['s1']))
        S.bullet('• %s (%.2f m): %s en el modelo.' % (eje, t['L'], 'VOLADO' if t['cond'] == 'VOLADO' else 'SIN APOYO en ninguno de sus extremos'))
    S.gap()
    S.section('F. CASTILLOS DEL PRIMER NIVEL QUE NO BAJAN A PLANTA BAJA (31)')
    S.secnote('Castillos K-15x15 del primer nivel (entrepiso N+6.00 del E2K) que no tienen castillo ni muro abajo: descargan sobre la losa N+3.00. Los marcados con (D1) bajan a cimentación (lista B).')
    r = S.r
    for k, t in enumerate(['Ubicación', 'x (m)', 'y (m)', '¿Baja a cimentación? (D1)', 'Nota', '', '', '']):
        hdr_cell(ws, r, 2 + k, t)
    S.r += 1
    props = {(round(c['x'], 2), round(c['y'], 2)): c['id'] for c in D['castillos_prop']}
    for c in D['castillos_sin_cont']:
        r = S.r; pid = props.get((round(c['x'], 2), round(c['y'], 2)))
        cell(ws, r, 2, c['ref'], al='left'); cell(ws, r, 3, c['x'], fmt='0.000'); cell(ws, r, 4, c['y'], fmt='0.000')
        cell(ws, r, 5, pid + ' (D1)' if pid else '—')
        ws.merge_cells('F%d:I%d' % (r, r))
        cell(ws, r, 6, 'Descarga sobre una nervadura principal o un tablero del N+3.00; su carga ya viene en el muro de arriba (wall load) de PRINCIPALES.' if not pid else 'Baja a cimentación (lista B, decisión D1).', al='left', sz=8)
        S.r += 1
    add_cf(ws, 'B%d:I%d' % (cf0, S.r))
    ws.freeze_panes = 'C1'
    ws.sheet_view.zoomScale = 85
    return ws


# =====================================================================
def sheet_cargas(wb):
    from gen_b import TC
    ws = wb.create_sheet('CARGAS'); S = Narr(ws, N)
    S.title('CARGAS — PESO PROPIO, SOBRECARGA, CARGA VIVA Y CARGAS LINEALES')
    S.subtitle('La sobrecarga de cada tablero se toma del modelo (varía de 160 a 315 kg/m²) y vive en TABLEROS; encima se cuentan los 20 + 20 kg/m² de 5.1.2.1 (decisión D4). Aquí está el desglose de referencia de Naim, los incrementos de la NTC y las cargas lineales del modelo.', 30)
    S.gap()
    S.section('1. PESO PROPIO DE LA LOSA')
    S.header()
    S.row('Peso propio geométrico LR-25 (con casetón 60×60×19 y capa de 6 cm)', '={PP}', 'kg/m²', 'Enlazado a DATOS sección 4.', 'DATOS', 'calc', '0.0')
    S.row('Incremento por losa colada en el lugar', '={inc_losa}', 'kg/m²', 'DATOS sección 5.', NTCCA + ' 5.1.2.1, p. 9', 'calc', '0')
    S.row('Incremento por capa de mortero', '={inc_mort}', 'kg/m²', 'DATOS sección 5. No se aplica donde la carga muerta es favorable (p. ej. el claro que equilibra un volado).', NTCCA + ' 5.1.2.1 y 5.1.2.2, p. 9', 'calc', '0')
    S.row('Peso propio de diseño', '={PP}+{inc_losa}+{inc_mort}', 'kg/m²', 'Suma de los tres renglones. Es el que usa TABLEROS en la LR-25 (casetón 60×60×19, capa de 6 cm, decisión D3).', NTCCA + ' 5.1.2.1, p. 9', 'key', '0.0', 'PPd')
    S.row('Peso propio de diseño de los volados (nervaduras de 20 cm, casetón 50×60×19)', '={PPv}+{inc_losa}+{inc_mort}', 'kg/m²', 'DATOS sección 4b + los dos incrementos. Decisión D10: más concreto que la LR-25 normal.', NTCCA + ' 5.1.2.1, p. 9 · Decisión D10', 'key', '0.0', 'PPdv')
    S.row('Peso propio de diseño de las losas macizas de azotea (h = 20 cm)', '={PPm}+{inc_losa}+{inc_mort}', 'kg/m²', 'DATOS sección 4b + los dos incrementos. Decisión D11.', NTCCA + ' 5.1.2.1, p. 9 · Decisión D11', 'key', '0.0', 'PPdm')
    S.row('Peso propio que trae el modelo', '={PPe2k}', 'kg/m²', 'Enlazado a DATOS sección 4. Es el de la sección vieja (20 + 5): se actualiza en ETABS (decisión D5).', E2K, 'calc', '0.0')
    S.gap()
    S.section('2. SOBRECARGA MUERTA — DESGLOSE DE REFERENCIA (NAIM)')
    S.secnote('Sólo de referencia: no entra en ningún cálculo. Sirve para juzgar si las sobrecargas del modelo (columna SCM de TABLEROS) son razonables. Ojo: la «carga adicional 0.04 t/m²» de Naim es el mismo incremento de 40 kg/m² de 5.1.2.1.')
    S.header(('Concepto', 'Carga', 'Unidad', 'Cómo se obtiene', 'Referencia'))
    for c, v, how in [('Entrepiso — relleno 8 cm', 0.13, '1.60 t/m³ × 0.08 m'), ('Entrepiso — piso 3 cm', 0.06, '2.00 t/m³ × 0.03 m'), ('Entrepiso — pegapiso 3 cm', 0.05, '1.50 t/m³ × 0.03 m'),
                      ('Entrepiso — yeso 3 cm', 0.05, '1.50 t/m³ × 0.03 m'), ('Entrepiso — instalaciones', 0.02, '—'), ('Entrepiso — carga adicional', 0.04, '—')]:
        S.row(c, v, 't/m²', how, NAIM24 + ' · pág. 5, CARGAS DE ENTREPISO', 'calc', '0.00')
    S.row('Entrepiso — TOTAL según Naim', 0.33, 't/m²', 'Naim reporta 0.33 (subtotal 0.29 + carga adicional 0.04); la suma de los renglones de arriba da 0.35 (0.338 sin redondear).', NAIM24 + ' · pág. 5', 'key', '0.00')
    for c, v, how in [('Azotea — relleno 10 cm', 0.16, '1.60 t/m³ × 0.10 m'), ('Azotea — yeso 3 cm', 0.05, '1.50 t/m³ × 0.03 m'), ('Azotea — impermeabilizante', 0.03, '1.40 t/m³ × 0.02 m'),
                      ('Azotea — instalaciones', 0.02, '—'), ('Azotea — carga adicional', 0.04, '—')]:
        S.row(c, v, 't/m²', how, NAIM24 + ' · pág. 5, CARGAS DE AZOTEA', 'calc', '0.00')
    S.row('Azotea — TOTAL según Naim', 0.29, 't/m²', 'Naim reporta 0.29 (la suma de renglones da 0.30).', NAIM24 + ' · pág. 5', 'key', '0.00')
    col = TC['scm']
    S.row('Sobrecarga del modelo — mínima y máxima en entrepiso', '=TEXT(MIN(TABLEROS!$%s$%d:$%s$%d),"0")&" a "&TEXT(MAX(TABLEROS!$%s$%d:$%s$%d),"0")' % (col, T0, col, T0 + 19, col, T0, col, T0 + 19), 'kg/m²', 'Rango de la columna SCM de TABLEROS para N+3.00.', E2K + ' · AREALOAD "Sobrecarga"', 'calc', '@')
    S.row('Sobrecarga del modelo — mínima y máxima en azotea', '=TEXT(MIN(TABLEROS!$%s$%d:$%s$%d),"0")&" a "&TEXT(MAX(TABLEROS!$%s$%d:$%s$%d),"0")' % (col, T0 + 20, col, T1, col, T0 + 20, col, T1), 'kg/m²', 'Rango de la columna SCM de TABLEROS para N+6.00.', E2K + ' · AREALOAD "Sobrecarga"', 'calc', '@')
    S.gap()
    S.section('3. COMBINACIONES DE DISEÑO')
    S.header()
    S.row('Carga muerta CM = PP de diseño + SCM', 'por tablero', '—', 'TABLEROS, columna CM.', NTCCA + ' 5.1.2.1, p. 9', 'calc', '@')
    S.row('w_u gravitacional = 1.3·CM + 1.5·W_m', 'por tablero', '—', 'TABLEROS, columna w_u. Grupo B. W_m uniforme en todos los tableros a la vez: se permite porque W_m ≤ 0.75·CM.', NTCCA + ' 2.3.1 a), p. 5 y 3.4.1 a), p. 7 · ' + NTCC + ' 3.3.2.2.3, p. 50', 'calc', '@')
    S.row('w_u con granizo (sólo azotea) = 1.1·(CM + W_g)', '={FCacc}*(' + '{PPd}' + '+270+{Wg})', 'kg/m²', 'Ejemplo con SCM = 270 kg/m². En TABLEROS se toma el mayor de las dos combinaciones; en azotea manda la gravitacional.', NTCCA + ' 3.4.1 b), p. 7 · 7.3.1.1 y 7.3.1.3, p. 17', 'calc', '0')
    S.row('w de servicio = CM + W_m', 'por tablero', '—', 'TABLEROS, columna w serv. Factor de carga 1.0.', NTCCA + ' 3.3.2, p. 6 y 3.4.1 d), p. 7', 'calc', '@')
    S.row('Carga sostenida para flechas diferidas = CM + W', 'por tramo', '—', 'PRINCIPALES, revisión de flechas.', NTCCA + ' 6.1.2.1 c), p. 10', 'calc', '@')
    S.gap()
    S.section('4. CARGAS CONCENTRADAS QUE PIDE LA NORMA')
    S.header()
    S.row('Capa entre nervaduras: carga concentrada mínima', 1000, 'kg', 'En un área de 10 × 10 cm en la posición más desfavorable. Se revisa en NERVADURA sección 11 (flexión, punzonamiento y acero mínimo de la capa).', NTCC + ' 6.7.8.1.6, p. 187 · 6.7.8.1.8, p. 188', 'input', '0', 'Pcap')
    S.row('Habitación: carga en lugar de W_m cuando sea más desfavorable', 500, 'kg', 'En 50 × 50 cm. La cubre la revisión de 1 000 kg en 10 × 10 cm.', NTCCA + ' Tabla 6.1.2.2 nota [1], p. 11', 'calc', '0')
    S.row('Azotea: carga concentrada en la posición más crítica', 100, 'kg', 'La cubre la revisión de 1 000 kg. Tinacos y equipos van aparte y deben indicarse en planos.', NTCCA + ' Tabla 6.1.2.2 notas [7] y [8], p. 12', 'calc', '0')
    S.row('Carga transitoria al colar la azotea sobre el entrepiso', 150, 'kg/m²', 'Más el peso del colado y la cimbra, y 150 kg concentrados. Indicar en planos el apuntalamiento a planta baja mientras se cuela N+6.00.', NTCCA + ' 6.1.3.1, p. 10', 'calc', '0')
    S.row('Temperatura en la estructura', 'se omite', '—', 'Dimensiones < 30 m, miembros de ejes rectos y azotea protegida (relleno + impermeabilizante; la norma pide «al menos un enladrillado»: confirmarlo).', NTCCA + ' 6.2.3.1, p. 13', 'calc', '@')
    S.gap()
    S.section('5. CARGAS LINEALES QUE TRAE EL MODELO (BARANDALES, MUROS BAJOS, PRETILES Y ESCALERA)')
    S.secnote('Se suman a la nervadura principal sobre la que caen (PRINCIPALES, columnas SCM y CV lineal). Las de 333.7 kg/m en las orillas libres de F80012 y F80019 entran como carga en la punta de cada nervadura del volado (TABLEROS) y, en los costados de F80012, en las vigas de orilla NP-09 y NP-14 (PRINCIPALES).')
    S.header(('Viga · nivel · caso', 'w', 'Unidad', 'Dónde (de – a) y largo cargado', 'Fuente'))
    for l in D['lineloads']:
        S.row('%s · %s · %s' % (l['beam'], l['level'], l['lc']), l['w'], 'kg/m', 'de (%.3f, %.3f) a (%.3f, %.3f) · %.2f m' % (l['x0'], l['y0'], l['x1'], l['y1'], l['L']),
              E2K + ' · FRAME OBJECT LOADS', 'calc', '0.0', h=19.5)
    add_cf(ws, 'C1:F%d' % S.r)
    return ws


# =====================================================================
# NTC-Concreto 2023, Tabla 3.4.2.1 (§3.4.2, pág. impresa 65), caso I — transcrita y cotejada contra la imagen de la tabla
COEF = {
    'INTERIOR': {'NIC': [998, 1160, 907, 736, 607, 503, 418], 'NIL': [516, 550, 517, 492, 466, 436, 418], 'NDC': [0] * 7, 'NDL': [0] * 7, 'PC': [630, 312, 331, 330, 298, 256, 209], 'PL': [175, 139, 174, 200, 213, 216, 209]},
    'BORDE-CORTO': {'NIC': [998, 1163, 965, 785, 652, 546, 460], 'NIL': [516, 561, 525, 498, 470, 440, 421], 'NDC': [0] * 7, 'NDL': [326, 258, 248, 236, 222, 206, 190], 'PC': [630, 329, 292, 331, 311, 276, 238], 'PL': [179, 232, 263, 266, 266, 263, 286]},
    'BORDE-LARGO': {'NIC': [1060, 1292, 981, 777, 628, 512, 421], 'NIL': [587, 566, 540, 513, 482, 483, 461], 'NDC': [651, 362, 321, 283, 250, 219, 190], 'NDL': [0] * 7, 'PC': [751, 458, 433, 383, 323, 293, 286], 'PL': [185, 217, 237, 248, 250, 243, 238]},
    'ESQUINA': {'NIC': [1060, 1382, 1046, 830, 675, 557, 465], 'NIL': [600, 580, 545, 516, 484, 486, 465], 'NDC': [651, 362, 321, 277, 250, 219, 190], 'NDL': [326, 333, 320, 298, 270, 241, 190], 'PC': [751, 452, 437, 397, 338, 306, 296], 'PL': [191, 298, 294, 292, 288, 274, 289]},
    'EXTREMO-LARGO': {'NIC': [1060, 1470, 1108, 882, 730, 650, 570], 'NIL': [0] * 7, 'NDC': [651, 370, 340, 310, 280, 250, 220], 'NDL': [220, 361, 337, 310, 280, 220, 220], 'PC': [751, 730, 670, 610, 550, 490, 430], 'PL': [185, 430, 430, 430, 430, 430, 430]},
    'EXTREMO-CORTO': {'NIC': [0] * 7, 'NIL': [570, 570, 570, 570, 570, 570, 570], 'NDC': [570, 480, 420, 370, 310, 270, 220], 'NDL': [330, 448, 402, 220, 313, 220, 220], 'PC': [1100, 960, 840, 730, 620, 540, 430], 'PL': [200, 430, 430, 430, 430, 430, 430]},
    'AISLADO': {'NIC': [0] * 7, 'NIL': [0] * 7, 'NDC': [570, 550, 530, 470, 430, 380, 330], 'NDL': [330, 446, 330, 330, 330, 330, 330], 'PC': [1100, 830, 800, 720, 640, 570, 500], 'PL': [200, 500, 500, 500, 500, 500, 500]},
}
CASODESC = {'INTERIOR': 'Interior — todos los bordes continuos', 'BORDE-CORTO': 'De borde — un lado corto discontinuo', 'BORDE-LARGO': 'De borde — un lado largo discontinuo',
            'ESQUINA': 'De esquina — dos lados adyacentes discontinuos', 'EXTREMO-LARGO': 'Extremo — tres bordes discontinuos, un lado largo continuo',
            'EXTREMO-CORTO': 'Extremo — tres bordes discontinuos, un lado corto continuo', 'AISLADO': 'Aislado — cuatro lados discontinuos'}
MOMDESC = {'NIC': 'negativo en bordes continuos, claro corto', 'NIL': 'negativo en bordes continuos, claro largo', 'NDC': 'negativo en bordes discontinuos, claro corto',
           'NDL': 'negativo en bordes discontinuos, claro largo', 'PC': 'positivo, claro corto', 'PL': 'positivo, claro largo'}


def sheet_tablas(wb):
    ws = wb.create_sheet('TABLAS')
    S = Narr(ws, N, {'A': 1.78, 'B': 22, 'C': 46, 'D': 8, 'E': 8, 'F': 8, 'G': 8, 'H': 8, 'I': 8, 'J': 8, 'K': 40}, last='K')
    S.title('TABLAS — COEFICIENTES, VARILLAS, MALLAS Y LISTAS')
    S.subtitle('Tablas auxiliares. Todas las hojas leen de aquí con INDEX/MATCH: si cambias un número aquí, cambia en todo el libro.', 18)
    assert S.r == 3
    S.r = TB_M - 1
    S.section('A. COEFICIENTES DE MOMENTO PARA TABLEROS RECTANGULARES, FRANJAS CENTRALES — CASO I (LOSA COLADA MONOLÍTICAMENTE CON SUS APOYOS) · ×10⁻⁴·w·a1²')
    r = TB_M
    hdr_cell(ws, r, 2, 'Clave'); hdr_cell(ws, r, 3, 'Caso — momento  (m = a1/a2 →)')
    for k, mv in enumerate([0, 0.5, 0.6, 0.7, 0.8, 0.9, 1.0]):
        x = hdr_cell(ws, r, 4 + k, mv); x.number_format = '0.0'
    hdr_cell(ws, r, 11, 'Referencia')
    rr = TB_C0
    for cs in CASOS:
        for mm in MOMS:
            cell(ws, rr, 2, '%s|%s' % (cs, mm), al='left', sz=8)
            cell(ws, rr, 3, '%s — %s' % (CASODESC[cs], MOMDESC[mm]), al='left', sz=8)
            for k, v in enumerate(COEF[cs][mm]):
                cell(ws, rr, 4 + k, v, 'calc', '0')
            cell(ws, rr, 11, NTCC + ' · Tabla 3.4.2.1, caso I (§3.4.2, pág. 65)', 'ref', al='left', wrap=False)
            rr += 1
    assert rr - 1 == TB_C1
    S.r = TB_C1 + 1
    S.notice('FUENTE: NTC-Concreto 2023, Tabla 3.4.2.1 «Coeficientes de momentos flexionantes para tableros rectangulares, franjas centrales», columna I (losa colada monolíticamente con sus apoyos), pág. 65. '
             'Transcrita y cotejada celda por celda contra la imagen de la tabla. Ojo: en 2023 cambiaron cerca de la mitad de los coeficientes respecto a 2017 (en promedio suben ~40%). '
             'Los renglones en cero son momentos que ese caso no tiene. Para m < 0.5 se interpola entre las columnas 0 y 0.5. Para franjas de columna se multiplican por 0.60 (nota 1 de la tabla); aquí no se usan porque cada nervadura se diseña con la franja central (conservador). '
             'Se conservan tal como vienen impresas dos rarezas de renglones NDL que parecen erratas (Extremo corto en m = 0.7: 220 entre 402 y 313; Aislado en m = 0.5: 446). Ningún tablero de esta casa usa esos tramos.', 66)
    S.r = TB_V0 - 2
    S.section('B. VARILLA CORRUGADA — ÁREAS, PESOS, DESARROLLO Y TRASLAPES')
    r = TB_V0 - 1
    for k, t in enumerate(['No.', 'db (cm)', 'As (cm²)', 'kg/m', 'L_d (cm)', 'L_dh gancho (cm)', 'Traslape B (cm)', 'L_d otros casos (cm)']):
        hdr_cell(ws, r, 2 + k, t)
    hdr_cell(ws, r, 10, 'Referencia'); ws.merge_cells('J%d:K%d' % (r, r))
    for i, (n_, d_, a_, w_) in enumerate(BARS):
        rr = TB_V0 + i
        cell(ws, rr, 2, n_, fmt='General'); cell(ws, rr, 3, d_, fmt='0.00'); cell(ws, rr, 4, a_, fmt='0.00'); cell(ws, rr, 5, w_, fmt='0.000')
        if n_ == 2:
            for cc in (6, 7, 8, 9):
                cell(ws, rr, cc, 'no aplica', al='center')
        else:
            cell(ws, rr, 6, '=MAX(%s/(IF(B%d<=6,6.6,5.3)*SQRT(%s))*C%d,30)' % (N['fy'], rr, N['fc'], rr), fmt='0')
            cell(ws, rr, 7, '=MAX(0.075*%s/SQRT(%s)*C%d,8*C%d,15)' % (N['fy'], N['fc'], rr, rr), fmt='0')
            cell(ws, rr, 8, '=MAX(1.3*F%d,30)' % rr, fmt='0')
            cell(ws, rr, 9, '=MAX(%s/(IF(B%d<=6,4.4,3.5)*SQRT(%s))*C%d,30)' % (N['fy'], rr, N['fc'], rr), fmt='0')
        ws.merge_cells('J%d:K%d' % (rr, rr))
        cell(ws, rr, 10, ('NMX-B-506 / ASTM A615M · L_d: NTC-Concreto 14.4.2.1 y Tabla 14.4.2.4, pp. 352–353 · L_dh: 14.4.3.1, p. 355 · traslape clase B: Tabla 14.5.2.1, p. 367' if n_ != 2 else
                               'Alambrón liso 1/4" (No. 2): la Tabla 14.4.2.4 y 14.4.3.1 son para barras corrugadas; no se usa como refuerzo longitudinal'), 'ref', al='left')
    S.r = TB_V1 + 1
    S.secnote('L_d (Tabla 14.4.2.4, p. 353) con separación libre ≥ 2d_b y recubrimiento ≥ d_b: fy·d_b/(6.6·√f\'c) para No. 6 y menores y fy·d_b/(5.3·√f\'c) para No. 7 a No. 11, ≥ 30 cm. '
              '«Otros casos» (separación libre < 2d_b o recubrimiento < d_b, sin estribos): fy·d_b/(4.4·√f\'c) y fy·d_b/(3.5·√f\'c); TABLEROS usa esta columna cuando la nervadura no da la separación. '
              'ψp = 1.0 en nervaduras secundarias (h = 25 cm) y en lechos inferiores; en el lecho superior de nervaduras principales con más de 30 cm de concreto debajo (aquí, h ≥ 40 cm) ψp = 1.3 (Tabla 14.4.2.6, p. 354): L_d y traslape × 1.3. '
              'Para 1.25 fy (barra inferior de integridad de cada nervadura, 6.7.8.1.7) se multiplica por 1.25. '
              'L_dh = el mayor de 0.075·fy·d_b/√f\'c, 8d_b y 15 cm, igual para el gancho estándar de 90° (extensión 12d_b) que para el de 180° (extensión 4d_b ≥ 6.5 cm) (14.4.3.1 a) a c), p. 355; Tabla 14.3.1, p. 349). '
              'En los bordes discontinuos de esta losa se pide el de 180°: el de 90° parado (16d_b de alto) no cabe en 25 cm con 4 cm de recubrimiento arriba y abajo (NERVADURA sección 10).', 54)
    S.r = TB_W0 - 2
    S.section('C. MALLA ELECTROSOLDADA (fy = 5000 kg/cm²; separación 15.24 cm ≤ 5t y ≤ 30 cm entre alambres soldados)')
    r = TB_W0 - 1
    for k, t in enumerate(['Malla', 'Ø alambre (mm)', 'As (cm²/m)', 'kg/m²']):
        hdr_cell(ws, r, 2 + k, t)
    hdr_cell(ws, r, 6, 'Referencia'); ws.merge_cells('F%d:K%d' % (r, r))
    for i, (nm, dmm) in enumerate(MALLAS):
        rr = TB_W0 + i
        cell(ws, rr, 2, nm); cell(ws, rr, 3, dmm, fmt='0.00')
        cell(ws, rr, 4, '=PI()*(C%d/10)^2/4*100/15.24' % rr, fmt='0.000')
        cell(ws, rr, 5, '=2*PI()*(C%d/1000)^2/4*7850/0.1524' % rr, fmt='0.00')
        ws.merge_cells('F%d:K%d' % (rr, rr))
        cell(ws, rr, 6, 'Calibre de alambre (galga) y separación de 6" = 15.24 cm en las dos direcciones · NMX-B-290 · NTC-Concreto 2.4.2.1.6.3, p. 44 y 6.8.7.3, p. 196 · traslape ≥ un cuadro + 5 cm (14.5.4.1, p. 369)', 'ref', al='left')
    S.r = TB_L0 - 1
    S.section('D. LISTAS DE LOS SELECTORES')
    for j, (nm, vals) in enumerate(LISTS):
        hdr_cell(ws, TB_L0, 2 + j, nm)
        for i, v in enumerate(vals):
            cell(ws, TB_L0 + 1 + i, 2 + j, v, al='left')
    ws.column_dimensions['B'].width = 24; ws.column_dimensions['C'].width = 58; ws.column_dimensions['D'].width = 12; ws.column_dimensions['E'].width = 12
    add_cf(ws, 'C1:K%d' % (TB_L0 + 10))
    return ws
