"""Estilos y ayudantes que copian el formato de CimentacionEspectacular_1_12.xlsx."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as CL

FONT = 'Arial'
TITLE_BG = '1F3864'; SEC_BG = '2E5C9A'; HDR_BG = 'D9E2F3'; HDR_FG = '1F3864'
NOTE_BG = 'F2F2F2'; WARN_BG = 'FFF0E0'; WARN_FG = '843C0C'
INP_BG = 'FFF2CC'; INP_FG = '0000CC'; TXT = '1A1A1A'; GRAY = '595959'; REF = '1F4E79'
thin = Side(style='thin', color='A6A6A6')
med = Side(style='medium', color='1F3864')
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HBOX = Border(left=thin, right=thin, top=thin, bottom=med)


def fill(c):
    return PatternFill('solid', fgColor=c)


def font(sz=10, b=False, i=False, color=TXT):
    return Font(name=FONT, size=sz, bold=b, italic=i, color=color)


def add_cf(ws, rng):
    """Mismas reglas del ejemplo: rojo para lo que falla, verde para lo que cumple; ámbar para avisos."""
    red = ('OR(ISNUMBER(FIND("NO CUMPLE",{c})),ISNUMBER(FIND("REVISAR",{c})),ISNUMBER(FIND("NO CABE",{c})),'
           'ISNUMBER(FIND("FUERA",{c})),ISNUMBER(FIND("SIN APOYO",{c})),ISNUMBER(FIND("INSUFICIENTE",{c})),ISNUMBER(FIND("PENDIENTE",{c})))')
    grn = 'OR(ISNUMBER(FIND("CUMPLE",{c})),ISNUMBER(FIND("SÍ aplica",{c})),ISNUMBER(FIND("MANDA",{c})))'
    amb = 'OR(ISNUMBER(FIND("MACIZAR",{c})),ISNUMBER(FIND("VOLADO",{c})),ISNUMBER(FIND("PROPUESTO",{c})),ISNUMBER(FIND("VER ",{c})))'
    first = rng.split(':')[0]
    ws.conditional_formatting.add(rng, FormulaRule(formula=[red.format(c=first)], stopIfTrue=True,
                                  font=Font(color='FFFFFF', bold=True), fill=PatternFill(bgColor='C00000', fill_type='solid')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[grn.format(c=first)], stopIfTrue=True,
                                  font=Font(color='FFFFFF', bold=True), fill=PatternFill(bgColor='00A650', fill_type='solid')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[amb.format(c=first)], stopIfTrue=True,
                                  font=Font(color='7F4F00', bold=True), fill=PatternFill(bgColor='FFD966', fill_type='solid')))


class Narr:
    """Hoja narrativa de columnas B..F: Concepto | Valor | Unidad | Cómo se calcula | Referencia."""

    def __init__(self, ws, names, widths=None, last='F'):
        self.ws = ws; self.N = names; self.r = 1; self.last = last
        ws.sheet_view.showGridLines = False
        w = widths or {'A': 1.78, 'B': 46, 'C': 17, 'D': 9, 'E': 54, 'F': 44}
        for k, v in w.items():
            ws.column_dimensions[k].width = v

    def merged(self, text, f, fl=None, h=None, al=None, border=False):
        ws = self.ws; r = self.r
        ws.merge_cells('B%d:%s%d' % (r, self.last, r))
        c = ws['B%d' % r]; c.value = text; c.font = f
        if fl:
            c.fill = fill(fl)
        c.alignment = al or Alignment(horizontal='left', vertical='top', wrap_text=True)
        if border:
            for col in range(2, ord(self.last) - 63):
                ws.cell(r, col).border = BOX
                if fl:
                    ws.cell(r, col).fill = fill(fl)
        if h:
            ws.row_dimensions[r].height = h
        self.r += 1
        return r

    def title(self, text):
        r = self.merged(text, font(13, True, color='FFFFFF'), TITLE_BG, 26, Alignment(horizontal='left', vertical='center'))
        return r

    def subtitle(self, text, h=30):
        return self.merged(text, font(9, i=True, color=GRAY), None, h)

    def notice(self, text, h=42):
        return self.merged(text, font(9, True, color=WARN_FG), WARN_BG, h, border=True)

    def gap(self, h=6):
        self.ws.row_dimensions[self.r].height = h; self.r += 1

    def section(self, text):
        return self.merged(text, font(10, True, color='FFFFFF'), SEC_BG, 21, Alignment(horizontal='left', vertical='center'))

    def secnote(self, text, h=30):
        return self.merged(text, font(9, i=True, color=GRAY), NOTE_BG, h, border=True)

    def header(self, cols=('Concepto', 'Valor', 'Unidad', 'Cómo se calcula / de dónde sale', 'Referencia'), h=30):
        ws = self.ws; r = self.r
        for k, t in enumerate(cols):
            c = ws.cell(r, 2 + k, t); c.font = font(9, True, color=HDR_FG); c.fill = fill(HDR_BG)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = HBOX
        ws.row_dimensions[r].height = h; self.r += 1
        return r

    def row(self, concept, value, unit='', how='', ref='', kind='calc', fmt='0.00', name=None, h=None, bold=False):
        """kind: input (amarillo/azul), calc (negro), key (resultado clave, azul pálido), text (texto calculado)."""
        ws = self.ws; r = self.r
        c = ws.cell(r, 2, concept); c.font = font(10, bold); c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True); c.border = BOX
        if isinstance(value, str) and value.startswith('='):
            value = value.format(**self.N)
        v = ws.cell(r, 3, value); v.border = BOX; v.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        if kind == 'input':
            v.font = font(10, True, color=INP_FG); v.fill = fill(INP_BG)
        elif kind == 'key':
            v.font = font(10, True, color=HDR_FG); v.fill = fill(HDR_BG)
        else:
            v.font = font(10)
        v.number_format = fmt
        u = ws.cell(r, 4, unit); u.font = font(9, color=GRAY); u.alignment = Alignment(horizontal='center', vertical='center'); u.border = BOX
        if isinstance(how, str) and how.startswith('='):
            how = how.format(**self.N)
        e = ws.cell(r, 5, how); e.font = font(9, color=GRAY); e.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True); e.border = BOX
        f = ws.cell(r, 6, ref); f.font = font(8, i=True, color=REF); f.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True); f.border = BOX
        vlen = 22.0 if (isinstance(value, str) and value.startswith('=') and kind == 'key' and '@' in fmt) else 0
        n = max(len(how) / 58.0, len(ref) / 52.0, len(concept) / 46.0, 1, vlen / 14.0)
        ws.row_dimensions[r].height = h or max(19.5, 13.5 * (int(n) + (1 if n > int(n) else 0)) + 3)
        if name:
            self.N[name] = "%s!$C$%d" % (q(ws.title), r)
        self.r += 1
        return r

    def bullet(self, text, h=None):
        n = len(text) / 150.0
        return self.merged(text, font(9, color=TXT), None, h or max(15, 13 * (int(n) + 1) + 2))


def q(title):
    return "'%s'" % title if any(ch in title for ch in ' -') else title


def hdr_cell(ws, r, c, text, w=None):
    x = ws.cell(r, c, text); x.font = font(9, True, color=HDR_FG); x.fill = fill(HDR_BG)
    x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); x.border = HBOX
    if w:
        ws.column_dimensions[CL(c)].width = w
    return x


def unit_cell(ws, r, c, text):
    x = ws.cell(r, c, text); x.font = font(8, i=True, color=GRAY); x.fill = fill('EEF2F9')
    x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); x.border = BOX
    return x


def cell(ws, r, c, value, kind='calc', fmt='General', al='center', sz=9, bold=False, wrap=False):
    x = ws.cell(r, c, value); x.border = BOX
    x.alignment = Alignment(horizontal=al, vertical='center', wrap_text=wrap)
    if kind == 'input':
        x.font = font(sz, True, color=INP_FG); x.fill = fill(INP_BG)
    elif kind == 'key':
        x.font = font(sz, True, color=HDR_FG); x.fill = fill(HDR_BG)
    elif kind == 'data':
        x.font = font(sz, color='006100')
    elif kind == 'ref':
        x.font = font(8, i=True, color=REF)
    else:
        x.font = font(sz, bold)
    x.number_format = fmt
    return x


def dv_list(ws, rng, formula):
    dv = DataValidation(type='list', formula1=formula, allow_blank=False)
    dv.error = 'Escoge un valor de la lista'; dv.errorTitle = 'Valor no válido'
    ws.add_data_validation(dv); dv.add(rng)
    return dv
