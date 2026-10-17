"""Parte B: TABLEROS, NERVADURA, PRINCIPALES, TRIB, TRANSFER, MUROS."""
from gen_a import *
from openpyxl.utils import column_index_from_string as CI

VOLB = {'F80012': 3.029, 'F80019': 4.0}   # claro de respaldo (m) del volado: tablero más corto detrás del muro raíz
for _t in TABS:
    if _t['modo'] == 'VOLADO':
        _t['volado'].setdefault('Lb', VOLB[_t['name']])


# ---------------------------------------------------------------- TABLEROS
TC = {}
BLKS = ('Pc', 'Nc', 'Pl', 'Nl')
OPPS = {'inferior': 'superior', 'superior': 'inferior', 'izquierdo': 'derecho', 'derecho': 'izquierdo'}
MEC = {'inferior': 'MeI', 'superior': 'MeS', 'izquierdo': 'MeL', 'derecho': 'MeR'}
_tcols = ['key', 'nivel', 'modo', 'dest', 'losa', 'Lx', 'Ly', 'ax', 'ay', 'a1', 'a2', 'm', 'dirc', 'ln1', 'ln2', 'hr', 'bwt', 'sb', 'bfc', 'bfl', 'bNc', 'bNl', 'rb', 'rt',
          'scm', 'cv', 'pp', 'cm', 'ws', 'wu',
          'nd', 'ndc', 'ndl', 'cauto', 'cforce', 'caso', 'idx', 'fr', 'cNIC', 'cNIL', 'cNDC', 'cNDL', 'cPC', 'cPL', 'MeI', 'MeS', 'MeL', 'MeR', 'kX', 'kY', 'MNcd', 'MNld', 'MPcd', 'MPld', 'dVv', 'f342', 'c342', 'pbor', 'MPc', 'MNc', 'MPl', 'MNl']
for _b in BLKS:
    _tcols += ['dq' + _b, 'As' + _b, 'v' + _b, 'n' + _b, 's' + _b, 'd' + _b, 'MR' + _b, 'FU' + _b, 'e' + _b]
_tcols += ['g678', 'Vu', 'VcR', 'cort', 'afm', 'hmin', 'heqv', 'peral', 'tdia', 'Lbc', 'Lbl', 'ncx', 'ncy', 'apoy', 'estado', 'nota']
for i, nm in enumerate(_tcols):
    TC[nm] = CL(2 + i)
TC_LAST = CL(2 + len(_tcols) - 1)


def tcol(nm):
    return "TABLEROS!$%s$%d:$%s$%d" % (TC[nm], T0, TC[nm], T1)


def asflex(M, b, d, neg):
    """As por flexión (cm²) para M en t·m; -1 si la sección no alcanza o rebasa p_max (ε_t ≥ 0.005)."""
    rad = '(1-2*%s*1E5/({FRf}*%s*%s^2*{fpc}))' % (M, b, d)
    asf = '(1-SQRT(%s))*{fpc}*%s*%s/{fy}' % (rad, b, d)
    return 'IF(%s<=0,0,IF(%s<0,-1,IF(%s>{pmax}*%s*%s,-1,%s)))' % (M, rad, asf, b, d, asf)


def asreq(M, b, d, neg, bmin='{bw}'):
    """As requerida = el mayor de la de flexión y el menor de (As mín de 6.3.5.1.2, 4/3 de la de flexión) — 6.3.5.1.3."""
    f = asflex(M, b, d, neg)
    return 'IF(%s<0,-1,MAX(%s,MIN({pmin}*%s*%s,4/3*%s)))' % (f, f, bmin, d, f)


def nfit(var, avail):
    return 'INT((%s+{slib})/(%s+{slib}))' % (avail, db(var))


def varsel2(As, pref, nmax, avail, nmin='1', as43=None):
    """Varilla más chica (desde la preferida) cuyas piezas necesarias (al menos nmin) caben en el ancho libre y no pasan del máximo.
    Con as43 (= 4/3 de la As de flexión): si las piezas no llegan a as43, 6.3.5.1.1 pide al menos 2 (6.3.5.1.3)."""
    i0 = 'MATCH(%s,%s,0)' % (pref, BARCOL)
    out = 'INDEX(%s,%d)' % (BARCOL, NBARS)
    for k in range(4, -1, -1):
        idx = 'MIN(%s+%d,%d)' % (i0, k, NBARS)
        a_ = 'INDEX(%s,%s)' % (ABCOL, idx); d_ = 'INDEX(%s,%s)' % (DBCOL, idx)
        npz = 'MAX(%s,CEILING(%s/%s,1))' % (nmin, As, a_)
        if as43:
            npz = 'IF(%s*%s>=%s-0.0001,%s,MAX(%s,2))' % (npz, a_, as43, npz, npz)
        ok = 'AND(%s>=0,%s<=MIN(%s,INT((%s+MAX({slib},%s))/(%s+MAX({slib},%s)))))' % (As, npz, nmax, avail, d_, d_, d_)
        out = 'IF(%s,INDEX(%s,%s),%s)' % (ok, BARCOL, idx, out)
    return out


def red391(d):
    """3.9.1 d) y 3.9.2: el peralte del lecho superior se reduce 20 mm si d ≤ 200 mm y no se anota la precaución en planos."""
    return '(%s-IF(AND(%s<=20,{ctl_bw}<>"SÍ"),2,0))' % (d, d)


def ldsel(var, n, r_, bw_='{bw}'):
    """L_d de la Tabla 14.4.2.4: «otros casos» si el recubrimiento es < d_b o la separación libre de n barras en la nervadura (ancho bw_) es < 2d_b."""
    d_ = db(var)
    return 'IF(OR(%s<%s,AND(%s>=2,(%s-2*%s-2*%s-%s*%s)/MAX(%s-1,1)<2*%s)),%s,%s)' % (r_, d_, n, bw_, r_, dbz('{est_sec}'), n, d_, n, d_, ldo(var), ld(var))


def sepmac(As, var, smax):
    """Separación (cm) de la varilla var en losa maciza para dar As (cm²/m): la menor de smax y la que da el área, redondeada hacia abajo a 2.5 cm."""
    return 'IF(%s>0,MIN(%s,FLOOR(100*%s/%s,2.5)),0)' % (As, smax, ab(var), As)


def varmac(As, pref, smax):
    """Varilla más chica (desde la preferida) cuya separación para dar As no baja de la mínima práctica de las losas macizas."""
    i0 = 'MATCH(%s,%s,0)' % (pref, BARCOL)
    out = 'INDEX(%s,MIN(%s+3,%d))' % (BARCOL, i0, NBARS)
    for k in range(2, -1, -1):
        idx = 'MIN(%s+%d,%d)' % (i0, k, NBARS)
        a_ = 'INDEX(%s,%s)' % (ABCOL, idx)
        ok = 'AND(%s>0,MIN(%s,FLOOR(100*%s/MAX(%s,1E-9),2.5))>={smin_mac})' % (As, smax, a_, As)
        out = 'IF(%s,INDEX(%s,%s),%s)' % (ok, BARCOL, idx, out)
    return out


def edge_af(t, side, k):
    """αf del borde (fórmula): muro completo = 99; nervadura principal = I_b/I_s; nada = 0."""
    e = t['bordes'][side]
    row = EDGE[(k, side)]
    perp = t['ay'] if side in ('inferior', 'superior') else t['ax']
    Is = '(%.3f*100*IF(APOYOS!$G$%d="CONTINUO",1,0.5)*IF({losa}{r}="MACIZA",{h_mac},{heq})^3/12)' % (perp, row)
    if e['apoyo'].startswith('MURO') and e['frac_muro'] >= 0.9:
        return '99'
    nps = e.get('nps') or []
    if nps and nps[0] in PROW:
        pr = PROW[nps[0]]
        return '(PRINCIPALES!$%s$%d*PRINCIPALES!$%s$%d^3/12/%s)' % ('{PCb}', pr, '{PCh}', pr, Is)
    return '0'


def sheet_tableros(wb):
    ws = wb.create_sheet('TABLEROS'); ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 1.78
    last = TC_LAST
    ws.merge_cells('B1:%s1' % last)
    c = ws['B1']; c.value = 'TABLEROS — NERVADURAS SECUNDARIAS DE CADA TABLERO (MÉTODO DE COEFICIENTES, NTC-CONCRETO 2023 TABLA 3.4.2.1)'
    c.font = font(13, True, color='FFFFFF'); c.fill = fill(TITLE_BG); c.alignment = Alignment(vertical='center'); ws.row_dimensions[1].height = 26
    ws.merge_cells('B2:%s2' % last)
    ws['B2'].value = ('Un renglón por tablero. Momentos = coeficiente × 10⁻⁴ × w_u × a1² por metro; por nervadura × s. Abajo trabaja la sección T con el patín efectivo de la Tabla 5.2.1.4.2; arriba sólo la nervadura (b_w de cálculo; 20 cm en las nervaduras de los volados, decisión D10). '
                      'Toda nervadura lleva al menos 1 No. 4 corrida abajo y 1 No. 4 corrida arriba (6.7.6.1.3); arriba, las demás varillas son bastones. Los renglones con «MACIZA» en la columna «Losa» (F80032 y F80010 de azotea, decisión D11) se diseñan por franja de 1 m (b = 100 cm, h = 20 cm, sin casetón ni malla) y su armado sale como «No. x @ y cm». '
                      'Lo AMARILLO se puede cambiar. El detalle paso a paso está en NERVADURA.')
    ws['B2'].font = font(9, i=True, color=GRAY); ws['B2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 42
    ws.merge_cells('B3:%s3' % last)
    ws['B3'].value = ('ESTADO: «CUMPLE» = flexión, ε_t ≥ 0.005, cortante, peralte, capa y apoyos bien. «VER PRINCIPALES» = el tablero pasa pero alguna nervadura principal que lo recibe no. «MACIZAR n hilera(s)» = quitar casetones junto al apoyo por cortante. '
                      '«6.7.8: NO» = la separación libre a_c pasa de 1/6 del claro corto (6.7.8.1.4): la losa se diseña como losa apoyada en vigas (6.7.8.1.9): nervaduras con recubrimiento de viga (r_np) y sin el ×1.1 de V_cR (decisión D7). '
                      '«3.4.2.1 c)» = los M⁻ del borde común con un vecino difieren más de 50 % del menor: los coeficientes quedan fuera de su campo; se distribuye 2/3 del desequilibrio (3.4.3.1) con rigideces d³/a1, cada tablero con su d (decisión D8), y se confirma con el modelo (D5).')
    ws['B3'].font = font(9, True, color=WARN_FG); ws['B3'].fill = fill(WARN_BG); ws['B3'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[3].height = 42
    groups = [('key', 'losa', 'TABLERO'), ('Lx', 'rt', 'GEOMETRÍA'), ('scm', 'wu', 'CARGAS (kg/m²)'), ('nd', 'caso', 'CASO DE LA TABLA'),
              ('idx', 'cPL', 'COEFICIENTES ×10⁻⁴ (Tabla 3.4.2.1, interpolados en m)'), ('MeI', 'c342', 'DISTRIBUCIÓN CON LOS VECINOS (3.4.3.1) Y 3.4.2.1 c)'), ('pbor', 'MNl', 'MOMENTOS POR NERVADURA (t·m; maciza: por metro)'),
              ('dqPc', 'ePc', 'INFERIOR — CLARO CORTO'), ('dqNc', 'eNc', 'SUPERIOR — CLARO CORTO'), ('dqPl', 'ePl', 'INFERIOR — CLARO LARGO'),
              ('dqNl', 'eNl', 'SUPERIOR — CLARO LARGO'), ('g678', 'cort', 'CORTANTE'), ('afm', 'tdia', 'PERALTE Y CAPA'), ('Lbc', 'nota', 'DETALLE Y RESUMEN')]
    for a, b, t in groups:
        ws.merge_cells('%s6:%s6' % (TC[a], TC[b]))
        x = ws['%s6' % TC[a]]; x.value = t; x.font = font(9, True, color='FFFFFF'); x.fill = fill(SEC_BG); x.alignment = Alignment(horizontal='center')
    heads = {'key': ('Tablero', 15), 'nivel': ('Nivel', 8), 'modo': ('Modo', 12), 'dest': ('Destino', 12), 'losa': ('Losa', 10), 'Lx': ('Lx', 6.5), 'Ly': ('Ly', 6.5), 'ax': ('a_x usado', 7.5), 'ay': ('a_y usado', 7.5),
             'a1': ('a1 corto', 7), 'a2': ('a2 largo', 7), 'm': ('m = a1/a2', 7), 'dirc': ('Corto ‖', 6), 'ln1': ('Claro libre corto', 7), 'ln2': ('Claro libre largo', 7),
             'hr': ('h de la losa', 6), 'bwt': ('b_w nervaduras ‖ corto (maciza: 100)', 8), 'sb': ('Ancho de diseño s (maciza: 100)', 8),
             'bfc': ('Patín b_f corto', 7), 'bfl': ('Patín b_f largo', 7), 'bNc': ('b a compresión M⁻ corto', 7.5), 'bNl': ('b a compresión M⁻ largo', 7.5), 'rb': ('r inferior', 6), 'rt': ('r superior', 6),
             'scm': ('SCM', 7), 'cv': ('CV W_m', 7), 'pp': ('PP diseño', 7), 'cm': ('CM', 7), 'ws': ('w serv.', 7), 'wu': ('w_u', 7),
             'nd': ('Bordes discont.', 7.5), 'ndc': ('… cortos', 6.5), 'ndl': ('… largos', 6.5), 'cauto': ('Caso automático', 15), 'cforce': ('Forzar caso', 12), 'caso': ('Caso usado', 15),
             'idx': ('col. m', 5), 'fr': ('interp.', 6), 'cNIC': ('M⁻ cont. corto', 7.5), 'cNIL': ('M⁻ cont. largo', 7.5), 'cNDC': ('M⁻ disc. corto', 7.5), 'cNDL': ('M⁻ disc. largo', 7.5),
             'cPC': ('M⁺ corto', 7.5), 'cPL': ('M⁺ largo', 7.5),
             'MeI': ('M⁻ borde inferior por metro', 8), 'MeS': ('M⁻ borde superior por metro', 8), 'MeL': ('M⁻ borde izquierdo por metro', 8), 'MeR': ('M⁻ borde derecho por metro', 8),
             'kX': ('Rigidez d³/a1 ‖X', 8), 'kY': ('Rigidez d³/a1 ‖Y', 8), 'dVv': ('ΔV por volado vecino', 8), 'f342': ('Vecinos con M⁻ > 50 % distinto', 14),
             'MNcd': ('M⁻ corto que deja el vecino', 8), 'MNld': ('M⁻ largo que deja el vecino', 8), 'MPcd': ('ΔM⁺ corto por distribución', 8), 'MPld': ('ΔM⁺ largo por distribución', 8), 'c342': ('3.4.2.1 c)', 16), 'pbor': ('Carga en orilla de volado', 8.5), 'MPc': ('M⁺ corto', 7.5), 'MNc': ('M⁻ corto', 7.5), 'MPl': ('M⁺ largo', 7.5), 'MNl': ('M⁻ largo', 7.5),
             'g678': ('6.7.8 cumple', 8), 'Vu': ('V_u por nerv.', 7.5), 'VcR': ('V_cR', 7.5), 'cort': ('Estado', 15),
             'afm': ('α_fm', 6.5), 'hmin': ('h mín. Tabla 6.7.3.1.2 / volado L/8', 9), 'heqv': ('h_eq / h', 7), 'peral': ('Estado', 12), 'tdia': ('Capa (diafragma)', 11),
             'Lbc': ('Bastón ‖ corto desde paño', 8), 'Lbl': ('Bastón ‖ largo desde paño', 8), 'ncx': ('Casetones ‖X', 7.5), 'ncy': ('Casetones ‖Y', 7.5), 'apoy': ('Apoyos', 15), 'estado': ('ESTADO', 26), 'nota': ('Nota', 60)}
    for blk in BLKS:
        heads['dq' + blk] = ('d para As (var. preferida)', 7); heads['As' + blk] = ('As req.', 7); heads['v' + blk] = ('Var. No.', 6); heads['n' + blk] = ('Pzas (maciza: por metro)', 6.5)
        heads['s' + blk] = ('Separación (maciza)', 6.5); heads['d' + blk] = ('d', 6); heads['MR' + blk] = ('M_R', 7); heads['FU' + blk] = ('Mu/M_R', 7); heads['e' + blk] = ('ε_t', 7)
    units = {'Lx': 'm', 'Ly': 'm', 'ax': 'm', 'ay': 'm', 'a1': 'm', 'a2': 'm', 'ln1': 'cm', 'ln2': 'cm', 'hr': 'cm', 'bwt': 'cm', 'sb': 'cm', 'bfc': 'cm', 'bfl': 'cm', 'bNc': 'cm', 'bNl': 'cm', 'rb': 'cm', 'rt': 'cm',
             'scm': 'kg/m²', 'cv': 'kg/m²', 'pp': 'kg/m²', 'cm': 'kg/m²', 'ws': 'kg/m²', 'wu': 'kg/m²', 'pbor': 'kg/m (serv.)',
             'MPc': 't·m', 'MNc': 't·m', 'MPl': 't·m', 'MNl': 't·m', 'MeI': 't·m/m', 'MeS': 't·m/m', 'MeL': 't·m/m', 'MeR': 't·m/m', 'kX': 'cm²', 'kY': 'cm²', 'dVv': 't', 'MNcd': 't·m', 'MNld': 't·m', 'MPcd': 't·m', 'MPld': 't·m', 'Vu': 't', 'VcR': 't', 'hmin': 'cm', 'heqv': 'cm', 'Lbc': 'cm', 'Lbl': 'cm'}
    for blk in BLKS:
        units['dq' + blk] = 'cm'; units['As' + blk] = 'cm² (maciza: /m)'; units['MR' + blk] = 't·m'; units['d' + blk] = 'cm'; units['e' + blk] = '≥ 0.005'; units['s' + blk] = 'cm'
    for nm, (t, w) in heads.items():
        hdr_cell(ws, 7, CI(TC[nm]), t, w)
        unit_cell(ws, 8, CI(TC[nm]), units.get(nm, ''))
    ws.row_dimensions[7].height = 42
    NN = dict(N); NN['PCb'] = PC['bq']; NN['PCh'] = PC['hq']
    F = lambda s, r: s.format(r=r, **{k: v for k, v in TC.items()}, **NN)
    for t in TABS:
        k = key(t); r = TROW[k]
        L = t['loads']
        scm = L.get('Sobrecarga', 270.0 if t['name'] == 'F2' else 0)
        cvm = L.get('Carga Viva Máxima', 100.0)
        dest = 'HABITACIÓN' if abs(cvm - 190) < 1 else 'AZOTEA'
        e = {s: EDGE[(k, s)] for s in SIDES}
        D_ = {s: '(APOYOS!$G$%d="DISCONTINUO")' % e[s] for s in SIDES}
        dbe = dbz('{est_sec}')
        vals = {
            'key': (k, 'calc', '@'), 'nivel': (t['level'], 'calc', '@'), 'modo': (t['modo'], 'input', '@'), 'dest': (dest, 'input', '@'),
            'losa': ('MACIZA' if t['seccion'] == 'Slab1' else 'NERVADA', 'input', '@'),
            'Lx': (t['Lx'], 'data', '0.000'), 'Ly': (t['Ly'], 'data', '0.000'), 'ax': (t['ax'], 'input', '0.000'), 'ay': (t['ay'], 'input', '0.000'),
            'a1': ('=MIN({ax}{r},{ay}{r})', 'calc', '0.000'), 'a2': ('=MAX({ax}{r},{ay}{r})', 'calc', '0.000'), 'm': ('=IF({a2}{r}>0,{a1}{r}/{a2}{r},0)', 'calc', '0.000'),
            'dirc': ('=IF({ax}{r}<={ay}{r},"X","Y")', 'calc', '@'),
            'ln1': ('=MAX(0,{a1}{r}*100-IF({modo}{r}="VOLADO",{bp}/2,{bp}))', 'calc', '0'), 'ln2': ('=MAX(0,{a2}{r}*100-{bp})', 'calc', '0'),
            # D10: en los volados las nervaduras de la dirección del volado (claro corto) son de b_w = bw_vol; D11: maciza por franja de 1 m
            'hr': ('=IF({losa}{r}="MACIZA",{h_mac},{h})', 'calc', '0'),
            'bwt': ('=IF({losa}{r}="MACIZA",100,IF({modo}{r}="VOLADO",{bw_vol},{bw}))', 'calc', '0'),
            'sb': ('=IF({losa}{r}="MACIZA",100,{s})', 'calc', '0'),
            'bfc': ('=IF({losa}{r}="MACIZA",100,MIN({s},{bwt}{r}+2*MIN(8*{t},({s}-{bwt}{r})/2,{ln1}{r}/8)))', 'calc', '0.0'),
            'bfl': ('=IF({losa}{r}="MACIZA",100,MIN({s},{bw}+2*MIN(8*{t},{ac}/2,{ln2}{r}/8)))', 'calc', '0.0'),
            # b de cálculo de la nervadura a compresión en M⁻: 3.9.1 e) quita 2 cm sólo si el ancho es < 200 mm y no se anota la precaución de 3.9.2
            'bNc': ('=IF({losa}{r}="MACIZA",100,IF(AND({ctl_bw}<>"SÍ",{bwt}{r}<20),{bwt}{r}-2,{bwt}{r}))', 'calc', '0.0'),
            'bNl': ('=IF({losa}{r}="MACIZA",100,{bwc})', 'calc', '0.0'),
            'rb': ('=IF({modo}{r}="VOLADO",{rec_ext},IF({g678}{r}="NO",{rec_np},{rec}))', 'calc', '0.0'),
            'rt': ('=IF(OR(AND({nivel}{r}="N+6.00",{exp_az}="SÍ"),AND({modo}{r}="VOLADO",{exp_vol}="SÍ")),{rec_ext},IF({g678}{r}="NO",{rec_np},{rec}))', 'calc', '0.0'),
            'scm': (scm, 'input', '0'), 'cv': ('=IF({dest}{r}="HABITACIÓN",{cv_hab},{cv_az})', 'calc', '0'),
            # peso propio del renglón: maciza h·γc; nervada con la b_w del renglón (casetón (s − b_w)·a_c·h_c por cada s²), más los incrementos de 5.1.2.1
            'pp': ('=IF({losa}{r}="MACIZA",{h_mac}/100*{gc}*1000,(({h}/100-({s}-{bwt}{r})/100*{ac}/100*{hc}/100/({s}/100)^2)*{gc}+({s}-{bwt}{r})/100*{ac}/100*{hc}/100/({s}/100)^2*{gpol})*1000)+{inc_losa}+{inc_mort}', 'calc', '0'),
            'cm': ('={pp}{r}+{scm}{r}', 'calc', '0'), 'ws': ('={cm}{r}+{cv}{r}', 'calc', '0'),
            'wu': ('=MAX({FCcm}*{cm}{r}+{FCcv}*{cv}{r},IF({dest}{r}="AZOTEA",{FCacc}*({cm}{r}+{Wg}),0))', 'key', '0'),
            'nd': ('=' + '+'.join(D_[s] for s in SIDES), 'calc', '0'),
            'ndc': ('=IF({dirc}{r}="X",%s+%s,%s+%s)' % (D_['inferior'], D_['superior'], D_['izquierdo'], D_['derecho']), 'calc', '0'),
            'ndl': ('={nd}{r}-{ndc}{r}', 'calc', '0'),
            'cauto': ('=IF({nd}{r}=0,"INTERIOR",IF({nd}{r}=1,IF({ndc}{r}=1,"BORDE-CORTO","BORDE-LARGO"),IF({nd}{r}=2,IF({ndc}{r}=1,"ESQUINA",IF({ndc}{r}=2,"EXTREMO-LARGO","EXTREMO-CORTO")),IF({nd}{r}=3,IF({ndc}{r}=2,"EXTREMO-LARGO","EXTREMO-CORTO"),"AISLADO"))))', 'calc', '@'),
            'cforce': ('AUTO', 'input', '@'), 'caso': ('=IF({cforce}{r}="AUTO",{cauto}{r},{cforce}{r})', 'key', '@'),
            'idx': ('=MATCH({m}{r},{MROW},1)', 'calc', '0'),
            'fr': ('=IF({idx}{r}>=7,0,({m}{r}-INDEX({MROW},{idx}{r}))/(INDEX({MROW},{idx}{r}+1)-INDEX({MROW},{idx}{r})))', 'calc', '0.00'),
            'pbor': (t.get('volado', {}).get('carga_borde', 0.0), 'input', '0.0'),
        }
        for mm in MOMS:
            idx = 'MATCH({caso}{r}&"|%s",{CKEY},0)' % mm
            vals['c' + mm] = ('=INDEX({CVAL},%s,{idx}{r})+{fr}{r}*(INDEX({CVAL},%s,MIN({idx}{r}+1,7))-INDEX({CVAL},%s,{idx}{r}))' % (idx, idx, idx), 'calc', '0')
        fac = '*{wu}{r}*{a1}{r}^2*1E-4*{sb}{r}/100/1000'     # por nervadura (s) o, en maciza, por franja de 1 m
        fm = '*{wu}{r}*{a1}{r}^2*1E-4/1000'                    # por metro: así se comparan nervadas y macizas en el borde común
        # 3.4.3.1: en cada borde continuo con vecino se reparten 2/3 del desequilibrio con rigideces d³/a1, cada tablero con su propio d (D8, D11).
        # 3.4.2.1 c): los M⁻ del borde común no deben diferir más de 50 % del menor.
        # M⁻ de este tablero en cada borde, según cómo lo clasifica él mismo en APOYOS (continuo o discontinuo), por metro
        for s_ in SIDES:
            dE = 'Y' if s_ in ('inferior', 'superior') else 'X'
            vals[MEC[s_]] = ('=IF({modo}{r}="PERIMETRAL",IF(APOYOS!$G$%d="CONTINUO",IF({dirc}{r}="%s",{cNIC}{r},{cNIL}{r}),IF({dirc}{r}="%s",{cNDC}{r},{cNDL}{r}))%s,0)' % (EDGE[(k, s_)], dE, dE, fm), 'calc', '0.000')
        # rigidez d³/a1 por dirección (d del lecho superior para dimensionar: sin referencia circular)
        for dX_ in ('X', 'Y'):
            vals['k%s' % dX_] = ('=IF({modo}{r}="PERIMETRAL",IF({dirc}{r}="%s",{dqNc}{r},{dqNl}{r})^3/({a1}{r}*100),0)' % dX_, 'calc', '0')
        gain = {'c': [], 'l': []}; redu = {'c': [], 'l': []}; fails = {}; asym = []; vols = []; dv = []
        for s_ in SIDES:
            e_ = t['bordes'][s_]
            dE = 'Y' if s_ in ('inferior', 'superior') else 'X'
            nbs = [n_.strip() for n_ in (e_.get('vecinos') or '').split(',') if n_.strip()]
            nbs = [n_ for n_ in nbs if (n_ + ' ' + t['level']) in TROW and (n_ + ' ' + t['level']) != k]
            if not nbs:
                continue
            gA_ = 'APOYOS!$G$%d' % EDGE[(k, s_)]
            MA = '{%s}{r}' % MEC[s_]
            reds = []
            for n_ in nbs:
                kb = n_ + ' ' + t['level']; rb_ = TROW[kb]
                tb = next(x for x in TABS if key(x) == kb)
                if tb['modo'] == 'VOLADO':
                    vols.append(('AND(%s="CONTINUO",{modo}%d="VOLADO")' % (gA_, rb_), n_))
                    dv.append('IF(AND(%s="CONTINUO",{modo}%d="VOLADO"),MAX(0,{MNc}%d*100/{sb}%d-%s)/MAX(IF({dirc}{r}="%s",{ln1}{r},{ln2}{r})/100,0.01),0)' % (gA_, rb_, rb_, rb_, MA, dE))
                    continue
                # borde recíproco: el del vecino que lista a este tablero (de preferencia el opuesto)
                sbs = [q for q in [OPPS[s_]] + SIDES if t['name'] in [x.strip() for x in (tb['bordes'][q].get('vecinos') or '').split(',')]]
                sb = sbs[0] if sbs else OPPS[s_]
                MB = '{%s}%d' % (MEC[sb], rb_)
                if sbs:
                    gB_ = 'APOYOS!$G$%d' % EDGE[(kb, sb)]
                    asym.append(('AND({modo}%d="PERIMETRAL",%s<>%s)' % (rb_, gA_, gB_), n_))
                    cont = 'OR(%s="CONTINUO",%s="CONTINUO")' % (gA_, gB_)
                else:
                    cont = '%s="CONTINUO"' % gA_
                cond = '%s,%s>0,%s>0' % (cont, MA, MB)
                kA = '{k%s}{r}/MAX({k%s}{r}+{k%s}%d,1E-9)' % (dE, dE, dE, rb_)
                gA = 'IF(AND(%s,%s>%s),%s+2/3*(%s-%s)*%s,0)' % (cond, MB, MA, MA, MB, MA, kA)
                rA = 'IF(AND(%s,%s>%s),2/3*(%s-%s)*%s,0)' % (cond, MA, MB, MA, MB, kA)
                for dd, op in (('c', '='), ('l', '<>')):
                    gain[dd].append('IF({dirc}{r}%s"%s",%s,0)' % (op, dE, gA))
                reds.append(rA)
                fails.setdefault(n_, []).append('AND(%s,ABS(%s-%s)>0.5*MIN(%s,%s))' % (cond, MA, MB, MA, MB))
            if reds:
                rmax = reds[0] if len(reds) == 1 else 'MAX(%s)' % ','.join(reds)
                for dd, op in (('c', '='), ('l', '<>')):
                    redu[dd].append('IF({dirc}{r}%s"%s",%s,0)' % (op, dE, rmax))
        for dd in ('c', 'l'):
            vals['MN%sd' % dd] = ('=IF({modo}{r}="PERIMETRAL",MAX(0,%s)*{sb}{r}/100,0)' % (','.join(gain[dd]) or '0'), 'calc', '0.000')
            vals['MP%sd' % dd] = ('=IF({modo}{r}="PERIMETRAL",(%s)/2*{sb}{r}/100,0)' % ('+'.join(redu[dd]) or '0'), 'calc', '0.000')
        vals['dVv'] = ('=IF({modo}{r}="PERIMETRAL",MAX(0,%s)*{sb}{r}/100,0)' % (','.join(dv) or '0'), 'calc', '0.000')
        lst = '&'.join('IF(%s," %s","")' % (fs[0] if len(fs) == 1 else 'OR(%s)' % ','.join(fs), n_) for n_, fs in fails.items())
        vals['f342'] = (('=IF({modo}{r}<>"PERIMETRAL","",%s)' % lst) if lst else '=""', 'calc', '@')
        seen = set(); tl = []
        for f_, n_ in asym:
            if n_ not in seen:
                seen.add(n_); tl.append('IF(%s," · continuidad asimétrica con %s (cada uno con su momento; revisar APOYOS)","")' % (f_, n_))
        tl += ['IF(%s," · junto al volado %s: se arma con su M⁻ y su cortante entra en V_u","")' % (f_, n_) for f_, n_ in vols]
        vals['c342'] = ('=IF({modo}{r}<>"PERIMETRAL","—",IF({f342}{r}="","CUMPLE","FUERA con"&{f342}{r})%s)' % ''.join('&' + x for x in tl), 'calc', '@')
        vals['MPc'] = ('=IF({modo}{r}="PERIMETRAL",{cPC}{r}%s+{MPcd}{r},0)' % fac, 'calc', '0.000')
        vals['MNc'] = ('=IF({modo}{r}="PERIMETRAL",MAX(MAX({cNIC}{r},{cNDC}{r})%s,{MNcd}{r}),IF({modo}{r}="VOLADO",({wu}{r}*({ln1}{r}/100)^2/2+{FCcm}*{pbor}{r}*{ln1}{r}/100)*{sb}{r}/100/1000,0))' % fac, 'calc', '0.000')
        vals['MPl'] = ('=IF({modo}{r}="PERIMETRAL",{cPL}{r}%s+{MPld}{r},0)' % fac, 'calc', '0.000')
        vals['MNl'] = ('=IF({modo}{r}="PERIMETRAL",MAX(MAX({cNIL}{r},{cNDL}{r})%s,{MNld}{r}),0)' % fac, 'calc', '0.000')
        # peraltes efectivos: el lecho del claro corto va por fuera; el del largo, encima (o debajo) de él.
        # Nervada: la varilla superior pasa debajo de la malla (2·Ø). Maciza (D11): sin malla ni estribo, h = h_mac.
        mac = '{losa}{r}="MACIZA"'
        mal = 'IF(%s,0,2*{fim})' % mac
        dbr = 'IF(%s,0,%s)' % (mac, dbe)
        pinf = 'IF(%s,{var_mac},{var_inf})' % mac
        psup = 'IF(%s,{var_mac},{var_sup})' % mac
        dq = {'Pc': '({hr}{r}-{rb}{r}-%s-%s/2)' % (dbr, db(pinf)),
              'Pl': '({hr}{r}-{rb}{r}-%s-%s-%s/2)' % (dbr, db(pinf), db(pinf)),
              'Nc': red391('({hr}{r}-{rt}{r}-%s-%s-%s/2)' % (mal, dbr, db(psup))),
              'Nl': red391('({hr}{r}-{rt}{r}-%s-%s-%s-%s/2)' % (mal, dbr, db(psup), db(psup)))}
        dact = {'Pc': '{hr}{r}-{rb}{r}-%s-%s/2' % (dbr, db('{vPc}{r}')),
                'Pl': '{hr}{r}-{rb}{r}-%s-%s-%s/2' % (dbr, db('{vPc}{r}'), db('{vPl}{r}')),
                'Nc': red391('({hr}{r}-{rt}{r}-%s-%s-%s/2)' % (mal, dbr, db('{vNc}{r}'))),
                'Nl': red391('({hr}{r}-{rt}{r}-%s-%s-%s-%s/2)' % (mal, dbr, db('{vNc}{r}'), db('{vNl}{r}')))}
        smx = 'MIN(2*{hr}{r},{smax_mac})'     # separación máxima en la maciza
        bwl = 'IF(%s,100,{bw})' % mac
        avs = {}
        for blk, M, neg in (('Pc', 'MPc', False), ('Nc', 'MNc', True), ('Pl', 'MPl', False), ('Nl', 'MNl', True)):
            b = {'Pc': '{bfc}{r}', 'Pl': '{bfl}{r}', 'Nc': '{bNc}{r}', 'Nl': '{bNl}{r}'}[blk]
            bwb = '{bwt}{r}' if blk[1] == 'c' else '{bw}'          # ancho de la nervadura de este lecho (D10: 20 cm en la dirección del volado)
            av = '(%s-2*{rt}{r}-2*%s)' % (bwb, dbe) if neg else '(%s-2*{rb}{r}-2*%s)' % (bwb, dbe)
            avs[blk] = av
            # 6.3.5.1.2: en el volado (viga isostática con el patín a tensión) b_w = el menor de b_f y 2b_w
            bmin = 'IF({modo}{r}="VOLADO",MIN({s},2*{bwt}{r}),{bwt}{r})' if blk == 'Nc' else bwb
            Mref = '{%s}{r}' % M
            vals['dq' + blk] = ('=' + dq[blk], 'calc', '0.00')
            dqr = '{dq%s}{r}' % blk
            fl = asflex(Mref, b, dqr, neg)
            # maciza: As mín de losa 0.0018·b·h en cada lecho (6.7.6.1.1); nervada: As mín de viga con la excepción de 4/3 (6.3.5.1.3)
            vals['As' + blk] = ('=IF({modo}{r}="HUECO",0,IF(%s,IF(%s<0,-1,MAX(%s,{Asm_mac})),%s))' % (mac, fl, fl, asreq(Mref, b, dqr, neg, bmin)), 'calc', '0.00')
            vals['v' + blk] = ('=IF(%s,%s,%s)' % (mac, varmac('{As%s}{r}' % blk, '{var_mac}', smx), varsel2('{As%s}{r}' % blk, '{var_sup}' if neg else '{var_inf}', '{nmax_sec}', av, '{nmin_sec}', '4/3*' + fl)), 'calc', 'General')
            ab_ = ab('{v%s}{r}' % blk)
            vals['s' + blk] = ('=IF(%s,%s,"—")' % (mac, sepmac('{As%s}{r}' % blk, '{v%s}{r}' % blk, smx)), 'calc', '0.0')
            n0 = 'MAX({nmin_sec},CEILING(MAX({As%s}{r},0)/%s,1))' % (blk, ab_)
            vals['n' + blk] = ('=IF({modo}{r}="HUECO",0,IF({As%s}{r}<0,0,IF(%s,IF({s%s}{r}>0,ROUND(100/{s%s}{r},2),0),IF(%s*%s>=4/3*%s-0.0001,%s,MAX(%s,2)))))' % (blk, mac, blk, blk, n0, ab_, fl, n0, n0), 'calc', 'General')
            vals['d' + blk] = ('=' + dact[blk], 'calc', '0.00')
            asp = '{n%s}{r}*%s' % (blk, ab_)
            vals['MR' + blk] = ('=IF({n%s}{r}=0,0,{FRf}*%s*{fy}*({d%s}{r}-%s*{fy}/({fpc}*%s)/2)/1E5)' % (blk, asp, blk, asp, b), 'calc', '0.000')
            vals['FU' + blk] = ('=IF({As%s}{r}<0,9.99,IF({MR%s}{r}>0,{%s}{r}/{MR%s}{r},0))' % (blk, blk, M, blk), 'key', '0.00')
            cexp = '(%s*{fy}/({fpc}*%s)/{beta1})' % (asp, b)
            vals['e' + blk] = ('=IF({n%s}{r}=0,"—",0.003*({d%s}{r}-%s)/%s)' % (blk, blk, cexp, cexp), 'calc', '0.0000')
        vals['g678'] = ('=IF({modo}{r}="HUECO","—",IF(%s,"—",IF(AND({modo}{r}="PERIMETRAL",{a1}{r}*100>=6*{ac},{ac}<=80,{bw}>=10,{h}<=3.5*{bw},{t}>=MAX({ac}/13,5)),"SÍ","NO")))' % mac, 'calc', '@')
        lam = 'MIN(1,SQRT(2/(1+0.04*{dNc}{r})))'
        pexp = '({nNc}{r}*%s/({bNc}{r}*{dNc}{r}))' % ab('{vNc}{r}')
        vals['Vu'] = ('=IF({modo}{r}="PERIMETRAL",MAX(0,({ln1}{r}/2-{dNc}{r})/100)*(0.95-0.5*{m}{r})*{wu}{r}*IF(AND({nd}{r}>0,{nd}{r}<4),{Vinc},1)*{sb}{r}/100/1000+{dVv}{r},'
                      'IF({modo}{r}="VOLADO",({wu}{r}*MAX(0,{ln1}{r}-{dNc}{r})/100+{FCcm}*{pbor}{r})*{sb}{r}/100/1000,0))', 'calc', '0.000')
        # ec. 5.5.3.2.1 sin estribos: nervadura (b_w de cálculo del renglón) o maciza (b = 100 cm), con la cuantía del acero superior en el apoyo
        vals['VcR'] = ('=IF({modo}{r}="HUECO",0,{FRv}*2*%s*%s^(1/3)*SQRT({fc})*{bNc}{r}*{dNc}{r}*IF({g678}{r}="SÍ",{VcRinc},1)/1000)' % (lam, pexp), 'calc', '0.000')
        # largo (cm desde el paño) donde V > V_cR; se compara contra el macizo de orilla que ya da el acomodo de casetones
        xper = 'IF({VcR}{r}<={dVv}{r},{ln1}{r},{ln1}{r}/2-({VcR}{r}-{dVv}{r})/MAX({Vu}{r}-{dVv}{r},1E-9)*({ln1}{r}/2-{dNc}{r}))'
        xvol = '({ln1}{r}-({VcR}{r}*1E5/{sb}{r}-{FCcm}*{pbor}{r})*100/{wu}{r})'
        ncs = 'INT(({a1}{r}*100-{bp}+{bw})/{s})'
        oril = 'IF(%s>0,({a1}{r}*100-{bp}-(%s*{ac}+(%s-1)*{bw}))/2,{a1}{r}*100/2)' % (ncs, ncs, ncs)
        hil = 'MAX(0,CEILING((IF({modo}{r}="VOLADO",%s,%s)-%s)/{s},1))' % (xvol, xper, oril)
        vals['cort'] = ('=IF({modo}{r}="HUECO","—",IF({Vu}{r}<={VcR}{r},"CUMPLE",IF(%s,"NO CUMPLE: V_u > V_cR (losa maciza)",IF(%s<=0,"CUMPLE (basta el macizo de orilla)","MACIZAR "&%s&" hilera(s)"))))' % (mac, hil, hil), 'calc', '@')
        afs = [edge_af(t, s_, k) for s_ in SIDES]
        vals['afm'] = ('=IF({modo}{r}="PERIMETRAL",(%s)/4,0)' % '+'.join(afs), 'calc', '0.00')
        disc_nobeam = 'OR(' + ','.join('AND(%s,%s<0.8)' % (D_[s_], a_) for s_, a_ in zip(SIDES, afs)) + ')'
        beta = '({ln2}{r}/MAX({ln1}{r},1))'
        k_ = '(0.8+{fy}/14000)'
        hd = 'MAX(9,{ln2}{r}*%s/(36+9*%s))' % (k_, beta)
        hb = 'MAX(13,{ln2}{r}*%s/(36+5*%s*({afm}{r}-0.2)))' % (k_, beta)
        ha = 'MAX(13,{ln2}{r}/IF({nd}{r}>0,30,33))'
        vals['hmin'] = ('=IF({modo}{r}="PERIMETRAL",IF({afm}{r}>2,%s,IF({afm}{r}>0.2,%s,%s))*IF(AND({afm}{r}>0.2,%s),1.1,1),IF({modo}{r}="VOLADO",MAX(25,{a1}{r}*100/8),0))' % (hd, hb, ha, disc_nobeam), 'calc', '0.0')
        # la maciza se compara con su h (Tabla 6.7.3.1.2 es para losas macizas); la nervada con h_eq; el volado con h (L/8)
        vals['heqv'] = ('=IF({modo}{r}="VOLADO",{h},IF(%s,{h_mac},{heq}))' % mac, 'calc', '0.0')
        vals['peral'] = ('=IF({modo}{r}="HUECO","—",IF({heqv}{r}>={hmin}{r},"CUMPLE",IF({modo}{r}="VOLADO","NO CUMPLE","CALCULAR FLECHA")))', 'calc', '@')
        vals['tdia'] = ('=IF({modo}{r}="HUECO","—",IF(IF(%s,{h_mac},{t})>=IF(IF({modo}{r}="VOLADO",{a1}{r},{a2}{r})>=6,6,5),"CUMPLE","CAPA ≥ 6 cm"))' % mac, 'calc', '@')
        # bastones: desde el paño, punto de inflexión + max(d, 12db, Ln/16), y no menos que L_d
        # punto de inflexión con momentos distintos en los extremos: M(ξ) = −M1(1−ξ) − M2·ξ + 4·M0·ξ(1−ξ), M0 = M⁺ + (M1+M2)/2.
        # M1 = M⁻ de diseño (incluye lo que deja el vecino), M2 = el del otro extremo sin distribuir, M⁺ sin el ΔM⁺ (lo más largo).
        # M0 (momento estático) con los M⁻ sin distribuir: al subir el M⁻ de un extremo por la distribución, el M⁺ baja y M0 no cambia.
        def xpi(ln, cP, cI, cD, MN):
            P_ = '({%s}{r}%s)' % (cP, fac); M1 = '{%s}{r}' % MN
            M1e = '(MAX({%s}{r},{%s}{r})%s)' % (cI, cD, fac)
            M2 = '(IF(AND({%s}{r}>0,{%s}{r}>0),MIN({%s}{r},{%s}{r}),MAX({%s}{r},{%s}{r}))%s)' % (cI, cD, cI, cD, cI, cD, fac)
            M0 = '(%s+(%s+%s)/2)' % (P_, M1e, M2)
            bq_ = '(4*%s+%s-%s)' % (M0, M1, M2)
            disc = '(%s^2-16*%s*%s)' % (bq_, M0, M1)
            return '{%s}{r}*(%s-SQRT(MAX(0,%s)))/(8*MAX(%s,1E-9))' % (ln, bq_, disc, M0), disc
        xpi_c, disc_c = xpi('ln1', 'cPC', 'cNIC', 'cNDC', 'MNc')
        xpi_l, disc_l = xpi('ln2', 'cPL', 'cNIL', 'cNDL', 'MNl')
        ext_c = 'MAX({dNc}{r},12*%s,{ln1}{r}/16)' % db('{vNc}{r}')
        ext_l = 'MAX({dNl}{r},12*%s,{ln2}{r}/16)' % db('{vNl}{r}')
        if t['modo'] == 'VOLADO':
            lb = t['volado']['Lb']
            # claro de respaldo (LR-25 normal) con 0.9·PP geométrico y sin carga viva (efecto favorable): donde se anula el momento negativo
            wD = '(0.9*{PP}*{sb}{r}/100/1000)'
            A_ = '(0.5*%s)' % wD; B_ = '(0.5*%s*%.3f+{MNc}{r}/%.3f)' % (wD, lb, lb); C_ = '{MNc}{r}'
            x0 = 'IF(%s^2-4*%s*%s<0,%.3f,MIN(%.3f,(%s-SQRT(%s^2-4*%s*%s))/(2*%s)))*100' % (B_, A_, C_, lb, lb, B_, B_, A_, C_, A_)
            vals['Lbc'] = ('=CEILING(MAX(%s+MAX({dNc}{r},12*%s,%.1f/16),%s),5)' % (x0, db('{vNc}{r}'), lb * 100, ldsel('{vNc}{r}', '{nNc}{r}', '{rt}{r}', '{bwt}{r}')), 'key', '0')
            vals['Lbl'] = (0, 'calc', '0')
        else:
            vals['Lbc'] = ('=IF({modo}{r}="HUECO",0,IF(%s<0,CEILING({ln1}{r},5),CEILING(MAX(%s+%s,%s),5)))' % (disc_c, xpi_c, ext_c, ldsel('{vNc}{r}', '{nNc}{r}', '{rt}{r}', '{bwt}{r}')), 'key', '0')
            vals['Lbl'] = ('=IF({modo}{r}="HUECO",0,IF(%s<0,CEILING({ln2}{r},5),CEILING(MAX(%s+%s,%s),5)))' % (disc_l, xpi_l, ext_l, ldsel('{vNl}{r}', '{nNl}{r}', '{rt}{r}', bwl)), 'key', '0')
        # casetones: a lo largo del muro del volado van de (s − b_w del volado) = 50 cm; en la otra dirección, a_c. La maciza no lleva.
        vals['ncx'] = ('=IF(%s,0,MAX(0,INT(({Lx}{r}*100-{bp}+IF(AND({modo}{r}="VOLADO",{dirc}{r}="Y"),{bw_vol},{bw}))/{s})))' % mac, 'calc', '0')
        vals['ncy'] = ('=IF(%s,0,MAX(0,INT(({Ly}{r}*100-{bp}+IF(AND({modo}{r}="VOLADO",{dirc}{r}="X"),{bw_vol},{bw}))/{s})))' % mac, 'calc', '0')
        libre = '+'.join('(APOYOS!$F$%d="LIBRE")' % e[s] for s in SIDES)
        vals['apoy'] = ('=IF(COUNTIFS(TRIB!$C$%d:$C$%d,{key}{r},TRIB!$K$%d:$K$%d,"NO CUMPLE*")+COUNTIFS(TRIB!$C$%d:$C$%d,{key}{r},TRIB!$K$%d:$K$%d,"SIN APOYO*")>0,"VER PRINCIPALES",IF(%s>0,"REVISAR: borde libre","OK"))'
                        % (R0, R1, R0, R1, R0, R1, R0, R1, libre), 'calc', '@')
        fus = 'MAX({FUPc}{r},{FUNc}{r},{FUPl}{r},{FUNl}{r})'
        # nervada: piezas que no caben en b_w − 2r; maciza: separación menor que la mínima práctica
        nn = 'OR(' + ','.join('IF(%s,{s%s}{r}<{smin_mac},{n%s}{r}>MIN({nmax_sec},%s))' % (mac, k_b, k_b, nfit('{v%s}{r}' % k_b, avs[k_b])) for k_b in BLKS) + ')'
        etl = 'OR(' + ','.join('AND(ISNUMBER({e%s}{r}),{e%s}{r}<0.005)' % (k_b, k_b) for k_b in BLKS) + ')'
        cfl = 'AND(%s,LEFT({cort}{r},6)<>"CUMPLE")' % mac
        vals['estado'] = ('=IF({modo}{r}="HUECO","HUECO — marco de nervaduras principales",IF(AND(%s<=1,NOT(%s),NOT(%s),NOT(%s),{peral}{r}="CUMPLE",{tdia}{r}="CUMPLE"),'
                          'IF({apoy}{r}="OK",IF(LEFT({cort}{r},6)="CUMPLE","CUMPLE","CUMPLE — "&{cort}{r})&IF(LEFT({c342}{r},5)="FUERA"," · fuera de 3.4.2.1 c)",""),IF({apoy}{r}="VER PRINCIPALES","VER PRINCIPALES","REVISAR APOYOS")),'
                          'IF(AND(%s<=1,NOT(%s),NOT(%s),NOT(%s),{peral}{r}="CALCULAR FLECHA"),"REVISAR: calcular flecha (h_eq < h mín.)",'
                          '"NO CUMPLE: "&IF(%s>1,"flexión ","")&IF(%s,IF(%s,"separación < mínima ","no caben "),"")&IF(%s,"ε_t < 0.005 ","")&IF(%s,"cortante ","")&IF({peral}{r}="NO CUMPLE","peralte ","")&IF({tdia}{r}<>"CUMPLE","capa < 6 cm",""))))'
                          % (fus, nn, etl, cfl, fus, nn, etl, cfl, fus, nn, mac, etl, cfl), 'key', '@')
        notes = []
        if t['rep_nota']:
            notes.append(t['rep_nota'])
        if t['seccion'] == 'Slab1':
            notes.append('LOSA MACIZA de 20 cm (decisión D11): así viene en el modelo («Slab1»). Se diseña por franja de 1 m con los coeficientes de la Tabla 3.4.2.1, r = 2 cm, As mín = 0.0018·b·h en cada lecho y armado «No. x @ y». '
                         'Concreto f\'c = 250 Clase 1A como el resto de la obra: en ETABS cambiar el de «Slab1» (4000 psi) a f\'c = 250 (D5). Cara superior al nivel de la LR-25 vecina y lecho inferior 5 cm más alto: anotar el desnivel en planos (D6).')
        if t['name'] == 'F2':
            notes.append('Sin cargas en el E2K: se toma la SCM de F80023 (270 kg/m²); en ETABS hay que ponérsela (decisión D5).')
        if t['modo'] == 'VOLADO':
            notes.append('Volado NERVADO de %.2f m desde el borde %s (muro) (decisión D10): LR-25 con las nervaduras de la dirección del volado de b_w = 20 cm (DATOS) y casetón 50×60×19 (s = 70 cm, alineado con la retícula del entrepiso); las paralelas al muro siguen de 10 cm. '
                         'Peso propio con esa b_w; orilla libre con %.1f kg/m del E2K. Cara inferior a la intemperie (r = r_ext); cara superior a la intemperie si DATOS lo indica (r = r_ext). Si la orilla lleva pretil de block u otro elemento frágil, el L/8 no basta: calcular flecha (13.4). '
                         'Bastón superior: se prolonga hacia el claro de respaldo (%.2f m) hasta pasar el punto donde se anula el momento con 0.9·PP y sin carga viva (NTC-Criterios 3.4.1 c y 5.1.2.2).' % (t['volado']['Lv'], t['volado']['raiz'], t['volado'].get('carga_borde', 0), t['volado']['Lb']))
        if t['modo'] == 'HUECO':
            notes.append('Tragaluz de 0.30 m: casi todo el tablero es hueco (quedan %.2f m² de losa).' % t['area'])
        txt = ' '.join(notes) or '—'
        txtf = '&'.join('"%s"' % txt[i_:i_ + 200].replace('"', '""') for i_ in range(0, len(txt), 200))
        vals['nota'] = ('=IF(AND({modo}{r}="PERIMETRAL",{losa}{r}="NERVADA",MIN({ncx}{r},{ncy}{r})=0),"Tablero macizo (no cabe un casetón en el claro corto): no aplica 6.7.8. La resistencia se revisa con una nervadura de 10 cm cada 70 cm y r_np (conservador); "&"como losa maciza lleva además As,mín = 0.0018·Ag en cada dirección cerca de la cara a tensión (6.7.6.1.1, p. 175). ",IF(AND({g678}{r}="NO",{modo}{r}="PERIMETRAL"),"No cumple 6.7.8.1.4 (a_c = "&{ac}&" cm > claro corto/6 = "&TEXT({a1}{r}*100/6,"0")&" cm): se diseña como losa apoyada en vigas (6.7.8.1.9): '
                        'nervaduras con recubrimiento de viga r_np = "&{rec_np}&" cm (Tabla 4.14.16.a), V_cR sin ×1.1 y cortante sin estribos por 5.5.3.8.1 (decisión D7). Para que cuente como encasetonada: casetón de "&TEXT(FLOOR(MIN(80,{a1}{r}*100/6),1),"0")&" cm o menor. ",""))'
                        '&IF(LEFT({c342}{r},5)="FUERA","M⁻ del borde común "&{c342}{r}&": difieren más de 50%% (3.4.2.1 c); se distribuyen 2/3 del desequilibrio (3.4.3.1, decisión D8) y se confirman los momentos con el modelo (D5). ","")&%s' % txtf, 'calc', '@')
        for nm in _tcols:
            v, kind, fmt = vals[nm]
            if isinstance(v, str) and v.startswith('='):
                v = F(v, r)
            x = cell(ws, r, CI(TC[nm]), v, kind, fmt, al='left' if nm in ('nota', 'estado', 'cort', 'apoy', 'key') else 'center', sz=9)
            if nm == 'nota':
                x.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[r].height = 30
        dv_list(ws, '%s%d' % (TC['modo'], r), '=' + N['LIST_1'])
        dv_list(ws, '%s%d' % (TC['dest'], r), '=' + N['LIST_2'])
        dv_list(ws, '%s%d' % (TC['cforce'], r), '=' + N['LIST_0'])
        dv_list(ws, '%s%d' % (TC['losa'], r), '=' + N['LIST_4'])
    ws.freeze_panes = 'C9'
    add_cf(ws, 'B9:%s%d' % (TC['estado'], T1))
    for c_ in ('g678', 'peral', 'tdia', 'cort', 'c342'):
        add_cf(ws, '%s9:%s%d' % (TC[c_], TC[c_], T1))
    ws.sheet_view.zoomScale = 80
    return ws


# ---------------------------------------------------------------- NERVADURA (detalle)
def sheet_nervadura(wb):
    ws = wb.create_sheet('NERVADURA'); S = Narr(ws, N)
    S.title('NERVADURA — DISEÑO PASO A PASO DE UN TABLERO')
    S.subtitle('Escoge el tablero en la celda amarilla. La hoja vuelve a hacer, con la fórmula a la vista, lo mismo que hace TABLEROS en un renglón; al final compara los dos resultados. La sección 11 revisa la capa de compresión (es igual para todos los tableros). '
               'En los volados (D10) las nervaduras del claro corto son de 20 cm; en las losas macizas (D11) todo es por franja de 1 m (b = 100 cm) y el número de varillas es por metro.', 42)
    S.gap()
    S.section('0. TABLERO')
    S.header()
    S.row('Tablero a revisar (y a dibujar en CROQUIS)', 'F80016 N+3.00', '—', 'SELECTOR. Lista de TABLEROS columna B.', '—', 'input', '@', 'sel')
    dv_list(ws, N['sel'].split('!')[1].replace('$', ''), '=' + N['TKEYS'])
    S.row('Renglón en TABLEROS', '=MATCH({sel},{TKEYS},0)', '—', 'Auxiliar.', '—', 'calc', '0', 'nr')

    def tv(c):
        return 'INDEX(%s,{nr})' % tcol(c)
    S.gap()
    S.section('1. GEOMETRÍA Y CARGAS')
    S.header()
    S.row('Modo', '=' + tv('modo'), '—', 'PERIMETRAL = apoyado en sus cuatro bordes; VOLADO; HUECO.', 'APOYOS / TABLEROS', 'key', '@', 'n_modo')
    S.row('Losa', '=' + tv('losa'), '—', 'NERVADA (LR-25) o MACIZA (F80032 y F80010 de azotea, losa maciza de 20 cm, decisión D11).', E2K + ' · SHELLPROP "LR-25" / "Slab1"', 'key', '@', 'n_losa')
    S.row('Nivel', '=' + tv('nivel'), '—', '—', E2K, 'calc', '@', 'n_niv')
    S.row('a_x — claro en X usado (a ejes)', '=' + tv('ax'), 'm', 'De TABLEROS (para irregulares es el rectángulo representativo). Usar claros a ejes queda del lado seguro: la tabla permite claros libres en el caso I.', DXFT + ' · TRIB_TABLERO · ' + NTCC + ' Tabla 3.4.2.1, nota final (Caso I: claros libres entre paños), p. 65', 'calc', '0.000', 'n_ax')
    S.row('a_y — claro en Y usado (a ejes)', '=' + tv('ay'), 'm', '—', DXFT, 'calc', '0.000', 'n_ay')
    S.row('a1 — claro corto', '=MIN({n_ax},{n_ay})', 'm', 'CALCULADA.', '—', 'key', '0.000', 'n_a1')
    S.row('a2 — claro largo', '=MAX({n_ax},{n_ay})', 'm', 'CALCULADA.', '—', 'calc', '0.000', 'n_a2')
    S.row('m = a1/a2', '={n_a1}/{n_a2}', '—', 'CALCULADA.', '—', 'calc', '0.000', 'n_m')
    S.row('Dirección del claro corto (ahí trabajan las nervaduras que más cargan)', '=IF({n_ax}<={n_ay},"‖ X","‖ Y")', '—', 'Su acero va por fuera; el del claro largo se apoya sobre él (d menor).', '—', 'key', '@', 'n_dir')
    S.row('Claros libres L_n corto / largo', '=TEXT(' + tv('ln1') + ',"0")&" / "&TEXT(' + tv('ln2') + ',"0")', 'cm', 'Claro a ejes menos el ancho de la nervadura principal o dala.', NTCC + ' 6.3.2.1, p. 117', 'calc', '@')
    S.row('h de la losa', '=' + tv('hr'), 'cm', 'LR-25: h_c + t; maciza: h de DATOS sección 1.', 'DATOS', 'calc', '0', 'n_h')
    S.row('b_w de las nervaduras ‖ claro corto (maciza: franja de 100 cm)', '=' + tv('bwt'), 'cm', 'b_w de DATOS; en los volados, el de sus nervaduras en la dirección del volado (decisión D10). Las nervaduras ‖ claro largo llevan la b_w general de DATOS.', 'DATOS sección 1 · Decisión D10', 'calc', '0', 'n_bwt')
    S.row('Ancho de diseño (s por nervadura; maciza: 100 cm)', '=' + tv('sb'), 'cm', 'Los momentos y cortantes de abajo son por este ancho.', 'Geometría', 'calc', '0', 'n_sb')
    S.row('Patín efectivo b_f claro corto / largo', '=TEXT(' + tv('bfc') + ',"0.0")&" / "&TEXT(' + tv('bfl') + ',"0.0")', 'cm', 'b_w + 2·el menor de 8t, (s − b_w)/2 y L_n/8 (no más que s). Maciza: 100 cm.', NTCC + ' Tabla 5.2.1.4.2, p. 89', 'calc', '@')
    S.row('SCM — sobrecarga muerta', '=' + tv('scm'), 'kg/m²', 'Del E2K, tablero por tablero.', E2K + ' · AREALOAD "Sobrecarga"', 'calc', '0', 'n_scm')
    S.row('W_m — carga viva', '=' + tv('cv'), 'kg/m²', 'Según el destino del tablero.', NTCCA + ' Tabla 6.1.2.2, p. 11', 'calc', '0', 'n_cv')
    S.row('PP de diseño (geométrico + 20 + 20)', '=' + tv('pp'), 'kg/m²', 'LR-25: (h − (s − b_w)·a_c·h_c/s²)·γc + casetón·γ_pol, con la b_w de este tablero (más concreto en los volados, D10); maciza: h·γc. Más 20 + 20 kg/m² (decisión D4). CARGAS sección 1.', NTCCA + ' 5.1.2.1, p. 9', 'calc', '0.0', 'n_pp')
    S.row('CM = PP + SCM', '={n_pp}+{n_scm}', 'kg/m²', 'CALCULADA. W_m/CM < 2.5: cumple 3.4.2.1 d). La condición c) (M⁻ de tableros vecinos que no difieran más de 50 %) se revisa en TABLEROS, columna «3.4.2.1 c)», y se reporta abajo.', NTCC + ' 3.4.2.1 c) y d), p. 64', 'calc', '0', 'n_cm')
    S.row('w — carga de servicio', '={n_cm}+{n_cv}', 'kg/m²', 'CM + W_m.', NTCCA + ' 3.4.1 d), p. 7', 'calc', '0', 'n_ws')
    S.row('w_u — carga última', '=MAX({FCcm}*{n_cm}+{FCcv}*{n_cv},IF(' + tv('dest') + '="AZOTEA",{FCacc}*({n_cm}+{Wg}),0))', 'kg/m²', 'El mayor de 1.3·CM + 1.5·W_m y, en azotea, 1.1·(CM + W_g).', NTCCA + ' 3.4.1 a) y b), p. 7 · 7.3.2.1, p. 17', 'key', '0', 'n_wu')
    S.gap()
    S.section('2. CASO Y COEFICIENTES')
    S.header()
    S.row('Bordes discontinuos (de APOYOS)', '=' + tv('nd'), '—', 'Cuenta de bordes con «DISCONTINUO» en APOYOS sección C.', 'APOYOS', 'calc', '0')
    S.row('Caso de la tabla', '=' + tv('caso'), '—', 'Automático según cuántos bordes discontinuos son cortos o largos; se puede forzar en TABLEROS. Caso I: losa colada monolíticamente con sus apoyos.', NTCC + ' Tabla 3.4.2.1, p. 65', 'key', '@', 'n_caso')
    S.row('Coef. M⁻ en bordes continuos, claro corto', '=' + tv('cNIC'), '×10⁻⁴', 'Interpolado en m.', 'TABLAS sección A', 'calc', '0', 'n_cNIC')
    S.row('Coef. M⁻ en bordes discontinuos, claro corto', '=' + tv('cNDC'), '×10⁻⁴', 'Interpolado en m.', 'TABLAS sección A', 'calc', '0', 'n_cNDC')
    S.row('Coef. M⁺ claro corto', '=' + tv('cPC'), '×10⁻⁴', 'Interpolado en m.', 'TABLAS sección A', 'calc', '0', 'n_cPC')
    S.row('Coef. M⁻ en bordes continuos, claro largo', '=' + tv('cNIL'), '×10⁻⁴', 'Interpolado en m.', 'TABLAS sección A', 'calc', '0', 'n_cNIL')
    S.row('Coef. M⁻ en bordes discontinuos, claro largo', '=' + tv('cNDL'), '×10⁻⁴', 'Interpolado en m.', 'TABLAS sección A', 'calc', '0', 'n_cNDL')
    S.row('Coef. M⁺ claro largo', '=' + tv('cPL'), '×10⁻⁴', 'Interpolado en m.', 'TABLAS sección A', 'calc', '0', 'n_cPL')
    S.row('Carga en la orilla del volado (servicio)', '=' + tv('pbor'), 'kg/m', 'Sólo para volados: barandal del E2K.', E2K + ' · FRAME OBJECT LOADS', 'calc', '0.0', 'n_pb')
    S.row('3.4.2.1 c): M⁻ del borde común con cada vecino', '=' + tv('c342'), '—', 'FUERA = con ese vecino (perimetral) los M⁻ difieren más de 50 % del menor: los coeficientes quedan fuera de su campo; se distribuye el desequilibrio (renglones de abajo) y conviene confirmar con el modelo (método reconocido, 3.4.1.1). Se avisa también si APOYOS marca distinto el borde común en los dos tableros y si el vecino es un volado.', NTCC + ' 3.4.2.1 c), p. 64', 'key', '@', 'n_c342')
    S.row('M⁻ que deja el vecino al distribuir (corto / largo)', '=TEXT(' + tv('MNcd') + ',"0.000")&" / "&TEXT(' + tv('MNld') + ',"0.000")', 't·m', 'En cada borde común que alguno de los dos tableros marque CONTINUO (APOYOS; cada uno entra con el M⁻ de su propia clasificación) con vecino de M⁻ mayor: M + 2/3·(M_vecino − M)·k/(k + k_vecino), con k = d³/a1 y d el peralte del lecho superior de cada tablero en esa dirección. Por nervadura. Los volados no entran: su momento es estático; se arma con el del volado y su desequilibrio se suma al cortante.', NTCC + ' 3.4.3.1, p. 64', 'calc', '@')
    S.row('ΔM⁺ por la distribución (corto / largo)', '=TEXT(' + tv('MPcd') + ',"0.000")&" / "&TEXT(' + tv('MPld') + ',"0.000")', 't·m', 'Donde el tablero cede negativo al vecino, el M⁺ sube la mitad de lo que baja el M⁻ en cada extremo.', NTCC + ' 3.4.3.1, p. 64', 'calc', '@')
    S.gap()
    S.section('3. MOMENTOS POR NERVADURA')
    S.secnote('Por metro: M = coef × 10⁻⁴ × w_u × a1². Por nervadura: × s (maciza: por franja de 1 m). En volado: M = (w_u·L_n²/2 + 1.3·P·L_n)·s, con L_n medido al paño del muro.')
    S.header()
    fac = '*{n_wu}*{n_a1}^2*1E-4*{n_sb}/100/1000'
    S.row('M⁺ claro corto', '=IF({n_modo}="PERIMETRAL",{n_cPC}%s+%s,0)' % (fac, tv('MPcd')), 't·m', 'Coef. M⁺ corto × 10⁻⁴ × w_u × a1² × s, más el ΔM⁺ de la distribución.', NTCC + ' 3.4.2.1, p. 64 y Tabla 3.4.2.1, p. 65 · 3.4.3.1, p. 64', 'key', '0.000', 'n_MPc')
    S.row('M⁻ claro corto', '=IF({n_modo}="PERIMETRAL",MAX(MAX({n_cNIC},{n_cNDC})%s,%s),IF({n_modo}="VOLADO",({n_wu}*(%s/100)^2/2+{FCcm}*{n_pb}*%s/100)*{n_sb}/100/1000,0))' % (fac, tv('MNcd'), tv('ln1'), tv('ln1')), 't·m', 'El mayor entre borde continuo, borde discontinuo y lo que deja el vecino al distribuir (3.4.3.1); en volado, el momento en el paño. En el borde común se arma con el mayor de los dos tableros.', NTCC + ' 3.4.3.1, p. 64', 'key', '0.000', 'n_MNc')
    S.row('M⁺ claro largo', '=IF({n_modo}="PERIMETRAL",{n_cPL}%s+%s,0)' % (fac, tv('MPld')), 't·m', 'Mismo procedimiento.', NTCC + ' Tabla 3.4.2.1, p. 65 · 3.4.3.1, p. 64', 'key', '0.000', 'n_MPl')
    S.row('M⁻ claro largo', '=IF({n_modo}="PERIMETRAL",MAX(MAX({n_cNIL},{n_cNDL})%s,%s),0)' % (fac, tv('MNld')), 't·m', 'Mismo procedimiento.', NTCC + ' Tabla 3.4.2.1, p. 65 · 3.4.3.1, p. 64', 'key', '0.000', 'n_MNl')
    sec = 4
    for blk, M, neg, lbl in (('Pc', 'n_MPc', False, 'LECHO INFERIOR — CLARO CORTO'), ('Nc', 'n_MNc', True, 'LECHO SUPERIOR — CLARO CORTO'),
                             ('Pl', 'n_MPl', False, 'LECHO INFERIOR — CLARO LARGO'), ('Nl', 'n_MNl', True, 'LECHO SUPERIOR — CLARO LARGO')):
        S.gap()
        S.section('%d. %s' % (sec, lbl)); sec += 1
        S.header()
        b = tv('b' + blk) if neg else tv('bf' + blk[1].lower())
        bwb = '{n_bwt}' if blk[1] == 'c' else '{bw}'
        S.row('b — ancho que trabaja a compresión', '=' + b, 'cm', ('Arriba en tensión: comprime la nervadura (b_w de cálculo; 20 cm en el volado). Maciza: 100 cm.' if neg else 'Arriba comprime el patín efectivo b_f (sección T). Maciza: 100 cm.'), NTCC + (' 3.9.1 e), p. 74' if neg else ' Tabla 5.2.1.4.2, p. 89 · 5.2.2.1.3.1, p. 90'), 'calc', '0.0', 'n_b' + blk)
        S.row('d — peralte efectivo con la varilla que se usa', '=' + tv('d' + blk), 'cm', 'h − r − d_b/2' + (' − 2·Ø de la malla (la varilla va bajo la malla)' if neg else '') + (' − d_b del lecho del claro corto' if blk[1] == 'l' else '') + '. r = ' + ('r superior (r_np si el tablero no cumple 6.7.8; r_ext si la cara queda expuesta); − 2 cm si d ≤ 20 cm sin la precaución de 3.9.2' if neg else 'r inferior (r_np si el tablero no cumple 6.7.8; r_ext en volados)') + '.', NTCC + ' Tabla 4.14.16.a, p. 85' + (' · 3.9.1 d) y 3.9.2, p. 74' if neg else ''), 'calc', '0.00', 'n_d' + blk)
        bmn = ('IF({n_modo}="VOLADO",MIN({s},2*{n_bwt}),{n_bwt})' if blk == 'Nc' else bwb)
        S.row("As mín = máx(0.80·√f'c, 14)/fy · b_w · d (maciza: 0.0018·b·h)", '=IF({n_losa}="MACIZA",{Asm_mac},{pmin}*%s*{n_d%s})' % (bmn, blk), 'cm²', 'Sobre b_w (en el volado, patín a tensión: el menor de b_f y 2b_w). Se dispensa si la As colocada es ≥ 4/3 de la requerida. En la maciza, 0.0018·b·h por metro en cada lecho, sin dispensa.', NTCC + ' 6.3.5.1.2 y 6.3.5.1.3, p. 122 · 6.7.6.1.1, p. 175', 'calc', '0.00', 'n_min' + blk)
        S.row('As requerida (TABLEROS)', '=' + tv('As' + blk), 'cm²', "q = 1 − √(1 − 2Mu/(F_R·b·d²·f''c)); As = q·f''c·b·d/fy; luego el mayor de esa y el menor de (As mín, 4/3·As). −1 = no alcanza o rebasa p_max.", NTCC + ' ecs. 5.2.2.1.1.1.a y b, p. 89 · 6.3.5.1.3, p. 122', 'key', '0.00', 'n_As' + blk)
        S.row('Varilla que se usa', '=' + tv('v' + blk), 'No.', 'La preferida (No. 4) o la siguiente que quepa.', NTCC + ' 6.7.6.1.3, p. 175', 'calc', 'General', 'n_v' + blk)
        S.row('Número de varillas (maciza: por metro)', '=' + tv('n' + blk), 'pzas', 'Al menos 1 (la No. 4 corrida). Si As colocada < 4/3 de la de flexión, al menos 2.' + (' Arriba: 1 corrida + el resto bastones.' if neg else ' Todas corridas.') + ' Maciza: 100/separación.', NTCC + ' 6.7.6.1.3, p. 175 · 6.3.5.1.1, p. 122', 'key', '0.00', 'n_n' + blk)
        S.row('Separación (sólo losa maciza)', '=' + tv('s' + blk), 'cm', 'La menor de la que da As (redondeada a 2.5 cm), 2h y el tope de DATOS sección 8; si sale menor que la mínima práctica, sube el número de varilla.', 'DATOS sección 8 · Criterio de esta hoja', 'calc', '0.0', 'n_s' + blk)
        asp = '{n_n%s}*%s' % (blk, ab('{n_v%s}' % blk))
        S.row('a — profundidad del bloque', '=%s*{fy}/({fpc}*{n_b%s})' % (asp, blk), 'cm', "a = As·fy/(f''c·b)." + ('' if neg else ' Debe quedar dentro del patín (a ≤ t).'), NTCC + ' 3.6.1 e), p. 68', 'calc', '0.00', 'n_a' + blk)
        S.row('M_R — momento resistente', '={FRf}*%s*{fy}*({n_d%s}-{n_a%s}/2)/1E5' % (asp, blk, blk), 't·m', "F_R·As·fy·(d − a/2).", NTCC + ' ec. 5.2.2.1.1.1.b, p. 89 · ' + NAIM24 + ' pág. 13', 'calc', '0.000', 'n_MR' + blk)
        S.row('Mu / M_R', '=IF({n_As%s}<0,9.99,IF({n_MR%s}>0,{%s}/{n_MR%s},0))' % (blk, blk, M, blk), '—', 'Debe ser ≤ 1. 9.99 = la sección no alcanza.', '—', 'key', '0.00', 'n_FU' + blk)
        S.row('ε_t — deformación del acero de tensión', '=0.003*({n_d%s}-{n_a%s}/{beta1})/({n_a%s}/{beta1})' % (blk, blk, blk), '—', 'ε_t = 0.003·(d − c)/c con c = a/β1. Debe ser ≥ 0.005 para usar F_R = 0.9 (sección controlada por tensión).', NTCC + ' Tabla 3.8.2.2, p. 70 · 6.7.3.3.1, p. 169', 'calc', '0.0000', 'n_e' + blk)
        S.row('¿Caben en la nervadura? (maciza: separación ≥ mínima)', '=IF({n_n%s}=0,"—",IF({n_losa}="MACIZA",IF({n_s%s}>={smin_mac},"CUMPLE","NO CABE"),IF({n_n%s}*%s+({n_n%s}-1)*MAX({slib},%s)<=%s-2*%s-2*%s+0.01,"CUMPLE","NO CABE")))' % (blk, blk, blk, db('{n_v%s}' % blk), blk, db('{n_v%s}' % blk), bwb, (tv('rt') if neg else tv('rb')), dbz('{est_sec}')),
              '—', 'n·d_b + (n−1)·s_libre ≤ b_w − 2r − 2·estribo, con la b_w de esa nervadura.', NTCC + ' 14.2.1, p. 347', 'calc', '@')
        S.row('Estado', '=IF(AND({n_FU%s}<=1,{n_e%s}>=0.005),"CUMPLE","NO CUMPLE")' % (blk, blk), '—', '—', '—', 'key', '@')
    S.gap()
    S.section('%d. CORTANTE EN LA NERVADURA' % sec); sec += 1
    S.secnote('V por metro a d del paño = (L_n/2 − d)·(0.95 − 0.5·m)·w_u, × 1.15 si hay bordes continuos y discontinuos (fórmula de la NTC 2017 que la 2023 ya no trae; se usa como método reconocido, 3.4.1.1). Por nervadura × s. '
              'V_cR sin estribos con la ec. 5.5.3.2.1 y la cuantía del acero superior.', 42)
    S.header()
    S.row('V_u por nervadura', '=' + tv('Vu'), 't', 'Incluye ΔV = (M⁻ del volado vecino − M⁻ propio)/L_n cuando las nervaduras rematan en un volado (si llegan por el claro largo se suma igual, del lado seguro); ese ΔV no baja a lo largo del claro y entra así en las hileras a macizar. En volado: V = (w_u·(L_n − d) + 1.3·P)·s.', NTCC + ' 3.4.1.1, p. 64 · 6.7.4.3.2, p. 172 (sección crítica a d del paño)', 'key', '0.000', 'n_Vu')
    S.row('p — cuantía del acero superior en el apoyo', '={n_nNc}*%s/({n_bNc}*{n_dNc})' % ab('{n_vNc}'), '—', 'As_sup / (b_w·d) (maciza: As por metro / (100·d)).', NTCC + ' 5.5.3.2.1, p. 95', 'calc', '0.0000', 'n_p')
    S.row('λs — factor de tamaño', '=MIN(1,SQRT(2/(1+0.04*{n_dNc})))', '—', '√(2/(1 + 0.04d)) ≤ 1, d en cm.', NTCC + ' ec. 5.5.3.2.1.b, p. 95', 'calc', '0.000', 'n_ls')
    S.row('¿Cumple 6.7.8 (losa encasetonada)?', '=' + tv('g678'), '—', 'Claro corto ≥ 6·a_c, a_c ≤ 80, b_w ≥ 10, h ≤ 3.5b_w, t ≥ máx(a_c/13, 5). Si NO, se diseña como losa apoyada en vigas: nervaduras con recubrimiento de viga (r_np) y sin el ×1.1.', NTCC + ' 6.7.8.1.2 a 6.7.8.1.4, p. 187 · 6.7.8.1.9, p. 188', 'key', '@', 'n_678')
    S.row('V_cR — cortante que toma el concreto', "={FRv}*2*{n_ls}*{n_p}^(1/3)*SQRT({fc})*{n_bNc}*{n_dNc}*IF({n_678}=\"SÍ\",{VcRinc},1)/1000", 't',
          "F_R·2·λs·p^(1/3)·√f'c·b_w·d (× 1.1 si cumple 6.7.8). Sin estribos: elemento en que no es práctico colocarlos. Maciza: misma expresión con b = 100 cm.", NTCC + ' 5.5.3.8.1, p. 97 · ec. 5.5.3.2.1.a, p. 95 · 6.7.8.1.5, p. 187 · Tabla 3.8.2.1, p. 70', 'key', '0.000', 'n_VcR')
    S.row('Estado', '=' + tv('cort'), '—', 'Si V_u > V_cR se mide desde el paño hasta donde V baja a V_cR y se compara con el macizo de orilla; lo que falte se cubre quitando hileras de casetones (macizo). '
          'Sin estribos por 5.5.3.8.1 (elemento en que no es práctico colocarlos, criterio del Proyectista); si el tablero no cumple 6.7.8 tampoco aplica la exención de la Tabla 6.3.5.4.1 c).', NTCC + ' 5.5.3.8.1, p. 97 · 5.5.3.1.1, p. 94 · Tabla 6.3.5.4.1 c), p. 123', 'key', '@')
    S.gap()
    S.section('%d. ESPESOR MÍNIMO (SE PUEDE OMITIR EL CÁLCULO DE FLECHAS)' % sec); sec += 1
    S.secnote('Tabla 6.7.3.1.2 con α_fm = promedio de α_f de los cuatro bordes: un muro bajo todo el borde cuenta como viga infinitamente rígida (α_f = ∞); una nervadura principal, como I_b/I_s. '
              'α_fm > 2: h = L_n·(0.8 + fy/14 000)/(36 + 9β) ≥ 9 cm; 0.2 < α_fm ≤ 2: h = L_n·(0.8 + fy/14 000)/(36 + 5β(α_fm − 0.2)) ≥ 13 cm; α_fm ≤ 0.2: Tabla 6.7.3.1.1 (L_n/30 o L_n/33, ≥ 13 cm). '
              '+10% con borde discontinuo sin viga de α_f ≥ 0.8. L_n = claro libre largo, β = L_n largo / L_n corto. En la losa nervada la hoja compara contra h_eq (losa maciza de la misma inercia), no contra los 25 cm; en las losas macizas (D11), contra su h.', 66)
    S.header()
    S.row('α_fm', '=' + tv('afm'), '—', 'Promedio de los cuatro bordes (99 = muro).', NTCC + ' Tabla 6.7.3.1.2 nota [1], p. 168 · comentario de 6.7.1, p. 165', 'calc', '0.00')
    S.row('h mínimo (o L/8 si es volado)', '=' + tv('hmin'), 'cm', 'Volado: las nervaduras son vigas, h ≥ L/8 y ≥ 25 cm (si no carga elementos frágiles).', NTCC + ' Tabla 6.7.3.1.2 y 6.7.3.1.2.1, p. 168 · Tabla 6.7.3.1.1, p. 167 · Tabla 6.3.3.1.1, p. 120', 'key', '0.00')
    S.row('h que se compara (h_eq; h si es volado o losa maciza)', '=' + tv('heqv'), 'cm', 'DATOS secciones 1 y 4.', 'Criterio de esta hoja', 'calc', '0.00')
    S.row('Estado', '=' + tv('peral'), '—', 'Si no cumple: calcular flechas (13.4) y revisar vibración (13.7).', NTCC + ' 6.7.3.2.1, p. 168 · 13.7.1.1, p. 343', 'key', '@')
    S.gap()
    S.section('%d. ACOMODO DE CASETONES, BASTONES Y DETALLE' % sec); sec += 1
    S.header()
    S.row('Lado del casetón ‖ X', '=IF(AND({n_modo}="VOLADO",{n_dir}="‖ Y"),{ac_vol},{ac})', 'cm', 'a_c; en el volado, el lado que queda entre sus nervaduras de 20 cm es s − b_w = 50 cm (casetón 50×60×19, decisión D10).', 'DATOS sección 1 · Decisión D10', 'calc', '0', 'n_acx')
    S.row('Lado del casetón ‖ Y', '=IF(AND({n_modo}="VOLADO",{n_dir}="‖ X"),{ac_vol},{ac})', 'cm', 'Mismo criterio.', 'DATOS sección 1 · Decisión D10', 'calc', '0', 'n_acy')
    S.row('Ancho de las nervaduras entre casetones ‖ X / ‖ Y', '=TEXT({s}-{n_acx},"0")&" / "&TEXT({s}-{n_acy},"0")', 'cm', 's − lado del casetón: la separación s = 70 cm no cambia, así la retícula queda alineada.', 'Geometría', 'calc', '@')
    S.row('Casetones en X', '=IF({n_losa}="MACIZA",0,MAX(0,INT(({n_ax}*100-{bp}+{s}-{n_acx})/{s})))', 'pzas', 'INT((a_x − b_p + b_w)/s). Lo que sobra se reparte como macizo en las dos orillas. Maciza: no lleva.', 'Geometría', 'key', '0', 'n_ncx')
    S.row('Macizo en cada orilla X', '=IF({n_ncx}>0,({n_ax}*100-{bp}-({n_ncx}*{n_acx}+({n_ncx}-1)*({s}-{n_acx})))/2,{n_ax}*100/2)', 'cm', 'Del paño de la nervadura principal al primer casetón.', 'Geometría', 'calc', '0.0', 'n_mx')
    S.row('Casetones en Y', '=IF({n_losa}="MACIZA",0,MAX(0,INT(({n_ay}*100-{bp}+{s}-{n_acy})/{s})))', 'pzas', 'Mismo cálculo.', 'Geometría', 'key', '0', 'n_ncy')
    S.row('Macizo en cada orilla Y', '=IF({n_ncy}>0,({n_ay}*100-{bp}-({n_ncy}*{n_acy}+({n_ncy}-1)*({s}-{n_acy})))/2,{n_ay}*100/2)', 'cm', 'Mismo cálculo.', 'Geometría', 'calc', '0.0', 'n_my')
    S.row('Nervaduras secundarias ‖ X (entre casetones)', '=MAX(0,{n_ncy}-1)', 'pzas', 'Una entre cada par de hileras de casetones.', 'Geometría', 'calc', '0', 'n_nrx')
    S.row('Nervaduras secundarias ‖ Y', '=MAX(0,{n_ncx}-1)', 'pzas', 'Mismo criterio.', 'Geometría', 'calc', '0', 'n_nry')
    S.row('Bastón superior ‖ claro corto: largo desde el paño', '=' + tv('Lbc'), 'cm', 'Punto de inflexión de M(ξ) = −M1(1−ξ) − M2·ξ + 4M0·ξ(1−ξ), con M1 = M⁻ de diseño (con la distribución), M2 = M⁻ del otro extremo y M0 = M⁺ + (M1+M2)/2 con los M⁻ sin distribuir (el momento estático no cambia), más el mayor de d, 12d_b y L_n/16, y no menos que L_d; si no hay punto de inflexión el bastón va de lado a lado '
          '(«otros casos» de la Tabla 14.4.2.4 si r < d_b o la separación libre < 2d_b). En volado: hacia el claro de respaldo.', NTCC + ' 6.3.7.3.4, p. 127 · 6.3.7.3.10.4, p. 129 · 14.4.2.1, p. 352', 'key', '0', 'n_Lbc')
    S.row('Bastón superior ‖ claro largo: largo desde el paño', '=' + tv('Lbl'), 'cm', 'Mismo criterio.', NTCC + ' 6.3.7.3.10.4, p. 129', 'key', '0', 'n_Lbl')
    ldh_ = lambda v_: 'INDEX({LDHCOL},MATCH(%s,{BARCOL},0))' % v_
    mxv = lambda col, nerv: 'SUMPRODUCT(MAX((%s="PERIMETRAL")%s*%s))' % (tcol('modo'), ('*(%s="NERVADA")' % tcol('losa')) if nerv else '', tcol(col))
    S.row('Ancho disponible para anclar en el macizo o dala de borde (decisión D9)', '={b_borde}-{rec_np}', 'cm', 'b_borde (40 cm, DATOS sección 1) menos el recubrimiento de la cara exterior. En todo borde discontinuo va un macizo o dala de borde de 40 cm de ancho en el peralte de la losa, apoyado en el muro de 15 cm; '
          'la nervadura nace en su paño interior. L_dh se mide del paño al lado exterior del doblez. Su peso extra (DATOS sección 4) baja directo al muro: en los momentos del tablero se desprecia porque está sobre el apoyo.', 'Decisión D9 · ' + NTCC + ' 14.4.3.1, p. 355', 'calc', '0.0', 'n_anc')
    S.row('Varilla inferior más gruesa de los tableros nervados (la que se ancla)', '=MAX(%s,%s)' % (mxv('vPc', True), mxv('vPl', True)), 'No.', 'La mayor de los lechos inferiores de todos los tableros nervados perimetrales de TABLEROS: la revisión vale para todos los bordes discontinuos.', 'TABLEROS', 'calc', 'General', 'n_vanc1')
    S.row('Varilla superior más gruesa de los tableros perimetrales (nervados y macizos)', '=MAX(%s,%s)' % (mxv('vNc', False), mxv('vNl', False)), 'No.', 'La mayor de los lechos superiores de todos los tableros perimetrales de TABLEROS (del lado seguro: no todas llegan a un borde discontinuo).', 'TABLEROS', 'calc', 'General', 'n_vanc2')
    S.row('Varilla inferior de integridad en bordes discontinuos (1.25 fy en el paño)', '=IF(AND(1.25*%s<={n_anc},{n_anc}>=15),"CUMPLE","NO CUMPLE: 1.25·L_dh = "&TEXT(1.25*%s,"0.0")&" cm y ≥ 15 cm dentro del apoyo; hay "&TEXT({n_anc},"0.0")&" cm")' % (ldh_('{n_vanc1}'), ldh_('{n_vanc1}')), '—',
          'Una varilla inferior por nervadura, continua y anclada para 1.25 fy en el paño, con ≥ 150 mm dentro del apoyo (el macizo o dala de borde de 40 cm, decisión D9). Con No. 4: 1.25·L_dh ≈ 32 cm ≤ 36 cm.',
          NTCC + ' 6.7.8.1.7, p. 188 · 6.7.7.4.1.1 a), p. 179 · 14.4.3.1, p. 355', 'key', '@', 'n_anc1')
    S.row('Bastón superior en bordes discontinuos (fy en el paño)', '=IF(%s<={n_anc},"CUMPLE","NO CUMPLE: L_dh = "&TEXT(%s,"0.0")&" cm > "&TEXT({n_anc},"0.0")&" cm disponibles en el macizo de borde")' % (ldh_('{n_vanc2}'), ldh_('{n_vanc2}')), '—',
          'El negativo se dobla con gancho de 90° y extensión de 12d_b dentro del macizo o dala de borde de 40 cm (decisión D9) para desarrollarse en el paño. Con No. 4: L_dh ≈ 25 cm; con No. 5, ≈ 32 cm.',
          NTCC + ' 6.7.7.4.1.1 b), p. 179 · 14.4.3.1, p. 355 · Tabla 14.3.1, p. 349 · NTC-Mampostería 2023 (CDMX) 4.1.17.2, p. 43', 'key', '@', 'n_anc2')
    S.row('Refuerzo de esquina exterior (tableros de esquina sobre muros)', 'arriba y abajo en L_largo/5', '—', 'Para el M⁺ máximo del tablero, en un cuadro de un quinto del claro largo desde la esquina.', NTCC + ' 6.7.7.3.1, p. 178', 'calc', '@')
    S.gap()
    S.section('%d. CAPA DE COMPRESIÓN — MALLA, DIAFRAGMA Y CARGA CONCENTRADA (IGUAL PARA TODOS LOS TABLEROS NERVADOS; DECISIÓN D3)' % sec); sec += 1
    S.header()
    capa0 = S.r
    S.row('As de la malla escogida', '=INDEX(TABLAS!$D$%d:$D$%d,MATCH({malla},TABLAS!$B$%d:$B$%d,0))' % (TB_W0, TB_W1, TB_W0, TB_W1), 'cm²/m', 'TABLAS sección C.', 'TABLAS', 'calc', '0.000', 'n_asm')
    S.row('Ø del alambre', '=INDEX(TABLAS!$C$%d:$C$%d,MATCH({malla},TABLAS!$B$%d:$B$%d,0))/10' % (TB_W0, TB_W1, TB_W0, TB_W1), 'cm', 'TABLAS sección C.', 'TABLAS', 'calc', '0.000', 'n_fim')
    S.row('Contracción y temperatura: As mín = 0.0018·t·100', '=0.0018*{t}*100', 'cm²/m', 'Ya no depende de fy ni de la exposición (la fórmula 660x₁/(fy(x₁+100)) era de 2017).', NTCC + ' 13.8.2.1.1, p. 345 · 6.7.8.1.8, p. 188', 'calc', '0.000', 'n_as1')
    S.row('Estado (temperatura)', '=IF({n_asm}>={n_as1},"CUMPLE","NO CUMPLE")', '—', '—', '—', 'key', '@')
    S.row('Diafragma: cuantía mínima por cortante = 0.0025·t·100', '=0.0025*{t}*100', 'cm²/m', 'Sólo cuenta el acero de la capa (el de cambios volumétricos sí contribuye; el de flexión de las nervaduras no).', NTCC + ' 6.8.6.4 y 6.8.6.3, p. 196 · 6.8.1.1 a), p. 188', 'calc', '0.000', 'n_asd')
    S.row('Estado (diafragma)', '=IF({n_asm}>={n_asd}-0.0005,"CUMPLE","NO CUMPLE: "&TEXT({n_asm}/{n_asd},"0%")&" — más capa o malla más pesada")', '—', 'Decisión D3: capa de 6 cm con 6x6-4/4. Con la capa de 5 cm del plano ninguna malla de la lista cumplía a la vez diafragma y recubrimiento.', '—', 'key', '@')
    S.row('Recubrimiento de la malla centrada en la capa (zonas no expuestas)', '=({t}-2*{n_fim})/2', 'cm', '(t − 2Ø)/2 arriba y abajo. Debe ser ≥ r (2 cm, interior A1).', NTCC + ' Tabla 4.14.16.a, p. 85 · 4.3.2, p. 76', 'calc', '0.00', 'n_rm')
    S.row('Estado (recubrimiento de la malla, zonas no expuestas)', '=IF({n_rm}>={rec}-0.005,"CUMPLE","NO CUMPLE")', '—', '—', '—', 'key', '@')
    S.row('Zonas con la cara superior expuesta (selectores de DATOS sección 1)', '=IF(AND({exp_vol}<>"SÍ",{exp_az}<>"SÍ"),"ninguna",IF({exp_vol}="SÍ","volados F80012 y F80019","")&IF(AND({exp_vol}="SÍ",{exp_az}="SÍ")," y ","")&IF({exp_az}="SÍ","azotea",""))', '—',
          'Ahí la malla baja a r_ext de la cara superior (así la toma TABLEROS para el peralte del lecho superior) y abajo, contra el casetón, le queda t − r_ext − 2Ø, que también debe ser ≥ r.', NTCC + ' Tabla 4.3.1 c) y 4.3.2, p. 76 · Tabla 4.14.16.a, p. 85', 'calc', '@', 'n_zexp')
    S.row('Recubrimiento abajo de la malla con la malla a r_ext de la cara superior', '={t}-{rec_ext}-2*{n_fim}', 'cm', 't − r_ext − 2Ø.', NTCC + ' Tabla 4.14.16.a, p. 85', 'calc', '0.00', 'n_rm2')
    S.row('Estado (recubrimiento de la malla, zonas expuestas)', '=IF({n_zexp}="ninguna","CUMPLE — no hay zonas expuestas",IF({n_rm2}>={rec}-0.005,"CUMPLE","NO CUMPLE en "&{n_zexp}&": con r_ext = "&TEXT({rec_ext},"0")&" cm arriba quedan "&TEXT({n_rm2},"0.0")&" cm abajo (mín. "&TEXT({rec},"0")&"): proteger esa cara (piso e impermeabilizante; selector en DATOS) o capa más gruesa ahí"))', '—',
          'Con capa de 6 cm no caben 4 cm arriba y 2 cm abajo de la malla. Lo resuelve el Responsable: proteger la cara superior de esas zonas con piso e impermeabilizante (y poner NO en el selector de DATOS, como en azotea) o engrosar la capa ahí.', '—', 'key', '@')
    S.row('Separación de la malla (15.24 cm) ≤ máx. permitida', '=IF(15.24<=MIN(5*{t},45,30),"CUMPLE","NO CUMPLE")', '—', 'El menor de 5t y 45 cm (diafragma) y 30 cm entre alambres soldados.', NTCC + ' 6.8.7.3, p. 196 · 2.4.2.1.6.3, p. 44', 'calc', '@')
    S.row('Espesor de pisos y rellenos que reparten la carga concentrada (45°)', 0, 'cm', 'CAPTURA. 0 = la carga actúa directo sobre la capa (lo más estricto). Si cuentas el piso y el pegapiso (≈ 6 cm), pon aquí su espesor.', 'Criterio de esta hoja', 'input', '0.0', 'tpz')
    S.row('P_u = FC_CV × 1 000 kg en 10 × 10 cm', '={FCcv}*' + N['Pcap'], 'kg', 'Carga concentrada mínima entre nervaduras.', NTCC + ' 6.7.8.1.6, p. 187 · ' + NTCCA + ' 3.4.1 a), p. 7', 'calc', '0', 'n_Pu')
    S.row('d de la capa (malla al centro)', '={t}/2', 'cm', '—', 'Geometría', 'calc', '0.00', 'n_dc')
    S.row('b₀ — perímetro crítico a d/2', '=4*(10+2*{tpz}+{n_dc})', 'cm', '4·(c + d) con c = 10 cm más el reparto.', NTCC + ' 5.6, p. 99–100', 'calc', '0.0', 'n_bo')
    S.row('V_cR punzonamiento de la capa', "={FRv}*MIN(1,0.5*(1+2/1),0.27*(2+40*{n_dc}/{n_bo}))*MIN(1,SQRT(2/(1+0.04*{n_dc})))*SQRT({fc})*{n_bo}*{n_dc}", 'kg', "v_cR = F_R·(el menor de 1.0, 0.5(1 + 2/β) y 0.27(2 + α_s·d/b₀))·λs·√f'c, con β = 1 y α_s = 40.", NTCC + ' Tabla 5.6.2.1.1, p. 100', 'calc', '0', 'n_vcp')
    S.row('Estado (punzonamiento)', '=IF({n_Pu}<={n_vcp},"CUMPLE","NO CUMPLE: P_u/V_cR = "&TEXT({n_Pu}/{n_vcp},"0.00")&" — más capa")', '—', 'Contar el reparto del piso (renglón de arriba) sólo ayuda aquí: no resuelve el diafragma ni el acero mínimo de abajo.', '—', 'key', '@')
    S.row('v_uv — esfuerzo cortante en la sección crítica', '={n_Pu}/({n_bo}*{n_dc})', 'kg/cm²', 'P_u/(b₀·d).', NTCC + ' 6.7.6.1.2, p. 175', 'calc', '0.00', 'n_vuv')
    S.row("Límite 0.5·F_R·λs·λ·√f'c", "=0.5*{FRv}*MIN(1,SQRT(2/(1+0.04*{n_dc})))*SQRT({fc})", 'kg/cm²', 'Si v_uv lo rebasa, se pide acero mínimo alrededor de la carga concentrada (ec. 6.7.6.1.2).', NTCC + ' 6.7.6.1.2, p. 175', 'calc', '0.00', 'n_vlim')
    S.row('b_losa — ancho en que se coloca ese acero', '=10+2*{tpz}+2*1.5*{t}', 'cm', 'Ancho cargado más 1.5h a cada lado.', NTCC + ' Tabla 6.7.4.2.2.3, p. 171', 'calc', '0.0', 'n_bls')
    S.row('As,mín alrededor de la carga = 5·v_uv·b_losa·b₀/(F_R·α_s·fy)', '=IF({n_vuv}>{n_vlim},5*{n_vuv}*{n_bls}*{n_bo}/({FRv}*40*{fym}),0)', 'cm²', 'α_s = 40 (carga interior); fy de la malla.', NTCC + ' ec. 6.7.6.1.2, p. 175 · Tabla 5.6.2.1.1, p. 100', 'calc', '0.000', 'n_asc')
    S.row('Estado (acero mínimo alrededor de la carga concentrada)', '=IF({n_asm}/100*{n_bls}>={n_asc}-0.0005,"CUMPLE","NO CUMPLE: la malla da "&TEXT({n_asm}/100*{n_bls},"0.00")&" de "&TEXT({n_asc},"0.00")&" cm² — malla más pesada")', '—', 'Malla dentro de b_losa contra As,mín.', NTCC + ' 6.7.6.1.2, p. 175', 'key', '@')
    S.row('m_R de la capa por unidad de ancho', "={FRf}*{n_asm}/100*{fym}*({n_dc}-{n_asm}/100*{fym}/({fpc})/2)", 'kg·cm/cm', "F_R·As·fy·(d − a/2) por cm.", NTCC + ' ec. 5.2.2.1.1.1.b, p. 89', 'calc', '0.0', 'n_mR')
    S.row('m requerido = P_u/(4π) (abanico de líneas de fluencia, m = m′)', '={n_Pu}/(4*PI())', 'kg·cm/cm', 'Mecanismo de abanico alrededor de la carga; la malla única al centro da m = m′. Ignorar el área cargada queda del lado seguro.', 'Criterio de esta hoja · ' + NTCC + ' 3.4.1.1, p. 64 (líneas de fluencia) · véase también ec. 3.4.5.1, p. 66 (se desprecia 2r/3R_b, del lado seguro) · 6.7.8.1.8, p. 188', 'calc', '0.0', 'n_mq')
    S.row('Estado (flexión de la capa)', '=IF({n_mq}<={n_mR},"CUMPLE","NO CUMPLE")', '—', '—', '—', 'key', '@')
    N['CAPA_RNG'] = 'NERVADURA!$C$%d:$C$%d' % (capa0, S.r - 1)
    S.gap()
    S.section('%d. COMPROBACIÓN CONTRA TABLEROS' % sec)
    S.header()
    S.row('Mu/M_R inferior corto (aquí / TABLEROS)', '=TEXT({n_FUPc},"0.000")&" / "&TEXT(%s,"0.000")' % tv('FUPc'), '—', 'Deben coincidir.', '—', 'calc', '@')
    S.row('Mu/M_R superior corto (aquí / TABLEROS)', '=TEXT({n_FUNc},"0.000")&" / "&TEXT(%s,"0.000")' % tv('FUNc'), '—', 'Deben coincidir.', '—', 'calc', '@')
    S.row('¿Coinciden?', '=IF(AND(ABS({n_FUPc}-%s)<0.0005,ABS({n_FUNc}-%s)<0.0005,ABS({n_FUPl}-%s)<0.0005,ABS({n_FUNl}-%s)<0.0005,ABS({n_VcR}-%s)<0.0005),"CUMPLE — coinciden","REVISAR")' % (tv('FUPc'), tv('FUNc'), tv('FUPl'), tv('FUNl'), tv('VcR')),
          '—', 'Control interno de la hoja.', '—', 'key', '@')
    S.row('ESTADO DEL TABLERO (TABLEROS)', '=' + tv('estado'), '—', '—', '—', 'key', '@')
    add_cf(ws, 'C1:F%d' % S.r)
    ws.sheet_view.zoomScale = 85
    return ws


# ---------------------------------------------------------------- PRINCIPALES
PC = {}
_pcols = ['np', 'nivel', 'eje', 'de', 'a', 'L', 'ei', 'ej', 'ci', 'cj', 'cond', 'raiz', 'bq', 'hq', 'cambio',
          'At', 'wt', 'Lm', 'wm', 'waz', 'scm', 'cv', 'wl', 'wpp', 'w', 'wD', 'wL', 'wW',
          'SP', 'SM', 'SV', 'SPm', 'SMs', 'Ri', 'Rj', 'Ris', 'Rjs', 'Mp', 'Mn', 'Vu', 'Ms',
          'd', 'AsP', 'vP', 'nPx', 'nP', 'MRP', 'FUP', 'eP', 'AsN', 'vN', 'nN', 'MRN', 'FUN', 'eN',
          'VcR', 'se', 'sext', 'cort', 'bfq', 'Ig', 'dr', 'nAs', 'kd', 'Icr', 'Ie', 'dlt', 'dlim', 'Pu', 'Kf', 'fcm', 'm418', 'htab', 'h418', 'hflex', 'hrec', 'estado', 'nota', 'nmax']
for i, nm in enumerate(_pcols):
    PC[nm] = CL(2 + i)
PC_LAST = CL(2 + len(_pcols) - 1)
XROW = {}
# vigas a lo largo del perímetro del edificio (fachadas y = 1.919, y = 21.446, x = 0, x = 8.85): integridad 6.3.7.7.1 y 2h de 6.3.7.7.7
PERIM = {t['np'] for t in NPS if t['cond'] != 'BORDE DE VOLADO' and ((t['ori'] == 'H' and min(abs(t['c'] - 1.919), abs(t['c'] - 21.446)) < 0.01) or (t['ori'] == 'V' and min(abs(t['c'] - 0.0), abs(t['c'] - 8.85)) < 0.01))}
_ntab = {}
for _tr in D['trib']:
    if _tr['tramo'].startswith('NP') and _tr['area'] > 0.05:
        _ntab.setdefault(_tr['tramo'], set()).add(_tr['tablero'])
ONESIDE = {t['np'] for t in NPS if t['np'] in PERIM or len(_ntab.get(t['np'], ())) < 2}
PROP = {}
try:
    PROP = json.load(open(W + 'work/propuesta_np.json'))
except Exception:
    PROP = {}


def trib_sum(col, tramo):
    return 'SUMIFS(TRIB!$%s$%d:$%s$%d,TRIB!$B$%d:$B$%d,"%s")' % (col, R0, col, R1, R0, R1, tramo)


def sheet_principales(wb):
    ws = wb.create_sheet('PRINCIPALES'); ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 1.78
    last = PC_LAST
    ws.merge_cells('B1:%s1' % last)
    c = ws['B1']; c.value = 'PRINCIPALES — NERVADURAS PRINCIPALES SIN MURO ABAJO (TRABAJAN COMO TRABE)'
    c.font = font(13, True, color='FFFFFF'); c.fill = fill(TITLE_BG); ws.row_dimensions[1].height = 26
    ws.merge_cells('B2:%s2' % last)
    ws['B2'].value = ('Cada renglón es un tramo de dala del modelo que NO tiene muro abajo, entre dos apoyos. Carga = áreas tributarias del DXF × w_u del tablero + muro de planta alta encima + la AZOTEA que ese muro recibe por su dala de N+6.00 + cargas lineales del E2K + peso extra de la trabe + reacciones de otras principales (TRANSFER, directas o por el muro). '
                      'Apoyada: M⁺ = wL²/8 + ΣPab/L, M⁻ = 0.667·M⁺. Volado: M⁻ = wL²/2 + ΣP·brazo. Recubrimiento de viga (4 cm), estribos cerrados No. 3 y A_v,mín en todo el claro; flecha con la inercia efectiva de la Tabla 13.4.3.2; con muro encima, además la NTC-Mampostería 4.1.18 (con H/L < 0.6, sin efecto de arco: la trabe toma todo el muro, decisión D2).')
    ws['B2'].font = font(9, i=True, color=GRAY); ws['B2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 42
    ws.merge_cells('B3:%s3' % last)
    ws['B3'].value = ('b y h (amarillo) son la SECCIÓN PROPUESTA: la hoja la dimensionó para que el tramo cumpla todo; donde dice «PERALTAR» la dala de 15×25 del modelo no alcanza y debe ir como trabe descolgada de esa sección. '
                      'Decisión D2: se aceptan estas secciones. Los extremos con castillo C1…C8 quedan firmes (aceptados, decisión D1; si alguno se rechaza en APOYOS sección B el tramo cambia). «h mín. por flexión, Tabla 6.3.3.1.1 y NTC-M 4.1.18» no incluye la flecha ni el acomodo de varillas; la flecha y el muro sobre la viga se revisan con la h capturada y el estado propone el h que falta.')
    ws['B3'].font = font(9, True, color=WARN_FG); ws['B3'].fill = fill(WARN_BG); ws['B3'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[3].height = 30
    groups = [('np', 'raiz', 'TRAMO Y APOYOS'), ('bq', 'cambio', 'SECCIÓN PROPUESTA'), ('At', 'wW', 'CARGA REPARTIDA (última; servicio D / W_m / W)'), ('SP', 'Rjs', 'CARGAS PUNTUALES Y REACCIONES'),
              ('Mp', 'Ms', 'ELEMENTOS MECÁNICOS'), ('d', 'eN', 'FLEXIÓN'), ('VcR', 'cort', 'CORTANTE'), ('bfq', 'dlim', 'FLECHA (13.4) — SECCIÓN T CON LA CAPA'), ('Pu', 'm418', 'MURO SOBRE VIGA (NTC-M 4.1.18)'), ('htab', 'nmax', 'PERALTE Y ESTADO')]
    for a, b, t in groups:
        ws.merge_cells('%s6:%s6' % (PC[a], PC[b]))
        x = ws['%s6' % PC[a]]; x.value = t; x.font = font(9, True, color='FFFFFF'); x.fill = fill(SEC_BG); x.alignment = Alignment(horizontal='center')
    heads = {'np': ('Tramo', 7), 'nivel': ('Nivel', 7.5), 'eje': ('Eje / línea', 16), 'de': ('De (x, y)', 13), 'a': ('A (x, y)', 13), 'L': ('L', 6.5), 'ei': ('Extremo i', 10.5), 'ej': ('Extremo j', 10.5),
             'ci': ('Castillo en i', 7), 'cj': ('Castillo en j', 7), 'cond': ('Condición', 13), 'raiz': ('Raíz del volado', 6.5), 'bq': ('b', 6), 'hq': ('h', 6), 'cambio': ('Respecto al modelo (15×25)', 16),
             'At': ('Área trib.', 7), 'wt': ('w losa', 7), 'Lm': ('Muro encima', 7), 'wm': ('w muro', 7), 'waz': ('w azotea por el muro', 7.5), 'scm': ('SCM lineal E2K', 7.5), 'cv': ('CV lineal E2K', 7.5),
             'wl': ('w lineal', 7), 'wpp': ('w peso extra', 7), 'w': ('w_u total', 7.5), 'wD': ('w_D serv.', 7), 'wL': ('w_L serv. (W_m)', 7), 'wW': ('w_L media (W)', 7),
             'SP': ('ΣP_u recibidas', 7.5), 'SM': ('ΣM_u por P', 7.5), 'SV': ('ΣV_u por P', 7.5), 'SPm': ('ΣP_u por el muro', 7.5), 'SMs': ('ΣM serv. por P', 7.5),
             'Ri': ('R_i', 7), 'Rj': ('R_j', 7), 'Ris': ('R_i serv.', 7), 'Rjs': ('R_j serv.', 7), 'Mp': ('Mu⁺', 7), 'Mn': ('Mu⁻', 7), 'Vu': ('Vu', 7), 'Ms': ('M serv.', 7), 'd': ('d', 6),
             'AsP': ('As⁺ req.', 7), 'vP': ('Var.', 5.5), 'nPx': ('Extra por flecha', 6), 'nP': ('Pzas', 5), 'MRP': ('M_R⁺', 7), 'FUP': ('Mu⁺/M_R', 7), 'eP': ('ε_t⁺', 6.5),
             'AsN': ('As⁻ req.', 7), 'vN': ('Var.', 5.5), 'nN': ('Pzas', 5), 'MRN': ('M_R⁻', 7), 'FUN': ('Mu⁻/M_R', 7), 'eN': ('ε_t⁻', 6.5),
             'VcR': ('V_cR', 7), 'se': ('Sep. estribo claro', 7), 'sext': ('Sep. en 2h extremos', 7), 'cort': ('Cortante', 20),
             'bfq': ('Patín b_f', 6.5), 'Ig': ('I_g sección T', 8.5), 'dr': ('d con la varilla real', 7), 'nAs': ('n·A_s a tensión', 7), 'kd': ('kd (rectángulo)', 7), 'Icr': ('I_cr agrietada (T real)', 8.5), 'Ie': ('I_e', 8), 'dlt': ('Flecha total', 7), 'dlim': ('Límite', 7), 'Pu': ('P_u del muro', 7), 'Kf': ('K', 6), 'fcm': ('f_CM', 7), 'm418': ('Estado muro', 16),
             'htab': ('h mín. (Tabla 6.3.3.1.1 / NTC-M 4.1.18.4)', 9), 'h418': ('h por NTC-M 4.1.18.3', 8), 'hflex': ('h req. flexión', 8), 'hrec': ('h mín. por flexión, Tabla 6.3.3.1.1 y NTC-M 4.1.18 (sin flecha ni acomodo)', 10), 'estado': ('ESTADO', 30),
             'nota': ('Nota', 50), 'nmax': ('Ancho libre para varillas', 7)}
    units = {'L': 'm', 'bq': 'cm', 'hq': 'cm', 'At': 'm²', 'wt': 't/m', 'Lm': 'm', 'wm': 't/m', 'waz': 't/m', 'scm': 'kg/m', 'cv': 'kg/m', 'wl': 't/m', 'wpp': 't/m', 'w': 't/m', 'wD': 't/m', 'wL': 't/m', 'wW': 't/m',
             'SP': 't', 'SM': 't·m', 'SV': 't', 'SPm': 't', 'SMs': 't·m', 'Ri': 't', 'Rj': 't', 'Ris': 't', 'Rjs': 't', 'Mp': 't·m', 'Mn': 't·m', 'Vu': 't', 'Ms': 't·m', 'd': 'cm',
             'AsP': 'cm²', 'MRP': 't·m', 'AsN': 'cm²', 'MRN': 't·m', 'eP': '≥ 0.005', 'eN': '≥ 0.005', 'VcR': 't', 'se': 'cm', 'sext': 'cm', 'bfq': 'cm', 'Ig': 'cm⁴', 'dr': 'cm', 'nAs': 'cm²', 'kd': 'cm', 'Icr': 'cm⁴', 'Ie': 'cm⁴', 'dlt': 'cm', 'dlim': 'cm',
             'Pu': 't', 'fcm': 'kg/cm²', 'htab': 'cm', 'h418': 'cm', 'hflex': 'cm', 'hrec': 'cm', 'nmax': 'cm'}
    for nm, (t, w) in heads.items():
        hdr_cell(ws, 7, CI(PC[nm]), t, w); unit_cell(ws, 8, CI(PC[nm]), units.get(nm, ''))
    ws.row_dimensions[7].height = 42
    for i, x in enumerate(D['transfer']):
        XROW[i] = X0 + i
    F = lambda s, r: s.format(r=r, **{k: v for k, v in PC.items()}, **N)
    sup = lambda c: 'OR({%s}{r}="FIRME",{%s}{r}="NERVADURA",{%s}{r}="PUNTA")' % (c, c, c)
    for t in NPS:
        r = PROW[t['np']]
        ext = t['ext']

        def etxt(k):
            e = ext[k]
            if e.get('prop'):
                return '=IF(%s="SÍ","FIRME","LIBRE")' % PROPROW[e['prop']]
            return e['tipo']
        ei = etxt(0); ej = etxt(1)
        base = t['cond']
        incoming = [(i, x) for i, x in enumerate(D['transfer']) if x['a'] == t['np']]
        sumcol = lambda col, via=None: ('+'.join('TRANSFER!$%s$%d' % (col, XROW[i]) for i, x in incoming if via is None or x['via'] == via)) or '0'
        bq0, hq0, nx0 = (PROP.get(t['np'], []) + [None, None, 0])[:3]
        waz = WAZ.get(t['np'], [])
        wz = lambda col: ('(' + '+'.join('%s*%s' % (trib_sum(col, dm), f) for dm, f in waz) + ')/{L}{r}') if waz else '0'
        wzD = ('(' + '+'.join('(%s-%s)*%s' % (trib_sum('I', dm), trib_sum('L', dm), f) for dm, f in waz) + ')/{L}{r}') if waz else '0'
        vals = {
            'np': (t['np'], 'key', '@'), 'nivel': (t['level'], 'calc', '@'), 'eje': (t['eje'], 'calc', '@'),
            'de': ('(%.3f, %.3f)' % (t['x0'], t['y0']), 'calc', '@'), 'a': ('(%.3f, %.3f)' % (t['x1'], t['y1']), 'calc', '@'),
            'L': (t['L'], 'input', '0.000'), 'ei': (ei, 'calc', '@'), 'ej': (ej, 'calc', '@'),
            'ci': (ext[0].get('prop') or '—', 'calc', '@'), 'cj': (ext[1].get('prop') or '—', 'calc', '@'),
            'cond': (('BORDE DE VOLADO', 'calc', '@') if base == 'BORDE DE VOLADO' else
                     ('=IF(AND(%s,%s),"APOYADA",IF(OR(%s,%s),"VOLADO","SIN APOYO"))' % (sup('ei'), sup('ej'), sup('ei'), sup('ej')), 'key', '@')),
            'raiz': ('=IF({cond}{r}="VOLADO",IF(%s,"i","j"),"—")' % sup('ei'), 'calc', '@'),
            'bq': (bq0 if bq0 else '={bp}', 'input', '0'), 'hq': (hq0 if hq0 else '={h}', 'input', '0'),
            'cambio': ('=IF({cond}{r}="BORDE DE VOLADO","—",IF(OR({bq}{r}<>{bp},{hq}{r}<>{h}),"PERALTAR: trabe "&TEXT({bq}{r},"0")&"×"&TEXT({hq}{r},"0")&IF({hq}{r}>{h}," (descuelga "&TEXT({hq}{r}-{h},"0")&" cm)",""),"Queda en la losa (15×25)"))', 'calc', '@'),
            'At': ('=' + trib_sum('E', t['np']), 'calc', '0.00'),
            'wt': ('=' + trib_sum('G', t['np']) + '/{L}{r}', 'calc', '0.000'),
            'Lm': (t['muro_arriba'], 'input', '0.00'), 'wm': ('={FCcm}*{Lm}{r}*{wmuro}/{L}{r}', 'calc', '0.000'),
            'waz': ('=' + wz('G'), 'calc', '0.000'),
            'scm': (t['e2k_scm'], 'input', '0.0'), 'cv': (t['e2k_cv'], 'input', '0.0'), 'wl': ('=({FCcm}*{scm}{r}+{FCcv}*{cv}{r})/1000', 'calc', '0.000'),
            'wpp': ('={FCcm}*({bq}{r}/100*{hq}{r}/100*{gc}-{bq}{r}/100*{PP}/1000)', 'calc', '0.000'),
            'w': ('={wt}{r}+{wm}{r}+{waz}{r}+{wl}{r}+{wpp}{r}', 'key', '0.000'),
            'wD': ('=(%s-%s)/{L}{r}+{Lm}{r}*{wmuro}/{L}{r}+%s+{scm}{r}/1000+{wpp}{r}/{FCcm}' % (trib_sum('I', t['np']), trib_sum('L', t['np']), wzD), 'calc', '0.000'),
            'wL': ('=%s/{L}{r}+%s+{cv}{r}/1000' % (trib_sum('L', t['np']), wz('L')), 'calc', '0.000'),
            'wW': ('=%s/{L}{r}+%s+{cv}{r}/1000*{cvw_hab}/{cv_hab}' % (trib_sum('M', t['np']), wz('M')), 'calc', '0.000'),
            'SP': ('=' + sumcol('F'), 'calc', '0.000'), 'SM': ('=' + sumcol('H'), 'calc', '0.000'), 'SV': ('=' + sumcol('I'), 'calc', '0.000'),
            'SPm': ('=' + sumcol('F', 'muro'), 'calc', '0.000'), 'SMs': ('=' + sumcol('N'), 'calc', '0.000'),
            'Ri': ('=IF({cond}{r}="APOYADA",{w}{r}*{L}{r}/2+%s,IF(AND({cond}{r}="VOLADO",{raiz}{r}="i"),{w}{r}*{L}{r}+{SP}{r},0))' % sumcol('J'), 'calc', '0.000'),
            'Rj': ('=IF({cond}{r}="APOYADA",{w}{r}*{L}{r}/2+%s,IF(AND({cond}{r}="VOLADO",{raiz}{r}="j"),{w}{r}*{L}{r}+{SP}{r},0))' % sumcol('K'), 'calc', '0.000'),
            'Ris': ('=IF({cond}{r}="APOYADA",({wD}{r}+{wL}{r})*{L}{r}/2+%s,IF(AND({cond}{r}="VOLADO",{raiz}{r}="i"),({wD}{r}+{wL}{r})*{L}{r}+%s,0))' % (sumcol('O'), sumcol('M')), 'calc', '0.000'),
            'Rjs': ('=IF({cond}{r}="APOYADA",({wD}{r}+{wL}{r})*{L}{r}/2+%s,IF(AND({cond}{r}="VOLADO",{raiz}{r}="j"),({wD}{r}+{wL}{r})*{L}{r}+%s,0))' % (sumcol('P'), sumcol('M')), 'calc', '0.000'),
            'Mp': ('=IF({cond}{r}="APOYADA",{w}{r}*{L}{r}^2/8+{SM}{r},0)', 'key', '0.000'),
            'Mn': ('=IF({cond}{r}="APOYADA",{Mrat}*{Mp}{r},IF({cond}{r}="VOLADO",{w}{r}*{L}{r}^2/2+{SM}{r},0))', 'key', '0.000'),
            'Vu': ('=IF({cond}{r}="APOYADA",{w}{r}*{L}{r}/2+{SV}{r},IF({cond}{r}="VOLADO",{w}{r}*{L}{r}+{SV}{r},0))', 'key', '0.000'),
            'Ms': ('=IF({cond}{r}="APOYADA",({wD}{r}+{wL}{r})*{L}{r}^2/8+{SMs}{r},IF({cond}{r}="VOLADO",({wD}{r}+{wL}{r})*{L}{r}^2/2+{SMs}{r},0))', 'calc', '0.000'),
            'd': ('={hq}{r}-{rec_np}-' + db('{est_np}') + '-' + db('{var_np}') + '/2', 'calc', '0.00'),
            'nmax': ('={bq}{r}-2*{rec_np}-2*' + db('{est_np}'), 'calc', '0.00'),
        }
        vals['nPx'] = (nx0 or 0, 'input', '0')
        for blk, M in (('P', 'Mp'), ('N', 'Mn')):
            # 6.3.5.1.2: en viga isostática con el patín a tensión (volado), b_w = el menor de b_f y 2b_w
            asr = asreq('{%s}{r}' % M, '{bq}{r}', '{d}{r}', True, 'IF({cond}{r}="VOLADO",MIN({bfq}{r},2*{bq}{r}),{bq}{r})' if blk == 'N' else '{bq}{r}')
            vals['As' + blk] = ('=IF(OR({cond}{r}="BORDE DE VOLADO",{cond}{r}="SIN APOYO"),0,%s)' % asr, 'calc', '0.00')
            vals['v' + blk] = ('=IF({As%s}{r}<0,0,' % blk + varsel2('{As%s}{r}' % blk, '{var_np}', '99', '{nmax}{r}', '{nmin_np}') + ')', 'calc', 'General')
            extra = '+{nPx}{r}' if blk == 'P' else ''
            vals['n' + blk] = ('=IF(OR({cond}{r}="BORDE DE VOLADO",{cond}{r}="SIN APOYO"),0,IF({As%s}{r}<0,0,MAX({nmin_np},CEILING({As%s}{r}/%s,1))%s))' % (blk, blk, ab('IF({v%s}{r}=0,%s,{v%s}{r})' % (blk, BARS[-1][0], blk)), extra), 'calc', '0')
            asp = '{n%s}{r}*%s' % (blk, ab('{v%s}{r}' % blk))
            dact = '({hq}{r}-{rec_np}-%s-%s/2)' % (db('{est_np}'), db('{v%s}{r}' % blk))
            vals['MR' + blk] = ('=IF({n%s}{r}=0,0,{FRf}*%s*{fy}*(%s-%s*{fy}/({fpc}*{bq}{r})/2)/1E5)' % (blk, asp, dact, asp), 'calc', '0.000')
            vals['FU' + blk] = ('=IF({As%s}{r}<0,9.99,IF({MR%s}{r}>0,{%s}{r}/{MR%s}{r},0))' % (blk, blk, M, blk), 'key', '0.00')
            cexp = '(%s*{fy}/({fpc}*{bq}{r})/{beta1})' % asp
            vals['e' + blk] = ('=IF({n%s}{r}=0,"—",0.003*(%s-%s)/%s)' % (blk, dact, cexp, cexp), 'calc', '0.0000')
        # cortante con A_v,mín: V_cR = F_R·máx(0.5, 2p^(1/3))·√f'c·b·d ≤ F_R·1.25·√f'c·b·d
        pmn = 'MAX(1E-6,MIN({nP}{r}*%s,{nN}{r}*%s)/({bq}{r}*{d}{r}))' % (ab('MAX({vP}{r},2)'), ab('MAX({vN}{r},2)'))
        vals['VcR'] = ('=IF({nN}{r}=0,0,{FRv}*MIN(1.25,MAX(0.5,2*%s^(1/3)))*SQRT({fc})*{bq}{r}*{d}{r}/1000)' % pmn, 'calc', '0.000')
        av = '2*' + ab('{est_np}')
        vs = 'MAX(0,{Vu}{r}-{VcR}{r})/{FRv}'
        smaxv = 'IF(%s<=1.1*SQRT({fc})*{bq}{r}*{d}{r}/1000,MIN({d}{r}/2,60),MIN({d}{r}/4,30))' % vs
        sreq = 'IF({Vu}{r}>{VcR}{r},{FRv}*%s*{fyt}*{d}{r}/(({Vu}{r}-{VcR}{r})*1000),999)' % av
        savm = '%s*{fyt}/(MAX(0.2*SQRT({fc}),3.5)*{bq}{r})' % av
        s147 = 'MIN(16*%s,48*%s)' % (db('MIN(MAX({vP}{r},2),MAX({vN}{r},2))'), db('{est_np}'))
        vals['se'] = ('=IF({Vu}{r}=0,0,FLOOR(MIN(%s,%s,%s,%s,{smax_np}),0.5))' % (sreq, smaxv, savm, s147), 'calc', '0.0')
        if t['np'] in PERIM:
            vals['sext'] = ('=IF({Vu}{r}=0,0,FLOOR(MIN({se}{r},{d}{r}/4,8*%s,24*%s,30),0.5))' % (db('MIN(MAX({vP}{r},2),MAX({vN}{r},2))'), db('{est_np}')), 'calc', '0.0')
        else:
            vals['sext'] = ('={se}{r}', 'calc', '0.0')
        vals['cort'] = ('=IF({Vu}{r}=0,"—",IF({Vu}{r}>{FRv}*({VcR}{r}/{FRv}+2.2*SQRT({fc})*{bq}{r}*{d}{r}/1000),"SECCIÓN INSUFICIENTE",'
                        '"E No. "&TEXT({est_np},"0")&IF({sext}{r}<{se}{r}," @ "&TEXT({sext}{r},"0.0")&" en 2h y"," @")&" "&TEXT({se}{r},"0.0")&" cm"))', 'calc', '@')
        # flecha: Bischoff (Tabla 13.4.3.2), inmediata + diferida, con la sección T que forma la capa de compresión (Tabla 5.2.1.4.2)
        if t['np'] in ONESIDE:
            vals['bfq'] = ('={bq}{r}+MIN(6*{t},{ac}/2,{L}{r}*100/12)', 'calc', '0')
        else:
            vals['bfq'] = ('={bq}{r}+2*MIN(8*{t},{ac}/2,{L}{r}*100/8)', 'calc', '0')
        AgT = '({bfq}{r}*{t}+{bq}{r}*({hq}{r}-{t}))'
        ybT = '(({bfq}{r}*{t}^2/2+{bq}{r}*({hq}{r}-{t})*({t}+({hq}{r}-{t})/2))/%s)' % AgT
        vals['Ig'] = ('={bfq}{r}*{t}^3/12+{bfq}{r}*{t}*(%s-{t}/2)^2+{bq}{r}*({hq}{r}-{t})^3/12+{bq}{r}*({hq}{r}-{t})*({t}+({hq}{r}-{t})/2-%s)^2' % (ybT, ybT), 'calc', '#,##0')
        Ig = '{Ig}{r}'
        AsT = 'IF({cond}{r}="VOLADO",{nN}{r}*%s,{nP}{r}*%s)' % (ab('MAX({vN}{r},2)'), ab('MAX({vP}{r},2)'))
        bc = 'IF({cond}{r}="VOLADO",{bq}{r},{bfq}{r})'
        vals['dr'] = ('={hq}{r}-{rec_np}-%s-%s/2' % (db('{est_np}'), db('MAX(IF({cond}{r}="VOLADO",{vN}{r},{vP}{r}),2)')), 'calc', '0.00')
        vals['nAs'] = ('={Es}/{Ec}*%s' % AsT, 'calc', '0.00')
        dr = '{dr}{r}'; nAs = '{nAs}{r}'
        rn = '(%s/(%s*%s))' % (nAs, bc, dr)
        vals['kd'] = ('=IF(%s<=0,0,(SQRT(2*%s+%s^2)-%s)*%s)' % (nAs, rn, rn, rn, dr), 'calc', '0.00')
        kd = '{kd}{r}'
        Icr_r = '(%s*%s^3/3+%s*(%s-%s)^2)' % (bc, kd, nAs, dr, kd)
        # eje neutro bajo la capa (sección T): (b_w/2)x² + [(b_f−b_w)t + nAs]x − [(b_f−b_w)t²/2 + nAs·d] = 0
        Bt = '(({bfq}{r}-{bq}{r})*{t}+%s)' % nAs
        Ct = '(({bfq}{r}-{bq}{r})*{t}^2/2+%s*%s)' % (nAs, dr)
        xT = '((-%s+SQRT(%s^2+2*{bq}{r}*%s))/{bq}{r})' % (Bt, Bt, Ct)
        Icr_t = '({bfq}{r}*{t}^3/12+{bfq}{r}*{t}*(%s-{t}/2)^2+{bq}{r}*(%s-{t})^3/3+%s*(%s-%s)^2)' % (xT, xT, nAs, dr, xT)
        vals['Icr'] = ('=IF(AND({cond}{r}<>"VOLADO",%s>{t}),%s,%s)' % (kd, Icr_t, Icr_r), 'calc', '#,##0')
        Icr = '{Icr}{r}'
        Mag = '({ffr}*%s/IF({cond}{r}="VOLADO",%s,{hq}{r}-%s))' % (Ig, ybT, ybT)
        Mse = '({Ms}{r}*1E5)'
        vals['Ie'] = ('=IF({Ms}{r}<=0,%s,IF(%s<=2/3*%s,%s,IF(%s<=0,0,%s/(1-(2/3*%s/%s)^2*(1-%s/%s)))))' % (Ig, Mse, Mag, Ig, nAs, Icr, Mag, Mse, Icr, Ig), 'calc', '#,##0')
        pcomp = 'IF({cond}{r}="VOLADO",{nP}{r}*%s,{nmin_np}*%s)/({bq}{r}*%s)' % (ab('MAX({vP}{r},2)'), ab('MAX({vN}{r},2)'), dr)
        rs = '(({wD}{r}+{wW}{r})/MAX({wD}{r}+{wL}{r},1E-6))'
        # volado: carga repartida → M·L²/(4EI); cargas puntuales (cota de la punta) → M·L²/(3EI)
        di = '(IF({cond}{r}="VOLADO",({Ms}{r}-{SMs}{r})/4+{SMs}{r}/3,5/48*{Ms}{r})*1E5*({L}{r}*100)^2/({Ec}*MAX({Ie}{r},1)))'
        vals['dlt'] = ('=IF(OR({cond}{r}="APOYADA",{cond}{r}="VOLADO"),%s*(1+2/(1+50*%s)*%s),0)' % (di, pcomp, rs), 'key', '0.00')
        vals['dlim'] = ('=IF({Lm}{r}>0,0.3+{L}{r}*100/480,0.5+{L}{r}*100/240)', 'calc', '0.00')
        # NTC-Mampostería 4.1.18: muro sobre viga
        ATm = '({tm}*100*{Lm}{r}*100)'
        a_ = 'MIN(1,{Lm}{r}/{L}{r})'
        Wp = '({Pu}{r}*1000/%s*{tm}*100*(%s^4-4*%s^3+8*%s))' % (ATm, a_, a_, a_)
        # w_u sobre la viga: losa + lineal + peso extra + reacciones que llegan DIRECTO de otras principales como carga repartida equivalente 8·M/L²
        # (las que llegan por el muro ya están en P_u)
        wu4 = '(({wt}{r}+{wl}{r}+{wpp}{r}+IF({cond}{r}="APOYADA",8*(%s)/{L}{r}^2,0))*10)' % sumcol('H', 'directo')
        fI = '(%s/(%s+5*%s))' % (Wp, Wp, wu4)
        vals['Pu'] = ('=IF({Lm}{r}>0,({wm}{r}+{waz}{r})*{L}{r}+{SPm}{r},0)', 'calc', '0.00')
        vals['Kf'] = ('=IF({Lm}{r}>0,({Em}*{tm}*100*({L}{r}*100)^3/({Ec}*%s*%s))^0.25,0)' % (Ig, fI), 'calc', '0.00')
        fce1 = '(1.6*{Kf}{r}-3.9)'; fce9 = '((3.2*{Kf}{r}-7.8)*1)'
        fce = 'MAX(1,IF(%s>=1,%s,IF(%s<0.9,%s,%s+(%s-%s)*(%s-0.9)/0.1)))' % (a_, fce1, a_, fce9, fce9, fce1, fce9, a_)
        vals['fcm'] = ('=IF({Lm}{r}>0,%s*{Pu}{r}*1000/%s,0)' % (fce, ATm), 'calc', '0.0')
        PD = '(({Lm}{r}*{wmuro}+(%s)*{L}{r}+{SPm}{r}/{FCcm})*1000)' % (wzD + '+' + wz('M'))
        WpD = '(%s/%s*{tm}*100*(%s^4-4*%s^3+8*%s))' % (PD, ATm, a_, a_, a_)
        wuD = '(({wD}{r}-{Lm}{r}*{wmuro}/{L}{r}-(%s))*10)' % wzD
        h4183 = '({L}{r}*100/((1/15*{Ec}*{bq}{r}/(%s+5*%s))^(1/3)))' % (WpD, wuD)
        I4183 = '(5/4*({L}{r}*100)^3/{Ec}*(%s+5*%s))' % (WpD, wuD)
        # h que da la misma I bruta con sección rectangular (para proponer peralte): h = (12·I/b)^(1/3)
        vals['h418'] = ('=IF({Lm}{r}>0,IF({Ig}{r}>=%s,0,MIN(%s,(12*%s/{bq}{r})^(1/3))),0)' % (I4183, h4183, I4183), 'calc', '0.0')
        # D2: con H/L < 0.6 el procedimiento de 4.1.18.2 no aplica; criterio del Responsable: sin efecto de arco, la trabe toma TODO el muro
        # (w muro + w azotea + P por el muro, sin reducción) y además se pide f_CM ≤ F_R·f'm con la misma expresión, como estimación fuera de su campo
        vals['m418'] = ('=IF({Lm}{r}=0,"—",IF({fcm}{r}<={FRm}*{fm},IF({hm}/{Lm}{r}<0.6,"CUMPLE — H/L < 0.6: sin arco, la trabe toma todo el muro; f_CM estimado fuera de campo (D2)","CUMPLE"),"NO CUMPLE f_CM > F_R·f\'m"&IF({hm}/{Lm}{r}<0.6," (estimación con H/L < 0.6)","")))', 'calc', '@')
        vals['htab'] = ('=IF({Lm}{r}>0,MAX(25,IF({Ig}{r}>=({L}{r}*100)^3/1600,0,MIN({L}{r}*100/14,(12*({L}{r}*100)^3/1600/{bq}{r})^(1/3)))),IF({cond}{r}="APOYADA",MAX(25,{L}{r}*100/16),IF({cond}{r}="VOLADO",MAX(25,{L}{r}*100/8),0)))', 'calc', '0.0')
        qm = '({pmax}*{fy}/{fpc})'
        vals['hflex'] = ('=SQRT(MAX({Mp}{r},{Mn}{r})*1E5/({FRf}*{bq}{r}*{fpc}*%s*(1-%s/2)))+({hq}{r}-{d}{r})' % (qm, qm), 'calc', '0.0')
        vals['hrec'] = ('=IF(OR({cond}{r}="BORDE DE VOLADO",{cond}{r}="SIN APOYO"),0,CEILING(MAX({htab}{r},{h418}{r},{hflex}{r},{h}),5))', 'key', '0')
        fitP = 'IF({nP}{r}=0,TRUE,{nP}{r}<=%s)' % nfit('MAX({vP}{r},2)', '{nmax}{r}'); fitN = 'IF({nN}{r}=0,TRUE,{nN}{r}<=%s)' % nfit('MAX({vN}{r},2)', '{nmax}{r}')
        etok = 'AND(OR(NOT(ISNUMBER({eP}{r})),{eP}{r}>=0.005),OR(NOT(ISNUMBER({eN}{r})),{eN}{r}>=0.005))'
        defok = '{dlt}{r}<={dlim}{r}'
        hok = 'AND({hq}{r}>={htab}{r},{hq}{r}>={h418}{r},{hq}{r}/{bq}{r}<=6)'
        mok = 'OR({Lm}{r}=0,LEFT({m418}{r},6)<>"NO CUM")'
        ok = 'AND({FUP}{r}<=1,{FUN}{r}<=1,%s,%s,%s,{cort}{r}<>"SECCIÓN INSUFICIENTE",%s,%s,%s)' % (fitP, fitN, etok, hok, defok, mok)
        vals['estado'] = ('=IF({cond}{r}="BORDE DE VOLADO","BORDE DE VOLADO — armado de orilla",IF({cond}{r}="SIN APOYO","SIN APOYO — REVISAR",IF(%s,IF(IF({Lm}{r}>0,{hm}/{Lm}{r}<0.6,FALSE),"CUMPLE — H/L < 0.6: sin arco, la trabe toma todo el muro (criterio del Responsable)","CUMPLE")&IF({cond}{r}="VOLADO"," (en VOLADO)",""),'
                          '"NO CUMPLE: "&IF(OR({AsP}{r}<0,{AsN}{r}<0),"sección insuficiente en flexión ",IF(OR({FUP}{r}>1,{FUN}{r}>1),"flexión ",""))&IF(NOT(AND(%s,%s)),"no caben ","")&IF(NOT(%s),"ε_t ","")'
                          '&IF({cort}{r}="SECCIÓN INSUFICIENTE","cortante ","")&IF(NOT(%s),"peralte ","")&IF(NOT(%s),"flecha ","")&IF(NOT(%s),"muro sobre viga ","")'
                          '&"→ h ≥ "&TEXT(MAX({hrec}{r},IF({dlt}{r}>{dlim}{r},MAX({hq}{r}+5,CEILING({hq}{r}*({dlt}{r}/{dlim}{r})^(1/3),5)),0)),"0")&" cm"&IF(NOT(AND(%s,%s))," y b ≥ "&TEXT({bq}{r}+5,"0")&" cm","")&" (verificar)")))'
                          % (ok, fitP, fitN, etok, hok, defok, mok, fitP, fitN), 'key', '@')
        notes = []
        for k, e in enumerate(ext):
            if e.get('prop'):
                notes.append('Extremo %s en el castillo %s (baja a cimentación, decisión D1).' % ('i' if k == 0 else 'j', e['prop']))
            if e['tipo'] == 'NERVADURA':
                notes.append('Extremo %s se apoya en %s.' % ('i' if k == 0 else 'j', e['recibe']))
            if e['tipo'] == 'PUNTA':
                notes.append('Extremo %s se apoya en la punta del volado corto %s.' % ('i' if k == 0 else 'j', e['recibe']))
        if incoming:
            notes.append('Recibe a ' + ', '.join(x['de'] + (' (por el muro de arriba)' if x['via'] == 'muro' else '') for _, x in incoming) + '.')
        if t['np'] in PERIM:
            notes.append('Viga perimetral: refuerzo de integridad continuo arriba y abajo (≥ 2 barras) y estribos cerrados en todo el claro, más cerrados en 2h junto a cada apoyo (NTC-Concreto 6.3.7.7.1, 6.3.7.7.7 y 6.3.7.7.8, p. 134).')
        if t['muro_arriba'] > 0:
            notes.append('Muro de planta alta encima (%.2f m)%s: viga que soporta muro estructural (NTC-Mampostería 4.1.18): la flecha se calcula y el límite es 0.3 cm + L/480.' % (t['muro_arriba'], (' con la azotea de ' + ', '.join(dm for dm, _ in waz)) if waz else ''))
        if t['np'] == 'NP-11':
            notes.append('Decisión D2: la dala bajo el muro de y = 14.247 se prolonga hasta la nervadura del eje F (x = 5.44); en el modelo remata suelta dentro de F80009: prolongarla también en ETABS (D5).')
        if t['level'] == 'N+3.00' and 8.8 < t['y0'] < 10 and 5.3 < t['x0'] < 7.4 and t['np'] in ('NP-07', 'NP-30', 'NP-25'):
            notes.append('Llegada de la escalera: revisar junto con el diseño de la escalera.')
        if base == 'BORDE DE VOLADO':
            notes.append('Orilla libre de tablero en volado: lleva el armado de la orilla, no trabaja como trabe.')
        txt = ' '.join(notes) or '—'
        if t['muro_arriba'] > 0:
            # la nota de D2 depende de H/L, que es fórmula (altura del muro en DATOS)
            d2 = (' H/L < 0.6 (NTC-M 4.1.18.2.1, p. 44): el procedimiento de 4.1.18.2 no aplica. Criterio del Responsable (decisión D2), lado seguro: sin efecto de arco, la trabe se diseña con el peso COMPLETO del muro y de la azotea que recibe, '
                  'repartido en el claro (w muro + w azotea + P por el muro, sin reducción); debe cumplir 4.1.18.4 (L/h ≤ 14 o I ≥ L³/1600) y la flecha de 4.1.18.3 (0.3 cm + L/480) (ver el estado); f_CM con la expresión de 4.1.18.2 se reporta como estimación fuera de su campo y debe ser ≤ F_R·f\'m. '
                  'Castillos en los dos extremos de ese muro.')
            lit = lambda x: '&'.join('"%s"' % x[i_:i_ + 200].replace('"', '""') for i_ in range(0, len(x), 200))
            vals['nota'] = ('=%s&IF({Lm}{r}>0,IF({hm}/{Lm}{r}<0.6,%s,""),"")' % (lit(txt), lit(d2)), 'calc', '@')
        else:
            vals['nota'] = (txt, 'calc', '@')
        for nm in _pcols:
            v, kind, fmt = vals[nm]
            if isinstance(v, str) and v.startswith('='):
                v = F(v, r)
            x = cell(ws, r, CI(PC[nm]), v, kind, fmt, al='left' if nm in ('nota', 'estado', 'eje', 'cort', 'cambio', 'm418') else 'center')
            if nm in ('nota', 'estado', 'cambio', 'm418'):
                x.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[r].height = 30
    ws.freeze_panes = 'C9'
    add_cf(ws, 'B9:%s%d' % (PC['estado'], P1))
    for c_ in ('cambio', 'm418', 'cort'):
        add_cf(ws, '%s9:%s%d' % (PC[c_], PC[c_], P1))
    ws.sheet_view.zoomScale = 80
    return ws


def sheet_trib(wb):
    ws = wb.create_sheet('TRIB'); ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 1.78
    ws.merge_cells('B1:M1'); c = ws['B1']; c.value = 'TRIB — ÁREAS TRIBUTARIAS DEL DXF POR TRAMO DE NERVADURA PRINCIPAL O DE MURO'
    c.font = font(13, True, color='FFFFFF'); c.fill = fill(TITLE_BG); ws.row_dimensions[1].height = 26
    ws.merge_cells('B2:M2'); ws['B2'].value = ('Cada polígono de la capa TRIB_AREA de ' + DXFT + ' se asignó al tramo sobre el que descansa su base (repartido por proyección cuando cae sobre dos tramos). '
                                              'Las áreas son datos; la carga se recalcula con el w del tablero. Las bajadas por áreas tributarias son admisibles (NTC-Mampostería 3.1.2.2.1, p. 31).')
    ws['B2'].font = font(9, i=True, color=GRAY); ws['B2'].alignment = Alignment(wrap_text=True); ws.row_dimensions[2].height = 30
    for k, (t, w) in enumerate([('Tramo', 9), ('Tablero', 15), ('Nivel', 8), ('Área (m²)', 9), ('w_u del tablero (kg/m²)', 10), ('Carga última (t)', 9), ('w serv. (kg/m²)', 9),
                                ('Carga servicio (t)', 9), ('Tipo', 9), ('Estado del tramo', 34), ('Carga viva W_m (t)', 9), ('Carga viva media W (t)', 9)]):
        hdr_cell(ws, 5, 2 + k, t, w)
    ws.row_dimensions[5].height = 30
    for i, tr in enumerate(TRIBROWS):
        r = R0 + i; k = '%s %s' % (tr['tablero'], tr['level']); tr_ = TROW[k]
        tipo = 'NP' if tr['tramo'].startswith('NP') else ('DM' if tr['tramo'].startswith('DM') else 'SIN DALA')
        cell(ws, r, 2, tr['tramo']); cell(ws, r, 3, k); cell(ws, r, 4, tr['level']); cell(ws, r, 5, tr['area'], 'data', '0.000')
        cell(ws, r, 6, '=TABLEROS!$%s$%d' % (TC['wu'], tr_), fmt='0'); cell(ws, r, 7, '=E%d*F%d/1000' % (r, r), fmt='0.000')
        cell(ws, r, 8, '=TABLEROS!$%s$%d' % (TC['ws'], tr_), fmt='0'); cell(ws, r, 9, '=E%d*H%d/1000' % (r, r), fmt='0.000')
        cell(ws, r, 10, tipo)
        cell(ws, r, 11, ('=PRINCIPALES!$%s$%d' % (PC['estado'], PROW[tr['tramo']])) if tipo == 'NP' else ('sobre muro' if tipo == 'DM' else 'borde sin dala en el E2K'), al='left')
        cell(ws, r, 12, '=E%d*TABLEROS!$%s$%d/1000' % (r, TC['cv'], tr_), fmt='0.000')
        cell(ws, r, 13, '=E%d*IF(TABLEROS!$%s$%d="HABITACIÓN",%s,%s)/1000' % (r, TC['dest'], tr_, N['cvw_hab'], N['cvw_az']), fmt='0.000')
    ws.freeze_panes = 'B6'
    add_cf(ws, 'K6:K%d' % R1)
    return ws


def sheet_transfer(wb):
    ws = wb.create_sheet('TRANSFER'); ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 1.78
    ws.merge_cells('B1:Q1'); c = ws['B1']; c.value = 'TRANSFER — NERVADURAS PRINCIPALES QUE SE APOYAN EN OTRAS (DIRECTO O A TRAVÉS DEL MURO DE ARRIBA)'
    c.font = font(13, True, color='FFFFFF'); c.fill = fill(TITLE_BG); ws.row_dimensions[1].height = 26
    ws.merge_cells('B2:Q2'); ws['B2'].value = ('La reacción del extremo de un tramo entra como carga puntual en el tramo que lo recibe, a la distancia «pos» de su extremo i. «Por el muro»: la principal de azotea descarga en un castillo o muro del primer nivel que está sobre una principal del N+3.00. '
                                              'Las referencias son directas (sin búsquedas) para que Excel no vea referencias circulares.')
    ws['B2'].font = font(9, i=True, color=GRAY); ws['B2'].alignment = Alignment(wrap_text=True); ws.row_dimensions[2].height = 30
    for k, (t, w) in enumerate([('Tramo que entrega', 10), ('Extremo', 8), ('Tramo que recibe', 10), ('pos desde i (m)', 9), ('P_u (t)', 8), ('L del que recibe (m)', 9),
                                ('Momento que aporta (t·m)', 10), ('Cortante que aporta (t)', 10), ('A R_i (t)', 8), ('A R_j (t)', 8), ('Condición del que recibe', 14),
                                ('P serv. (t)', 8), ('M serv. que aporta (t·m)', 10), ('A R_i serv. (t)', 8), ('A R_j serv. (t)', 8), ('Vía', 9)]):
        hdr_cell(ws, 5, 2 + k, t, w)
    ws.row_dimensions[5].height = 42
    for i, x in enumerate(D['transfer']):
        r = XROW[i]; rd = PROW[x['de']]; ra = PROW[x['a']]
        cell(ws, r, 2, x['de']); cell(ws, r, 3, 'i' if x['extremo'] == 0 else 'j'); cell(ws, r, 4, x['a']); cell(ws, r, 5, x['pos'], 'data', '0.000')
        cell(ws, r, 6, '=PRINCIPALES!$%s$%d' % (PC['Ri'] if x['extremo'] == 0 else PC['Rj'], rd), fmt='0.000')
        cell(ws, r, 7, '=PRINCIPALES!$%s$%d' % (PC['L'], ra), fmt='0.000')
        cond = 'PRINCIPALES!$%s$%d' % (PC['cond'], ra); raiz = 'PRINCIPALES!$%s$%d' % (PC['raiz'], ra)
        R = lambda f: f.replace('{r}', str(r))
        cell(ws, r, 8, R('=IF(%s="APOYADA",F{r}*E{r}*(G{r}-E{r})/G{r},IF(%s="VOLADO",F{r}*IF(%s="i",E{r},G{r}-E{r}),0))') % (cond, cond, raiz), fmt='0.000')
        cell(ws, r, 9, R('=IF(%s="APOYADA",F{r}*MAX(E{r},G{r}-E{r})/G{r},IF(%s="VOLADO",F{r},0))') % (cond, cond), fmt='0.000')
        cell(ws, r, 10, R('=IF(%s="APOYADA",F{r}*(G{r}-E{r})/G{r},0)') % cond, fmt='0.000')
        cell(ws, r, 11, R('=IF(%s="APOYADA",F{r}*E{r}/G{r},0)') % cond, fmt='0.000')
        cell(ws, r, 12, '=' + cond)
        cell(ws, r, 13, '=PRINCIPALES!$%s$%d' % (PC['Ris'] if x['extremo'] == 0 else PC['Rjs'], rd), fmt='0.000')
        cell(ws, r, 14, R('=IF(%s="APOYADA",M{r}*E{r}*(G{r}-E{r})/G{r},IF(%s="VOLADO",M{r}*IF(%s="i",E{r},G{r}-E{r}),0))') % (cond, cond, raiz), fmt='0.000')
        cell(ws, r, 15, R('=IF(%s="APOYADA",M{r}*(G{r}-E{r})/G{r},0)') % cond, fmt='0.000')
        cell(ws, r, 16, R('=IF(%s="APOYADA",M{r}*E{r}/G{r},0)') % cond, fmt='0.000')
        cell(ws, r, 17, 'por el muro' if x['via'] == 'muro' else 'directo')
    return ws


def sheet_muros(wb):
    ws = wb.create_sheet('MUROS'); S = Narr(ws, N, {'A': 1.78, 'B': 22, 'C': 9, 'D': 22, 'E': 9, 'F': 10, 'G': 11, 'H': 11, 'I': 40}, last='I')
    S.title('MUROS — LO QUE LA LOSA LE ENTREGA A LOS MUROS Y A LOS CASTILLOS')
    S.subtitle('Para la revisión de muros con la NTC-Mampostería 2023: carga lineal de la losa sobre cada tramo de dala que sí tiene muro abajo, y reacciones concentradas de las nervaduras principales en sus apoyos.', 30)
    S.gap()
    S.section('A. CARGA LINEAL DE LOSA SOBRE MUROS (por tramo de dala sobre muro)')
    S.secnote('Sólo la losa (áreas tributarias del DXF). No incluye el peso del propio muro, ni el muro de arriba, ni las cargas lineales del E2K: eso se suma en la revisión del muro.')
    r = S.r
    for k, t in enumerate(['Tramo · eje', 'Nivel', 'De – a', 'L (m)', 'Área trib. (m²)', 'w servicio (t/m)', 'w última (t/m)', 'Nota']):
        hdr_cell(ws, r, 2 + k, t)
    S.r += 1
    for t in DMS:
        r = S.r
        cell(ws, r, 2, '%s · %s' % (t['np'], t['eje']), al='left'); cell(ws, r, 3, t['level'])
        cell(ws, r, 4, '(%.2f, %.2f) – (%.2f, %.2f)' % (t['x0'], t['y0'], t['x1'], t['y1']), al='left', sz=8)
        cell(ws, r, 5, t['L'], fmt='0.000')
        cell(ws, r, 6, '=SUMIFS(TRIB!$E$%d:$E$%d,TRIB!$B$%d:$B$%d,"%s")' % (R0, R1, R0, R1, t['np']), fmt='0.00')
        cell(ws, r, 7, '=SUMIFS(TRIB!$I$%d:$I$%d,TRIB!$B$%d:$B$%d,"%s")/E%d' % (R0, R1, R0, R1, t['np'], r), fmt='0.000')
        cell(ws, r, 8, '=SUMIFS(TRIB!$G$%d:$G$%d,TRIB!$B$%d:$B$%d,"%s")/E%d' % (R0, R1, R0, R1, t['np'], r), 'key', '0.000')
        cell(ws, r, 9, 'Muro del entrepiso de abajo (planta baja)' if t['level'] == 'N+3.00' else 'Muro del primer nivel', al='left', sz=8)
        S.r += 1
    S.gap()
    S.section('B. REACCIONES DE LAS NERVADURAS PRINCIPALES EN SUS APOYOS FIRMES')
    S.secnote('Carga concentrada que baja por el castillo o el muro del extremo. Las que caen en uno de los castillos C1…C8 (aceptados, decisión D1) son las que ese castillo tiene que llevar a la cimentación; donde no hay muro abajo en planta baja es columna de 25×25 como mínimo (NTC-Concreto 2023 (CDMX) 6.4.2.1.1, p. 138).')
    r = S.r
    for k, t in enumerate(['Tramo · extremo', 'Nivel', 'Apoyo (x, y)', 'Castillo', 'Tipo', 'R última (t)', 'R servicio aprox. (t)', 'Nota']):
        hdr_cell(ws, r, 2 + k, t)
    S.r += 1
    for t in NPS:
        for k in (0, 1):
            e = t['ext'][k]
            if e['tipo'] != 'FIRME':
                continue
            r = S.r; pr = PROW[t['np']]
            cell(ws, r, 2, '%s · %s' % (t['np'], 'i' if k == 0 else 'j'), al='left'); cell(ws, r, 3, t['level'])
            cell(ws, r, 4, '(%.3f, %.3f)' % ((t['x0'], t['y0']) if k == 0 else (t['x1'], t['y1'])), sz=8)
            cell(ws, r, 5, e.get('prop') or '—')
            cell(ws, r, 6, '=PRINCIPALES!$%s$%d' % (PC['ei'] if k == 0 else PC['ej'], pr))
            cell(ws, r, 7, '=PRINCIPALES!$%s$%d' % (PC['Ri'] if k == 0 else PC['Rj'], pr), 'key', '0.000')
            cell(ws, r, 8, '=PRINCIPALES!$%s$%d' % (PC['Ris'] if k == 0 else PC['Rjs'], pr), fmt='0.000')
            cell(ws, r, 9, 'Servicio = cargas sin factorizar (D + W_m).', al='left', sz=8)
            S.r += 1
    S.gap()
    S.notice('Para el muro: P_u total = (w_losa + peso propio del muro + muro de arriba + macizo de borde)·L + reacciones concentradas, y se compara con la resistencia a compresión del muro confinado de la NTC-Mampostería 2023. '
             'El macizo o dala de borde de 40 cm de los bordes discontinuos (decisión D9) agrega su peso (DATOS sección 4, «Carga muerta extra del macizo de borde») directo al muro de ese borde. Eso no está en este libro.', 42)
    add_cf(ws, 'B1:I%d' % S.r)
    return ws
