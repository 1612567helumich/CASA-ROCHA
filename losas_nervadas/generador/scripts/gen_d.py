"""Parte D: RESUMEN, REFERENCIAS y armado del libro."""
import sys
from gen_c import *
from openpyxl import Workbook


def fx(*parts):
    """Texto vivo: los literales se parten en trozos de ≤ 200 caracteres (Excel no acepta literales de más de 255); las tuplas son expresiones."""
    out = []
    for p in parts:
        if isinstance(p, tuple):
            out.append(p[0].format(**N))
        else:
            out += ['"%s"' % p[i:i + 200].replace('"', '""') for i in range(0, len(p), 200)]
    return '=' + '&'.join(out)


def endlist(x, none='ninguno.'):
    """Lista armada con «; » entre elementos: cambia el último «; » por punto (o «ninguno.» si está vacía)."""
    return 'SUBSTITUTE(SUBSTITUTE(%s&"|","; |","."),"|","%s")' % (x, none)


def fxlen(*parts):
    return sum(len(p) if not isinstance(p, tuple) else 12 for p in parts)


def sheet_resumen(wb):
    ws = wb.create_sheet('RESUMEN')
    widths = {'A': 1.78, 'B': 16, 'C': 12, 'D': 10, 'E': 9, 'F': 7, 'G': 7, 'H': 8, 'I': 8, 'J': 9, 'K': 8, 'L': 9, 'M': 9, 'N': 30}
    S = Narr(ws, N, widths, last='N')

    def lbullet(*parts):
        n = fxlen(*parts) / 150.0
        return S.bullet(fx(*parts), max(15, 13 * (int(n) + 1) + 2))
    T_ = lambda c, k: 'TABLEROS!$%s$%d' % (TC[c], TROW[k])
    T2 = T_
    S.title('RESUMEN DE LAS LOSAS NERVADAS — CASA ROCHA')
    S.merged(fx('Losa reticular LR-25 (casetón ', ('TEXT({ac},"0")',), '×', ('TEXT({ac},"0")',), '×', ('TEXT({hc},"0")',), ', nervadura de ', ('TEXT({bw},"0")',), ' cm, capa de ', ('TEXT({t},"0")',), ' cm con malla ', ('{malla}',), ' al centro de la capa',
                ') en entrepiso N+3.00 y azotea N+6.00; volados F80012 y F80019 nervados con nervaduras de ', ('TEXT({bw_vol},"0")',), ' cm (casetón ', ('TEXT({ac_vol},"0")',), '×', ('TEXT({ac},"0")',), '×', ('TEXT({hc},"0")',),
                '); F80032 y F80010 de azotea: losa maciza de ', ('TEXT({h_mac},"0")',), ' cm. Apoyada en muros de block confinados y en nervaduras principales.'), font(9, i=True, color=GRAY), None, 30)
    S.gap()
    S.section('1. REVISIONES — TODO EN UN SOLO LUGAR')
    r = S.r
    for k, t in enumerate(['Revisión', '', '', 'Cuántos', 'de', '', '', '', '', '', '', '', 'Estado']):
        hdr_cell(ws, r, 2 + k, t)
    ws.merge_cells('B%d:D%d' % (r, r)); ws.merge_cells('G%d:M%d' % (r, r)); S.r += 1
    est = 'TABLEROS!$%s$%d:$%s$%d' % (TC['estado'], T0, TC['estado'], T1)
    pest = 'PRINCIPALES!$%s$%d:$%s$%d' % (PC['estado'], P0, PC['estado'], P1)
    tc_ = lambda c: 'TABLEROS!$%s$%d:$%s$%d' % (TC[c], T0, TC[c], T1)
    rows = [('Tableros que CUMPLEN (nervados, volados y macizos)', '=COUNTIF(%s,"CUMPLE*")' % est, '=COUNTA(%s)-COUNTIF(%s,"HUECO*")' % (est, est), 'ok'),
            ('Volados F80012 y F80019, nervados con nervaduras de 20 cm, que CUMPLEN (decisión D10)', '=COUNTIFS(%s,"VOLADO",%s,"CUMPLE*")' % (tc_('modo'), est), '=COUNTIF(%s,"VOLADO")' % tc_('modo'), 'ok'),
            ('Losas macizas de 20 cm de azotea (F80032 y F80010) que CUMPLEN (decisión D11)', '=COUNTIFS(%s,"MACIZA",%s,"CUMPLE*")' % (tc_('losa'), est), '=COUNTIF(%s,"MACIZA")' % tc_('losa'), 'ok'),
            ('Tableros que piden MACIZAR casetones junto al apoyo (lista en la sección 7, notas para planos)', '=COUNTIF(%s,"*MACIZAR*")' % est, '=COUNTA(%s)' % est, 'info'),
            ('Tableros que esperan a una nervadura principal', '=COUNTIF(%s,"VER PRINCIPALES*")' % est, '=COUNTA(%s)' % est, 'zero'),
            ('Tableros que NO CUMPLEN', '=COUNTIF(%s,"NO CUMPLE*")' % est, '=COUNTA(%s)' % est, 'zero'),
            ('Tableros donde h_eq < espesor mínimo (calcular flecha)', '=COUNTIF(%s,"REVISAR: calcular flecha*")' % est, '=COUNTA(%s)' % est, 'zero'),
            ('Tableros nervados que no cumplen 6.7.8.1.4: losa sobre vigas, recubrimiento de viga y sin ×1.1 (decisión D7)', '=COUNTIFS(%s,"NO",%s,"PERIMETRAL",%s,"NERVADA")' % (tc_('g678'), tc_('modo'), tc_('losa')), '=COUNTIFS(%s,"PERIMETRAL",%s,"NERVADA")' % (tc_('modo'), tc_('losa')), 'info'),
            ('Tableros fuera de 3.4.2.1 c): distribución de 3.4.3.1 (decisión D8), se confirma con ETABS (D5)', '=COUNTIF(%s,"FUERA*")' % tc_('c342'), '=COUNTIF(%s,"PERIMETRAL")' % tc_('modo'), 'info'),
            ('Nervaduras principales que CUMPLEN con la sección propuesta (decisión D2)', '=COUNTIF(%s,"CUMPLE*")' % pest, '=COUNTA(%s)-COUNTIF(%s,"BORDE*")' % (pest, pest), 'ok'),
            ('Nervaduras principales que hay que PERALTAR respecto a la dala 15×25 del modelo (trabes descolgadas, D2)', '=COUNTIF(PRINCIPALES!$%s$%d:$%s$%d,"PERALTAR*")' % (PC['cambio'], P0, PC['cambio'], P1), '=COUNTA(%s)-COUNTIF(%s,"BORDE*")' % (pest, pest), 'info'),
            ('Principales con muro encima y H/L < 0.6 (NP-15, NP-26): sin arco, la trabe toma todo el muro (criterio del Responsable, D2)', '=COUNTIF(%s,"CUMPLE*H/L < 0.6*")' % pest, '=COUNTIF(PRINCIPALES!$%s$%d:$%s$%d,"*H/L < 0.6*")' % (PC['m418'], P0, PC['m418'], P1), 'ok'),
            ('Nervaduras principales que NO CUMPLEN', '=COUNTIF(%s,"NO CUMPLE*")' % pest, '=COUNTA(%s)-COUNTIF(%s,"BORDE*")' % (pest, pest), 'zero'),
            ('Respaldos de los volados (D10): el M⁻ y el ΔV del volado en las nervaduras de esa dirección, con nervaduras de 20 cm junto al muro raíz', '=COUNTIF(%s,"NO CUMPLE*")' % tc_('vest'), '=COUNTIF(%s,"<>—")' % tc_('vest'), 'zero'),
            ('Orillas laterales del volado F80012 (NP-09 y NP-14) como vigas de orilla en volado (D10) que CUMPLEN', '=%s' % '+'.join('COUNTIF(PRINCIPALES!$%s$%d,"CUMPLE*")' % (PC['estado'], PROW[np_]) for np_ in sorted(VOLLAT)), len(VOLLAT), 'ok'),
            ('Capa de compresión de 6 cm con 6x6-4/4 al centro de la capa (D3, D12): temperatura, diafragma, recubrimiento (zonas no expuestas y expuestas), separación, punzonamiento, acero mínimo junto a la carga (6.7.6.1.2) y flexión con 1 000 kg', '=COUNTIF(%s,"NO CUMPLE*")' % N['CAPA_RNG'], 8, 'zero'),
            ('Anclaje en bordes discontinuos con la dala de borde armada (D9): barra inferior a 1.25 fy, bastón superior, gancho de 180° y losas macizas', '=COUNTIF(%s,"NO CUMPLE*")+COUNTIF(%s,"NO CUMPLE*")+COUNTIF(%s,"NO CUMPLE*")+COUNTIF(%s,"NO CUMPLE*")' % (N['n_anc1'], N['n_anc2'], N['n_g180'], N['n_ancm']), 4, 'zero'),
            ('Nervaduras principales sin apoyo', '=COUNTIF(%s,"SIN APOYO*")' % pest, '=COUNTA(%s)' % pest, 'zero'),
            ('Castillos C1…C8 aceptados: bajan a cimentación (decisión D1)', '=COUNTIF(APOYOS!$C$1:$C$400,"SÍ")', len(D['castillos_prop']), 'info')]
    for c, f1, f2, kind in rows:
        r = S.r
        ws.merge_cells('B%d:D%d' % (r, r)); cell(ws, r, 2, c, al='left', sz=10, wrap=True)
        cell(ws, r, 5, f1, 'key', '0'); cell(ws, r, 6, f2, 'calc', '0')
        ws.merge_cells('G%d:M%d' % (r, r)); cell(ws, r, 7, '', 'calc')
        if kind == 'ok':
            st = '=IF(E%d=F%d,"CUMPLE — todos","REVISAR — "&(F%d-E%d)&" no cumple(n)")' % (r, r, r, r)
        elif kind == 'zero':
            st = '=IF(E%d=0,"CUMPLE — ninguno","REVISAR — ver lista")' % r
        else:
            st = '="informativo"'
        cell(ws, r, 14, st, 'key', '@', al='left')
        ws.row_dimensions[r].height = 15 * max(1, -(-len(c) // 48))
        S.r += 1
    S.gap()
    S.section('2. DECISIONES TOMADAS (propuesta del calculista auxiliar — revisar y firmar el Responsable)')
    S.secnote('El usuario pidió decidir del lado conservador. Cada decisión queda aplicada en el libro (fórmulas vivas) y cada resultado de abajo se recalcula solo. Lo que todavía no cumple se dice tal cual.', 18)
    pr = lambda np_, c: 'PRINCIPALES!$%s$%d' % (PC[c], PROW[np_])
    lbullet('D1) CASTILLOS C1…C8 (APOYOS sección B): ACEPTADOS, bajan a cimentación (', ('COUNTIF(APOYOS!$C$1:$C$400,"SÍ")',), ' de ', str(len(D['castillos_prop'])), ' con SÍ). Sin ellos las principales de la cochera, del vacío de doble altura y de la escalera quedan en volado de 4 a 8 m. '
            'C7 (llegada de la escalera) es nuevo: no tiene castillo arriba y se coordina con arquitectura (2a). Toda bajada que no tenga muro abajo en planta baja es columna de 25×25 como mínimo (NTC-Concreto 2023 (CDMX) 6.4.2.1.1, p. 138).')
    fl = lambda np_: (pr(np_, 'bq') + '&"×"&' + pr(np_, 'hq') + '&" cm, "&' + pr(np_, 'estado') + '&"; flecha "&TEXT(' + pr(np_, 'dlt') + ',"0.00")&" de "&TEXT(' + pr(np_, 'dlim') + ',"0.00")&" cm"&IF(' + pr(np_, 'dlt') + '>0.97*' + pr(np_, 'dlim') + '," (EN EL LÍMITE)","")',)
    lbullet('D2) PRINCIPALES CON MURO ENCIMA: se aceptan las secciones propuestas (trabes descolgadas, tabla 4). Cumplen NTC-Mampostería 2023 (CDMX) 4.1.18.4 (L/h ≤ 14 o I ≥ L³/1600, p. 46), la flecha de 4.1.18.3 (0.3 cm + L/480) y f_CM ≤ F_R·f\'m (4.1.18.2, p. 44). '
            'En NP-15 y NP-26 (H/L < 0.6, donde 4.1.18.2 no aplica): criterio del Responsable, lado seguro — sin efecto de arco, la trabe se diseña con el peso completo del muro y de la azotea que recibe, repartido en el claro, sin reducción; f_CM se sigue calculando como estimación fuera de su campo; castillos en los dos extremos de esos muros. '
            'La flecha diferida cuenta con las varillas superiores corridas en todo el claro (ρ′): así se piden en planos. NP-15: ', fl('NP-15'), ' · NP-26: ', fl('NP-26'), '.')
    lbullet('D3) CAPA DE COMPRESIÓN: capa de ', ('TEXT({t},"0")',), ' cm con casetón 60×60×', ('TEXT({hc},"0")',), ' (h sigue en ', ('TEXT({h},"0")',), ' cm, LR-25) y malla ', ('{malla}',), ' al centro de la capa (≈ ', ('TEXT(({t}-2*{fim})/2,"0.0")',), ' cm por cara, sobre silletas: así la revisa NERVADURA sección 11 y así entra en el d superior de TABLEROS). '
            'Con la capa de 5 cm del plano ninguna malla cumplía a la vez diafragma (6.8.6.4, p. 196), punzonamiento y acero mínimo junto a la carga (6.7.6.1.2, p. 175). '
            'La losa pesa más (peso propio geométrico ', ('TEXT({PP},"0")',), ' kg/m²; el modelo trae ', ('TEXT({PPe2k},"0")',), '): las cargas cambiaron y las secciones de las principales se volvieron a iterar. Corregir la nota 2 del plano de cortes y la malla 6x6-10/10 de EST-03.')
    lbullet('D4) SOBRECARGA: se cuentan los ', ('TEXT({inc_losa},"0")',), ' + ', ('TEXT({inc_mort},"0")',), ' kg/m² de NTC-Criterios y Acciones 2023 (CDMX) 5.1.2.1, p. 9, además de la sobrecarga del modelo (160 a 315 kg/m²): no se puede confirmar que ésta ya traiga la «carga adicional 0.04 t/m²» de Naim, así que se cuenta (lado seguro).')
    lbullet('D5) MODELO DE ETABS: no se puede correr aquí. La lista de cambios está en la sección 2b; hasta que se corra, las reacciones (MUROS) y los momentos de los tableros fuera de 3.4.2.1 c) son los de esta hoja.')
    lbullet('D6) NOTAS PARA PLANOS: sección 7 de esta hoja.')
    lbullet('D7) TABLEROS CHICOS (NTC-Concreto 2023 (CDMX) 6.7.8.1.4, p. 187): donde a_c = ', ('TEXT({ac},"0")',), ' cm pasa de 1/6 del claro corto, la losa se diseña como losa apoyada en vigas (6.7.8.1.9, p. 188): recubrimiento de viga de 4 cm (Tabla 4.14.16.a, p. 85) y V_cR sin ×1.1. '
            'Es el camino que da la propia norma; se acepta y cada tablero lo dice en su nota (', ('COUNTIFS(%s,"NO",%s,"PERIMETRAL",%s,"NERVADA")' % (tc_('g678'), tc_('modo'), tc_('losa')),), ' tableros).')
    lbullet('D8) MÉTODO DE COEFICIENTES fuera de 3.4.2.1 c) (NTC-Concreto 2023 (CDMX) 3.4.2.1 c), p. 64): se acepta la distribución de 3.4.3.1 que ya hace la hoja (2/3 del desequilibrio con rigideces d³/a1, cada tablero con su propio d; la maciza con el suyo). Se confirma con ETABS (D5). Informativo (',
            ('COUNTIF(%s,"FUERA*")' % tc_('c342'),), ' tableros).')
    lbullet('D9) ANCLAJE EN BORDES DISCONTINUOS: en todo borde discontinuo apoyado va una DALA DE BORDE ARMADA de ', ('TEXT({b_borde},"0")',), ' cm de ancho en el peralte de la losa (refuerzo longitudinal y estribos cerrados; armado mínimo propuesto en DATOS sección 1), en lugar de la primera franja junto al apoyo: '
            'sobre muro se apoya en el muro de 15 cm; sobre nervadura principal es la principal ensanchada a ese ancho; en las losas macizas baja ', ('TEXT({h}-{h_mac},"0")',), ' cm bajo la maciza. Un macizo sin armar no sirve: sería losa y el apoyo sería el paño del muro. Quedan ', ('TEXT({b_borde}-{rec_np},"0")',),
            ' cm para anclar: ≥ 1.25·L_dh de la barra inferior (6.7.8.1.7, p. 188) y ≥ L_dh del bastón (14.4.3.1, p. 355), con gancho de 180° (el de 90° parado no cabe en 25 cm). TABLEROS acomoda los casetones con ella y CUANTIFICACION la cuenta; su peso extra (',
            ('TEXT({Wborde},"0.000")',), ' t/m) baja al muro o a la principal (PRINCIPALES ya lo suma). Anclaje: ',
            ('IF(COUNTIF(%s,"NO CUMPLE*")+COUNTIF(%s,"NO CUMPLE*")+COUNTIF(%s,"NO CUMPLE*")+COUNTIF(%s,"NO CUMPLE*")=0,"CUMPLE","NO CUMPLE (NERVADURA sección 10)")' % (N['n_anc1'], N['n_anc2'], N['n_g180'], N['n_ancm']),), '.')
    k12, k19 = 'F80012 N+3.00', 'F80019 N+3.00'
    vtxt = lambda k: ('"%s: "&%s&" No. "&TEXT(%s,"0.#")&" arriba por nervadura (M_u = "&TEXT(%s,"0.00")&" t·m con los macizos de la punta, Mu/M_R = "&TEXT(%s,"0.00")&", V_u/V_cR = "&TEXT(%s/%s,"0.00")&"); 1 corrida hasta el apoyo lejano del respaldo y las demás hasta "&TEXT(%s/100,"0.00")&" m dentro del respaldo; "&%s' %
                      (k.split()[0], T_('nNc', k), T_('vNc', k), T_('MNc', k), T_('FUNc', k), T_('Vu', k), T_('VcR', k), T_('Lbc', k), T_('estado', k)),)
    rsp = endlist('&'.join('"%s: "&IF(LEFT(%s,2)="NO",%s,IF(%s>0,%s&" hilera(s) de 20 cm",IF(LEFT(%s,8)="CUMPLE —",MID(%s,10,40),%s))&IF(%s<0," (levanta en su apoyo lejano con 0.9·PP: la corrida se ancla ahí)",""))&"; "'
                           % (kk, T_('vest', kk), T_('vest', kk), T_('vnh', kk), T_('vnh', kk), T_('vest', kk), T_('vest', kk), T_('vest', kk), T_('vRl', kk)) for kk in BACK))
    lat = '&" · "&'.join('"%s: "&%s' % (np_, pr(np_, 'estado')) for np_ in sorted(VOLLAT))
    lbullet('D10) VOLADOS F80012 Y F80019 (entrepiso): «el entrepiso es nervado»: siguen NERVADOS (LR-25, h = 25 cm, casetón de ', ('TEXT({hc},"0")',), ' y capa de ', ('TEXT({t},"0")',), ' cm); en la dirección del volado las nervaduras son de ', ('TEXT({bw_vol},"0")',), ' cm con casetón ', ('TEXT({ac_vol},"0")',), '×60×', ('TEXT({hc},"0")',),
            ' (s = 70 cm). Con 10 cm el M⁻ pedía una No. 8 que no cabía. ', vtxt(k12), ' · ', vtxt(k19), '. Su M⁻ y su ΔV siguen dentro de los tableros de respaldo (los que están detrás del muro raíz; F80020 y F80021 quedan a los costados de F80019 y no los reciben): '
            'junto al muro raíz las nervaduras siguen de 20 cm hasta donde la de 10 cm resiste con 1 varilla del volado — ', (rsp,), ' Los costados de F80012, que corren en la dirección del volado y llevan barandal, son vigas de orilla en volado (PRINCIPALES): ', (lat,), '. '
            'El área de la orilla libre de F80019 que el DXF daba a NP-16 y el par del volado (M/L_b) bajan al muro raíz (MUROS).')
    sup_ = lambda k, b, L, ln: '"@ "&TEXT(%s,"0.0")&" cm"&IF(2*%s>=%s," corridas"," (L = "&TEXT(%s/100,"0.00")&" m desde el paño)")' % (T_('s' + b, k), T_(L, k), T_(ln, k), T_(L, k))
    mtxt = lambda k: ('"%s: inf. corto No. "&TEXT(%s,"0.#")&" @ "&TEXT(%s,"0.0")&" cm, sup. corto No. "&TEXT(%s,"0.#")&" "&%s&", inf. largo No. "&TEXT(%s,"0.#")&" @ "&TEXT(%s,"0.0")&" cm, sup. largo No. "&TEXT(%s,"0.#")&" "&%s&"; "&%s' %
                      (k.split()[0], T_('vPc', k), T_('sPc', k), T_('vNc', k), sup_(k, 'Nc', 'Lbc', 'ln1'), T_('vPl', k), T_('sPl', k), T_('vNl', k), sup_(k, 'Nl', 'Lbl', 'ln2'), T_('estado', k)),)
    lbullet('D11) AZOTEA F80032 Y F80010 (N+6.00): se respeta el modelo («Slab1»): LOSA MACIZA de ', ('TEXT({h_mac},"0")',), ' cm perimetralmente apoyada, con los mismos coeficientes de la Tabla 3.4.2.1 (son de losa maciza), por franja de 1 m, r = 2 cm, As mín = 0.0018·b·h y f\'c = 250 Clase 1A. '
            'Su lecho inferior queda ', ('TEXT({h}-{h_mac},"0")',), ' cm más alto que el de la LR-25 vecina; en sus bordes discontinuos remata en la dala de borde de D9 y sus dalas de orilla (NP-33, NP-34) descuelgan 5 cm bajo ella. ', mtxt('F80032 N+6.00'), ' · ', mtxt('F80010 N+6.00'), '.')
    lbullet('D12) VOLADOS — CARA SUPERIOR PROTEGIDA: F80012 y F80019 son terrazas con impermeabilizante y piso, igual que la azotea (selector de DATOS sección 1 en «', ('{exp_vol}',), '»; en planos: «terraza con impermeabilizante y piso»). '
            'La exposición ya no manda arriba: la malla va al centro de la capa y la capa cumple en todas las zonas (sección 1). El lecho superior de las nervaduras del volado conserva el recubrimiento de viga r_np = ', ('TEXT({rec_np},"0")',),
            ' cm porque no cumplen 6.7.8.1.4 (losa sobre vigas, D7), del lado seguro; abajo sigue r_ext = ', ('TEXT({rec_ext},"0")',), ' cm (cara exterior). Con eso las nervaduras de 20 cm siguen cumpliendo (D10).')
    S.gap(4)
    S.section('2a. LO QUE TE QUEDA A TI')
    lbullet('(1) Actualizar y volver a correr el modelo de ETABS con la lista de la sección 2b (D5) y comparar reacciones contra MUROS y momentos de los tableros fuera de 3.4.2.1 c).')
    lbullet('(2) Pasar a planos las notas de la sección 7 (D6).')
    lbullet('(3) Coordinar con arquitectura el castillo nuevo C7 (llegada de la escalera, sin castillo arriba) y revisar NP-07 y NP-30 con el diseño de la escalera, que está fuera del alcance (DATOS, ALCANCE).')
    lbullet('(4) Revisar y firmar las decisiones D1 a D12: son una propuesta del calculista auxiliar, del lado conservador.')
    S.gap(4)
    S.section('2b. CAMBIOS AL MODELO DE ETABS (D5)')
    nps_txt = endlist('&'.join('IF(%s="BORDE DE VOLADO","","%s "&%s&"×"&%s&"; ")' % (pr(t['np'], 'cond'), t['np'], pr(t['np'], 'bq'), pr(t['np'], 'hq')) for t in NPS if t['cond'] != 'BORDE DE VOLADO' and t['np'] not in VOLLAT))
    lbullet('(1) Losa «LR-25» con la sección nueva (casetón 60×60×', ('TEXT({hc},"0")',), ', capa de ', ('TEXT({t},"0")',), ' cm, h = ', ('TEXT({h},"0")',), ' cm): modificadores de membrana F11, F22, F12 y de peso WMOD = ', ('TEXT({wmod},"0.0000")',),
            ' (hoy 0.6688) y de flexión M11, M22, M12 = ', ('TEXT({Irat},"0.0000")',), ' (hoy 0.29915), con el mismo material «LR-25 equiv» (DATOS sección 4).')
    lbullet('(2) F2 de azotea: ponerle la sobrecarga de F80023 (270 kg/m²) y la carga viva de azotea (', ('TEXT({cv_az},"0")',), ' kg/m²); hoy no tiene cargas.')
    lbullet('(3) Secciones propuestas de las nervaduras principales (b×h en cm, PRINCIPALES; NP-11 además se prolonga hasta la nervadura del eje F, x = 5.44; NP-09 y NP-14 siguen como dalas de 15×25 en el modelo: su franja maciza es losa): ', (nps_txt,), '')
    lbullet('(4) Volados F80012 y F80019: crear una losa «LR-25V» (nervaduras de ', ('TEXT({bw_vol},"0")',), ' cm, casetón ', ('TEXT({ac_vol},"0")',), '×60×', ('TEXT({hc},"0")',), ') con WMOD y F = ', ('TEXT({wmodv},"0.0000")',), '. Con el eje local 1 en X (el de ETABS por omisión): F80012 (volado en X) M11 = ', ('TEXT({Iratv},"0.0000")',),
            ' y M22 = ', ('TEXT({Irat},"0.0000")',), '; F80019 (volado en Y) M11 = ', ('TEXT({Irat},"0.0000")',), ' y M22 = ', ('TEXT({Iratv},"0.0000")',), '; M12 = ', ('TEXT({Irat},"0.0000")',), ' en los dos (el de la LR-25, el menor; DATOS sección 4b).')
    lbullet('(5) F80032 y F80010 de azotea: se quedan «Slab1» (losa maciza de ', ('TEXT({h_mac},"0")',), ' cm, sin modificadores) pero con concreto f\'c = 250 kg/cm² Clase 1A (hoy 4000 psi).')
    lbullet('(6) Volver a correr y comparar: reacciones contra la hoja MUROS (incluidos los muros raíz de los volados), momentos de los tableros fuera de 3.4.2.1 c) (TABLEROS) y elementos mecánicos de las principales (PRINCIPALES).')
    S.gap()
    S.section('3. TABLA DE NERVADURAS SECUNDARIAS (COMO LA DE NAIM)')
    r = S.r
    heads = ['Tablero', 'Nervadura (maciza: franja de 1 m)', 'Lecho', 'Mu (t·m)', 'b (cm)', 'h (cm)', 'Cant. var. (maciza: @ cm)', 'Var. No.', 'M_R (t·m)', 'Mu/M_R', 'Vu (t)', 'V_cR (t)', 'Estado del tablero']
    for k, t in enumerate(heads):
        hdr_cell(ws, r, 2 + k, t)
    ws.row_dimensions[r].height = 30; S.r += 1
    for t in TABS:
        tr = TROW[key(t)]
        T_ = lambda c: 'TABLEROS!$%s$%d' % (TC[c], tr)
        for i, (lab, lecho, blk, M, b) in enumerate([('‖ claro corto', 'Inferior', 'Pc', 'MPc', '=' + T_('bfc')), ('‖ claro corto', 'Superior', 'Nc', 'MNc', '=' + T_('bNc')),
                                                     ('‖ claro largo', 'Inferior', 'Pl', 'MPl', '=' + T_('bfl')), ('‖ claro largo', 'Superior', 'Nl', 'MNl', '=' + T_('bNl'))]):
            r = S.r
            cell(ws, r, 2, key(t) if i == 0 else '', al='left', bold=True)
            cell(ws, r, 3, lab); cell(ws, r, 4, lecho)
            cell(ws, r, 5, '=' + T_(M), fmt='0.000')
            cell(ws, r, 6, b.format(**N), fmt='0'); cell(ws, r, 7, '=' + T_('hr'), fmt='0')
            cell(ws, r, 8, '=IF(%s="MACIZA","@ "&TEXT(%s,"0.0"),%s)' % (T_('losa'), T_('s' + blk), T_('n' + blk)), 'key', '0'); cell(ws, r, 9, '=' + T_('v' + blk), 'key', 'General')
            cell(ws, r, 10, '=' + T_('MR' + blk), fmt='0.000'); cell(ws, r, 11, '=' + T_('FU' + blk), fmt='0.00')
            cell(ws, r, 12, ('=' + T_('Vu')) if i == 1 else '', fmt='0.000'); cell(ws, r, 13, ('=' + T_('VcR')) if i == 1 else '', fmt='0.000')
            cell(ws, r, 14, ('=' + T_('estado')) if i == 0 else '', al='left', sz=8)
            S.r += 1
    S.gap()
    S.section('4. TABLA DE NERVADURAS PRINCIPALES')
    r = S.r
    heads = ['Tramo', 'Nivel', 'Condición', 'L (m)', 'b (cm)', 'h (cm)', 'Inferior (corridas)', '', 'Superior (corridas)', '', 'Estribos', 'h mín. sin flecha (cm)', 'Estado']
    for k, t in enumerate(heads):
        hdr_cell(ws, r, 2 + k, t)
    ws.merge_cells('H%d:I%d' % (r, r)); ws.merge_cells('J%d:K%d' % (r, r)); ws.row_dimensions[r].height = 30; S.r += 1
    for t in NPS:
        r = S.r; pr = PROW[t['np']]
        P_ = lambda c: 'PRINCIPALES!$%s$%d' % (PC[c], pr)
        cell(ws, r, 2, t['np'] + ' · ' + t['eje'], al='left', sz=8); cell(ws, r, 3, t['level'])
        cell(ws, r, 4, '=' + P_('cond'), sz=8); cell(ws, r, 5, '=' + P_('L'), fmt='0.00')
        cell(ws, r, 6, '=' + P_('bq'), 'key', '0'); cell(ws, r, 7, '=' + P_('hq'), 'key', '0')
        ws.merge_cells('H%d:I%d' % (r, r)); cell(ws, r, 8, '=IF(%s=0,"—",%s&" No. "&TEXT(%s,"0.#"))' % (P_('nP'), P_('nP'), P_('vP')))
        ws.merge_cells('J%d:K%d' % (r, r)); cell(ws, r, 10, '=IF(%s=0,"—",%s&" No. "&TEXT(%s,"0.#"))' % (P_('nN'), P_('nN'), P_('vN')))
        cell(ws, r, 12, '=' + P_('cort'), sz=8); cell(ws, r, 13, '=' + P_('hrec'), 'key', '0')
        cell(ws, r, 14, '=' + P_('estado'), al='left', sz=8)
        S.r += 1
    S.secnote('Todas las varillas de los dos lechos van corridas en todo el claro: así las cuenta CUANTIFICACION y así entra el acero de compresión de la flecha diferida (ρ′). «h mín. sin flecha» = el mínimo por flexión, Tabla 6.3.3.1.1 y NTC-M 4.1.18; '
              'la h propuesta (columna h) ya incluye la flecha de 13.4 y el acomodo de varillas: no se baja la trabe a ese valor. NP-09 y NP-14 son las vigas de orilla del volado F80012: b es la franja maciza (dala + macizo de orilla).')
    S.gap()
    S.section('5. CANTIDADES')
    for k, (c, u) in enumerate([('Concreto de losas (con desperdicio)', 'm³'), ('Casetones 60×60×19 (LR-25)', 'pzas'), ('Casetones 50×60×19 (volados y respaldos)', 'pzas'), ('Malla electrosoldada', 'm²'),
                                ('Acero de nervaduras secundarias y losas macizas', 'kg'), ('Acero de nervaduras principales', 'kg')]):
        r = S.r
        ws.merge_cells('B%d:D%d' % (r, r)); cell(ws, r, 2, c, al='left', sz=10)
        cell(ws, r, 5, '=' + N['qT%d' % k], 'key', '#,##0'); cell(ws, r, 6, u)
        S.r += 1
    S.gap()
    S.section('6. CONDICIONES DE OBRA')
    for txt in ['• Castillos que se bajan a cimentación: refuerzo anclado en el cimiento con escuadra a 90° hacia adentro, empotramiento ≥ 40 cm en cimiento de piedra; traslapes sólo en el tercio medio de la planta baja con estribos @ 7.5 cm (NTC-M 4.1.16.3 a 4.1.16.5, p. 43; 4.2.7.5.2, p. 53). Dalas con 4 varillas y As ≥ 0.2(f\'c/fy)·b·h: la D-15x20 pide 3.57 cm² (2 No. 4 + 2 No. 3) (NTC-M 6.2.2.2, p. 62).',
                '• Nervaduras principales sin muro abajo: se arman como trabe con recubrimiento de 4 cm; TODAS las varillas de la tabla 4 van corridas en todo el claro (2 No. 4 como mínimo arriba y abajo; la flecha diferida cuenta con las superiores, ρ′) y estribos cerrados No. 3 a 135° (gancho 8 cm) con la separación de PRINCIPALES (columna «Cortante»). '
                'En las vigas perimetrales (NP-01 y NP-33) van más cerrados en 2h junto a cada apoyo (NTC-Concreto 6.3.5.1.1, p. 122; 6.3.7.7.7 y 6.3.7.7.8, p. 134; 14.7.3.3, p. 375). Las que piden más de 25 cm quedan descolgadas.',
                '• Donde la hoja pide MACIZAR, se quita la hilera de casetones junto al apoyo en todo el ancho del tablero. En esquinas exteriores sobre muros se refuerza arriba y abajo en un quinto del claro largo (6.7.7.3.1, p. 178).',
                '• En el borde común de dos tableros se arma con el mayor de los dos bastones; los momentos ya incluyen la distribución de 2/3 del desequilibrio (3.4.3.1, p. 64).',
                '• Traslapes clase B: secundarias No. 3 ≈ 50 cm y No. 4 ≈ 66 cm; principales, lecho inferior No. 4 ≈ 66, No. 5 ≈ 83 y No. 8 ≈ 166 cm; principales con h ≥ 40 cm, lecho superior No. 4 ≈ 86 y No. 5 ≈ 108 cm (ψp = 1.3, Tabla 14.4.2.6, p. 354). '
                'Fuera de las secciones de esfuerzo máximo y alternados; malla traslapada un cuadro + 5 cm (14.5.1.3, Tabla 14.5.2.1, 14.5.4.1).']:
        S.bullet(txt)
    lbullet('• Toda nervadura secundaria lleva 1 No. 4 corrida abajo y 1 No. 4 corrida arriba; la inferior se ancla para 1.25 fy en el paño (6.7.6.1.3, p. 175; 6.7.8.1.7, p. 188). Bastones con el largo de TABLEROS medido desde el paño. '
            'En bordes discontinuos las dos rematan con gancho de 180° dentro de la dala de borde armada de ', ('TEXT({b_borde},"0")',), ' cm (decisión D9; 6.7.7.4.1.1, p. 179; 14.4.3.1, p. 355).')
    S.gap()
    S.section('7. NOTAS PARA PLANOS (DECISIÓN D6)')
    sup2 = lambda k, b, L, ln: '" @ "&TEXT(%s,"0.0")&IF(2*%s>=%s," cm corridas"," cm (L = "&TEXT(%s/100,"0.00")&" m desde el paño)")' % (T2('s' + b, k), T2(L, k), T2(ln, k), T2(L, k))
    mac = lambda k: ('"%s de azotea: inf. corto No. "&TEXT(%s,"0.#")&" @ "&TEXT(%s,"0.0")&" cm, sup. corto No. "&TEXT(%s,"0.#")&%s&", inf. largo No. "&TEXT(%s,"0.#")&" @ "&TEXT(%s,"0.0")&" cm, sup. largo No. "&TEXT(%s,"0.#")&%s' %
                     (k.split()[0], T2('vPc', k), T2('sPc', k), T2('vNc', k), sup2(k, 'Nc', 'Lbc', 'ln1'), T2('vPl', k), T2('sPl', k), T2('vNl', k), sup2(k, 'Nl', 'Lbl', 'ln2')),)
    hil = endlist('&'.join('IF(LEFT(%s,7)="MACIZAR","%s: "&LOWER(%s)&"; ","")' % (T2('cort', key(t)), key(t), T2('cort', key(t))) for t in TABS), 'ninguna.')
    rsp20 = endlist('&'.join('IF(%s>0,"%s: "&%s&" hilera(s) junto al muro raíz de %s; ","")' % (T2('vnh', kk), kk, T2('vnh', kk), bb['vol'].split()[0]) for kk, bb in BACK.items()), 'ninguna.')
    lbullet('• Concreto Clase 1A, f\'c = ', ('TEXT({fc},"0")',), ' kg/cm², con agregado grueso CALIZO de 1/2" (TMA ≤ 2/3 del recubrimiento y de la separación libre; Ec = 14 000·√f\'c) (NTC-Concreto 2023 (CDMX) 2.2.2 a), p. 30; Tabla 2.2.1, p. 31; 2.2.5.2.1, p. 32).')
    lbullet('• Precauciones de 3.9.2 (NTC-Concreto 2023 (CDMX) 3.9.1 d) y e) y 3.9.2, p. 74): casetón de medida fija que garantiza b_w y silletas que fijan el peralte del lecho superior.')
    lbullet('• Losa reticular LR-25: casetón de poliestireno 60×60×', ('TEXT({hc},"0")',), ', nervaduras de ', ('TEXT({bw},"0")',), ' cm @ ', ('TEXT({s},"0")',), ' cm y capa de compresión de ', ('TEXT({t},"0")',), ' cm con malla ', ('{malla}',),
            ' AL CENTRO de la capa, sobre silletas (≈ ', ('TEXT(({t}-2*{fim})/2,"0.0")',), ' cm arriba y abajo) (corrige la nota 2 del plano de cortes, que dice 60×60×20 con capa de 5 cm, y la 6x6-10/10 de EST-03; decisión D3).')
    lbullet('• Volados F80012 y F80019: ', ('IF({exp_vol}="SÍ","cara superior a la intemperie: recubrimiento de "&TEXT({rec_ext},"0")&" cm abajo y arriba","TERRAZA CON IMPERMEABILIZANTE Y PISO (cara superior protegida, decisión D12): recubrimiento de "&TEXT({rec_ext},"0")&" cm abajo (cara exterior) y "&TEXT({rec_np},"0")&" cm arriba (recubrimiento de viga, 6.7.8.1.9)")',),
            '; nervaduras de ', ('TEXT({bw_vol},"0")',), ' cm en la dirección del volado con casetón ', ('TEXT({ac_vol},"0")',), '×60×', ('TEXT({hc},"0")',), ' (decisión D10), alineadas con las de sus tableros de respaldo; armado de TABLEROS (sección 3 de esta hoja): 1 varilla superior corrida hasta el apoyo lejano del respaldo y las demás hasta el fin de las nervaduras de 20 cm del respaldo. '
            'Orillas libres (NP-16 y NP-17) como dala de orilla de 15×25 con 4 No. 4 y estribos No. 3 @ 20 cm, donde rematan las varillas inferiores de las nervaduras del volado.')
    lbullet('• Tableros de respaldo de los volados: siguen con nervaduras de ', ('TEXT({bw_vol},"0")',), ' cm (casetón ', ('TEXT({ac_vol},"0")',), '×60×', ('TEXT({hc},"0")',), ') en las hileras junto al muro raíz: ', (rsp20,), ' (TABLEROS, columnas RESPALDO, decisión D10).')
    lbullet('• Vigas de orilla del volado F80012 (NP-09 y NP-14, sus costados): franja maciza de la sección de la tabla 4 (descuelga lo que pase de 25 cm), con sus varillas superiores corridas desde la punta hasta L_d dentro de la principal que sigue del otro lado del muro raíz (NP-10 y NP-15) y los estribos de PRINCIPALES (decisión D10).')
    lbullet('• Losas macizas de ', ('TEXT({h_mac},"0")',), ' cm en F80032 y F80010 de azotea (decisión D11), con su lecho inferior ', ('TEXT({h}-{h_mac},"0")',), ' cm más alto que el de la LR-25 vecina (desnivel en cimbra y plafón); en sus bordes discontinuos rematan en la dala de borde de ',
            ('TEXT({b_borde},"0")',), '×', ('TEXT({h},"0")',), ' (baja ', ('TEXT({h}-{h_mac},"0")',), ' cm bajo la maciza). ', mac('F80032 N+6.00'), ' · ', mac('F80010 N+6.00'), '.')
    lbullet('• Dala de borde ARMADA de ', ('TEXT({b_borde},"0")',), ' cm de ancho en el peralte de la losa en todo borde discontinuo apoyado (decisión D9): sobre muro se apoya en el muro de 15 cm; sobre nervadura principal es la principal ensanchada a ese ancho con estribos que lo abarcan. '
            'Armado mínimo propuesto: 4 No. 4 corridas y estribos cerrados No. 3 @ 20 cm de todo el ancho. Ahí rematan con GANCHO DE 180° (extensión 4d_b ≥ 6.5 cm) la varilla inferior de cada nervadura (1.25 fy en el paño interior de la dala) y el bastón superior; '
            'el gancho de 90° parado no cabe en 25 cm (se permite el de 90° acostado, dentro del armado de la dala). Un macizo sin armar no sirve.')
    lbullet('• Castillos C1…C8 hasta la cimentación (APOYOS sección B, decisión D1; C7 es nuevo, sin castillo arriba: coordinar con arquitectura); donde no haya muro abajo en planta baja, columna de 25×25 como mínimo (NTC-Concreto 2023 (CDMX) 6.4.2.1.1, p. 138). Castillos en los dos extremos de los muros que cargan NP-15 y NP-26 (decisión D2).')
    lbullet('• Trabes descolgadas de las nervaduras principales con la sección de la tabla 4 y todas sus varillas corridas en todo el claro (decisión D2).')
    lbullet('• Apuntalamiento del entrepiso N+3.00 hasta planta baja mientras se cuela la azotea: 150 kg/m² más el peso del colado y la cimbra (NTC-Criterios y Acciones 2023 (CDMX) 6.1.3.1, p. 10).')
    lbullet('• Hileras de casetones a macizar junto al apoyo del claro corto, en todo el ancho del tablero: ', (hil,), '')
    add_cf(ws, 'C1:N%d' % S.r)
    ws.sheet_view.zoomScale = 85
    return ws


def sheet_refs(wb):
    ws = wb.create_sheet('REFERENCIAS')
    S = Narr(ws, N, {'A': 1.78, 'B': 58, 'C': 40, 'D': 46, 'E': 10, 'F': 10})
    S.title('REFERENCIAS — QUÉ DOCUMENTO Y QUÉ PARTE CHECAR')
    S.subtitle('Los archivos del proyecto se citan por su tabla o capa; los documentos de Naim por página impresa; las NTC 2023 por numeral y página impresa de cada norma, cotejados contra el texto oficial del compendio.', 30)

    def block(title, items):
        S.section(title)
        S.header(('Tema', 'Documento', 'Tabla / capa / página / sección', '—', '—'))
        for a, b, c in items:
            r = S.r
            cell(ws, r, 2, a, al='left', sz=9, wrap=True); cell(ws, r, 3, b, al='left', sz=9, wrap=True); cell(ws, r, 4, c, 'ref', al='left', wrap=True)
            cell(ws, r, 5, ''); cell(ws, r, 6, '')
            ws.row_dimensions[r].height = max(18, 13 * (1 + max(len(a) // 60, len(c) // 46)))
            S.r += 1
    block('MODELO ETABS — ' + E2K, [
        ('Geometría de tableros, dalas, castillos y muros por nivel', E2K, 'POINT COORDINATES / LINE CONNECTIVITIES / AREA CONNECTIVITIES / AREA ASSIGNS'),
        ('Losa LR-25: espesor 0.25 y modificadores F 0.6688, M 0.29915 (sección vieja 60×60×20 con capa de 5 cm; con la de la decisión D3 cambian: DATOS secciones 4 y 4b)', E2K, 'SLAB PROPERTIES — SHELLPROP "LR-25"'),
        ('Losa maciza «Slab1» de 0.20 m (F80032 y F80010 de azotea), concreto 4000 psi: se respeta como losa maciza (decisión D11) con f\'c = 250', E2K, 'SLAB PROPERTIES — SHELLPROP "Slab1" · AREA ASSIGNS'),
        ("Concreto f'c = 250 kg/cm², E = 14,000√f'c, γ = 2440 kg/m³", E2K, "MATERIAL PROPERTIES — \"Concreto f'c=250kg/cm2. Clase 1A\""),
        ('Peso equivalente de la LR-25: 1503.75 kg/m³', E2K, 'MATERIAL "Concreto f\'c=250kg/cm2. Clase 1A (LR-25 equiv)"'),
        ('Varilla fy = 4200 kg/cm²', E2K, 'MATERIAL "Fy=4200"'),
        ('Sobrecarga y cargas vivas por tablero', E2K, 'SHELL OBJECT LOADS — AREALOAD'),
        ('Barandales, muros bajos, pretiles y reacción de la escalera', E2K, 'FRAME OBJECT LOADS — LINELOAD'),
        ('Combinaciones 1.3 CM + 1.5 CV (Comb2.2) y 1.5/1.7 (Comb2.1, grupo A)', E2K, 'LOAD COMBINATIONS'),
        ('Dalas D-15x20 y vigas VP-15x40 del pergolado', E2K, 'FRAME SECTIONS / LINE ASSIGNS'),
    ])
    block('PLANOS DXF', [
        ('Casetón 60×60×20, nervadura 10 cm, capa 5 cm (h = 0.25); la hoja usa 60×60×19 con capa de 6 cm (decisión D3): corregir la nota', DXFA, 'NOTAS DE CORTES, nota 2 (capa 010_Notas)'),
        ('Azotea: losa reticular, peralte según cálculo', DXFA, 'NOTAS DE CORTES, nota 3'),
        ('Vacío de doble altura, vacío de escalera, tragaluces', DXFA, 'Capa 004_Vacios, plantas PRIMER NIVEL y AZOTEA'),
        ('Escalera curva (cubo de escalera)', DXFA, 'Capa 005_Escalera'),
        ('Cochera de 8.70 m libres (muros sólo en A y J)', DXFA, 'Capa 001_Muros, PLANTA BAJA'),
        ('Tableros y áreas tributarias a 45°', DXFT, 'Capas TRIB_TABLERO, TRIB_AREA, TRIB_HUECO'),
    ])
    block('DOCUMENTOS DE NAIM (página impresa)', [
        ('Normas usadas: NTC-2023 concreto, mampostería, acero; MOC-CFE sismo 2015', NAIM24, 'pág. 3'),
        ("Materiales: f'c 250, fy 4200, f'm 15 (valor del Cap. 12, sólo para Tipo II; esta casa es Tipo I)", NAIM24, 'pág. 4'),
        ('Cargas muertas de entrepiso (0.33 t/m²) y azotea (0.29 t/m²); carga viva de habitación', NAIM24, 'pág. 5'),
        ('Carga viva de azotea con pendiente ≤ 5%', NAIM24, 'pág. 6'),
        ('Factores 1.3 permanentes / 1.5 variables, estructura grupo B', NAIM24, 'pág. 8 (lista de combinaciones en pág. 9)'),
        ('Tabla de nervaduras: recubrimiento 2 cm, MR y VcR de cada nervadura (con eso se verificaron las fórmulas)', NAIM24, 'pág. 13'),
        ('FR = 0.9 en flexión y As mín = 0.7√f\'c/fy·b·d (fórmula anterior; la 2023 pide máx(0.80√f\'c, 14)/fy)', NAIM21, 'pág. 13'),
        ('Detalle de nervadura principal/secundaria, malla 6x6-10/10, tabla de nervaduras', NAIMPL, 'EST-03, EST-04 y EST-05'),
        ('Mallas con Fy = 5000 kg/cm²; recubrimiento de 2 cm en losas y dalas', 'Naim — EST CASA YÉPEZ', 'Notas 1.2 y 1.2.2'),
    ])
    block(NTCC + ' — NUMERAL Y PÁGINA IMPRESA', [
        ("Concreto Clase 1A, f'c ≥ 250; Ec = 14 000√f'c (agregado calizo); f̄f = 2√f'c", NTCC, '2.2.2 a), p. 30 · Tabla 2.2.1, p. 31 · Tabla 2.2.6.3, p. 33 · 2.2.8.4, p. 34'),
        ('Tamaño máximo del agregado ≤ 2/3 del recubrimiento y de la separación libre', NTCC, '2.2.5.2.1, p. 32'),
        ('Barras Grado 42 y malla NMX-B-290 (fy de diseño ≤ 7 000 en losas; estribos corrugados, fy ≤ 4 200)', NTCC, '2.4.2.1.1, p. 43 · 2.4.2.1.6 y 2.4.2.1.7, p. 44 · Tabla 2.4.2.2.4.a, p. 46'),
        ('Métodos reconocidos y líneas de fluencia; losa como diafragma', NTCC, '3.4.1.1 y 3.4.1.2, p. 64'),
        ('Método de coeficientes: condiciones a) a d) y distribución de momentos', NTCC, '3.4.2.1 y 3.4.3.1, p. 64'),
        ('Coeficientes de momento, franjas centrales, caso I (losa monolítica)', NTCC, 'Tabla 3.4.2.1, p. 65 (≈50% de valores cambiaron respecto a 2017, comentario p. 66)'),
        ('Cargas lineales de muros dentro de un tablero', NTCC, '3.4.4.1 a 3.4.4.3 y Tabla 3.4.4.2, pp. 65–66'),
        ("Bloque equivalente: f''c = 0.85f'c, β1 = 0.85 hasta f'c = 300, ε_cu = 0.003", NTCC, '3.6.1 d) y e), ec. 3.6.1, p. 68'),
        ('Factores de resistencia: flexión 0.65–0.90 según ε_t; cortante 0.75', NTCC, 'Tablas 3.8.2.1 y 3.8.2.2, p. 70 · 3.8.2.4, p. 71'),
        ('Reducción de 20 mm en anchos < 200 mm y en el peralte del lecho superior ≤ 200 mm, salvo precauciones anotadas en planos', NTCC, '3.9.1 d) y e) y 3.9.2, p. 74'),
        ('Clase de exposición (cara superior de volados sin techo: A2) y recubrimientos (losas y viguetas 20 mm; vigas 40 mm; intemperie 40 mm)', NTCC, 'Tabla 4.3.1 c) y 4.3.2, p. 76 · 4.13.1.1 y 4.13.2.2, p. 83 · Tabla 4.13.3.1 y 4.13.3.2, p. 84 · Tabla 4.14.16.a, p. 85'),
        ('Cuantía balanceada; ancho efectivo del patín; MR de sección rectangular y T', NTCC, 'ec. 5.2.1.3.1, p. 88 · Tabla 5.2.1.4.2, p. 89 · ecs. 5.2.2.1.1.1.a y b, p. 89 · 5.2.2.1.3.1, p. 90'),
        ("Cortante: V_cR con A_v,mín (0.5√f'c o 2p^(1/3)√f'c ≤ 1.25√f'c); sin refuerzo en el alma (2λs·p^(1/3)√f'c); límite de la sección; V_sR", NTCC, '5.5.2.1 y 5.5.2.2, p. 94 · 5.5.3.1.1 y 5.5.3.1.2, pp. 94–95 · 5.5.3.2.1, p. 95 · 5.5.3.6.1 y 5.5.3.8.1, p. 97'),
        ('Punzonamiento de la capa con la carga concentrada', NTCC, 'Tabla 5.6.2.1.1, p. 100'),
        ('Vigas (incluye nervaduras): claro, h/b ≤ 6, peralte mínimo L/16–L/8 y h ≥ 25 cm', NTCC, '6.3.1.2 b), 6.3.2.1 y 6.3.2.2, p. 117 · Tabla 6.3.3.1.1, p. 120 · 6.3.3.2.1 y 6.3.3.3.1, p. 121'),
        ("Vigas: As mín = máx(0.80√f'c, 14)/fy·b_w·d (con patín a tensión en viga isostática —volados— b_w = el menor de b_f y 2b_w), 2 barras corridas, excepción 4/3; As máx 0.90 As,b", NTCC, '6.3.5.1.1 a 6.3.5.1.3 y 6.3.5.2.1, p. 122'),
        ('Vigas: A_v,mín (casos exentos) y su valor', NTCC, '6.3.5.4.1 y Tabla 6.3.5.4.1, p. 123 · Tabla 6.3.5.4.4, p. 124'),
        ('Corte y anclaje del refuerzo de flexión; punto de inflexión', NTCC, '6.3.7.3.4 a 6.3.7.3.6, p. 127 · 6.3.7.3.10, p. 129'),
        ('Separación máxima de estribos; integridad estructural (2h en extremos)', NTCC, 'Tabla 6.3.7.6.2.2, p. 132 · 6.3.7.7, p. 134'),
        ('Losas en dos direcciones: alcance (nervaduras en dos direcciones)', NTCC, '6.7.1.1 e) y comentario de 6.7.1, p. 165'),
        ('Espesor mínimo: Tabla 6.7.3.1.1 (sin vigas) y Tabla 6.7.3.1.2 (α_fm), +10% en borde discontinuo; flechas si no cumple', NTCC, 'Tabla 6.7.3.1.1, p. 167 · Tabla 6.7.3.1.2, 6.7.3.1.2.1 y 6.7.3.2.1, p. 168'),
        ('Losas controladas por tensión', NTCC, '6.7.3.3.1, p. 169'),
        ('Sección crítica de cortante a d del paño', NTCC, '6.7.4.3.2, p. 172'),
        ('As mín de losa 0.0018·A_g; acero mínimo alrededor de una carga concentrada (ec. 6.7.6.1.2, ancho b_losa); 1 No. 4 abajo y 1 No. 4 arriba corridas en toda nervadura', NTCC, '6.7.6.1.1 a 6.7.6.1.3, p. 175 · Tabla 6.7.4.2.2.3, p. 171'),
        ('Refuerzo de esquinas exteriores; bordes discontinuos', NTCC, '6.7.7.3.1, p. 178 · 6.7.7.4.1.1, p. 179'),
        ('Losas encasetonadas: b_w ≥ 10 cm, h ≤ 3.5b_w, separación libre ≤ 80 cm y ≤ claro/6; V_cR × 1.1; carga de 1 000 kg en 10×10 cm', NTCC, '6.7.8.1.2 a 6.7.8.1.6, p. 187'),
        ('Losas encasetonadas: barra inferior continua para 1.25 fy; refuerzo de la capa; si no cumple, losa sobre vigas; capa ≥ máx(a_c/13, 50 mm)', NTCC, '6.7.8.1.7 a 6.7.8.1.9 y 6.7.8.3.1, p. 188'),
        ('Losas macizas de azotea (F80032 y F80010): mismo método de coeficientes, As mín 0.0018·b·h, cortante con la ec. 5.5.3.2.1 y b = 100 cm, espesor mínimo con la Tabla 6.7.3.1.2 comparado con h', NTCC, '3.4.2.1, p. 64 · Tabla 3.4.2.1, p. 65 · 5.5.3.2.1, p. 95 · 6.7.6.1.1, p. 175 · Tabla 6.7.3.1.2, p. 168 · 13.8.2.1.1, p. 345'),
        ('Diafragma: firme ≥ 50 mm (≥ 60 mm con claros ≥ 6 m); V_R; cuantía mínima 0.0025; separación ≤ 5t y 45 cm', NTCC, '6.8.1.1 a), p. 188 · 6.8.3.1.2, p. 189 · ec. 6.8.5.3.2, p. 192 · 6.8.6.3, 6.8.6.4 y 6.8.7.3, p. 196'),
        ('Columnas de ductilidad baja (si una bajada queda aislada): dimensión ≥ 25 cm', NTCC, 'ec. 6.4.1.1, p. 137 · 6.4.2.1.1, p. 138'),
        ('Flechas: límites 0.5 cm + L/240 y 0.3 cm + L/480; inercia efectiva (Bischoff); diferida λΔ = ξ/(1 + 50p′)', NTCC, '13.4.1.1, p. 336 · Tabla 13.4.3.2 y ec. 13.4.3.3, p. 338 · ec. 13.4.4.1, p. 339'),
        ('Vibración de pisos que no cumplen el espesor mínimo', NTCC, '13.7.1.1, p. 343'),
        ('Contracción y temperatura: cuantía 0.0018 (ya no 660x₁/(fy(x₁+100)))', NTCC, '13.8.2.1.1, p. 345'),
        ('Separación libre mínima: máx(25 mm, d_b, 1.5·TMA)', NTCC, '14.2.1, p. 347'),
        ('Ganchos y dobleces (90°: 12d_b; estribos 135°: 6d_b ≥ 80 mm)', NTCC, 'Tabla 14.3.1, p. 349 · 14.3.4 y Tabla 14.3.2, p. 350'),
        ('Longitud de desarrollo L_d (No. 6 y menores / No. 7 a 11; «otros casos»), factor ψp de barra superior, gancho L_dh y traslapes clase B; traslape de malla', NTCC, '14.4.2.1, p. 352 · Tabla 14.4.2.4, p. 353 · Tabla 14.4.2.6, p. 354 · 14.4.3.1 a) a c), p. 355 · 14.4.10.2, p. 366 · Tabla 14.5.2.1, p. 367 · 14.5.4.1, p. 369'),
        ('Estribo cerrado mínimo No. 3; separación ≤ 16d_b long. y 48d_b del estribo', NTCC, '14.7.3.2 y 14.7.3.3, p. 375'),
    ])
    block(NTCCA + ' — NUMERAL Y PÁGINA IMPRESA', [
        ('Combinación de cargas permanentes + variables con W_m uniforme', NTCCA, '2.3.1 a), p. 5'),
        ('Factores de carga: 1.3 / 1.5 (grupo B), 1.5 / 1.7 (grupo A); 1.1 accidentales; 0.9 favorables; 1.0 servicio', NTCCA, '3.4.1 a) a d), p. 7'),
        ('Límites de desplazamiento: remite a la norma del material', NTCCA, '4.1.1, p. 8'),
        ('Incremento de 20 kg/m² por losa colada en sitio y 20 kg/m² por capa de mortero', NTCCA, '5.1.2.1 y 5.1.2.2, p. 9'),
        ('Uso de W_m, W_a y W; muros divisorios aparte; cargas transitorias de obra', NTCCA, '6.1.1.2, 6.1.1.3, 6.1.2.1 y 6.1.3.1, p. 10'),
        ('Cargas vivas unitarias: habitación 190/100/80; azotea ≤ 5%: 100/70/15; cargas concentradas', NTCCA, 'Tabla 6.1.2.2 incisos a) y h), p. 11 · notas [1], [7] y [8], pp. 11–12'),
        ('Temperatura: se puede omitir en estructuras de menos de 30 m con azotea protegida', NTCCA, '6.2.3.1, p. 13'),
        ('Granizo: W_g = 900·H_s·C_s = 112.5 kg/m² (grupo B); flecha ≤ L/240 con pendiente ≤ 2%', NTCCA, '7.3.1.1 y 7.3.2.1, p. 17 · Tablas 7.3.2.2 y 7.3.2.3, p. 18 · 7.3.4.2, p. 19'),
    ])
    block(NTCM + ' — NUMERAL Y PÁGINA IMPRESA', [
        ('Estructura Tipo I: tiene muros estructurales desplantados sobre vigas o sobre la losa', NTCM, '1.6.1 c), p. 15 · 1.6.2, p. 16'),
        ("f'm de block de concreto f'p = 100 con mortero tipo I = 50 kg/cm²; Em = 800 f'm", NTCM, 'Tabla 2.8.1.4.1, p. 24 · ec. 2.8.5.4.a, p. 27'),
        ('Bajada de cargas por áreas tributarias', NTCM, '3.1.2.2.1, p. 31'),
        ('Factor de resistencia en compresión F_R = 0.6', NTCM, '4.1.6.2, p. 39'),
        ('Anclaje de castillos en la cimentación; muros que parten de otro nivel; barras del piso ancladas en los muros', NTCM, '4.1.16.3 a 4.1.16.5 y comentario, 4.1.17.2, p. 43'),
        ('Muros sobre vigas: esfuerzo en los extremos (f_CM ≤ F_R f\'m, H/L ≥ 0.6), flecha (ec. 4.1.18.3.1.a), L/h ≤ 14 y h ≥ 25 cm', NTCM, '4.1.18.1 a 4.1.18.2, pp. 44–45 · 4.1.18.3 y 4.1.18.4, p. 46'),
        ('Traslapes de castillos en el primer nivel (tercio medio, estribos @ h_c/2)', NTCM, '4.2.7.5.2 y 4.2.7.5.3, p. 53'),
        ('Castillos y dalas: dimensiones, f\'c ≥ 200, 4 barras con As ≥ 0.2(f\'c/fy)·b·h, estribos', NTCM, '6.1.3 y 6.2.1.1 a 6.2.1.3, pp. 60–61 · 6.2.2.2, 6.2.2.3 y 6.2.3.2, p. 62'),
        ('Resistencia a compresión del muro P_R = F_R·F_E·(f\'m·A_T + ΣAs·fy)', NTCM, 'ec. 6.4.1.1, p. 65 · F_E: 3.1.2.3.1, p. 32'),
    ])
    S.notice('ACI 318-19 ya no se usa en esta hoja: todo lo que antes se tomaba de ahí (nervaduras, peraltes de vigas, estribos, separación de barras, bastones) lo resuelve la NTC-Concreto 2023 con los numerales de arriba.', 30)
    return ws


def build(path):
    wb = Workbook(); wb.remove(wb.active)
    sheet_datos(wb)
    sheet_apoyos(wb)
    sheet_cargas(wb)
    sheet_tableros(wb)
    sheet_nervadura(wb)
    sheet_principales(wb)
    sheet_muros(wb)
    wsK = sheet_croquis(wb)
    sheet_cuant(wb)
    sheet_resumen(wb)
    sheet_refs(wb)
    sheet_tablas(wb)
    sheet_trib(wb)
    sheet_transfer(wb)
    build_croquis_datos(wb, wsK)
    from openpyxl.worksheet.properties import PageSetupProperties
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    wsK.page_setup.fitToHeight = 1; wsK.page_setup.orientation = 'portrait'
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return path


if __name__ == '__main__':
    print(build(sys.argv[1]))
