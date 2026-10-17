"""Parte C: CROQUIS + CROQUIS_DATOS (gráficas de dispersión que se recalculan solas), CUANTIFICACION, RESUMEN, REFERENCIAS."""
from gen_b import *
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.series import SeriesLabel, StrRef
from openpyxl.chart.label import DataLabelList, DataLabel
from openpyxl.chart.marker import Marker
from openpyxl.chart.legend import LegendEntry
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import cm_to_EMU
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DFont

LV = ['N+3.00', 'N+6.00']
CD = 'CROQUIS_DATOS'
NPT = 60           # puntos por serie dinámica
Y0 = 8             # primera fila de datos
LAB0 = 400         # fila donde viven los textos de las leyendas


def colpair(i):
    return CL(i), CL(i + 1)


# ---------------------------------------------------------------- CROQUIS (parámetros)
def sheet_croquis(wb):
    ws = wb.create_sheet('CROQUIS'); S = Narr(ws, N)
    S.title('CROQUIS DE LAS NERVADURAS — PLANTAS, TABLERO, CORTES Y ALZADO')
    S.subtitle('Seis dibujos: las dos plantas con apoyos, nervaduras principales y vacíos; la planta del tablero escogido con sus casetones y nervaduras; el corte y el alzado de la nervadura secundaria; '
               'y el corte de la nervadura principal escogida. Todo sale de las tablas de abajo, que están enlazadas a NERVADURA y PRINCIPALES: se redibuja solo.', 42)
    S.notice('Para cambiar de tablero usa el selector de la hoja NERVADURA (celda amarilla de arriba). Para cambiar de nervadura principal usa el selector de esta hoja (sección 2). '
             'Las coordenadas viven en CROQUIS_DATOS, al final del libro: se generan solas, no se tocan.', 30)
    S.gap()
    S.section('1. TABLERO QUE SE DIBUJA — TODO ENLAZADO')
    S.header()
    S.row('Tablero', '={sel}', '—', 'Enlazado a NERVADURA.', 'NERVADURA', 'key', '@', 'k_sel')
    S.row('Nivel', '={n_niv}', '—', '—', 'NERVADURA', 'calc', '@', 'k_niv')
    S.row('Modo', '={n_modo}', '—', '—', 'NERVADURA', 'calc', '@', 'k_modo')
    S.row('Losa', '={n_losa}', '—', 'NERVADA (LR-25) o MACIZA (decisión D11).', 'NERVADURA', 'calc', '@', 'k_losa')
    S.row('h de la losa', '={n_h}', 'cm', '—', 'NERVADURA', 'calc', '0', 'k_h')
    S.row('Ancho del alma en el corte', '=IF({k_losa}="MACIZA",{s},{n_bwt})', 'cm', 'b_w de la nervadura del claro corto (20 cm en los volados, D10); en la maciza se dibuja una franja de ancho s.', 'NERVADURA', 'calc', '0', 'k_bwd')
    S.row('a_x', '={n_ax}*100', 'cm', 'NERVADURA × 100.', 'NERVADURA', 'calc', '0', 'k_ax')
    S.row('a_y', '={n_ay}*100', 'cm', 'NERVADURA × 100.', 'NERVADURA', 'calc', '0', 'k_ay')
    S.row('Claro corto paralelo a', '=IF({k_ax}<={k_ay},"X","Y")', '—', '—', 'NERVADURA', 'calc', '@', 'k_dir')
    S.row('Casetones en X / en Y', '={n_ncx}&" × "&{n_ncy}', 'pzas', '—', 'NERVADURA sección 10', 'calc', '@')
    S.row('Escala del dibujo de planta', '=82/MAX({k_ax},{k_ay})', 'u/cm', 'El tablero se ajusta a un cuadro de 82 unidades.', 'Sólo dibujo', 'calc', '0.0000', 'k_sc')
    S.row('Origen X del dibujo', '=12+(82-{k_ax}*{k_sc})/2', 'u', 'Centra el tablero.', 'Sólo dibujo', 'calc', '0.00', 'k_ox')
    S.row('Origen Y del dibujo', '=12+(82-{k_ay}*{k_sc})/2', 'u', 'Centra el tablero.', 'Sólo dibujo', 'calc', '0.00', 'k_oy')
    S.row('Inferior ‖ claro corto', '=IF({k_losa}="MACIZA","No. "&TEXT({n_vPc},"0.#")&" @ "&TEXT({n_sPc},"0.0")&" cm corridas",{n_nPc}&" No. "&TEXT({n_vPc},"0.#")&" corrida(s)")', '—', 'Enlazado a NERVADURA sección 4.', 'NERVADURA', 'calc', '@', 'k_tPc')
    S.row('Superior ‖ claro corto', '=IF({k_losa}="MACIZA","No. "&TEXT({n_vNc},"0.#")&" @ "&TEXT({n_sNc},"0.0")&" cm, L = "&TEXT({n_Lbc}/100,"0.00")&" m desde el paño","1 No. "&TEXT({n_vNc},"0.#")&" corrida"&IF({n_nNc}>1," + "&({n_nNc}-1)&" bastón(es) No. "&TEXT({n_vNc},"0.#")&", L = "&TEXT({n_Lbc}/100,"0.00")&" m desde el paño"," (no requiere bastones)"))', '—', 'Largo desde el paño: punto de inflexión + máx(d, 12d_b, L_n/16), y ≥ L_d (NERVADURA sección 10).', 'NERVADURA secciones 5 y 10', 'calc', '@', 'k_tNc')
    S.row('Inferior ‖ claro largo', '=IF({k_losa}="MACIZA","No. "&TEXT({n_vPl},"0.#")&" @ "&TEXT({n_sPl},"0.0")&" cm corridas",{n_nPl}&" No. "&TEXT({n_vPl},"0.#")&" corrida(s)")', '—', '—', 'NERVADURA sección 6', 'calc', '@', 'k_tPl')
    S.row('Superior ‖ claro largo', '=IF({k_losa}="MACIZA","No. "&TEXT({n_vNl},"0.#")&" @ "&TEXT({n_sNl},"0.0")&" cm, L = "&TEXT({n_Lbl}/100,"0.00")&" m desde el paño","1 No. "&TEXT({n_vNl},"0.#")&" corrida"&IF({n_nNl}>1," + "&({n_nNl}-1)&" bastón(es) No. "&TEXT({n_vNl},"0.#")&", L = "&TEXT({n_Lbl}/100,"0.00")&" m desde el paño"," (no requiere bastones)"))', '—', '—', 'NERVADURA secciones 7 y 10', 'calc', '@', 'k_tNl')
    S.gap()
    S.section('2. NERVADURA PRINCIPAL QUE SE DIBUJA')
    S.header()
    S.row('Tramo de nervadura principal', 'NP-23', '—', 'SELECTOR. Lista de PRINCIPALES.', 'PRINCIPALES', 'input', '@', 'k_np')
    dv_list(ws, N['k_np'].split('!')[1].replace('$', ''), '=' + N['NPKEYS'])
    S.row('Renglón en PRINCIPALES', '=MATCH({k_np},{NPKEYS},0)', '—', 'Auxiliar.', '—', 'calc', '0', 'k_npr')

    def pv(c):
        return 'INDEX(PRINCIPALES!$%s$%d:$%s$%d,{k_npr})' % (PC[c], P0, PC[c], P1)
    S.row('Nivel', '=' + pv('nivel'), '—', '—', 'PRINCIPALES', 'calc', '@', 'k_npniv')
    S.row('Eje', '=' + pv('eje'), '—', '—', 'PRINCIPALES', 'calc', '@', 'k_npeje')
    S.row('Condición', '=' + pv('cond'), '—', '—', 'PRINCIPALES', 'key', '@', 'k_npcond')
    S.row('L', '=' + pv('L'), 'm', '—', 'PRINCIPALES', 'calc', '0.00', 'k_npL')
    S.row('b', '=' + pv('bq'), 'cm', '—', 'PRINCIPALES', 'calc', '0', 'k_b')
    S.row('h', '=' + pv('hq'), 'cm', '—', 'PRINCIPALES', 'calc', '0', 'k_hh')
    S.row('Varillas inferiores', '=' + pv('nP'), 'pzas', '—', 'PRINCIPALES', 'calc', '0', 'k_nP')
    S.row('No. inferior', '=' + pv('vP'), '—', '—', 'PRINCIPALES', 'calc', 'General', 'k_vP')
    S.row('Varillas superiores', '=' + pv('nN'), 'pzas', '—', 'PRINCIPALES', 'calc', '0', 'k_nN')
    S.row('No. superior', '=' + pv('vN'), '—', '—', 'PRINCIPALES', 'calc', 'General', 'k_vN')
    S.row('Separación de estribos (resto del claro)', '=' + pv('se'), 'cm', '—', 'PRINCIPALES', 'calc', '0.0', 'k_se')
    S.row('Separación de estribos en 2h junto a cada apoyo', '=' + pv('sext'), 'cm', '—', 'PRINCIPALES · NTC-Concreto 6.3.7.7.7, p. 134', 'calc', '0.0', 'k_sext')
    S.row('Estado', '=' + pv('estado'), '—', '—', 'PRINCIPALES', 'key', '@', 'k_npest')
    S.row('Paso del achurado de los apoyos', 4, 'u', 'CAPTURA. Sólo dibujo.', 'Sólo dibujo', 'input', '0', 'k_hatch')
    S.gap()
    S.section('3. DIBUJOS')
    S.secnote('Arriba: plantas N+3.00 y N+6.00 (azul = tablero escogido, naranja = nervadura principal escogida, rojo = principales sin muro, gris grueso = muros que cargan, morado = vacíos, rombos rojos = castillos C1…C8 que bajan a cimentación, decisión D1). '
              'Abajo: planta del tablero, corte y alzado de la nervadura secundaria, corte de la principal.', 42)
    N['k_chartrow'] = S.r
    add_cf(ws, 'C1:F%d' % S.r)
    ws.sheet_view.zoomScale = 70
    return ws


# ---------------------------------------------------------------- CROQUIS_DATOS + charts
def build_croquis_datos(wb, wsK):
    ws = wb.create_sheet(CD); ws.sheet_view.showGridLines = False
    ws['A1'] = 'COORDENADAS DE LAS GRÁFICAS DE LA HOJA CROQUIS'; ws['A1'].font = font(12, True, color=HDR_FG)
    ws['A2'] = ('La genera el libro solo. Cada par de columnas es un trazo CONTINUO: cuando el trazo ya terminó, las filas que sobran repiten el último punto. '
                'Las plantas (columnas CW en adelante) son datos fijos del modelo y del DXF; se separan con renglones vacíos para levantar la pluma.')
    ws['A2'].font = font(9, i=True, color=GRAY)
    ws.column_dimensions['A'].width = 46
    labels = {}
    lab_r = [LAB0]

    def label(name, formula):
        r = lab_r[0]; ws.cell(r, 1, formula.format(**N) if formula.startswith('=') else formula); labels[name] = "'%s'!$A$%d" % (CD, r) if ' ' in CD else '%s!$A$%d' % (CD, r)
        lab_r[0] += 1
    series_cols = {}
    col = [2]

    def dyn(name, xf, yf, npt=NPT, title=None):
        """Serie dinámica: xf, yf son funciones de la celda k (columna A) que devuelven fórmula."""
        cx, cy = colpair(col[0]); col[0] += 2
        ws['%s4' % cx] = name + ' x'; ws['%s4' % cy] = name + ' y'
        for i in range(npt):
            r = Y0 + i
            ws['%s%d' % (cx, r)] = ('=' + xf('$A%d' % r)).format(**N)
            ws['%s%d' % (cy, r)] = ('=' + yf('$A%d' % r)).format(**N)
        series_cols[name] = (cx, cy, Y0, Y0 + npt - 1)
    for i in range(300):
        ws['A%d' % (Y0 + i)] = i + 1
    ws['A4'] = 'k'
    # ------------------------------------------------ Ch3: planta del tablero
    W = '{k_ax}*{k_sc}'; H = '{k_ay}*{k_sc}'; ox = '{k_ox}'; oy = '{k_oy}'; sc = '{k_sc}'
    rect = lambda k, x0, y0, x1, y1: 'CHOOSE(MIN(MAX(%s,1),5),%s,%s,%s,%s,%s)' % (k, x0, x1, x1, x0, x0), lambda k, x0, y0, x1, y1: 'CHOOSE(MIN(MAX(%s,1),5),%s,%s,%s,%s,%s)' % (k, y0, y0, y1, y1, y0)
    rx, ry = rect
    dyn('Tablero', lambda k: rx(k, ox, oy, '%s+%s' % (ox, W), '%s+%s' % (oy, H)), lambda k: ry(k, ox, oy, '%s+%s' % (ox, W), '%s+%s' % (oy, H)))
    b2 = '{bp}/2*%s' % sc
    dyn('Paño de principales', lambda k: rx(k, '%s+%s' % (ox, b2), 0, '%s+%s-%s' % (ox, W, b2), 0), lambda k: ry(k, 0, '%s+%s' % (oy, b2), 0, '%s+%s-%s' % (oy, H, b2)))
    # casetones: aristas verticales en serpentina
    def serp(ncell, mcell, along_len, cross_m, cross_len, vertical, acell):
        # j = indice de linea, w = 0/1 dentro de la linea; acell = lado del casetón en esa dirección (50 cm a lo largo del muro en los volados, D10)
        def xf(k):
            j = 'MIN(INT((%s-1)/2),MAX(2*%s-1,0))' % (k, ncell)
            pos = '({bp}/2+%s+INT(%s/2)*{s}+MOD(%s,2)*%s)' % (mcell, j, j, acell)
            return 'IF(%s=0,%s,%s)' % (ncell, '%s+%s/2' % (ox, W) if vertical else '%s+%s/2' % (oy, H), ('%s+%s*%s' % (ox, pos, sc)) if vertical else ('%s+%s*%s' % (oy, pos, sc)))
        def yf(k):
            j = 'MIN(INT((%s-1)/2),MAX(2*%s-1,0))' % (k, ncell)
            w = 'IF(INT((%s-1)/2)>MAX(2*%s-1,0),1,MOD(%s-1,2))' % (k, ncell, k)
            lo = '({bp}/2+%s)' % cross_m
            hi = '(%s-{bp}/2-%s)' % (cross_len, cross_m)
            v = 'IF(MOD(%s,2)=0,IF(%s=0,%s,%s),IF(%s=0,%s,%s))' % (j, w, lo, hi, w, hi, lo)
            base = oy if vertical else ox
            return 'IF(%s=0,%s,%s+%s*%s)' % (ncell, '%s+%s/2' % (oy, H) if vertical else '%s+%s/2' % (ox, W), base, v, sc)
        return (xf, yf) if vertical else (yf, xf)
    fx, fy = serp('{n_ncx}', '{n_mx}', '{k_ax}', '{n_my}', '{k_ay}', True, '{n_acx}')
    dyn('Casetones (aristas ‖ Y)', fx, fy)
    fx, fy = serp('{n_ncy}', '{n_my}', '{k_ay}', '{n_mx}', '{k_ax}', False, '{n_acy}')
    dyn('Casetones (aristas ‖ X)', fx, fy)
    # nervadura representativa del claro corto (al centro) y del largo
    def bar(dirv, frac0, frac1, off):
        """barra horizontal (dirv='X') o vertical a lo largo del tablero, de frac0 a frac1 del claro, desplazada off unidades del centro."""
        if dirv == 'X':
            xf = lambda k: 'IF(%s<=1,%s+(%s)*%s,%s+(%s)*%s)' % (k, ox, frac0, W, ox, frac1, W)
            yf = lambda k: '%s+%s/2+%s' % (oy, H, off)
        else:
            xf = lambda k: '%s+%s/2+%s' % (ox, W, off)
            yf = lambda k: 'IF(%s<=1,%s+(%s)*%s,%s+(%s)*%s)' % (k, oy, frac0, H, oy, frac1, H)
        return xf, yf
    vol = '{k_modo}="VOLADO"'
    for tag, dsel_short in (('corto', True), ('largo', False)):
        dX = '{k_dir}="X"' if dsel_short else '{k_dir}="Y"'
        # el trazo corre en X si (claro corto es X y es la nervadura corta) o (claro corto es Y y es la larga)
        fb = 'MIN(0.5,({bp}/2+{n_Lb%s})/({n_a%s}*100))' % ('c' if dsel_short else 'l', '1' if dsel_short else '2')
        nsup = '{n_nNc}' if dsel_short else '{n_nNl}'
        for part, f0, f1, off in (('inf', '0.02', '0.98', '-1.2'), ('sup', '0.02', '0.98', '1.2'), ('bI', '0', 'IF(%s,1,%s)' % (vol, fb), '2.6'),
                                  ('bD', 'IF(%s,1,1-%s)' % (vol, fb), '1', '2.6')):
            xX, yX = bar('X', f0, f1, off if dsel_short else '(%s)+0.0' % off)
            xY, yY = bar('Y', f0, f1, off)
            hide = '' if part in ('inf', 'sup') else ('%s<2' % nsup if part == 'bI' else 'OR(%s<2,%s)' % (nsup, vol))
            fx = (lambda k, xX=xX, xY=xY, dX=dX: 'IF(%s,%s,%s)' % (dX, xX(k), xY(k)))
            if hide:
                fx = (lambda k, xX=xX, xY=xY, dX=dX, hide=hide: 'IF(%s,-50,IF(%s,%s,%s))' % (hide, dX, xX(k), xY(k)))
            dyn('n%s_%s' % (tag, part), fx, (lambda k, yX=yX, yY=yY, dX=dX: 'IF(%s,%s,%s)' % (dX, yX(k), yY(k))), 6)
    # cotas
    def cota_h(x0, x1, y):
        xs = [x0, x0, x0, '(%s+%s)/2' % (x0, x1), x1, x1, x1]
        ys = ['%s-1.5' % y, '%s+1.5' % y, y, y, y, '%s+1.5' % y, '%s-1.5' % y]
        return (lambda k: 'CHOOSE(MIN(MAX(%s,1),7),%s)' % (k, ','.join(xs))), (lambda k: 'CHOOSE(MIN(MAX(%s,1),7),%s)' % (k, ','.join(ys)))
    def cota_v(y0, y1, x):
        f, g = cota_h(y0, y1, x)
        return g, f
    dyn('Cota a_x', *cota_h(ox, '%s+%s' % (ox, W), '%s-5' % oy), 7)
    dyn('Cota a_y', *cota_v(oy, '%s+%s' % (oy, H), '%s-5' % ox), 7)
    dyn('Cota s', *cota_h('%s+({bp}/2+{n_mx})*%s' % (ox, sc), '%s+({bp}/2+{n_mx}+{s})*%s' % (ox, sc), '%s+%s+4' % (oy, H)), 7)
    label('Tablero', '="Tablero "&{k_sel}&" — "&TEXT({k_ax}/100,"0.00")&" × "&TEXT({k_ay}/100,"0.00")&" m"')
    label('Paño de principales', '="Paño de nervaduras principales / dalas (b = "&{bp}&" cm)"')
    label('Casetones (aristas ‖ Y)', '=IF({k_losa}="MACIZA","Losa maciza de "&{k_h}&" cm: sin casetones (D11)","Casetones "&{n_acx}&"×"&{n_acy}&"×"&{hc}&": "&{n_ncx}&" × "&{n_ncy}&" pzas")')
    label('Casetones (aristas ‖ X)', '=IF({k_losa}="MACIZA","(maciza)","Macizo en orillas: "&TEXT({n_mx},"0")&" cm en X, "&TEXT({n_my},"0")&" cm en Y")')
    label('ncorto_inf', '="Claro corto · inferior: "&{k_tPc}')
    label('ncorto_sup', '="Claro corto · superior: "&{k_tNc}')
    label('ncorto_bI', '=IF({n_nNc}<2,"Claro corto · bastones: no lleva","Claro corto · bastones: "&({n_nNc}-1)&" No. "&TEXT({n_vNc},"0.#")&" en cada apoyo")')
    label('ncorto_bD', '="(bastón del otro apoyo)"')
    label('nlargo_inf', '="Claro largo · inferior: "&{k_tPl}')
    label('nlargo_sup', '="Claro largo · superior: "&{k_tNl}')
    label('nlargo_bI', '=IF({n_nNl}<2,"Claro largo · bastones: no lleva","Claro largo · bastones: "&({n_nNl}-1)&" No. "&TEXT({n_vNl},"0.#")&" en cada apoyo")')
    label('nlargo_bD', '="(bastón del otro apoyo)"')
    label('Cota a_x', '="a_x = "&TEXT({k_ax}/100,"0.00")&" m"')
    label('Cota a_y', '="a_y = "&TEXT({k_ay}/100,"0.00")&" m"')
    label('Cota s', '="s = "&{s}&" cm"')
    # ------------------------------------------------ Ch4: corte nervadura secundaria (cm reales)
    hs = '{s}/2'
    # nervada: sección T con la b_w del claro corto (20 cm en volados); maciza: rectángulo de ancho s y h = h_mac (el alma ocupa todo el ancho)
    poly = ['-%s' % hs, hs, hs, '{k_bwd}/2', '{k_bwd}/2', '-{k_bwd}/2', '-{k_bwd}/2', '-%s' % hs, '-%s' % hs]
    polyy = ['{k_h}', '{k_h}', '{hc}', '{hc}', '0', '0', '{hc}', '{hc}', '{k_h}']
    dyn('Concreto', lambda k: 'CHOOSE(MIN(MAX(%s,1),9),%s)' % (k, ','.join(poly)), lambda k: 'CHOOSE(MIN(MAX(%s,1),9),%s)' % (k, ','.join(polyy)), 9)
    hidem = lambda f: 'IF({k_losa}="MACIZA",-100,%s)' % f
    dyn('Casetón izq.', lambda k: hidem(rx(k, '-%s' % hs, 0, '-{k_bwd}/2', 0)), lambda k: ry(k, 0, '0', 0, '{hc}'), 5)
    dyn('Casetón der.', lambda k: hidem(rx(k, '{k_bwd}/2', 0, hs, 0)), lambda k: ry(k, 0, '0', 0, '{hc}'), 5)
    dyn('Malla', lambda k: hidem('IF(%s<=1,-%s+1,%s-1)' % (k, hs, hs)), lambda k: '{h}-{rec}-{fim}', 2)
    dbe = dbz('{est_sec}')
    def barpos(n, var, ytop, sep):
        dbb = db(var)
        x = lambda k: 'IF({k_losa}="MACIZA",IF(%s<=1,-1,1)*%s/2,IF(%s=1,0,IF(MIN(%s,%s)=1,-1,1)*(MAX({slib},%s)/2+%s/2)))' % (k, sep, n, k, n, dbb, dbb)
        y = lambda k: 'IF(%s=0,-100,%s)' % (n, '{n_dNc}' if ytop else '{k_h}-{n_dPc}')
        return x, y
    dyn('Varillas inferiores', *barpos('{n_nPc}', '{n_vPc}', False, '{n_sPc}'), 2)
    dyn('Bastones', *barpos('{n_nNc}', '{n_vNc}', True, '{n_sNc}'), 2)
    dyn('Estribo / grapa', lambda k: rx(k, '-({bw}/2-{rec})', 0, '({bw}/2-{rec})', 0), lambda k: 'IF({est_sec}=0,-100,%s)' % ry(k, 0, '{rec}', 0, '{h}-{rec}'), 5)
    dyn('Cota b_w', *cota_h('-{k_bwd}/2', '{k_bwd}/2', '-4'), 7)
    dyn('Cota s (corte)', *cota_h('-%s' % hs, hs, '{k_h}+4'), 7)
    dyn('Cota h', *cota_v('0', '{k_h}', '%s+4' % hs), 7)
    label('Concreto', '=IF({k_losa}="MACIZA","Losa maciza h = "&{k_h}&" cm (D11), franja de "&{s}&" cm","Nervadura "&{k_bwd}&" cm + capa "&{t}&" cm (h = "&{h}&" cm)")')
    label('Casetón izq.', '=IF({k_losa}="MACIZA","(sin casetón)","Casetón de poliestireno "&{n_acx}&"×"&{n_acy}&"×"&{hc})')
    label('Casetón der.', '="(casetón)"')
    label('Malla', '=IF({k_losa}="MACIZA","(sin malla: parrilla de varillas)","Malla "&{malla})')
    label('Varillas inferiores', '="Inferior claro corto: "&{k_tPc}')
    label('Bastones', '="Superior claro corto: "&{k_tNc}')
    label('Estribo / grapa', '=IF({est_sec}=0,"Sin estribo: el cortante lo toma el concreto (5.5.3.8.1); si no alcanza, se maciza","Estribo No. "&TEXT({est_sec},"0.#"))')
    label('Cota b_w', '=IF({k_losa}="MACIZA","franja de "&{s}&" cm","b_w = "&{k_bwd}&" cm")')
    label('Cota s (corte)', '="s = "&{s}&" cm"')
    label('Cota h', '="h = "&{k_h}&" cm"')
    # ------------------------------------------------ Ch5: alzado de la nervadura del claro corto (x 0..100, alto exagerado)
    Hy = '12'
    dyn('Nervadura (alzado)', lambda k: rx(k, '0', 0, '100', 0), lambda k: ry(k, 0, '0', 0, Hy), 5)
    dyn('Apoyo izquierdo', lambda k: rx(k, '-3', 0, '0', 0), lambda k: ry(k, 0, '-10', 0, Hy), 5)
    dyn('Apoyo derecho', lambda k: rx(k, 'IF(%s,-3,100)' % vol, 0, 'IF(%s,0,103)' % vol, 0), lambda k: ry(k, 0, '-10', 0, Hy), 5)
    dyn('Inferior (alzado)', lambda k: 'CHOOSE(MIN(MAX(%s,1),4),1.5,1.5,98.5,98.5)' % k, lambda k: 'CHOOSE(MIN(MAX(%s,1),4),4,1.5,1.5,4)' % k, 4)
    fbz = 'MIN(50,({bp}/2+{n_Lbc})/({n_a1}*100)*100)'
    dyn('Superior corrido (alzado)', lambda k: 'CHOOSE(MIN(MAX(%s,1),4),-1.5,-1.5,101.5,101.5)' % k, lambda k: 'CHOOSE(MIN(MAX(%s,1),4),6,9.5,9.5,IF(%s,9.5,6))' % (k, vol), 4)
    dyn('Bastón izq. (alzado)', lambda k: 'IF({n_nNc}<2,-50,CHOOSE(MIN(MAX(%s,1),3),-1.5,-1.5,IF(%s,98.5,%s)))' % (k, vol, fbz), lambda k: 'CHOOSE(MIN(MAX(%s,1),3),6,10.7,10.7)' % k, 3)
    dyn('Bastón der. (alzado)', lambda k: 'IF(OR({n_nNc}<2,%s),-50,CHOOSE(MIN(MAX(%s,1),3),101.5,101.5,100-%s))' % (vol, k, fbz), lambda k: 'CHOOSE(MIN(MAX(%s,1),3),6,10.7,10.7)' % k, 3)
    dyn('Cota a1', *cota_h('0', '100', '-14'), 7)
    label('Nervadura (alzado)', '=IF({k_losa}="MACIZA","Franja de 1 m del claro corto","Nervadura del claro corto")&" ("&IF({k_modo}="VOLADO","VOLADO","apoyada en sus dos extremos")&") — alto exagerado"')
    label('Apoyo izquierdo', '=IF({k_modo}="VOLADO","Apoyo del volado: nervadura principal o muro (punta libre)","Apoyos: nervadura principal o muro")')
    label('Apoyo derecho', '=IF({k_modo}="VOLADO","(volado: sin apoyo en la punta)","(apoyo)")')
    label('Inferior (alzado)', '="Inferior corrido: "&{k_tPc}')
    label('Superior corrido (alzado)', '=IF({k_losa}="MACIZA","Superior: "&{k_tNc},"Superior corrido: 1 No. "&TEXT({n_vNc},"0.#")&" (6.7.6.1.3)")')
    label('Bastón izq. (alzado)', '=IF({k_losa}="MACIZA","(bastones de la maciza: ver arriba)",IF({n_nNc}<2,"Sin bastones: basta la corrida superior","Bastones: "&({n_nNc}-1)&" No. "&TEXT({n_vNc},"0.#")&", "&TEXT({n_Lbc}/100,"0.00")&" m desde el paño"))')
    label('Bastón der. (alzado)', '="(bastón del otro apoyo)"')
    label('Cota a1', '="a1 = "&TEXT({n_a1},"0.00")&" m"')
    # ------------------------------------------------ Ch6: corte de la nervadura principal
    hb = '{k_b}/2'; hh = '{k_hh}'
    poly = ['-35', '35', '35', hb, hb, '-' + hb, '-' + hb, '-35', '-35']
    polyy = [hh, hh, '%s-{t}' % hh, '%s-{t}' % hh, '0', '0', '%s-{t}' % hh, '%s-{t}' % hh, hh]
    dyn('Sección principal', lambda k: 'CHOOSE(MIN(MAX(%s,1),9),%s)' % (k, ','.join(poly)), lambda k: 'CHOOSE(MIN(MAX(%s,1),9),%s)' % (k, ','.join(polyy)), 9)
    dyn('Casetón izq. (principal)', lambda k: rx(k, '-35', 0, '-%s' % hb, 0), lambda k: ry(k, 0, '%s-{t}-{hc}' % hh, 0, '%s-{t}' % hh), 5)
    dyn('Casetón der. (principal)', lambda k: rx(k, hb, 0, '35', 0), lambda k: ry(k, 0, '%s-{t}-{hc}' % hh, 0, '%s-{t}' % hh), 5)
    dbn = db('{est_np}')
    dyn('Estribo principal', lambda k: rx(k, '-(%s-{rec_np}-%s/2)' % (hb, dbn), 0, '(%s-{rec_np}-%s/2)' % (hb, dbn), 0), lambda k: ry(k, 0, '{rec_np}+%s/2' % dbn, 0, '%s-{rec_np}-%s/2' % (hh, dbn)), 5)

    def barsp(n, var, top):
        dbb = db('MAX(%s,3)' % var)
        span = '({k_b}-2*{rec_np}-2*%s-%s)' % (dbn, dbb)
        x = lambda k: 'IF(%s<=1,0,-%s/2+(MIN(%s,%s)-1)*%s/(%s-1))' % (n, span, k, n, span, n)
        y = lambda k: 'IF(%s=0,-100,%s)' % (n, ('%s-{rec_np}-%s-%s/2' % (hh, dbn, dbb)) if top else ('{rec_np}+%s+%s/2' % (dbn, dbb)))
        return x, y
    dyn('Inferiores principal', *barsp('{k_nP}', '{k_vP}', False), 8)
    dyn('Superiores principal', *barsp('{k_nN}', '{k_vN}', True), 8)
    dyn('Cota b principal', *cota_h('-' + hb, hb, '-5'), 7)
    dyn('Cota h principal', *cota_v('0', hh, '38'), 7)
    label('Sección principal', '="Nervadura principal "&{k_np}&" ("&{k_npeje}&", L = "&TEXT({k_npL},"0.00")&" m, "&{k_npcond}&")"')
    label('Casetón izq. (principal)', '="Casetones de los tableros vecinos"')
    label('Casetón der. (principal)', '="(casetón)"')
    label('Estribo principal', '="Estribos cerrados No. "&TEXT({est_np},"0.#")&" a 135°: @ "&TEXT({k_sext},"0.0")&" cm en 2h y @ "&TEXT({k_se},"0.0")&" cm resto"')
    label('Inferiores principal', '=IF({k_nP}=0,"Inferior: la sección no alcanza — sube h","Inferior: "&{k_nP}&" var. No. "&TEXT({k_vP},"0.#"))')
    label('Superiores principal', '=IF({k_nN}=0,"Superior: la sección no alcanza — sube h","Superior: "&{k_nN}&" var. No. "&TEXT({k_vN},"0.#"))')
    label('Cota b principal', '="b = "&{k_b}&" cm"')
    label('Cota h principal', '="h = "&{k_hh}&" cm"')
    # ------------------------------------------------ plantas (datos fijos)
    c0 = 101  # columna CW (deja libres B..CV para las series dinámicas)
    static = {}

    def put(name, pts_lists, marker=False):
        nonlocal c0
        cx, cy = colpair(c0); c0 += 2
        ws['%s4' % cx] = name + ' x'; ws['%s4' % cy] = name + ' y'
        r = Y0
        for pts in pts_lists:
            for (x, y) in pts:
                ws['%s%d' % (cx, r)] = round(x, 3); ws['%s%d' % (cy, r)] = round(y, 3); r += 1
            if not marker:
                r += 1
        static[name] = (cx, cy, Y0, max(r - 1, Y0))
    for lvl in LV:
        tabs = [t for t in TABS if t['level'] == lvl]
        put('tab' + lvl, [t['pts'] + [t['pts'][0]] for t in tabs])
        put('mur' + lvl, D['plantas'][lvl]['muros'])
        put('np' + lvl, [[(t['x0'], t['y0']), (t['x1'], t['y1'])] for t in NPS if t['level'] == lvl])
        vv = []
        for v in D['vacios'][lvl]:
            cs = v['coords']; xs = [p[0] for p in cs]; ys = [p[1] for p in cs]
            vv.append(cs)
            vv.append([(min(xs), min(ys)), (max(xs), max(ys))]); vv.append([(min(xs), max(ys)), (max(xs), min(ys))])
        put('vac' + lvl, vv)
        put('cas' + lvl, [[p] for p in D['plantas'][lvl]['castillos']], marker=True)
    # tablas de búsqueda: polígono de cada tablero y línea de cada NP
    tb_c = c0; ws.cell(3, tb_c, 'Polígonos de tableros (búsqueda)')
    for i, t in enumerate(TABS):
        r = Y0 + i
        ws.cell(r, tb_c, key(t)); ws.cell(r, tb_c + 1, t['level'])
        pts = t['pts'] + [t['pts'][0]]
        pts = pts + [pts[-1]] * (18 - len(pts))
        for j, (x, y) in enumerate(pts[:18]):
            ws.cell(r, tb_c + 2 + 2 * j, round(x, 3)); ws.cell(r, tb_c + 3 + 2 * j, round(y, 3))
    TBK = '$%s$%d:$%s$%d' % (CL(tb_c), Y0, CL(tb_c), Y0 + len(TABS) - 1)
    np_c = tb_c + 2 + 36 + 1
    ws.cell(3, np_c, 'Nervaduras principales (búsqueda)')
    for i, t in enumerate(NPS):
        r = Y0 + i
        for j, v in enumerate([t['np'], t['level'], t['x0'], t['y0'], t['x1'], t['y1']]):
            ws.cell(r, np_c + j, v)
    NPK = '$%s$%d:$%s$%d' % (CL(np_c), Y0, CL(np_c), Y0 + len(NPS) - 1)
    hl_c = np_c + 7
    for lvl_i, lvl in enumerate(LV):
        cx, cy = colpair(hl_c + 4 * lvl_i)
        ws['%s4' % cx] = 'sel ' + lvl
        for j in range(18):
            r = Y0 + j
            m_ = 'MATCH({k_sel},%s,0)' % TBK
            ws['%s%d' % (cx, r)] = ('=IF(INDEX($%s$%d:$%s$%d,%s)="%s",INDEX($%s$%d:$%s$%d,%s),-100)' % (CL(tb_c + 1), Y0, CL(tb_c + 1), Y0 + len(TABS) - 1, m_, lvl,
                                     CL(tb_c + 2 + 2 * j), Y0, CL(tb_c + 2 + 2 * j), Y0 + len(TABS) - 1, m_)).format(**N)
            ws['%s%d' % (cy, r)] = ('=IF(INDEX($%s$%d:$%s$%d,%s)="%s",INDEX($%s$%d:$%s$%d,%s),-100)' % (CL(tb_c + 1), Y0, CL(tb_c + 1), Y0 + len(TABS) - 1, m_, lvl,
                                     CL(tb_c + 3 + 2 * j), Y0, CL(tb_c + 3 + 2 * j), Y0 + len(TABS) - 1, m_)).format(**N)
        static['sel' + lvl] = (cx, cy, Y0, Y0 + 17)
        cx2, cy2 = colpair(hl_c + 4 * lvl_i + 2)
        ws['%s4' % cx2] = 'np ' + lvl
        m_ = 'MATCH({k_np},%s,0)' % NPK
        for j in range(2):
            r = Y0 + j
            ws['%s%d' % (cx2, r)] = ('=IF(INDEX($%s$%d:$%s$%d,%s)="%s",INDEX($%s$%d:$%s$%d,%s),-100)' % (CL(np_c + 1), Y0, CL(np_c + 1), Y0 + len(NPS) - 1, m_, lvl,
                                      CL(np_c + 2 + 2 * j), Y0, CL(np_c + 2 + 2 * j), Y0 + len(NPS) - 1, m_)).format(**N)
            ws['%s%d' % (cy2, r)] = ('=IF(INDEX($%s$%d:$%s$%d,%s)="%s",INDEX($%s$%d:$%s$%d,%s),-100)' % (CL(np_c + 1), Y0, CL(np_c + 1), Y0 + len(NPS) - 1, m_, lvl,
                                      CL(np_c + 3 + 2 * j), Y0, CL(np_c + 3 + 2 * j), Y0 + len(NPS) - 1, m_)).format(**N)
        static['npsel' + lvl] = (cx2, cy2, Y0, Y0 + 1)
    # etiquetas de tableros y castillos propuestos (una celda por punto)
    lb_c = hl_c + 9
    for i, t in enumerate(TABS):
        xs = [p[0] for p in t['pts']]; ys = [p[1] for p in t['pts']]
        from shapely.geometry import Polygon as _P
        c = _P(t['pts']).representative_point()
        ws.cell(Y0 + i, lb_c, t['name']); ws.cell(Y0 + i, lb_c + 1, round(c.x, 3)); ws.cell(Y0 + i, lb_c + 2, round(c.y, 3)); ws.cell(Y0 + i, lb_c + 3, t['level'])
    for i, c in enumerate(D['castillos_prop']):
        r = Y0 + 40 + i
        ws.cell(r, lb_c, c['id']); ws.cell(r, lb_c + 1, c['x']); ws.cell(r, lb_c + 2, c['y']); ws.cell(r, lb_c + 3, c['level'])
        ws.cell(r, lb_c + 4, '=IF(%s="SÍ",%s&" (D1)",%s&" RECHAZADO")' % (PROPROW[c['id']], '"%s"' % c['id'], '"%s"' % c['id']))
    for name, f in [('k', None)]:
        pass
    ws['A%d' % (LAB0 - 1)] = 'TEXTOS DE LAS LEYENDAS'; ws['A%d' % (LAB0 - 1)].font = font(10, True, color=HDR_FG)
    for r in range(LAB0, lab_r[0]):
        ws['A%d' % r].font = font(9)
    # ================================================== gráficas
    from openpyxl.chart.layout import Layout, ManualLayout

    def mk_chart(title, xmin, xmax, ymin, ymax, w_cm, h_cm, legend=True, hfrac=None, ytop=0.11):
        ch = ScatterChart(); ch.title = title; ch.style = 13; ch.scatterStyle = 'lineMarker'
        ch.x_axis.scaling.min = xmin; ch.x_axis.scaling.max = xmax; ch.y_axis.scaling.min = ymin; ch.y_axis.scaling.max = ymax
        ch.x_axis.delete = True; ch.y_axis.delete = True
        ch.x_axis.majorGridlines = None; ch.y_axis.majorGridlines = None
        ch.width = w_cm; ch.height = h_cm
        if legend:
            ch.legend.position = 'b'
        else:
            ch.legend = None
        ch.display_blanks = 'gap'
        # plot area con la misma escala en X y en Y (que los cuadrados salgan cuadrados)
        if hfrac:
            wfrac = (xmax - xmin) / (ymax - ymin) * hfrac * h_cm / w_cm
            if wfrac > 0.94:
                hfrac = hfrac * 0.94 / wfrac; wfrac = 0.94
            ch.plot_area.layout = Layout(manualLayout=ManualLayout(layoutTarget='inner', xMode='edge', yMode='edge', x=(1 - wfrac) / 2, y=ytop, w=wfrac, h=hfrac))
        return ch

    def add(ch, ref_cols, title_ref=None, title_txt=None, color='1F4E79', w=1.5, dash='solid', marker=None, msize=7, lbl_idx=None, sheet=ws, noleg=False):
        cx, cy, r0, r1 = ref_cols
        xr = Reference(sheet, min_col=CI(cx), min_row=r0, max_row=r1)
        yr = Reference(sheet, min_col=CI(cy), min_row=r0, max_row=r1)
        s = Series(yr, xr)
        if title_ref:
            s.tx = SeriesLabel(strRef=StrRef(title_ref))
        elif title_txt:
            s.tx = SeriesLabel(v=title_txt)
        s.smooth = False
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = int(w * 12700)
        s.graphicalProperties.line.prstDash = dash
        if marker:
            s.marker = Marker(symbol=marker, size=msize)
            s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
            s.marker.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line = LineProperties(noFill=True)
        else:
            s.marker = Marker(symbol='none')
        if lbl_idx is not None:
            s.dLbls = DataLabelList(dLbl=[DataLabel(idx=lbl_idx, showSerName=True, showVal=False, showCatName=False, showLegendKey=False, showPercent=False, dLblPos='t')])
            s.dLbls.showVal = False; s.dLbls.showSerName = False; s.dLbls.showLegendKey = False; s.dLbls.showCatName = False
        if noleg and ch.legend is not None:
            ch.legend.legendEntry.append(LegendEntry(idx=len(ch.series), delete=True))
        ch.series.append(s)
        return s
    anchor_row = N['k_chartrow']
    # plantas
    for i, lvl in enumerate(LV):
        ch = mk_chart('PLANTA %s — APOYOS, PRINCIPALES Y VACÍOS' % lvl, -2.6, 9.6, -0.6, 25.0, 11.5, 22.5, legend=False, hfrac=0.85)
        add(ch, static['tab' + lvl], title_txt='Tableros', color='8EA9DB', w=0.75)
        add(ch, static['mur' + lvl], title_txt='Muros que cargan', color='7F7F7F', w=3.5)
        add(ch, static['vac' + lvl], title_txt='Vacíos', color='7030A0', w=1.0, dash='dash')
        add(ch, static['np' + lvl], title_txt='Nervaduras principales sin muro', color='C00000', w=2.0)
        add(ch, static['cas' + lvl], title_txt='Castillos', color='000000', marker='square', msize=4)
        add(ch, static['sel' + lvl], title_txt='Tablero escogido', color='0070C0', w=3.0)
        add(ch, static['npsel' + lvl], title_txt='Principal escogida', color='ED7D31', w=4.5)
        for j, t in enumerate(TABS):
            if t['level'] != lvl:
                continue
            r = Y0 + j
            add(ch, (CL(lb_c + 1), CL(lb_c + 2), r, r), title_ref="%s!$%s$%d" % (CD, CL(lb_c), r), color='1F3864', marker='none', lbl_idx=0)
        if lvl == 'N+3.00':
            for j, c in enumerate(D['castillos_prop']):
                r = Y0 + 40 + j
                add(ch, (CL(lb_c + 1), CL(lb_c + 2), r, r), title_ref="%s!$%s$%d" % (CD, CL(lb_c + 4), r), color='C00000', marker='diamond', msize=9, lbl_idx=0)
        wsK.add_chart(ch, 'B%d' % anchor_row if i == 0 else 'E%d' % anchor_row)
    r2 = anchor_row + 48
    def beside_tab(ch, row):
        # a la derecha de la planta del tablero (17 cm), que es más ancha que B:D: se recorre desde E lo que falte
        px = sum(int(wsK.column_dimensions[k].width * 7 + 5) for k in 'BCD')
        off = max(0.0, 17.3 - px * 2.54 / 96)
        return OneCellAnchor(_from=AnchorMarker(col=4, colOff=cm_to_EMU(off), row=row - 1, rowOff=0),
                             ext=XDRPositiveSize2D(cm_to_EMU(ch.width), cm_to_EMU(ch.height)))
    ch = mk_chart('PLANTA DEL TABLERO — CASETONES Y NERVADURAS', 0, 100, 0, 100, 17, 23, hfrac=0.56, ytop=0.08)
    add(ch, series_cols['Tablero'], labels['Tablero'], color='1F3864', w=2.2)
    add(ch, series_cols['Paño de principales'], labels['Paño de principales'], color='C00000', w=1.0, dash='dash')
    add(ch, series_cols['Casetones (aristas ‖ Y)'], labels['Casetones (aristas ‖ Y)'], color='A6A6A6', w=0.75)
    add(ch, series_cols['Casetones (aristas ‖ X)'], labels['Casetones (aristas ‖ X)'], color='A6A6A6', w=0.75)
    add(ch, series_cols['ncorto_inf'], labels['ncorto_inf'], color='0070C0', w=2.5)
    add(ch, series_cols['ncorto_sup'], labels['ncorto_sup'], color='7F4F00', w=1.5)
    add(ch, series_cols['ncorto_bI'], labels['ncorto_bI'], color='C00000', w=2.5)
    add(ch, series_cols['ncorto_bD'], labels['ncorto_bD'], noleg=True, color='C00000', w=2.5)
    add(ch, series_cols['nlargo_inf'], labels['nlargo_inf'], color='00B050', w=2.0)
    add(ch, series_cols['nlargo_sup'], labels['nlargo_sup'], color='375623', w=1.5)
    add(ch, series_cols['nlargo_bI'], labels['nlargo_bI'], color='ED7D31', w=2.0)
    add(ch, series_cols['nlargo_bD'], labels['nlargo_bD'], noleg=True, color='ED7D31', w=2.0)
    add(ch, series_cols['Cota a_x'], labels['Cota a_x'], color='7030A0', w=0.75, lbl_idx=3)
    add(ch, series_cols['Cota a_y'], labels['Cota a_y'], color='7030A0', w=0.75, lbl_idx=3)
    add(ch, series_cols['Cota s'], labels['Cota s'], color='7030A0', w=0.75, lbl_idx=3)
    wsK.add_chart(ch, 'B%d' % r2)
    ch = mk_chart('CORTE DE LA NERVADURA SECUNDARIA (cm)', -42, 42, -10, 34, 15, 13.5, hfrac=0.42, ytop=0.12)
    add(ch, series_cols['Concreto'], labels['Concreto'], color='1F3864', w=2.2)
    add(ch, series_cols['Casetón izq.'], labels['Casetón izq.'], color='A6A6A6', w=1.0, dash='dash')
    add(ch, series_cols['Casetón der.'], labels['Casetón der.'], noleg=True, color='A6A6A6', w=1.0, dash='dash')
    add(ch, series_cols['Malla'], labels['Malla'], color='00B050', w=1.0, dash='sysDash')
    add(ch, series_cols['Estribo / grapa'], labels['Estribo / grapa'], color='7030A0', w=1.0)
    add(ch, series_cols['Varillas inferiores'], labels['Varillas inferiores'], color='0070C0', marker='circle', msize=8)
    add(ch, series_cols['Bastones'], labels['Bastones'], color='C00000', marker='circle', msize=8)
    add(ch, series_cols['Cota b_w'], labels['Cota b_w'], color='7F7F7F', w=0.75, lbl_idx=3)
    add(ch, series_cols['Cota s (corte)'], labels['Cota s (corte)'], color='7F7F7F', w=0.75, lbl_idx=3)
    add(ch, series_cols['Cota h'], labels['Cota h'], color='7F7F7F', w=0.75, lbl_idx=3)
    wsK.add_chart(ch, beside_tab(ch, r2))
    ch = mk_chart('ALZADO DE LA NERVADURA DEL CLARO CORTO', -8, 108, -20, 20, 15, 11, hfrac=0.36, ytop=0.13)
    add(ch, series_cols['Nervadura (alzado)'], labels['Nervadura (alzado)'], color='1F3864', w=2.0)
    add(ch, series_cols['Apoyo izquierdo'], labels['Apoyo izquierdo'], color='7F7F7F', w=2.5)
    add(ch, series_cols['Apoyo derecho'], labels['Apoyo derecho'], noleg=True, color='7F7F7F', w=2.5)
    add(ch, series_cols['Inferior (alzado)'], labels['Inferior (alzado)'], color='0070C0', w=2.5)
    add(ch, series_cols['Superior corrido (alzado)'], labels['Superior corrido (alzado)'], color='7F4F00', w=2.0)
    add(ch, series_cols['Bastón izq. (alzado)'], labels['Bastón izq. (alzado)'], color='C00000', w=2.5)
    add(ch, series_cols['Bastón der. (alzado)'], labels['Bastón der. (alzado)'], noleg=True, color='C00000', w=2.5)
    add(ch, series_cols['Cota a1'], labels['Cota a1'], color='7030A0', w=0.75, lbl_idx=3)
    wsK.add_chart(ch, beside_tab(ch, r2 + 30))
    ch = mk_chart('CORTE DE LA NERVADURA PRINCIPAL ESCOGIDA (cm)', -42, 42, -10, 62, 15, 15, hfrac=0.5, ytop=0.1)
    add(ch, series_cols['Sección principal'], labels['Sección principal'], color='1F3864', w=2.2)
    add(ch, series_cols['Casetón izq. (principal)'], labels['Casetón izq. (principal)'], color='A6A6A6', w=1.0, dash='dash')
    add(ch, series_cols['Casetón der. (principal)'], labels['Casetón der. (principal)'], noleg=True, color='A6A6A6', w=1.0, dash='dash')
    add(ch, series_cols['Estribo principal'], labels['Estribo principal'], color='7030A0', w=1.25)
    add(ch, series_cols['Inferiores principal'], labels['Inferiores principal'], color='0070C0', marker='circle', msize=8)
    add(ch, series_cols['Superiores principal'], labels['Superiores principal'], color='C00000', marker='circle', msize=8)
    add(ch, series_cols['Cota b principal'], labels['Cota b principal'], color='7F7F7F', w=0.75, lbl_idx=3)
    add(ch, series_cols['Cota h principal'], labels['Cota h principal'], color='7F7F7F', w=0.75, lbl_idx=3)
    wsK.add_chart(ch, 'B%d' % (r2 + 50))
    return ws


# ---------------------------------------------------------------- CUANTIFICACION
def sheet_cuant(wb):
    ws = wb.create_sheet('CUANTIFICACION')
    S = Narr(ws, N, {'A': 1.78, 'B': 18, 'C': 9, 'D': 10, 'E': 11, 'F': 11, 'G': 12, 'H': 12, 'I': 12, 'J': 40}, last='J')
    S.title('CUANTIFICACIÓN — CONCRETO, CASETONES, MALLA Y ACERO DE NERVADURAS')
    S.subtitle('Cantidades aproximadas para presupuesto. Los casetones se cuentan con el rectángulo envolvente y se escalan por el área real del tablero; el acero de nervaduras secundarias '
               'supone una nervadura por cada hilera de casetones más las de orilla, con el armado de TABLEROS. Volados (D10): casetón 50×60×19 y nervaduras de 20 cm. Losas macizas de azotea (D11): concreto h·área, sin casetón ni malla, y acero «No. x @ y» por m².', 42)
    S.gap()
    S.section('1. DATOS DE CUANTIFICACIÓN')
    S.header(('Concepto', 'Valor', 'Unidad', 'Cómo se usa', 'Referencia'))
    ws.column_dimensions['E'].width = 11
    for k, (c, v, u, how, nm) in enumerate([('Desperdicio de concreto', 1.05, '—', 'Multiplica el volumen.', 'q_dc'), ('Desperdicio de acero', 1.07, '—', 'Multiplica el peso.', 'q_da'),
                                            ('Traslape / anclaje por extremo de varilla', 0.30, 'm', 'Se suma a cada varilla.', 'q_tr')]):
        r = S.r
        cell(ws, r, 2, c, al='left'); cell(ws, r, 3, v, 'input', '0.00'); cell(ws, r, 4, u)
        ws.merge_cells('E%d:I%d' % (r, r)); cell(ws, r, 5, how, al='left'); cell(ws, r, 10, 'Práctica de obra', 'ref', al='left')
        N[nm] = 'CUANTIFICACION!$C$%d' % r; S.r += 1
    S.gap()
    S.section('2. POR TABLERO')
    r = S.r
    heads = ['Tablero', 'Área (m²)', 'Casetones (pzas)', 'Concreto (m³)', 'Malla (m²)', 'Acero ‖ corto (kg)', 'Acero ‖ largo (kg)', 'Acero total (kg)', 'Nota']
    for k, t in enumerate(heads):
        hdr_cell(ws, r, 2 + k, t)
    S.r += 1
    first = S.r
    for t in TABS:
        r = S.r; tr = TROW[key(t)]
        T_ = lambda c: 'TABLEROS!$%s$%d' % (TC[c], tr)
        cell(ws, r, 2, key(t), al='left')
        cell(ws, r, 3, t['area'], 'data', '0.00')
        cell(ws, r, 4, '=IF(%s="HUECO",0,ROUND(%s*%s*C%d/(%s*%s),0))' % (T_('modo'), T_('ncx'), T_('ncy'), r, T_('Lx'), T_('Ly')), fmt='0')
        # concreto: maciza h·área; nervada h·área − casetones (en el volado el casetón es (s − b_w del volado) × a_c × h_c)
        cell(ws, r, 5, ('=IF(%s="HUECO",0,IF(%s="MACIZA",C%d*{h_mac}/100,C%d*{h}/100-D%d*IF(%s="VOLADO",{ac_vol},{ac})/100*{ac}/100*({hc}/100)))' % (T_('modo'), T_('losa'), r, r, r, T_('modo'))).format(**N), fmt='0.000')
        cell(ws, r, 6, '=IF(OR(%s="HUECO",%s="MACIZA"),0,C%d)' % (T_('modo'), T_('losa'), r), fmt='0.00')
        # nervaduras ‖ claro corto: tantas como hileras de casetones en la otra dirección + 1; largo = a1 + 2 traslapes
        ncross_c = 'IF(%s="X",%s,%s)' % (T_('dirc'), T_('ncy'), T_('ncx'))
        ncross_l = 'IF(%s="X",%s,%s)' % (T_('dirc'), T_('ncx'), T_('ncy'))
        kg = lambda v: 'INDEX(%s,MATCH(%s,%s,0))' % (KGCOL, v, BARCOL)
        # inferiores corridas; arriba 1 corrida + (n − 1) bastones en cada apoyo (en volado: todas hasta el respaldo)
        topc = 'IF(%s="VOLADO",%s*(%s+%s/100+{q_tr})*%s,(%s+2*{q_tr})*%s+MAX(%s-1,0)*2*(%s/100+{bp}/200+{q_tr})*%s)' % (T_('modo'), T_('nNc'), T_('a1'), T_('Lbc'), kg(T_('vNc')), T_('a1'), kg(T_('vNc')), T_('nNc'), T_('Lbc'), kg(T_('vNc')))
        topl = '((%s+2*{q_tr})*%s+MAX(%s-1,0)*2*(%s/100+{bp}/200+{q_tr})*%s)' % (T_('a2'), kg(T_('vNl')), T_('nNl'), T_('Lbl'), kg(T_('vNl')))
        corto = '(%s+1)*(%s*(%s+2*{q_tr})*%s+%s)' % (ncross_c, T_('nPc'), T_('a1'), kg(T_('vPc')), topc)
        largo = '(%s+1)*(%s*(%s+2*{q_tr})*%s+%s)' % (ncross_l, T_('nPl'), T_('a2'), kg(T_('vPl')), topl)
        # maciza: n = varillas por metro (100/separación); acero por m² × área: inferiores corridas + bastones en los dos apoyos
        mc = '(%s*(%s+2*{q_tr})*%s+%s*2*(%s/100+{bp}/200+{q_tr})*%s)/%s*C%d' % (T_('nPc'), T_('a1'), kg(T_('vPc')), T_('nNc'), T_('Lbc'), kg(T_('vNc')), T_('a1'), r)
        ml = '(%s*(%s+2*{q_tr})*%s+%s*2*(%s/100+{bp}/200+{q_tr})*%s)/%s*C%d' % (T_('nPl'), T_('a2'), kg(T_('vPl')), T_('nNl'), T_('Lbl'), kg(T_('vNl')), T_('a2'), r)
        cell(ws, r, 7, ('=IF(%s="HUECO",0,IF(%s="MACIZA",%s,%s))' % (T_('modo'), T_('losa'), mc, corto)).format(**N), fmt='0.0')
        cell(ws, r, 8, ('=IF(%s="HUECO",0,IF(%s="MACIZA",%s,%s))' % (T_('modo'), T_('losa'), ml, largo)).format(**N), fmt='0.0')
        cell(ws, r, 9, '=G%d+H%d' % (r, r), 'key', '0.0')
        rep_ = ' · rectángulo representativo' if t['rep_nota'] else ''
        cell(ws, r, 10, '=IF(%s="VOLADO","Volado: casetón 50×60×19, nervaduras de "&{bw_vol}&" cm (D10)%s",IF(%s="MACIZA","Losa maciza h = "&{h_mac}&" cm (D11)%s","%s"))'.format(**N) % (T_('modo'), rep_, T_('losa'), rep_, 'Rectángulo representativo' if t['rep_nota'] else '—'), al='left', sz=8)
        S.r += 1
    last = S.r - 1
    r = S.r
    cell(ws, r, 2, 'TOTAL', bold=True, al='left'); cell(ws, r, 3, '=SUM(C%d:C%d)' % (first, last), 'key', '0.00'); cell(ws, r, 4, '=SUM(D%d:D%d)' % (first, last), 'key', '0')
    cell(ws, r, 5, '=SUM(E%d:E%d)' % (first, last), 'key', '0.00'); cell(ws, r, 6, '=SUM(F%d:F%d)' % (first, last), 'key', '0.0')
    cell(ws, r, 7, '=SUM(G%d:G%d)' % (first, last), 'key', '0'); cell(ws, r, 8, '=SUM(H%d:H%d)' % (first, last), 'key', '0'); cell(ws, r, 9, '=SUM(I%d:I%d)' % (first, last), 'key', '0')
    N['q_area'] = 'CUANTIFICACION!$C$%d' % r; N['q_cas'] = 'CUANTIFICACION!$D$%d' % r; N['q_conc'] = 'CUANTIFICACION!$E$%d' % r
    N['q_malla'] = 'CUANTIFICACION!$F$%d' % r; N['q_acs'] = 'CUANTIFICACION!$I$%d' % r
    S.r += 2
    S.section('3. NERVADURAS PRINCIPALES')
    r = S.r
    for k, t in enumerate(['Tramo', 'L (m)', 'b × h (cm)', 'Concreto extra (m³)', 'Long. (kg)', 'Estribos (kg)', 'Total (kg)', '', 'Nota']):
        hdr_cell(ws, r, 2 + k, t)
    S.r += 1
    f2 = S.r
    for t in NPS:
        r = S.r; pr = PROW[t['np']]
        P_ = lambda c: 'PRINCIPALES!$%s$%d' % (PC[c], pr)
        kg = lambda v: 'INDEX(%s,MATCH(%s,%s,0))' % (KGCOL, v, BARCOL)
        cell(ws, r, 2, t['np']); cell(ws, r, 3, '=' + P_('L'), fmt='0.00'); cell(ws, r, 4, '=%s&" × "&%s' % (P_('bq'), P_('hq')))
        cell(ws, r, 5, ('=%s*(%s/100*%s/100-%s/100*{Vm2})' % (P_('L'), P_('bq'), P_('hq'), P_('bq'))).format(**N), fmt='0.000')
        cell(ws, r, 6, ('=(%s*%s+%s*%s)*(%s+2*{q_tr})' % (P_('nP'), kg('MAX(%s,2)' % P_('vP')), P_('nN'), kg('MAX(%s,2)' % P_('vN')), P_('L'))).format(**N), fmt='0.0')
        nst = '(2*ROUNDUP(2*%s/%s,0)+ROUNDUP(MAX(0,%s*100-4*%s)/%s,0)+1)' % (P_('hq'), P_('sext'), P_('L'), P_('hq'), P_('se'))
        cell(ws, r, 7, ('=IF(AND(%s>0,%s>0),%s*(2*((%s-2*{rec_np})+(%s-2*{rec_np}))+16)/100*%s,0)' % (P_('se'), P_('sext'), nst, P_('bq'), P_('hq'), kg('{est_np}'))).format(**N), fmt='0.0')
        cell(ws, r, 8, '=F%d+G%d' % (r, r), 'key', '0.0')
        ws.merge_cells('I%d:J%d' % (r, r)); cell(ws, r, 9, '=' + P_('estado'), al='left', sz=8)
        S.r += 1
    l2 = S.r - 1
    r = S.r
    cell(ws, r, 2, 'TOTAL', bold=True, al='left'); cell(ws, r, 5, '=SUM(E%d:E%d)' % (f2, l2), 'key', '0.00'); cell(ws, r, 8, '=SUM(H%d:H%d)' % (f2, l2), 'key', '0')
    N['q_cnp'] = 'CUANTIFICACION!$E$%d' % r; N['q_anp'] = 'CUANTIFICACION!$H$%d' % r
    S.r += 2
    S.section('4. TOTALES A PEDIR')
    S.header(('Concepto', 'Cantidad', 'Unidad', 'Cómo se calcula', 'Referencia'))
    volc = 'SUMIF(J%d:J%d,"Volado*",D%d:D%d)' % (first, last, first, last)
    for qi, (c, f, u, how) in enumerate([('Concreto de losas (con desperdicio)', '=({q_conc}+{q_cnp})*{q_dc}', 'm³', 'Tableros (nervados y macizos) + exceso de principales, × desperdicio.'),
                         ('Casetones de poliestireno 60×60×19 (LR-25, decisión D3)', '={q_cas}-' + volc, 'pzas', 'Suma por tablero, sin los volados.'),
                         ('Casetones de poliestireno 50×60×19 (volados, decisión D10)', '=' + volc, 'pzas', 'Suma de los tableros en volado.'),
                         ('Malla electrosoldada', '={q_malla}*1.10', 'm²', 'Área de losa nervada + 10% de traslapes (las macizas no llevan malla).'),
                         ('Acero de nervaduras secundarias y losas macizas (con desperdicio)', '={q_acs}*{q_da}', 'kg', '—'),
                         ('Acero de nervaduras principales (con desperdicio)', '={q_anp}*{q_da}', 'kg', '—')]):
        r = S.r
        cell(ws, r, 2, c, al='left'); ws.cell(r, 2).font = font(10)
        cell(ws, r, 3, f.format(**N), 'key', '#,##0.0'); cell(ws, r, 4, u)
        ws.merge_cells('E%d:I%d' % (r, r)); cell(ws, r, 5, how, al='left')
        cell(ws, r, 10, '—', 'ref'); S.r += 1
        N['qT%d' % qi] = 'CUANTIFICACION!$C$%d' % r
    add_cf(ws, 'I1:J%d' % S.r)
    return ws
