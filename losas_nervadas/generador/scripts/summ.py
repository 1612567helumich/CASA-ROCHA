import sys, openpyxl
sys.path.insert(0, '/tmp/claude-0/-home-user-CASA-ROCHA/faa149e0-9f48-5f0d-b5ee-28b4c7156d8c/scratchpad/scripts')
from openpyxl.utils import column_index_from_string as CI
import gen_b
TC, PC = gen_b.TC, gen_b.PC
f = lambda x: ('%.2f' % x) if isinstance(x, (int, float)) else str(x)
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
which = sys.argv[2] if len(sys.argv) > 2 else 'TP'
if 'T' in which:
    ws = wb['TABLEROS']
    for r in range(gen_b.T0, gen_b.T1 + 1):
        g = lambda k: ws.cell(r, CI(TC[k])).value
        print(g('key'), g('caso')[:6], 'wu', f(g('wu')), 'FU', '/'.join(f(g('FU' + b)) for b in gen_b.BLKS), 'n', '/'.join('%s#%s' % (g('n' + b), g('v' + b)) for b in gen_b.BLKS),
              'et', f(min([g('e' + b) for b in gen_b.BLKS if isinstance(g('e' + b), (int, float))] or [9])), '678', g('g678'), 'V', f(g('Vu')), f(g('VcR')), g('cort'),
              'afm', f(g('afm')), 'hmin', f(g('hmin')), g('peral'), g('tdia'), 'Lb', g('Lbc'), g('Lbl'), '|', g('estado'))
if 'P' in which:
    ws = wb['PRINCIPALES']
    for r in range(gen_b.P0, gen_b.P1 + 1):
        g = lambda k: ws.cell(r, CI(PC[k])).value
        print(g('np'), g('cond')[:7], 'L', f(g('L')), '%sx%s' % (g('bq'), g('hq')), 'w', f(g('w')), 'waz', f(g('waz')), 'M', f(g('Mp')), f(g('Mn')), 'FU', f(g('FUP')), f(g('FUN')),
              'n', '%s#%s/%s#%s' % (g('nP'), g('vP'), g('nN'), g('vN')), 'et', f(g('eP')), f(g('eN')), 'V', f(g('Vu')), f(g('VcR')), g('cort'),
              'd', f(g('dlt')), '/', f(g('dlim')), 'm418', g('m418'), 'fcm', f(g('fcm')), 'K', f(g('Kf')), 'htab', f(g('htab')), 'h418', f(g('h418')), 'hrec', g('hrec'), '|', g('estado'))
if 'N' in which:
    ws = wb['NERVADURA']
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if v and ws.cell(r, 3).value is not None and any(k in str(v) for k in ('Estado', 'V_cR', 'As de la malla', 'm_R', 'm requerido', 'P_u', 'b₀', 'Recubrimiento', '¿Coinciden', 'ESTADO')):
            print(r, str(v)[:70], '|', ws.cell(r, 3).value)
