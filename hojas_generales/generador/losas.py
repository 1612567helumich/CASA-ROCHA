"""NERVADAS y MACIZAS: un renglón por tablero (hasta 40), con las fórmulas de gen_b.py (Casa Rocha, HEAD)
sin la parte que depende del modelo de ese proyecto (vecinos, principales, áreas tributarias)."""
import re
from openpyxl.styles import Alignment
from openpyxl.comments import Comment
from common import *            # noqa: F401,F403
from openpyxl.utils import column_index_from_string as CI

NC = {}    # NERVADAS: clave -> letra de columna
MC = {}    # MACIZAS
BLKS = ('Pc', 'Nc', 'Pl', 'Nl')
BLKNAME = {'Pc': 'INFERIOR ‖ CLARO CORTO', 'Nc': 'SUPERIOR ‖ CLARO CORTO', 'Pl': 'INFERIOR ‖ CLARO LARGO', 'Nl': 'SUPERIOR ‖ CLARO LARGO'}
MOMOF = {'Pc': 'MPc', 'Nc': 'MNc', 'Pl': 'MPl', 'Nl': 'MNl'}


def mk_F(COL):
    def F(expr, r):
        s = expr.format(**N) if '{' in expr else expr
        return re.sub(r'\[(\w+)\]', lambda m: '%s%d' % (COL[m.group(1)], r), s)
    return F


def G(expr):
    """Renglón vacío o incompleto: celda en blanco (sin ceros ni errores)."""
    return '=IF([ok]<>"SÍ","",%s)' % expr


def varsel2(As, pref, nmax, avail, nmin, as43):
    """gen_b.varsel2 sin estribo: varilla más chica (desde la preferida) cuyas piezas caben y no pasan del máximo."""
    i0 = 'MATCH(%s,%s,0)' % (pref, BARCOL)
    out = 'INDEX(%s,%d)' % (BARCOL, NBARS)
    for k in range(4, -1, -1):
        idx = 'MIN(%s+%d,%d)' % (i0, k, NBARS)
        a_ = 'INDEX(%s,%s)' % (ABCOL, idx); d_ = 'INDEX(%s,%s)' % (DBCOL, idx)
        npz = 'MAX(%s,CEILING(%s/%s,1))' % (nmin, As, a_)
        npz = 'IF(%s*%s>=%s-0.0001,%s,MAX(%s,2))' % (npz, a_, as43, npz, npz)
        ok = 'AND(%s>=0,%s<=MIN(%s,INT((%s+MAX({slib},%s))/(%s+MAX({slib},%s)))))' % (As, npz, nmax, avail, d_, d_, d_)
        out = 'IF(%s,INDEX(%s,%s),%s)' % (ok, BARCOL, idx, out)
    return out


def macsel(As, pref, smax):
    """Losa maciza: varilla más chica (desde la preferida) con separación ≥ la mínima práctica."""
    i0 = 'MATCH(%s,%s,0)' % (pref, BARCOL)
    out = 'INDEX(%s,%d)' % (BARCOL, NBARS)
    for k in range(4, -1, -1):
        idx = 'MIN(%s+%d,%d)' % (i0, k, NBARS)
        a_ = 'INDEX(%s,%s)' % (ABCOL, idx)
        ok = 'AND(%s>0,FLOOR(MIN(%s,100*%s/MAX(%s,0.0001)),{spaso})>={smin})' % (As, smax, a_, As)
        out = 'IF(%s,INDEX(%s,%s),%s)' % (ok, BARCOL, idx, out)
    return out


# --------------------------------------------------------------- piezas comunes a las dos hojas
def common_geo(vals, kind):
    """Geometría, cargas y caso (iguales en nervadas y macizas)."""
    vals['emp'] = (G('IF([tipo]<>"VOLADO","—",IF(COUNTIF([bI]:[bR],"CONTINUO")<>1,"?",IF([bI]="CONTINUO","INFERIOR",IF([bS]="CONTINUO","SUPERIOR",IF([bL]="CONTINUO","IZQUIERDO","DERECHO")))))'), 'calc', '@')
    vals['dirc'] = (G('IF([tipo]="VOLADO",IF(OR([emp]="IZQUIERDO",[emp]="DERECHO"),"X","Y"),IF([ax]<=[ay],"X","Y"))'), 'calc', '@')
    vals['a1'] = (G('IF([tipo]="VOLADO",IF([dirc]="X",[ax],[ay]),MIN([ax],[ay]))'), 'calc', '0.000')
    vals['a2'] = (G('IF([tipo]="VOLADO",IF([dirc]="X",[ay],[ax]),MAX([ax],[ay]))'), 'calc', '0.000')
    vals['m'] = (G('MIN([a1],[a2])/MAX([a1],[a2])'), 'calc', '0.000')
    vals['ln1'] = (G('MAX(1,[a1]*100-IF([tipo]="VOLADO",{bp}/2,{bp}))'), 'calc', '0.0')
    vals['ln2'] = (G('MAX(1,[a2]*100-{bp})'), 'calc', '0.0')
    vals['azo'] = (G('IFERROR(INDEX({CV_AZ},MATCH([dest],{CV_NAMES},0)),"NO")'), 'calc', '@')
    vals['cv'] = (G('IF(ISNUMBER([cvo]),[cvo],IFERROR(INDEX({CV_WM},MATCH([dest],{CV_NAMES},0)),0))'), 'calc', '0')
    vals['cm'] = (G('[pp]+N([scm])'), 'calc', '0')
    vals['wu'] = (G('MAX({FCcm}*[cm]+{FCcv}*[cv],IF([azo]="SÍ",{FCacc}*([cm]+{Wg}),0))'), 'key', '0')
    vals['nd'] = (G('([bI]="DISCONTINUO")+([bS]="DISCONTINUO")+([bL]="DISCONTINUO")+([bR]="DISCONTINUO")'), 'calc', '0')
    vals['ndc'] = (G('IF([dirc]="X",([bI]="DISCONTINUO")+([bS]="DISCONTINUO"),([bL]="DISCONTINUO")+([bR]="DISCONTINUO"))'), 'calc', '0')
    vals['cauto'] = (G('IF([nd]=0,"INTERIOR",IF([nd]=1,IF([ndc]=1,"BORDE-CORTO","BORDE-LARGO"),IF([nd]=2,IF([ndc]=1,"ESQUINA",IF([ndc]=2,"EXTREMO-LARGO","EXTREMO-CORTO")),IF([nd]=3,IF([ndc]=2,"EXTREMO-LARGO","EXTREMO-CORTO"),"AISLADO"))))'), 'calc', '@')
    vals['caso'] = (G('IF(OR([cforce]="",[cforce]="AUTO"),[cauto],[cforce])'), 'key', '@')
    vals['idx'] = (G('MATCH([m],{MROW},1)'), 'calc', '0')
    vals['fr'] = (G('IF([idx]>=7,0,([m]-INDEX({MROW},[idx]))/(INDEX({MROW},[idx]+1)-INDEX({MROW},[idx])))'), 'calc', '0.00')
    for mm in MOMS:
        i_ = 'MATCH([caso]&"|%s",{CKEY},0)' % mm
        vals['c' + mm] = (G('INDEX({CVAL},%s,[idx])+[fr]*(INDEX({CVAL},%s,MIN([idx]+1,7))-INDEX({CVAL},%s,[idx]))' % (i_, i_, i_)), 'calc', '0')


def c342_formula(fac):
    own_c = 'MAX([cNIC],[cNDC])%s' % fac
    own_l = 'MAX([cNIL],[cNDL])%s' % fac
    fc_ = 'AND(N([Mvc])>0,ABS(%s-N([Mvc]))>0.5*MIN(%s,N([Mvc])))' % (own_c, own_c)
    fl_ = 'AND(N([Mvl])>0,ABS(%s-N([Mvl]))>0.5*MIN(%s,N([Mvl])))' % (own_l, own_l)
    return G('IF([tipo]<>"PERIMETRAL","—",IF(AND(N([Mvc])=0,N([Mvl])=0),"— (sin M⁻ del vecino)",IF(OR(%s,%s),"FUERA: difieren > 50 %%"&IF(%s," ‖ corto","")&IF(%s," ‖ largo",""),"CUMPLE")))' % (fc_, fl_, fc_, fl_))


def xpi(ln, cP, cI, cD, MN, fac):
    """Punto de inflexión (gen_b.xpi) sin la distribución con vecinos: M1 = M⁻ de diseño, M2 = el del otro extremo, M0 con los M⁻ de la tabla."""
    P_ = '([%s]%s)' % (cP, fac); M1 = '[%s]' % MN
    M1e = '(MAX([%s],[%s])%s)' % (cI, cD, fac)
    M2 = '(IF(AND([%s]>0,[%s]>0),MIN([%s],[%s]),MAX([%s],[%s]))%s)' % (cI, cD, cI, cD, cI, cD, fac)
    M0 = '(%s+(%s+%s)/2)' % (P_, M1e, M2)
    bq_ = '(4*%s+%s-%s)' % (M0, M1, M2)
    disc = '(%s^2-16*%s*%s)' % (bq_, M0, M1)
    return '[%s]*(%s-SQRT(MAX(0,%s)))/(8*MAX(%s,1E-9))' % (ln, bq_, disc, M0), disc


def volado_x0(wD):
    """Distancia (cm) desde el apoyo, dentro del claro de respaldo L_b, donde se anula el M⁻ del volado (gen_b, volado)."""
    Lb = 'IF(ISNUMBER([Lb]),[Lb],[a1])'
    A_ = '(0.5*%s)' % wD; B_ = '(0.5*%s*%s+[MNc]/%s)' % (wD, Lb, Lb); C_ = '[MNc]'
    return 'IF(%s^2-4*%s*%s<0,%s,MIN(%s,(%s-SQRT(%s^2-4*%s*%s))/(2*%s)))*100' % (B_, A_, C_, Lb, Lb, B_, B_, A_, C_, A_), Lb


def write_table(ws, COL, cols, groups, notes, vals_by_row, inputs, examples, title, subtitle, legend, last_key):
    """Encabezados, notas «CÓMO SE CALCULA» por grupo, unidades y los 40 renglones."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3.2
    last = COL[last_key]
    for rr, txt, f_, h in ((1, title, font(13, True, color='FFFFFF'), 26), (2, subtitle, font(9, i=True, color=GRAY), 30), (3, legend, font(9, True, color=WARN_FG), 42)):
        ws.merge_cells('B%d:%s%d' % (rr, last, rr))
        c = ws['B%d' % rr]; c.value = txt; c.font = f_
        c.alignment = Alignment(wrap_text=True, vertical='center' if rr == 1 else 'top')
        if rr == 1:
            c.fill = fill(TITLE_BG)
        if rr == 3:
            c.fill = fill(WARN_BG)
        ws.row_dimensions[rr].height = h
    ws.row_dimensions[4].height = 6
    for a, b, t in groups:
        ws.merge_cells('%s6:%s6' % (COL[a], COL[b]))
        x = ws['%s6' % COL[a]]; x.value = t; x.font = font(9, True, color='FFFFFF'); x.fill = fill(SEC_BG if t[:5] != 'DATOS' else '7F6000')
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.merge_cells('%s5:%s5' % (COL[a], COL[b]))
        y = ws['%s5' % COL[a]]; y.value = notes.get(a, ''); y.font = font(7.5, i=True, color=GRAY); y.fill = fill(NOTE_BG)
        y.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[5].height = 118
    ws.row_dimensions[6].height = 18
    for k, (t, w, u) in cols.items():
        hdr_cell(ws, 7, CI(COL[k]), t, w)
        unit_cell(ws, 8, CI(COL[k]), u)
    ws.row_dimensions[7].height = 54
    hdr_cell(ws, 7, 1, '#'); unit_cell(ws, 8, 1, 'cuenta')
    F = mk_F(COL)
    for r in range(R0, R1 + 1):
        x = ws.cell(r, 1, '=IF(B%d="",0,1)' % r if r == R0 else '=IF(B%d="",0,1)+A%d' % (r, r - 1))
        x.font = font(7, color='A6A6A6'); x.alignment = Alignment(horizontal='center')
        ex = examples[r - R0] if r - R0 < len(examples) else {}
        for k in inputs:
            v = ex.get(k)
            fmt = '@' if k in ('key', 'tipo', 'dest', 'bI', 'bS', 'bL', 'bR', 'cforce', 'col') else ('0.000' if k in ('ax', 'ay', 'Lb', 'Lb2') else ('0.000' if k in ('Mvc', 'Mvl') else '0.0' if k in ('pbor', 'afo', 'ho') else 'General'))
            c = cell(ws, r, CI(COL[k]), v, 'input', fmt, al='left' if k in ('key', 'dest') else 'center', sz=9)
        vals = vals_by_row
        for k in cols:
            if k in inputs:
                continue
            v, kind, fmt = vals[k]
            v = F(v, r)
            c = cell(ws, r, CI(COL[k]), v, kind, fmt, al='left' if k in ('estado', 'nota', 'cort', 'c342', 'anc') or k.startswith('t') and k[1:] in BLKS else 'center', sz=9)
            if k in ('nota',) or (k.startswith('t') and k[1:] in BLKS):
                c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
        ws.row_dimensions[r].height = 33
    for k, lst in (('tipo', 'LST_TIPO'), ('dest', 'LST_DEST'), ('cforce', 'LST_CASO'), ('col', 'LST_COLADO')):
        dv_list(ws, '%s%d:%s%d' % (COL[k], R0, COL[k], R1), '=' + lst)
    dv_list(ws, '%s%d:%s%d' % (COL['bI'], R0, COL['bR'], R1), '=LST_BORDE')
    ws.freeze_panes = 'C%d' % R0
    add_cf(ws, '%s%d:%s%d' % (COL['ok'], R0, COL['ok'], R1))
    for k in ('estado', 'cort', 'peral', 'c342', 'anc') + (('g678', 'tdia', 'rest') if 'g678' in COL else ()):
        add_cf(ws, '%s%d:%s%d' % (COL[k], R0, COL[k], R1))
    ws.sheet_view.zoomScale = 80
    print_setup(ws, 'A1:%s%d' % (last, R1), landscape=True)
    ws.print_title_rows = '5:8'
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID


def _mkcols(spec):
    COL = {}
    cols = {}
    for i, (k, t, w, u) in enumerate(spec):
        COL[k] = CL(2 + i)
        cols[k] = (t, w, u)
    return COL, cols


# =====================================================================
INP_N = ['key', 'tipo', 'dest', 'ax', 'ay', 'bI', 'bS', 'bL', 'bR', 'cforce', 'col', 'scm', 'cvo', 'pbor', 'Lb', 'Lb2', 'bwo', 'afo', 'Mvc', 'Mvl']
SPEC_N = [
    ('key', 'Tablero', 22, ''), ('tipo', 'Tipo', 11.5, ''), ('dest', 'Destino (lista de DATOS)', 19, ''), ('ax', 'a_x — claro en X a ejes', 7.5, 'm'), ('ay', 'a_y — claro en Y a ejes', 7.5, 'm'),
    ('bI', 'Borde inferior (‖ X)', 11, 'CONT./DISC.'), ('bS', 'Borde superior (‖ X)', 11, 'CONT./DISC.'), ('bL', 'Borde izquierdo (‖ Y)', 11, 'CONT./DISC.'), ('bR', 'Borde derecho (‖ Y)', 11, 'CONT./DISC.'),
    ('cforce', 'Forzar caso', 12, 'AUTO = sale solo'), ('col', 'Colado con los apoyos', 13, 'caso I / II'), ('scm', 'SCM — sobrecarga muerta', 7.5, 'kg/m²'), ('cvo', 'CV propia (opcional)', 7.5, 'kg/m²'),
    ('pbor', 'Carga en la orilla del volado', 8, 'kg/m serv.'), ('Lb', 'L_b — claro de respaldo (volado)', 8, 'm'), ('Lb2', 'Claro del respaldo ‖ al apoyo (opcional; vacío = losa en una dirección, del lado seguro)', 8, 'm'), ('bwo', 'b_w propio (opcional)', 7, 'cm'), ('afo', 'α_fm propio (opcional)', 7, '—'),
    ('Mvc', 'M⁻ del vecino ‖ corto (opcional)', 8.5, 't·m/nerv.'), ('Mvl', 'M⁻ del vecino ‖ largo (opcional)', 8.5, 't·m/nerv.'),
    ('ok', '¿Datos completos?', 7.5, ''), ('emp', 'Borde empotrado (volado)', 10, ''), ('dirc', 'Claro corto ‖', 5.5, ''), ('a1', 'a1 corto (volado: largo)', 7, 'm'), ('a2', 'a2 largo', 7, 'm'),
    ('m', 'm = a1/a2', 6.5, ''), ('ln1', 'L_n corto', 6.5, 'cm'), ('ln2', 'L_n largo', 6.5, 'cm'), ('bwr', 'b_w nerv. ‖ corto', 6, 'cm'), ('acr', 'Casetón entre nerv. ‖ corto = s − b_w', 7, 'cm'), ('bwc', 'b_w cálculo ‖ corto (M⁻ y V)', 7, 'cm'), ('bwcl', 'b_w cálculo ‖ largo (M⁻)', 7, 'cm'),
    ('bfc', 'Patín b_f ‖ corto', 7, 'cm'), ('bfl', 'Patín b_f ‖ largo', 7, 'cm'), ('g678', '6.7.8 cumple', 7, ''), ('rb', 'r inferior', 6, 'cm'), ('rt', 'r superior', 6, 'cm'), ('rtm', 'r hasta la malla (al centro de la capa)', 6.5, 'cm'),
    ('azo', '¿Azotea?', 6.5, ''), ('cv', 'CV W_m', 6.5, 'kg/m²'), ('pp', 'PP de diseño', 7, 'kg/m²'), ('cm', 'CM', 6.5, 'kg/m²'), ('wu', 'w_u', 7, 'kg/m²'),
    ('nd', 'Bordes discont.', 6.5, ''), ('ndc', '… cortos', 6, ''), ('cauto', 'Caso automático', 14, ''), ('caso', 'Caso usado', 14, ''),
    ('idx', 'col. m', 5, ''), ('fr', 'interp.', 5.5, ''), ('cNIC', 'M⁻ cont. corto', 7, '×10⁻⁴'), ('cNIL', 'M⁻ cont. largo', 7, '×10⁻⁴'), ('cNDC', 'M⁻ disc. corto', 7, '×10⁻⁴'),
    ('cNDL', 'M⁻ disc. largo', 7, '×10⁻⁴'), ('cPC', 'M⁺ corto', 7, '×10⁻⁴'), ('cPL', 'M⁺ largo', 7, '×10⁻⁴'),
    ('Mtip', 'M⁻ extra de los macizos del volado', 7.5, 't·m/nerv.'), ('Vtip', 'V extra de los macizos del volado', 7.5, 't/nerv.'),
    ('MPc', 'M⁺ corto', 7, 't·m/nerv.'), ('MNc', 'M⁻ corto', 7, 't·m/nerv.'), ('MPl', 'M⁺ largo', 7, 't·m/nerv.'), ('MNl', 'M⁻ largo', 7, 't·m/nerv.'), ('c342', '3.4.2.1 c) con el vecino', 15, ''),
]
for _b in BLKS:
    SPEC_N += [('dq' + _b, 'd para As (var. preferida)', 7, 'cm'), ('As' + _b, 'As req.', 6.5, 'cm²'), ('v' + _b, 'Var. No.', 5.5, ''), ('n' + _b, 'Pzas', 5, ''), ('d' + _b, 'd real', 6, 'cm'),
               ('MR' + _b, 'M_R', 6.5, 't·m'), ('FU' + _b, 'Mu/M_R', 6.5, '≤ 1'), ('e' + _b, 'ε_t', 6.5, '≥ 0.005')]
SPEC_N += [('Vu', 'V_u por nerv.', 7, 't'), ('VcR', 'V_cR', 7, 't'), ('nh', 'Hileras a macizar', 6.5, 'por apoyo'), ('cort', 'Estado cortante', 17, ''),
           ('yb', 'ȳ de la T', 6, 'cm'), ('IT', 'I_T de la T', 8, 'cm⁴'), ('heq', 'h_eq', 6.5, 'cm'), ('afm', 'α_fm', 6, ''), ('hmin', 'h mín. (Tabla 6.7.3.1.2 / volado L/8)', 8.5, 'cm'),
           ('peral', 'Estado peralte', 13, ''), ('tdia', 'Capa (diafragma)', 10, ''),
           ('Lbc', 'Bastón ‖ corto desde el paño', 8, 'cm'), ('Lbl', 'Bastón ‖ largo desde el paño', 8, 'cm'), ('bdc', 'Bastón también en borde disc. ‖ corto', 8, ''), ('bdl', 'Bastón también en borde disc. ‖ largo', 8, ''),
           ('anc', 'Anclaje en bordes discontinuos', 16, ''), ('ncx', 'Casetones ‖ X', 6.5, 'pzas'), ('ncy', 'Casetones ‖ Y', 6.5, 'pzas'), ('mx0', 'Macizo de orilla izq.', 6.5, 'cm'), ('mx1', 'Macizo de orilla der.', 6.5, 'cm'), ('my0', 'Macizo de orilla inf.', 6.5, 'cm'), ('my1', 'Macizo de orilla sup.', 6.5, 'cm'),
           ('rn1', 'Varillas del volado que siguen en la nervadura de b_w de DATOS (la corrida)', 8, 'pzas'), ('rd', 'd de esa nervadura', 6.5, 'cm'), ('rM1', 'M_R de esa nervadura', 7, 't·m'), ('rV1', 'V_cR de esa nervadura', 7, 't'),
           ('rxM', 'Donde M baja a ese M_R', 7, 'm desde el paño'), ('rxV', 'Donde V baja a ese V_cR', 7, 'm desde el paño'), ('rVw', 'V_u del respaldo junto al apoyo (nerv. ancha)', 7.5, 't'),
           ('rxr', 'Nervadura ancha en el respaldo hasta', 7.5, 'cm desde el paño'), ('rnh', 'Hileras del respaldo con b_w del volado', 7, 'hileras'), ('rest', 'Respaldo del volado', 24, ''),
           ('tPc', 'ARMADO inferior ‖ corto (por nervadura)', 16, ''), ('tNc', 'ARMADO superior ‖ corto', 28, ''), ('tPl', 'ARMADO inferior ‖ largo', 16, ''), ('tNl', 'ARMADO superior ‖ largo', 28, ''),
           ('estado', 'ESTADO', 30, ''), ('nota', 'NOTA', 70, '')]
NC_, COLS_N = _mkcols(SPEC_N)
NC.update(NC_)

GROUPS_N = [('key', 'Mvl', 'DATOS DEL TABLERO — CAPTURA (amarillo)'), ('ok', 'rtm', 'GEOMETRÍA'), ('azo', 'wu', 'CARGAS'), ('nd', 'caso', 'CASO DE LA TABLA'),
            ('idx', 'cPL', 'COEFICIENTES ×10⁻⁴ (Tabla 3.4.2.1)'), ('Mtip', 'c342', 'MOMENTOS POR NERVADURA'),
            ('dqPc', 'ePc', 'LECHO INFERIOR ‖ CLARO CORTO'), ('dqNc', 'eNc', 'LECHO SUPERIOR ‖ CLARO CORTO'), ('dqPl', 'ePl', 'LECHO INFERIOR ‖ CLARO LARGO'), ('dqNl', 'eNl', 'LECHO SUPERIOR ‖ CLARO LARGO'),
            ('Vu', 'cort', 'CORTANTE'), ('yb', 'tdia', 'ESPESOR MÍNIMO Y CAPA'), ('Lbc', 'my1', 'BASTONES, ANCLAJE Y CASETONES'), ('rn1', 'rest', 'RESPALDO DE UN VOLADO CON b_w PROPIO'), ('tPc', 'nota', 'RESULTADO')]
NOTES_N = {
    'key': 'CAPTURA. Claros A EJES (del lado seguro; la tabla admite claros libres en el caso I). Borde CONTINUO = sigue losa del otro lado; DISCONTINUO = orilla, hueco o losa que no sigue. '
           'VOLADO: marca CONTINUO sólo el borde empotrado. CV propia, b_w propio y α_fm propio: vacío = el de DATOS. M⁻ del vecino: por nervadura, en el borde común; la hoja arma con el mayor.',
    'ok': 'a1 = claro corto, a2 = largo (volado: a1 = largo perpendicular al borde empotrado). L_n = a − b_p (volado a − b_p/2) (NTC-Concreto 6.3.2.1, p. 117). Patín b_f = b_w + 2·mín(8t, a_c/2, L_n/8) ≤ s (Tabla 5.2.1.4.2, p. 89). '
          '6.7.8: claro corto ≥ 6·a_c, a_c ≤ 80, b_w ≥ 10, h ≤ 3.5·b_w, t ≥ máx(a_c/13, 5) (6.7.8.1.2 a 6.7.8.1.4, p. 187); si NO: losa sobre vigas, r de viga, sin ×1.1 (6.7.8.1.9, p. 188). b_w − 2 cm si b_w < 20 sin 3.9.2 (3.9.1 e), p. 74).',
    'azo': 'PP = (h − a_c²·h_c/s²)·γc + casetón + 20 + 20 (NTC-Criterios 5.1.2.1, p. 9). CM = PP + SCM. w_u = máx(1.3·CM + 1.5·W_m; azotea: 1.1·(CM + W_g)) (NTC-Criterios 3.4.1 a) y b), p. 7; ec. 7.3.2.1, p. 17).',
    'nd': 'Se cuentan los bordes DISCONTINUOS; «corto» = borde de largo a1. Mismo criterio que la hoja de Casa Rocha; dos bordes opuestos discontinuos se toman como «extremo» (del lado seguro). Tabla 3.4.2.1, p. 65.',
    'idx': 'Interpolación lineal en m entre las columnas de la Tabla 3.4.2.1, caso I (TABLAS sección A).',
    'Mtip': 'M = coef·10⁻⁴·w_u·a1²·s, por nervadura (NTC-Concreto 3.4.2.1, p. 64). M⁻ = el mayor de borde continuo, discontinuo y el del vecino capturado. Volado: M⁻ = (w_u·L_n²/2 + 1.3·P·L_n)·s + M de los macizos '
            '(el PP promedio no cuenta el concreto que llena los casetones en el macizo de la punta, de ancho macizo de orilla + b_p/2, ni en el de la raíz: 1.3·γc·s·(vacío/m²)·ancho·brazo, como en la hoja de Casa Rocha). '
            '3.4.2.1 c): M⁻ propio y del vecino no deben diferir más de 50 % (p. 64).',
    'dqPc': 'b = patín b_f. d con la varilla preferida para As; con la real para M_R. q = 1 − √(1 − 2Mu/(F_R·b·d²·f\'\'c)), As = q·f\'\'c·b·d/fy (ecs. 5.2.2.1.1.1, p. 89); As ≥ mín(p_min·b_w·d, 4/3·As) (6.3.5.1.3, p. 122); −1 = no alcanza o rebasa p_max. '
            'Varilla: la preferida o la siguiente que quepa (14.2.1, p. 347). ε_t = 0.003(d − c)/c ≥ 0.005 (Tabla 3.8.2.2, p. 70).',
    'dqNc': 'b = b_w de cálculo. d = h − máx(r, (t − 2Ø)/2) − 2·Ø malla − d_b/2 (la malla va al centro de la capa y la varilla pasa debajo), − 2 cm si d ≤ 20 cm sin la precaución de 3.9.2 (3.9.1 d), p. 74). 1 varilla corrida + bastones (6.7.6.1.3, p. 175). En volado el patín trabaja a tensión: As mín sobre 2·b_w.',
    'dqPl': 'Igual que el lecho inferior corto; el lecho del claro largo va encima del corto: d = h − r − d_b,corto − d_b/2.',
    'dqNl': 'Igual que el lecho superior corto; va debajo del lecho superior corto.',
    'Vu': 'V_u a d del paño = (L_n/2 − d)(0.95 − 0.5m)·w_u·s, ×1.15 con bordes continuos y discontinuos (método reconocido, 3.4.1.1, p. 64; 6.7.4.3.2, p. 172). Volado: (w_u(L_n − d) + 1.3P)·s. '
          'V_cR = F_R·2·λs·p^⅓·√f\'c·b_w·d (ec. 5.5.3.2.1, p. 95), ×1.1 si cumple 6.7.8 (6.7.8.1.5, p. 187). Sin estribos (5.5.3.8.1, p. 97): lo que falte se cubre macizando hileras de casetones.',
    'yb': 'Tabla 6.7.3.1.2 (p. 168), α_fm > 2: L_n(0.8 + fy/14 000)/(36 + 9β) ≥ 9; 0.2–2: …/(36 + 5β(α_fm − 0.2)) ≥ 13; ≤ 0.2: Tabla 6.7.3.1.1 (L_n/30 ó /33, ≥ 13). +10 % con borde discontinuo sin viga (6.7.3.1.2.1). '
          'Se compara h_eq = ∛(12·I_T/s). Volado: L/8 contra h (Tabla 6.3.3.1.1, p. 120). Capa ≥ 6 cm si el claro es ≥ 6 m (6.8.3.1.2, p. 189).',
    'Lbc': 'Bastón desde el paño = punto de inflexión + máx(d, 12d_b, L_n/16), ≥ L_d (6.3.7.3.4, p. 127; 6.3.7.3.10.4, p. 129; Tabla 14.4.2.4, p. 353). Volado: entra al respaldo hasta M = 0 con 0.9·PP y sin CV. '
           'Anclaje en bordes discontinuos (dala de borde armada, D9 de Casa Rocha): 1.25·L_dh ≤ ancho − r_np y ≥ 15 cm (6.7.8.1.7, p. 188); L_dh del bastón (6.7.7.4.1.1 b), p. 179); gancho de 180° con alto 8·d_b ≤ peralte de la dala − 2·r_np (Tabla 14.3.1, p. 349). '
           'Respaldo de un volado con b_w propio: sus nervaduras anchas siguen en el respaldo hasta donde la nervadura de b_w de DATOS resiste el M⁻ (con 0.9·PP, sin CV) y el cortante del volado.',
    'rn1': 'Sólo volados con b_w propio mayor que el de DATOS (decisión D10 de Casa Rocha). Del otro lado del apoyo el M⁻ del volado lo toman las nervaduras del tablero de respaldo: sus primeras hileras siguen con el b_w del volado '
           '(casetón s − b_w) hasta donde la nervadura normal (con r de viga, del lado seguro), con la varilla corrida del volado, resiste M(x) = M⁻(1 − x/L_b) − 0.9·PP·s·x(L_b − x)/2 (sin CV) y el cortante V(x) = M⁻/L_n + (0.95 − 0.5·m_b)·1.15·w_u·s·(L_n/2 − x), con L_n = L_b − b_p y m_b del respaldo (sin su otro claro, m_b = 0); sin ×1.1 y con ×1.15 porque aquí no se sabe si el respaldo cumple 6.7.8 ni si tiene bordes discontinuos, del lado seguro. '
           'Las demás varillas del volado son bastones que terminan ahí, más máx(d, 12d_b) y no menos que L_d. La nervadura ancha se revisa a cortante junto al apoyo con el V_cR del volado (ec. 5.5.3.2.1, p. 95; 6.3.7.3.10.4, p. 129).',
    'tPc': 'Armado por NERVADURA. Superior: 1 corrida + bastones a cada lado de los bordes continuos (y de los discontinuos si la columna lo pide). ESTADO junta flexión, cupo, ε_t, cortante, peralte, capa y anclaje.',
}


def vals_nervadas():
    v = {}
    v['ok'] = ('=IF([key]="","",IF(AND(OR([tipo]="PERIMETRAL",[tipo]="VOLADO"),ISNUMBER([ax]),ISNUMBER([ay]),[ax]>0,[ay]>0,OR(ISNUMBER([cvo]),ISNUMBER(MATCH([dest],{CV_NAMES},0))),OR([tipo]<>"VOLADO",N([Lb])>0)),"SÍ","NO"))', 'calc', '@')
    common_geo(v, 'N')
    v['bwr'] = (G('IF(ISNUMBER([bwo]),[bwo],{bw})'), 'calc', '0.0')
    v['acr'] = (G('MAX(0,{s}-[bwr])'), 'calc', '0.0')
    v['bwc'] = (G('IF(OR({ctl_bw}="SÍ",[bwr]>=20),[bwr],[bwr]-2)'), 'calc', '0.0')
    v['bwcl'] = (G('IF(OR({ctl_bw}="SÍ",{bw}>=20),{bw},{bw}-2)'), 'calc', '0.0')
    v['bfc'] = (G('MIN({s},[bwr]+2*MIN(8*{t},[acr]/2,[ln1]/8))'), 'calc', '0.0')
    v['bfl'] = (G('MIN({s},{bw}+2*MIN(8*{t},{ac}/2,[ln2]/8))'), 'calc', '0.0')
    v['g678'] = (G('IF(AND([tipo]="PERIMETRAL",[a1]*100>=6*MAX({ac},[acr]),MAX({ac},[acr])<=80,MIN({bw},[bwr])>=10,{h}<=3.5*MIN({bw},[bwr]),{t}>=MAX(MAX({ac},[acr])/13,5)),"SÍ","NO")'), 'calc', '@')
    v['rb'] = (G('IF([tipo]="VOLADO",{rec_ext},IF([g678]="NO",{rec_np},{rec}))'), 'calc', '0.0')
    v['rt'] = (G('IF(OR(AND([azo]="SÍ",{exp_az}="SÍ"),AND([tipo]="VOLADO",{exp_vol}="SÍ")),{rec_ext},IF([g678]="NO",{rec_np},{rec}))'), 'calc', '0.0')
    # la malla va al centro de la capa ((t − 2Ø)/2 por cara) y la varilla superior pasa debajo; si r es mayor, manda r
    v['rtm'] = (G('MAX([rt],({t}-2*{fim})/2)'), 'calc', '0.00')
    vf = '([acr]*{ac}/{s}^2)*{hc}/100'
    v['pp'] = (G('IF(MIN([ncx],[ncy])=0,{h}/100*{gc}*1000,({h}/100-%s)*{gc}*1000+%s*{gpol}*1000)+{inc_losa}+{inc_mort}' % (vf, vf)), 'calc', '0')
    fac = '*[wu]*[a1]^2*1E-4*{s}/100/1000'
    v['MPc'] = (G('IF([tipo]="PERIMETRAL",[cPC]%s,0)' % fac), 'calc', '0.000')
    # volado: el concreto que llena los casetones en el macizo de la punta (macizo de orilla + b_p/2) y en el de la raíz no lo cuenta el PP promedio
    otp = 'IF([dirc]="X",[mx0],[my0])'
    Ax = '(([acr]*{ac}/{s}^2)*{hc}/100)'
    v['Mtip'] = (G('IF(OR([tipo]<>"VOLADO",MIN([ncx],[ncy])=0),0,{FCcm}*{gc}*{s}/100*(%s*(%s+{bp}/2)/100*([ln1]-(%s+{bp}/2)/2)/100+%s*%s/100*%s/200))' % (Ax, otp, otp, Ax, otp, otp)), 'calc', '0.000')
    v['Vtip'] = (G('IF(OR([tipo]<>"VOLADO",MIN([ncx],[ncy])=0),0,{FCcm}*{gc}*{s}/100*%s*(%s+{bp}/2)/100)' % (Ax, otp)), 'calc', '0.000')
    v['MNc'] = (G('IF([tipo]="PERIMETRAL",MAX(MAX([cNIC],[cNDC])%s,N([Mvc])),([wu]*([ln1]/100)^2/2+{FCcm}*N([pbor])*[ln1]/100)*{s}/100/1000+[Mtip])' % fac), 'key', '0.000')
    v['MPl'] = (G('IF([tipo]="PERIMETRAL",[cPL]%s,0)' % fac), 'calc', '0.000')
    v['MNl'] = (G('IF([tipo]="PERIMETRAL",MAX(MAX([cNIL],[cNDL])%s,N([Mvl])),0)' % fac), 'calc', '0.000')
    v['c342'] = (c342_formula(fac), 'calc', '@')
    dq = {'Pc': '([h0]-[rb]-%s/2)' % db('MAX(4,{var_inf})'),
          'Pl': '([h0]-[rb]-%s-%s/2)' % (db('MAX(4,{var_inf})'), db('MAX(4,{var_inf})')),
          'Nc': red391('([h0]-[rtm]-2*{fim}-%s/2)' % db('MAX(4,{var_sup})')),
          'Nl': red391('([h0]-[rtm]-2*{fim}-%s-%s/2)' % (db('MAX(4,{var_sup})'), db('MAX(4,{var_sup})')))}
    dact = {'Pc': '[h0]-[rb]-%s/2' % db('[vPc]'),
            'Pl': '[h0]-[rb]-%s-%s/2' % (db('[vPc]'), db('[vPl]')),
            'Nc': red391('([h0]-[rtm]-2*{fim}-%s/2)' % db('[vNc]')),
            'Nl': red391('([h0]-[rtm]-2*{fim}-%s-%s/2)' % (db('[vNc]'), db('[vNl]')))}
    for blk in BLKS:
        neg = blk[0] == 'N'
        M = '[%s]' % MOMOF[blk]
        # nervaduras ‖ corto con b_w del renglón (b_w propio); nervaduras ‖ largo con la b_w de DATOS (como D10 de Casa Rocha)
        bwb = '[bwr]' if blk[1] == 'c' else '{bw}'
        b = ('[bwc]' if blk[1] == 'c' else '[bwcl]') if neg else '[bf%s]' % blk[1]
        pref = 'MAX(4,{var_sup})' if neg else 'MAX(4,{var_inf})'
        av = '(%s-2*%s)' % (bwb, '[rt]' if neg else '[rb]')
        bmin = 'IF([tipo]="VOLADO",MIN([bfc],2*[bwr]),[bwr])' if blk == 'Nc' else bwb
        dqr = '[dq%s]' % blk
        f = asflex(M, b, dqr)
        v['dq' + blk] = (G(dq[blk].replace('[h0]', '{h}')), 'calc', '0.00')
        v['As' + blk] = (G('IF(%s<0,-1,MAX(%s,MIN({pmin}*%s*%s,4/3*%s)))' % (f, f, bmin, dqr, f)), 'calc', '0.00')
        v['v' + blk] = (G(varsel2('[As%s]' % blk, pref, '{nmax_sec}', av, '{nmin_sec}', '4/3*' + f)), 'calc', 'General')
        ab_ = ab('[v%s]' % blk)
        n0 = 'MAX({nmin_sec},CEILING(MAX([As%s],0)/%s,1))' % (blk, ab_)
        v['n' + blk] = (G('IF([As%s]<0,0,IF(%s*%s>=4/3*%s-0.0001,%s,MAX(%s,2)))' % (blk, n0, ab_, f, n0, n0)), 'calc', '0')
        v['d' + blk] = (G(dact[blk].replace('[h0]', '{h}')), 'calc', '0.00')
        asp = '[n%s]*%s' % (blk, ab_)
        v['MR' + blk] = (G('IF([n%s]=0,0,{FRf}*%s*{fy}*([d%s]-%s*{fy}/({fpc}*%s)/2)/1E5)' % (blk, asp, blk, asp, b)), 'calc', '0.000')
        v['FU' + blk] = (G('IF([As%s]<0,9.99,IF([MR%s]>0,%s/[MR%s],0))' % (blk, blk, M, blk)), 'key', '0.00')
        cexp = '(%s*{fy}/({fpc}*%s)/{beta1})' % (asp, b)
        v['e' + blk] = (G('IF([n%s]=0,"—",0.003*([d%s]-%s)/%s)' % (blk, blk, cexp, cexp)), 'calc', '0.0000')
    lam = 'MIN(1,SQRT(2/(1+0.04*[dNc])))'
    pexp = '([nNc]*%s/([bwc]*[dNc]))' % ab('[vNc]')
    v['Vu'] = (G('IF([tipo]="PERIMETRAL",MAX(0,([ln1]/2-[dNc])/100)*(0.95-0.5*[m])*[wu]*IF(AND([nd]>0,[nd]<4),{Vinc},1)*{s}/100/1000,'
                 '([wu]*MAX(0,[ln1]-[dNc])/100+{FCcm}*N([pbor]))*{s}/100/1000+[Vtip])'), 'calc', '0.000')
    v['VcR'] = (G('{FRv}*2*%s*%s^(1/3)*SQRT({fc})*[bwc]*[dNc]*IF([g678]="SÍ",{VcRinc},1)/1000' % (lam, pexp)), 'calc', '0.000')
    xper = 'IF([VcR]<=0,[ln1],[ln1]/2-[VcR]/MAX([Vu],1E-9)*([ln1]/2-[dNc]))'
    xvol = '([ln1]-(([VcR]-[Vtip])*1E5/{s}-{FCcm}*N([pbor]))*100/[wu])'
    # a lo largo del claro corto los casetones miden a_c y los separan nervaduras ‖ largo de b_w de DATOS
    # macizo de orilla del claro corto (el menor de los dos bordes; ya incluye la dala de borde de D9)
    oril = 'IF([dirc]="X",MIN([mx0],[mx1]),MIN([my0],[my1]))'
    v['nh'] = (G('IF([Vu]<=[VcR],0,CEILING(MAX(0,IF([tipo]="VOLADO",%s,%s)-%s)/{s},1))' % (xvol, xper, oril)), 'calc', '0')
    v['cort'] = (G('IF([Vu]<=[VcR],"CUMPLE",IF([nh]<=0,"CUMPLE (basta el macizo de orilla)",IF([nh]*{s}*IF([tipo]="VOLADO",1,2)>=[a1]*100-{bp},"NO CUMPLE: no alcanza ni macizando todo","MACIZAR "&[nh]&" hilera(s)")))'), 'key', '@')
    bwq = 'MIN({bw},[bwr])'
    v['yb'] = (G('({s}*{t}^2/2+%s*{hc}*({t}+{hc}/2))/({s}*{t}+%s*{hc})' % (bwq, bwq)), 'calc', '0.00')
    v['IT'] = (G('{s}*{t}^3/12+{s}*{t}*([yb]-{t}/2)^2+%s*{hc}^3/12+%s*{hc}*({t}+{hc}/2-[yb])^2' % (bwq, bwq)), 'calc', '#,##0')
    v['heq'] = (G('(12*[IT]/{s})^(1/3)'), 'calc', '0.00')
    v['afm'] = (G('IF([tipo]="PERIMETRAL",IF(ISNUMBER([afo]),[afo],{afm_def}),0)'), 'calc', '0.00')
    beta = '([ln2]/MAX([ln1],1))'; k_ = '(0.8+{fy}/14000)'
    hd = 'MAX(9,[ln2]*%s/(36+9*%s))' % (k_, beta)
    hb = 'MAX(13,[ln2]*%s/(36+5*%s*([afm]-0.2)))' % (k_, beta)
    ha = 'MAX(13,[ln2]/IF([nd]>0,30,33))'
    v['hmin'] = (G('IF([tipo]="PERIMETRAL",IF([afm]>2,%s,IF([afm]>0.2,%s,%s))*IF(AND([afm]>0.2,[nd]>0,{borde_viga}="NO"),1.1,1),MAX(25,[a1]*100/8*(0.4+{fy}/7000)))' % (hd, hb, ha)), 'calc', '0.0')
    v['peral'] = (G('IF(IF([tipo]="VOLADO",{h},[heq])>=[hmin],"CUMPLE",IF([tipo]="VOLADO","NO CUMPLE","CALCULAR FLECHA"))'), 'key', '@')
    v['tdia'] = (G('IF({t}>=IF(IF([tipo]="VOLADO",[a1],[a2])>=6,6,5),"CUMPLE","CAPA ≥ 6 cm")'), 'calc', '@')
    xc, dc_ = xpi('ln1', 'cPC', 'cNIC', 'cNDC', 'MNc', fac)
    xl, dl_ = xpi('ln2', 'cPL', 'cNIL', 'cNDL', 'MNl', fac)

    def ldsel(var, n, r_, bw_='[bwr]'):
        d_ = db(var)
        return 'IF(OR(%s<%s,AND(%s>=2,(%s-2*%s-%s*%s)/MAX(%s-1,1)<2*%s)),%s,%s)' % (r_, d_, n, bw_, r_, n, d_, n, d_, ldo(var), ld(var))
    # el M⁻ del volado baja dentro del tablero de respaldo con su PP (sección normal de DATOS), no con el del volado
    x0, Lb = volado_x0('(0.9*{PP}*{s}/100/1000)')
    # si con 0.9·PP el M⁻ no se anula dentro del respaldo, la corrida llega a su apoyo lejano y ahí se ancla con gancho, con su recubrimiento
    v['Lbc'] = (G('IF([tipo]="VOLADO",MIN(FLOOR(%s*100-{rec_np},5),CEILING(MAX(%s+MAX([dNc],12*%s,%s*100/16),%s),5)),IF(%s<0,CEILING([ln1],5),CEILING(MAX(%s+MAX([dNc],12*%s,[ln1]/16),%s),5)))'
                  % (Lb, x0, db('[vNc]'), Lb, ldsel('[vNc]', '[nNc]', '[rt]'), dc_, xc, db('[vNc]'), ldsel('[vNc]', '[nNc]', '[rt]'))), 'key', '0')
    v['Lbl'] = (G('IF([tipo]="VOLADO",0,IF(%s<0,CEILING([ln2],5),CEILING(MAX(%s+MAX([dNl],12*%s,[ln2]/16),%s),5)))'
                  % (dl_, xl, db('[vNl]'), ldsel('[vNl]', '[nNl]', '[rt]', '{bw}'))), 'key', '0')
    mr1 = lambda blk: '{FRf}*%s*{fy}*([d%s]-%s*{fy}/({fpc}*%s)/2)/1E5' % (ab('[v%s]' % blk), blk, ab('[v%s]' % blk), '[bwc]' if blk[1] == 'c' else '[bwcl]')
    v['bdc'] = (G('IF(AND([tipo]="PERIMETRAL",[nNc]>=2,[nd]-[ndc]>0,[cNDC]%s>%s),"SÍ","NO")' % (fac, mr1('Nc'))), 'calc', '@')
    v['bdl'] = (G('IF(AND([tipo]="PERIMETRAL",[nNl]>=2,[ndc]>0,[cNDL]%s>%s),"SÍ","NO")' % (fac, mr1('Nl'))), 'calc', '@')
    # D9 de Casa Rocha: dala de borde armada; la varilla inferior (1.25 fy, ≥ 15 cm dentro) y el bastón (fy) rematan con gancho de 180°, que debe caber parado (8·d_b)
    disp = '({b_borde}-{rec_np})'
    vi = 'MAX([vPc],[vPl])'; vs = 'MAX([vNc],[vNl])'
    wok = 'AND(1.25*%s<=%s,%s>=15,%s<=%s)' % (ldh(vi), disp, disp, ldh(vs), disp)
    gok = '8*%s<={h_borde}-2*{rec_np}' % db('MAX(%s,%s)' % (vi, vs))
    v['anc'] = (G('IF(OR([tipo]="VOLADO",[nd]=0),"—",IF(AND(%s,%s),"CUMPLE (gancho de 180°)",IF(%s,"NO CUMPLE: el gancho de 180° (8·d_b) no cabe en el peralte de la dala de borde",'
                  '"NO CUMPLE: dala de borde ≥ "&TEXT(CEILING(MAX(1.25*%s,15,%s)+{rec_np},5),"0")&" cm de ancho"&IF(%s,""," y más peralte para el gancho"))))'
                  % (wok, gok, wok, ldh(vi), ldh(vs), gok)), 'calc', '@')
    # casetón a_c (DATOS) a lo largo del claro corto y s − b_w (del renglón) a lo largo del largo: la retícula s no cambia
    acx = 'IF([dirc]="X",{ac},[acr])'; acy = 'IF([dirc]="Y",{ac},[acr])'
    # D9: en cada borde discontinuo de un perimetral la dala de borde ocupa b_borde desde la cara exterior del apoyo, b_borde − b_p más que en los demás
    ded = lambda b_: 'IF(AND([tipo]="PERIMETRAL",[%s]="DISCONTINUO"),MAX(0,{b_borde}-{bp}),0)' % b_
    for dX_, a_, s0_, s1_, ac_ in (('x', 'ax', 'bL', 'bR', acx), ('y', 'ay', 'bI', 'bS', acy)):
        v['nc' + dX_] = (G('MAX(0,INT(([%s]*100-{bp}-%s-%s+{s}-%s)/{s}))' % (a_, ded(s0_), ded(s1_), ac_)), 'calc', '0')
        left = '([%s]*100-{bp}-%s-%s-([nc%s]*%s+([nc%s]-1)*({s}-%s)))' % (a_, ded(s0_), ded(s1_), dX_, ac_, dX_, ac_)
        v['m%s0' % dX_] = (G('IF([nc%s]>0,%s+%s/2,[%s]*100/2)' % (dX_, ded(s0_), left, a_)), 'calc', '0.0')
        v['m%s1' % dX_] = (G('IF([nc%s]>0,%s+%s/2,[%s]*100/2)' % (dX_, ded(s1_), left, a_)), 'calc', '0.0')
    # respaldo de un volado con b_w propio (D10 de Casa Rocha): hasta dónde siguen sus nervaduras anchas dentro del tablero de respaldo
    wide = 'AND([tipo]="VOLADO",[bwr]>{bw})'
    dbv = db('[vNc]')
    # el respaldo puede no cumplir 6.7.8.1.4 (claro corto < 6·a_c): del lado seguro, su nervadura lleva r de viga
    v['rn1'] = (G('IF(NOT(%s),0,IF(%s<={bw}-2*{rec_np},1,0))' % (wide, dbv)), 'calc', '0')
    v['rd'] = (G('IF(NOT(%s),0,%s)' % (wide, red391('({h}-MAX({rec_np},({t}-2*{fim})/2)-2*{fim}-%s-%s/2)' % (db('MAX(4,{var_sup})'), dbv)))), 'calc', '0.00')
    As1 = '([rn1]*%s)' % ab('[vNc]'); c1 = '(%s*{fy}/({fpc}*[bwcl])/{beta1})' % As1
    v['rM1'] = (G('IF(OR(NOT(%s),[rn1]<1),0,IF(0.003*([rd]-%s)/%s<0.005,0,{FRf}*%s*{fy}*([rd]-%s*{fy}/({fpc}*[bwcl])/2)/1E5))' % (wide, c1, c1, As1, As1)), 'calc', '0.000')
    v['rV1'] = (G('IF(OR(NOT(%s),[rn1]<1),0,{FRv}*2*MIN(1,SQRT(2/(1+0.04*[rd])))*(%s/([bwcl]*[rd]))^(1/3)*SQRT({fc})*[bwcl]*[rd]/1000)' % (wide, As1)), 'calc', '0.000')
    Lbm = 'IF(ISNUMBER([Lb]),[Lb],[a1])'
    wD = '(0.9*{PP}*{s}/100/1000)'
    # sin el otro claro del respaldo capturado se toma m = 0 (losa en una dirección), del lado seguro
    mb = 'IF(N([Lb2])>0,MIN(%s,[Lb2])/MAX(%s,[Lb2]),0)' % (Lbm, Lbm)
    Lnb = '(%s*100-{bp})' % Lbm
    kV = '((0.95-0.5*%s)*{Vinc}*[wu]*{s}/100/1000)' % mb
    dVb = '([MNc]/(%s/100))' % Lnb
    A_ = '(%s/2)' % wD; B_ = '(%s*%s/2+[MNc]/%s)' % (wD, Lbm, Lbm); C_ = '([MNc]-[rM1])'
    v['rxM'] = (G('IF(NOT(%s),0,IF(%s<=0,0,IF([rM1]<=0,%s,IF(%s^2-4*%s*%s<0,%s,MIN(%s,(%s-SQRT(%s^2-4*%s*%s))/(2*%s))))))' % (wide, C_, Lbm, B_, A_, C_, Lbm, Lbm, B_, B_, A_, C_, A_)), 'calc', '0.000')
    v['rxV'] = (G('IF(NOT(%s),0,IF([rV1]<=%s,%s,MIN(%s,MAX(0,%s/200-([rV1]-%s)/%s))))' % (wide, dVb, Lbm, Lbm, Lnb, dVb, kV)), 'calc', '0.000')
    v['rVw'] = (G('IF(NOT(%s),0,%s+%s*MAX(0,%s/200-[dNc]/100))' % (wide, dVb, kV, Lnb)), 'calc', '0.000')
    v['rxr'] = (G('IF(NOT(%s),0,MIN(FLOOR(100*%s-{rec_np},5),CEILING(100*MIN(%s,MAX(IF([rxM]>0,[rxM]+IF([nNc]>=2,MAX([rd],12*%s)/100,0),0),[rxV],IF([nNc]>=2,%s/100,0))),5)))' % (wide, Lbm, Lbm, dbv, ld('[vNc]'))), 'calc', '0')
    v['rnh'] = (G('IF(NOT(%s),0,IF([rxr]<=0,0,CEILING([rxr]/{s},1)))' % wide), 'key', '0')
    v['rest'] = (G('IF([tipo]<>"VOLADO","—",IF([bwr]<={bw},"— (mismo b_w que el respaldo)",IF([rVw]>[VcR],"NO CUMPLE: cortante de la nervadura ancha junto al apoyo ("&TEXT([rVw],"0.00")&" > "&TEXT([VcR],"0.00")&" t)",'
                   'IF([rxr]>=100*%s,"CUMPLE — todas las nervaduras del respaldo ‖ al volado con b_w = "&[bwr]&" cm","CUMPLE — "&[rnh]&" hilera(s) del respaldo con nervaduras de b_w = "&[bwr]&" cm (casetón "&[acr]&" cm), hasta "&TEXT([rxr]/100,"0.00")&" m del paño"))))' % Lbm), 'key', '@')
    v['tPc'] = (G('IF([nPc]=0,"NO ALCANZA",[nPc]&" No. "&[vPc]&IF([nPc]>1," corridas"," corrida"))'), 'key', '@')
    v['tNc'] = (G('IF([nNc]=0,"NO ALCANZA",IF(AND([tipo]="VOLADO",[bwr]>{bw},[nNc]>=2),"1 No. "&[vNc]&" corrida de la orilla al respaldo ("&IF([Lbc]>=FLOOR(100*IF(ISNUMBER([Lb]),[Lb],[a1])-{rec_np},5),"hasta su apoyo lejano, con gancho","entra "&TEXT([Lbc]/100,"0.00")&" m desde el paño")&") + "&([nNc]-1)&" bastón No. "&[vNc]&" hasta "&TEXT([rxr]/100,"0.00")&" m del paño (fin de la nervadura ancha)",'
                  'IF([tipo]="VOLADO",[nNc]&" No. "&[vNc]&" de la orilla al respaldo ("&IF([Lbc]>=FLOOR(100*IF(ISNUMBER([Lb]),[Lb],[a1])-{rec_np},5),"hasta su apoyo lejano, con gancho","entran "&TEXT([Lbc]/100,"0.00")&" m desde el paño")&")",'
                  '"1 No. "&[vNc]&" corrida"&IF([nNc]>1," + "&([nNc]-1)&" bastón No. "&[vNc]&" L = "&TEXT([Lbc]/100,"0.00")&" m desde el paño",""))))'), 'key', '@')
    v['tPl'] = (G('IF([nPl]=0,"NO ALCANZA",[nPl]&" No. "&[vPl]&IF([nPl]>1," corridas"," corrida"))'), 'key', '@')
    v['tNl'] = (G('IF([nNl]=0,"NO ALCANZA",IF([tipo]="VOLADO","1 No. "&[vNl]&" corrida (repartición)","1 No. "&[vNl]&" corrida"&IF([nNl]>1," + "&([nNl]-1)&" bastón No. "&[vNl]&" L = "&TEXT([Lbl]/100,"0.00")&" m desde el paño","")))'), 'key', '@')
    fus = 'MAX([FUPc],[FUNc],[FUPl],[FUNl])'
    nn = 'OR(' + ','.join('[n%s]>MIN({nmax_sec},INT((%s+MAX({slib},%s))/(%s+MAX({slib},%s))))' % (b_, '(%s-2*%s)' % ('[bwr]' if b_[1] == 'c' else '{bw}', '[rt]' if b_[0] == 'N' else '[rb]'), db('[v%s]' % b_), db('[v%s]' % b_), db('[v%s]' % b_)) for b_ in BLKS) + ')'
    etl = 'OR(' + ','.join('AND(ISNUMBER([e%s]),[e%s]<0.005)' % (b_, b_) for b_ in BLKS) + ')'
    sfx = ('IF(LEFT([c342],5)="FUERA"," · fuera de 3.4.2.1 c)","")&IF(LEFT([anc],2)="NO"," · REVISAR anclaje en borde discontinuo","")'
           '&IF([rnh]>0," · respaldo: "&[rnh]&" hilera(s) de b_w "&[bwr]&" cm","")')
    rsp = 'LEFT([rest],2)="NO"'
    v['estado'] = ('=IF([key]="","",IF([ok]<>"SÍ","FALTAN DATOS: tipo, a_x, a_y, destino y, en volados, L_b",IF([col]="NO MONOLÍTICO (caso II)","FUERA DE ALCANCE: caso II no transcrito",'
                   'IF(AND([tipo]="VOLADO",[emp]="?"),"REVISAR: el volado lleva UN solo borde CONTINUO (el empotrado)",'
                   'IF(MIN([ncx],[ncy])=0,"REVISAR: no cabe un casetón (tablero macizo): diséñalo en MACIZAS",'
                   'IF(AND(%s<=1,NOT(%s),NOT(%s),NOT(%s),[peral]="CUMPLE",[tdia]="CUMPLE",LEFT([cort],2)<>"NO",LEFT({capa},6)="CUMPLE"),IF(LEFT([cort],6)="CUMPLE","CUMPLE","CUMPLE — "&[cort])&%s,'
                   'IF(AND(%s<=1,NOT(%s),NOT(%s),NOT(%s),[peral]="CALCULAR FLECHA",LEFT([cort],2)<>"NO",LEFT({capa},6)="CUMPLE"),"REVISAR: calcular flecha (h_eq < h mín.)"&%s,'
                   '"NO CUMPLE: "&IF(%s>1,"flexión ","")&IF(%s,"no caben ","")&IF(%s,"ε_t < 0.005 ","")&IF(LEFT([cort],2)="NO","cortante ","")&IF([peral]="NO CUMPLE","peralte ","")&IF([tdia]<>"CUMPLE","capa < 6 cm ","")&IF(LEFT({capa},6)<>"CUMPLE","capa o malla (DATOS sección 7) ","")&IF(%s,"respaldo del volado",""))))))))'
                   % (fus, nn, etl, rsp, sfx, fus, nn, etl, rsp, sfx, fus, nn, etl, rsp), 'key', '@')
    # con 0.9·PP y sin CV, si el M⁻ del volado no se anula en el respaldo su apoyo lejano levanta
    Rf = '(%s*%s/2-[MNc]/%s)' % (wD, Lbm, Lbm)
    upl = ('&IF(AND([tipo]="VOLADO",%s<0),"Con 0.9·PP y sin carga viva el apoyo lejano del respaldo levanta (R = "&TEXT(-%s,"0.00")&" t por nervadura): '
           'ancla ahí la corrida con gancho y revisa ese apoyo a tensión. ","")' % (Rf, Rf))
    v['nota'] = (G('IF(LEFT([key],7)="EJEMPLO","EJEMPLO de la hoja de Casa Rocha: bórralo o sobrescríbelo. ","")'
                   '&IF(AND([tipo]="PERIMETRAL",MIN([ncx],[ncy])=0),"No cabe un casetón en el claro corto: es un tablero macizo, diséñalo en MACIZAS. ","")'
                   '&IF(AND([g678]="NO",[tipo]="PERIMETRAL"),"No cumple 6.7.8.1.4 (a_c = "&MAX({ac},[acr])&" cm > claro corto/6 = "&TEXT([a1]*100/6,"0")&" cm): se diseña como losa apoyada en vigas (6.7.8.1.9): r de viga = "&{rec_np}&" cm y V_cR sin ×1.1. Para que cuente como encasetonada: casetón de "&TEXT(FLOOR(MIN(80,[a1]*100/6),1),"0")&" cm o menor. ","")'
                   '&IF([tipo]="VOLADO","Volado de "&TEXT([a1],"0.00")&" m empotrado en el borde "&LOWER([emp])&IF(N([pbor])>0,", con "&TEXT([pbor],"0.0")&" kg/m en la orilla","")&". Lecho inferior con r_ext. Si la orilla carga pretil u otro elemento frágil, calcula la flecha (13.4). El acero superior entra al respaldo de "&TEXT(IF(ISNUMBER([Lb]),[Lb],[a1]),"0.00")&" m"&" hasta pasar M = 0 con 0.9·PP sin carga viva; si ahí no se anula, llega al apoyo lejano y se ancla con gancho (NTC-Criterios 3.4.1 c y 5.1.2.2). "&IF([bwr]>{bw},"Respaldo: "&[rest]&". ",""),"")'
                   + upl +
                   '&IF(LEFT([c342],5)="FUERA","El M⁻ del vecino difiere más de 50 % (3.4.2.1 c): distribuye el desequilibrio (3.4.3.1) o confirma con un modelo; se arma con el mayor. ","")'
                   '&IF(OR([bdc]="SÍ",[bdl]="SÍ"),"El M⁻ del borde discontinuo pide bastón también ahí (con gancho). ","")'
                   '&IF(LEFT([anc],2)="NO","Anclaje: "&[anc]&" (6.7.8.1.7). ","")'
                   '&IF(AND([tipo]="PERIMETRAL",[ndc]>=1,[nd]-[ndc]>=1),"Esquinas exteriores sobre muro: refuerzo de esquina arriba y abajo en L_largo/5 (6.7.7.3.1). ","")'), 'calc', '@')
    return v


EX_N = [
    dict(key='EJEMPLO F80014 (interior)', tipo='PERIMETRAL', dest='HABITACIÓN', ax=3.029, ay=1.6, bI='CONTINUO', bS='CONTINUO', bL='CONTINUO', bR='CONTINUO', cforce='AUTO', col='MONOLÍTICO (caso I)', scm=230),
    dict(key='EJEMPLO F80016 (de borde)', tipo='PERIMETRAL', dest='HABITACIÓN', ax=4.426, ay=4.0, bI='CONTINUO', bS='CONTINUO', bL='DISCONTINUO', bR='CONTINUO', cforce='AUTO', col='MONOLÍTICO (caso I)', scm=160),
    dict(key='EJEMPLO F80024 (esquina, azotea)', tipo='PERIMETRAL', dest='AZOTEA (pendiente ≤ 5 %)', ax=4.327, ay=4.95, bI='CONTINUO', bS='DISCONTINUO', bL='DISCONTINUO', bR='CONTINUO', cforce='AUTO', col='MONOLÍTICO (caso I)', scm=315),
    dict(key='EJEMPLO F80012 (volado)', tipo='VOLADO', dest='HABITACIÓN', ax=2.0, ay=3.425, bI='DISCONTINUO', bS='DISCONTINUO', bL='DISCONTINUO', bR='CONTINUO', cforce='AUTO', col='MONOLÍTICO (caso I)', scm=210,
         pbor=333.7, Lb=3.029, bwo=20),
]


def sheet_nervadas(wb):
    ws = wb.create_sheet('NERVADAS')
    write_table(ws, NC, COLS_N, GROUPS_N, NOTES_N, vals_nervadas(), INP_N, EX_N,
                'NERVADAS — LOSA NERVADA (ENCASETONADA), UN RENGLÓN POR TABLERO · MÉTODO DE COEFICIENTES NTC-CONCRETO 2023, TABLA 3.4.2.1 (CASO I)',
                'Momentos por nervadura = coef·10⁻⁴·w_u·a1²·s. Abajo trabaja la sección T (patín b_f); arriba sólo la nervadura (b_w de cálculo). Toda nervadura lleva al menos 1 varilla corrida abajo y 1 arriba (6.7.6.1.3); '
                'arriba las demás son bastones. Sólo se captura lo AMARILLO; los renglones vacíos quedan en blanco. La memoria paso a paso de un tablero está en DETALLE y su dibujo en CROQUIS.',
                'ESTADO: «CUMPLE» = flexión, cupo, ε_t ≥ 0.005, cortante, peralte, capa y anclaje bien. «MACIZAR n hilera(s)» = quitar casetones junto a cada apoyo por cortante. «REVISAR: calcular flecha» = h_eq menor que el mínimo de la Tabla 6.7.3.1.2. '
                '«No cumple 6.7.8.1.4» (nota) = a_c > claro corto/6: losa apoyada en vigas (6.7.8.1.9). «FUERA de 3.4.2.1 c)» = el M⁻ del vecino capturado difiere más de 50 %.',
                'nota')
    return ws


# =====================================================================
INP_M = ['key', 'tipo', 'dest', 'ax', 'ay', 'bI', 'bS', 'bL', 'bR', 'cforce', 'col', 'ho', 'scm', 'cvo', 'pbor', 'Lb', 'afo', 'varo', 'Mvc', 'Mvl']
SPEC_M = [
    ('key', 'Tablero', 22, ''), ('tipo', 'Tipo', 11.5, ''), ('dest', 'Destino (lista de DATOS)', 19, ''), ('ax', 'a_x — claro en X a ejes', 7.5, 'm'), ('ay', 'a_y — claro en Y a ejes', 7.5, 'm'),
    ('bI', 'Borde inferior (‖ X)', 11, 'CONT./DISC.'), ('bS', 'Borde superior (‖ X)', 11, 'CONT./DISC.'), ('bL', 'Borde izquierdo (‖ Y)', 11, 'CONT./DISC.'), ('bR', 'Borde derecho (‖ Y)', 11, 'CONT./DISC.'),
    ('cforce', 'Forzar caso', 12, 'AUTO = sale solo'), ('col', 'Colado con los apoyos', 13, 'caso I / II'), ('ho', 'h propio (vacío = DATOS)', 7, 'cm'), ('scm', 'SCM — sobrecarga muerta', 7.5, 'kg/m²'),
    ('cvo', 'CV propia (opcional)', 7.5, 'kg/m²'), ('pbor', 'Carga en la orilla del volado', 8, 'kg/m serv.'), ('Lb', 'L_b — claro de respaldo (volado)', 8, 'm'), ('afo', 'α_fm propio (opcional)', 7, '—'),
    ('varo', 'No. de varilla preferido (opcional)', 7, 'No.'), ('Mvc', 'M⁻ del vecino ‖ corto (opcional)', 8.5, 't·m/m'), ('Mvl', 'M⁻ del vecino ‖ largo (opcional)', 8.5, 't·m/m'),
    ('ok', '¿Datos completos?', 7.5, ''), ('hr', 'h usado', 6, 'cm'), ('emp', 'Borde empotrado (volado)', 10, ''), ('dirc', 'Claro corto ‖', 5.5, ''), ('a1', 'a1 corto (volado: largo)', 7, 'm'),
    ('a2', 'a2 largo', 7, 'm'), ('m', 'm = a1/a2', 6.5, ''), ('ln1', 'L_n corto', 6.5, 'cm'), ('ln2', 'L_n largo', 6.5, 'cm'), ('rb', 'r inferior', 6, 'cm'), ('rt', 'r superior', 6, 'cm'),
    ('azo', '¿Azotea?', 6.5, ''), ('cv', 'CV W_m', 6.5, 'kg/m²'), ('pp', 'PP de diseño', 7, 'kg/m²'), ('cm', 'CM', 6.5, 'kg/m²'), ('wu', 'w_u', 7, 'kg/m²'),
    ('nd', 'Bordes discont.', 6.5, ''), ('ndc', '… cortos', 6, ''), ('cauto', 'Caso automático', 14, ''), ('caso', 'Caso usado', 14, ''),
    ('idx', 'col. m', 5, ''), ('fr', 'interp.', 5.5, ''), ('cNIC', 'M⁻ cont. corto', 7, '×10⁻⁴'), ('cNIL', 'M⁻ cont. largo', 7, '×10⁻⁴'), ('cNDC', 'M⁻ disc. corto', 7, '×10⁻⁴'),
    ('cNDL', 'M⁻ disc. largo', 7, '×10⁻⁴'), ('cPC', 'M⁺ corto', 7, '×10⁻⁴'), ('cPL', 'M⁺ largo', 7, '×10⁻⁴'),
    ('MPc', 'M⁺ corto', 7, 't·m/m'), ('MNc', 'M⁻ corto', 7, 't·m/m'), ('MPl', 'M⁺ largo', 7, 't·m/m'), ('MNl', 'M⁻ largo', 7, 't·m/m'), ('c342', '3.4.2.1 c) con el vecino', 15, ''),
    ('Asm', 'As mín = 0.0018·b·h', 7, 'cm²/m'), ('smax', 'Separación máx.', 6.5, 'cm'),
]
for _b in BLKS:
    SPEC_M += [('dq' + _b, 'd para As (var. preferida)', 7, 'cm'), ('Asf' + _b, 'As flexión', 6.5, 'cm²/m'), ('As' + _b, 'As req. (≥ mín.)', 6.5, 'cm²/m'), ('v' + _b, 'Var. No.', 5.5, ''),
               ('sp' + _b, 'Separación', 6.5, 'cm'), ('Asc' + _b, 'As colocada', 6.5, 'cm²/m'), ('d' + _b, 'd real', 6, 'cm'), ('MR' + _b, 'M_R', 6.5, 't·m/m'), ('FU' + _b, 'Mu/M_R', 6.5, '≤ 1'), ('e' + _b, 'ε_t', 6.5, '≥ 0.005')]
SPEC_M += [('Vu', 'V_u', 7, 't/m'), ('p', 'p superior en el apoyo', 7, ''), ('VcR', 'V_cR', 7, 't/m'), ('cort', 'Estado cortante', 15, ''),
           ('afm', 'α_fm', 6, ''), ('hmin', 'h mín. (Tabla 6.7.3.1.2 / volado L/10)', 8.5, 'cm'), ('peral', 'Estado peralte', 13, ''),
           ('Lbc', 'Bastón ‖ corto desde el paño', 8, 'cm'), ('Lbl', 'Bastón ‖ largo desde el paño', 8, 'cm'), ('bdc', 'Bastón también en borde disc. ‖ corto', 8, ''), ('bdl', 'Bastón también en borde disc. ‖ largo', 8, ''),
           ('anc', 'Anclaje del bastón en bordes discontinuos', 16, ''),
           ('tPc', 'ARMADO inferior ‖ corto', 16, ''), ('tNc', 'ARMADO superior ‖ corto', 28, ''), ('tPl', 'ARMADO inferior ‖ largo', 16, ''), ('tNl', 'ARMADO superior ‖ largo', 28, ''),
           ('estado', 'ESTADO', 30, ''), ('nota', 'NOTA', 70, '')]
MC_, COLS_M = _mkcols(SPEC_M)
MC.update(MC_)
GROUPS_M = [('key', 'Mvl', 'DATOS DEL TABLERO — CAPTURA (amarillo)'), ('ok', 'rt', 'GEOMETRÍA'), ('azo', 'wu', 'CARGAS'), ('nd', 'caso', 'CASO DE LA TABLA'),
            ('idx', 'cPL', 'COEFICIENTES ×10⁻⁴ (Tabla 3.4.2.1)'), ('MPc', 'smax', 'MOMENTOS POR METRO Y ACERO MÍNIMO'),
            ('dqPc', 'ePc', 'LECHO INFERIOR ‖ CLARO CORTO'), ('dqNc', 'eNc', 'LECHO SUPERIOR ‖ CLARO CORTO'), ('dqPl', 'ePl', 'LECHO INFERIOR ‖ CLARO LARGO'), ('dqNl', 'eNl', 'LECHO SUPERIOR ‖ CLARO LARGO'),
            ('Vu', 'cort', 'CORTANTE'), ('afm', 'peral', 'ESPESOR MÍNIMO'), ('Lbc', 'anc', 'BASTONES Y ANCLAJE'), ('tPc', 'nota', 'RESULTADO')]
NOTES_M = {
    'key': 'CAPTURA. Claros A EJES. Borde CONTINUO = sigue losa del otro lado; DISCONTINUO = orilla, hueco o losa que no sigue. VOLADO: marca CONTINUO sólo el borde empotrado. '
           'h propio, CV propia, α_fm y No. de varilla: vacío = el de DATOS. M⁻ del vecino por metro, en el borde común; la hoja arma con el mayor.',
    'ok': 'Franja de 1 m. a1 = claro corto, a2 = largo (volado: a1 = largo perpendicular al borde empotrado). L_n = a − b_p (volado a − b_p/2). r inferior: r (r_ext en volados); r superior: r_ext si la cara queda expuesta (DATOS).',
    'azo': 'PP = h·γc + 20 + 20 (NTC-Criterios 5.1.2.1, p. 9). CM = PP + SCM. w_u = máx(1.3·CM + 1.5·W_m; azotea: 1.1·(CM + W_g)) (NTC-Criterios 3.4.1 a) y b), p. 7; ec. 7.3.2.1, p. 17).',
    'nd': 'Se cuentan los bordes DISCONTINUOS; «corto» = borde de largo a1. Tabla 3.4.2.1, p. 65 (mismo criterio que nervadas).',
    'idx': 'Interpolación lineal en m entre las columnas de la Tabla 3.4.2.1, caso I (TABLAS sección A).',
    'MPc': 'M = coef·10⁻⁴·w_u·a1² por metro (NTC-Concreto 3.4.2.1, p. 64). M⁻ = el mayor de borde continuo, discontinuo y el del vecino capturado. Volado: M⁻ = w_u·L_n²/2 + 1.3·P·L_n. '
           'As mín = 0.0018·100·h en cada lecho (6.7.6.1.1, p. 175; 13.8.2.1.1, p. 345). Separación máx. = criterio de esta hoja (DATOS sección 8).',
    'dqPc': 'b = 100 cm. d = h − r − d_b/2 (preferida para As, real para M_R). As = q·f\'\'c·b·d/fy (ecs. 5.2.2.1.1.1, p. 89), ≥ As mín; −1 = no alcanza o rebasa p_max (ε_t ≥ 0.005). '
            'Varilla: la preferida o la siguiente cuya separación (redondeada hacia abajo) no baje de la mínima; separación ≤ máx. ε_t ≥ 0.005 (Tabla 3.8.2.2, p. 70).',
    'dqNc': 'd = h − r − d_b/2, − 2 cm si d ≤ 20 cm sin la precaución de 3.9.2 (3.9.1 d), p. 74). Perimetral: bastones en los bordes; volado: acero principal de la orilla al respaldo.',
    'dqPl': 'El lecho del claro largo va encima del corto: d = h − r − d_b,corto − d_b/2.',
    'dqNl': 'Va debajo del lecho superior corto. En volados es el acero de repartición (As mín).',
    'Vu': 'V_u a d del paño = (L_n/2 − d)(0.95 − 0.5m)·w_u, ×1.15 con bordes continuos y discontinuos (3.4.1.1, p. 64; 6.7.4.3.2, p. 172). Volado: w_u(L_n − d) + 1.3P. '
          'V_cR = F_R·2·λs·p^⅓·√f\'c·b·d con b = 100 cm y p del acero superior (ec. 5.5.3.2.1, p. 95). Losa sin estribos: si no alcanza, sube h.',
    'afm': 'Tabla 6.7.3.1.2 (p. 168) con α_fm (igual que nervadas), contra h. Volado: L/10 (criterio de esta hoja, DATOS sección 8).',
    'Lbc': 'Bastón desde el paño = punto de inflexión + máx(d, 12d_b, L_n/16), ≥ L_d (6.3.7.3.4, p. 127; 6.3.7.3.10.4, p. 129; Tabla 14.4.2.4, p. 353). Volado: entra al respaldo hasta M = 0 con 0.9·PP y sin CV. '
           'Anclaje en bordes discontinuos (dala de borde armada): L_dh del bastón ≤ ancho − r_np y ≥ 15 cm para la varilla inferior (6.7.7.4.1.1, p. 179); gancho de 180° con alto 8·d_b ≤ peralte de la dala − 2·r_np (Tabla 14.3.1, p. 349).',
    'tPc': 'Armado por metro: «No. 3 @ 20 cm». Superior perimetral: bastones en los bordes continuos (y en los discontinuos con M⁻ de tabla, con gancho). ESTADO junta flexión, separación, ε_t, cortante y peralte.',
}


def vals_macizas():
    v = {}
    v['ok'] = ('=IF([key]="","",IF(AND(OR([tipo]="PERIMETRAL",[tipo]="VOLADO"),ISNUMBER([ax]),ISNUMBER([ay]),[ax]>0,[ay]>0,OR(ISNUMBER([cvo]),ISNUMBER(MATCH([dest],{CV_NAMES},0))),OR([tipo]<>"VOLADO",N([Lb])>0)),"SÍ","NO"))', 'calc', '@')
    v['hr'] = (G('IF(ISNUMBER([ho]),[ho],{h_mac})'), 'calc', '0.0')
    common_geo(v, 'M')
    v['rb'] = (G('IF([tipo]="VOLADO",{rec_ext},{rec})'), 'calc', '0.0')
    v['rt'] = (G('IF(OR(AND([azo]="SÍ",{exp_az}="SÍ"),AND([tipo]="VOLADO",{exp_vol}="SÍ")),{rec_ext},{rec})'), 'calc', '0.0')
    v['pp'] = (G('[hr]/100*{gc}*1000+{inc_losa}+{inc_mort}'), 'calc', '0')
    fac = '*[wu]*[a1]^2*1E-4/1000'
    v['MPc'] = (G('IF([tipo]="PERIMETRAL",[cPC]%s,0)' % fac), 'calc', '0.000')
    v['MNc'] = (G('IF([tipo]="PERIMETRAL",MAX(MAX([cNIC],[cNDC])%s,N([Mvc])),([wu]*([ln1]/100)^2/2+{FCcm}*N([pbor])*[ln1]/100)/1000)' % fac), 'key', '0.000')
    v['MPl'] = (G('IF([tipo]="PERIMETRAL",[cPL]%s,0)' % fac), 'calc', '0.000')
    v['MNl'] = (G('IF([tipo]="PERIMETRAL",MAX(MAX([cNIL],[cNDL])%s,N([Mvl])),0)' % fac), 'calc', '0.000')
    v['c342'] = (c342_formula(fac), 'calc', '@')
    v['Asm'] = (G('{pmac}*100*[hr]'), 'calc', '0.00')
    v['smax'] = (G('MIN(IF([tipo]="VOLADO",{kotr},{kcrit})*[hr],{stope})'), 'calc', '0.0')
    pref = 'IF(ISNUMBER([varo]),[varo],{var_mac})'
    dq = {'Pc': '([hr]-[rb]-%s/2)' % db(pref), 'Pl': '([hr]-[rb]-1.5*%s)' % db(pref),
          'Nc': red391('([hr]-[rt]-%s/2)' % db(pref)), 'Nl': red391('([hr]-[rt]-1.5*%s)' % db(pref))}
    dact = {'Pc': '[hr]-[rb]-%s/2' % db('[vPc]'), 'Pl': '[hr]-[rb]-%s-%s/2' % (db('[vPc]'), db('[vPl]')),
            'Nc': red391('([hr]-[rt]-%s/2)' % db('[vNc]')), 'Nl': red391('([hr]-[rt]-%s-%s/2)' % (db('[vNc]'), db('[vNl]')))}
    for blk in BLKS:
        M = '[%s]' % MOMOF[blk]
        dqr = '[dq%s]' % blk
        v['dq' + blk] = (G(dq[blk]), 'calc', '0.00')
        v['Asf' + blk] = (G(asflex(M, '100', dqr)), 'calc', '0.00')
        v['As' + blk] = (G('IF([Asf%s]<0,-1,MAX([Asf%s],[Asm]))' % (blk, blk)), 'calc', '0.00')
        v['v' + blk] = (G(macsel('[As%s]' % blk, pref, '[smax]')), 'calc', 'General')
        ab_ = ab('[v%s]' % blk)
        v['sp' + blk] = (G('IF([As%s]<=0,0,FLOOR(MIN([smax],100*%s/[As%s]),{spaso}))' % (blk, ab_, blk)), 'key', '0.0')
        v['Asc' + blk] = (G('IF([sp%s]<=0,0,%s*100/[sp%s])' % (blk, ab_, blk)), 'calc', '0.00')
        v['d' + blk] = (G(dact[blk]), 'calc', '0.00')
        v['MR' + blk] = (G('IF([Asc%s]=0,0,{FRf}*[Asc%s]*{fy}*([d%s]-[Asc%s]*{fy}/({fpc}*100)/2)/1E5)' % (blk, blk, blk, blk)), 'calc', '0.000')
        v['FU' + blk] = (G('IF([As%s]<0,9.99,IF([MR%s]>0,%s/[MR%s],0))' % (blk, blk, M, blk)), 'key', '0.00')
        cexp = '([Asc%s]*{fy}/({fpc}*100)/{beta1})' % blk
        v['e' + blk] = (G('IF([Asc%s]=0,"—",0.003*([d%s]-%s)/%s)' % (blk, blk, cexp, cexp)), 'calc', '0.0000')
    v['Vu'] = (G('IF([tipo]="PERIMETRAL",MAX(0,([ln1]/2-[dNc])/100)*(0.95-0.5*[m])*[wu]*IF(AND([nd]>0,[nd]<4),{Vinc},1)/1000,([wu]*MAX(0,[ln1]-[dNc])/100+{FCcm}*N([pbor]))/1000)'), 'calc', '0.000')
    v['p'] = (G('[AscNc]/(100*[dNc])'), 'calc', '0.0000')
    v['VcR'] = (G('{FRv}*2*MIN(1,SQRT(2/(1+0.04*[dNc])))*[p]^(1/3)*SQRT({fc})*100*[dNc]/1000'), 'calc', '0.000')
    v['cort'] = (G('IF([Vu]<=[VcR],"CUMPLE","NO CUMPLE: sube h")'), 'key', '@')
    v['afm'] = (G('IF([tipo]="PERIMETRAL",IF(ISNUMBER([afo]),[afo],{afm_def}),0)'), 'calc', '0.00')
    beta = '([ln2]/MAX([ln1],1))'; k_ = '(0.8+{fy}/14000)'
    hd = 'MAX(9,[ln2]*%s/(36+9*%s))' % (k_, beta)
    hb = 'MAX(13,[ln2]*%s/(36+5*%s*([afm]-0.2)))' % (k_, beta)
    ha = 'MAX(13,[ln2]/IF([nd]>0,30,33))'
    v['hmin'] = (G('IF([tipo]="PERIMETRAL",IF([afm]>2,%s,IF([afm]>0.2,%s,%s))*IF(AND([afm]>0.2,[nd]>0,{borde_viga}="NO"),1.1,1),[a1]*100/{kvol}*(0.4+{fy}/7000))' % (hd, hb, ha)), 'calc', '0.0')
    v['peral'] = (G('IF([hr]>=[hmin]-0.001,"CUMPLE",IF([tipo]="VOLADO","NO CUMPLE","CALCULAR FLECHA"))'), 'key', '@')
    xc, dc_ = xpi('ln1', 'cPC', 'cNIC', 'cNDC', 'MNc', fac)
    xl, dl_ = xpi('ln2', 'cPL', 'cNIL', 'cNDL', 'MNl', fac)

    def ldsel(blk):
        d_ = db('[v%s]' % blk)
        return 'IF(OR([rt]<%s,[sp%s]-%s<2*%s),%s,%s)' % (d_, blk, d_, d_, ldo('[v%s]' % blk), ld('[v%s]' % blk))
    x0, Lb = volado_x0('(0.9*[hr]/100*{gc})')
    # si con 0.9·PP el M⁻ no se anula dentro del respaldo, el acero llega a su apoyo lejano (y ahí se ancla con gancho)
    v['Lbc'] = (G('IF([tipo]="VOLADO",MIN(FLOOR(%s*100-{rec_np},5),CEILING(MAX(%s+MAX([dNc],12*%s,%s*100/16),%s),5)),IF(%s<0,CEILING([ln1],5),CEILING(MAX(%s+MAX([dNc],12*%s,[ln1]/16),%s),5)))'
                  % (Lb, x0, db('[vNc]'), Lb, ldsel('Nc'), dc_, xc, db('[vNc]'), ldsel('Nc'))), 'key', '0')
    v['Lbl'] = (G('IF([tipo]="VOLADO",0,IF(%s<0,CEILING([ln2],5),CEILING(MAX(%s+MAX([dNl],12*%s,[ln2]/16),%s),5)))' % (dl_, xl, db('[vNl]'), ldsel('Nl'))), 'key', '0')
    v['bdc'] = (G('IF(AND([tipo]="PERIMETRAL",[nd]-[ndc]>0,[cNDC]>0),"SÍ","NO")'), 'calc', '@')
    v['bdl'] = (G('IF(AND([tipo]="PERIMETRAL",[ndc]>0,[cNDL]>0),"SÍ","NO")'), 'calc', '@')
    # bastón superior (fy) y varilla inferior (≥ 15 cm dentro) en la dala de borde, con gancho de 180° que quepa parado (8·d_b)
    disp = '({b_borde}-{rec_np})'
    vs = 'MAX([vNc],[vNl])'
    wok = 'AND(%s<=%s,%s>=15)' % (ldh(vs), disp, disp)
    gok = '8*%s<={h_borde}-2*{rec_np}' % db('MAX([vPc],[vPl],[vNc],[vNl])')
    v['anc'] = (G('IF(OR([tipo]="VOLADO",[nd]=0),"—",IF(AND(%s,%s),"CUMPLE (gancho de 180°)",IF(%s,"NO CUMPLE: el gancho de 180° (8·d_b) no cabe en el peralte de la dala de borde",'
                  '"NO CUMPLE: dala de borde ≥ "&TEXT(CEILING(MAX(%s,15)+{rec_np},5),"0")&" cm de ancho"&IF(%s,""," y más peralte para el gancho"))))'
                  % (wok, gok, wok, ldh(vs), gok)), 'calc', '@')
    at = lambda b_: '"No. "&[v%s]&" @ "&[sp%s]&" cm"' % (b_, b_)
    v['tPc'] = (G('IF([sp%s]=0,"NO ALCANZA",%s&IF([tipo]="VOLADO"," (temperatura)",""))' % ('Pc', at('Pc'))), 'key', '@')
    v['tNc'] = (G('IF([spNc]=0,"NO ALCANZA",IF([tipo]="VOLADO",%s&" de la orilla al respaldo ("&IF([Lbc]>=FLOOR(100*IF(ISNUMBER([Lb]),[Lb],[a1])-{rec_np},5),"hasta su apoyo lejano, con gancho","entran "&TEXT([Lbc]/100,"0.00")&" m desde el paño")&")","Bastones "&%s&", L = "&TEXT([Lbc]/100,"0.00")&" m desde el paño"))' % (at('Nc'), at('Nc'))), 'key', '@')
    v['tPl'] = (G('IF([spPl]=0,"NO ALCANZA",%s&IF([tipo]="VOLADO"," (repartición)",""))' % at('Pl')), 'key', '@')
    v['tNl'] = (G('IF([spNl]=0,"NO ALCANZA",IF([tipo]="VOLADO",%s&" (repartición)","Bastones "&%s&", L = "&TEXT([Lbl]/100,"0.00")&" m desde el paño"))' % (at('Nl'), at('Nl'))), 'key', '@')
    fus = 'MAX([FUPc],[FUNc],[FUPl],[FUNl])'
    sepf = 'OR(' + ','.join('[sp%s]<{smin}' % b_ for b_ in BLKS) + ')'
    etl = 'OR(' + ','.join('AND(ISNUMBER([e%s]),[e%s]<0.005)' % (b_, b_) for b_ in BLKS) + ')'
    sfx = 'IF(LEFT([c342],5)="FUERA"," · fuera de 3.4.2.1 c)","")&IF(LEFT([anc],2)="NO"," · REVISAR anclaje en borde discontinuo","")'
    v['estado'] = ('=IF([key]="","",IF([ok]<>"SÍ","FALTAN DATOS: tipo, a_x, a_y, destino y, en volados, L_b",IF([col]="NO MONOLÍTICO (caso II)","FUERA DE ALCANCE: caso II no transcrito",'
                   'IF(AND([tipo]="VOLADO",[emp]="?"),"REVISAR: el volado lleva UN solo borde CONTINUO (el empotrado)",'
                   'IF(AND(%s<=1,NOT(%s),NOT(%s),[cort]="CUMPLE",[peral]="CUMPLE"),"CUMPLE"&%s,'
                   'IF(AND(%s<=1,NOT(%s),NOT(%s),[cort]="CUMPLE",[peral]="CALCULAR FLECHA"),"REVISAR: calcular flecha (h < h mín.)"&%s,'
                   '"NO CUMPLE: "&IF(%s>1,"flexión ","")&IF(%s,"separación < mínima ","")&IF(%s,"ε_t < 0.005 ","")&IF([cort]<>"CUMPLE","cortante ","")&IF([peral]="NO CUMPLE","peralte","")))))))'
                   % (fus, sepf, etl, sfx, fus, sepf, etl, sfx, fus, sepf, etl), 'key', '@')
    Lbm = 'IF(ISNUMBER([Lb]),[Lb],[a1])'; Rf = '(0.9*[hr]/100*{gc}*%s/2-[MNc]/%s)' % (Lbm, Lbm)
    upl = ('&IF(AND([tipo]="VOLADO",%s<0),"Con 0.9·PP y sin carga viva el apoyo lejano del respaldo levanta (R = "&TEXT(-%s,"0.00")&" t/m): '
           'ancla ahí el acero con gancho y revisa ese apoyo a tensión. ","")' % (Rf, Rf))
    v['nota'] = (G('IF(LEFT([key],7)="EJEMPLO","EJEMPLO de tablero típico: bórralo o sobrescríbelo. ","")'
                   '&IF([tipo]="VOLADO","Volado de "&TEXT([a1],"0.00")&" m empotrado en el borde "&LOWER([emp])&IF(N([pbor])>0,", con "&TEXT([pbor],"0.0")&" kg/m en la orilla","")&". Lecho inferior con r_ext. Si la orilla carga pretil u otro elemento frágil, calcula la flecha (13.4). El acero superior entra al respaldo de "&TEXT(IF(ISNUMBER([Lb]),[Lb],[a1]),"0.00")&" m"&" hasta pasar M = 0 con 0.9·PP sin carga viva; si ahí no se anula, llega al apoyo lejano y se ancla con gancho. ","")'
                   + upl +
                   '&IF(LEFT([c342],5)="FUERA","El M⁻ del vecino difiere más de 50 % (3.4.2.1 c): distribuye el desequilibrio (3.4.3.1) o confirma con un modelo; se arma con el mayor. ","")'
                   '&IF(OR([bdc]="SÍ",[bdl]="SÍ"),"Bastones también en los bordes discontinuos (M⁻ de tabla), con gancho. ","")'
                   '&IF(LEFT([anc],2)="NO","Anclaje: "&[anc]&" (6.7.7.4.1.1). ","")'
                   '&IF(AND([tipo]="PERIMETRAL",[ndc]>=1,[nd]-[ndc]>=1),"Esquinas exteriores sobre muro: refuerzo de esquina arriba y abajo en L_largo/5 (6.7.7.3.1). ","")'), 'calc', '@')
    return v


EX_M = [
    dict(key='EJEMPLO M-1 (interior 3.0 × 4.0)', tipo='PERIMETRAL', dest='HABITACIÓN', ax=3.0, ay=4.0, bI='CONTINUO', bS='CONTINUO', bL='CONTINUO', bR='CONTINUO', cforce='AUTO', col='MONOLÍTICO (caso I)', ho=12, scm=160),
    dict(key='EJEMPLO M-2 (esquina 2.0 × 3.4)', tipo='PERIMETRAL', dest='HABITACIÓN', ax=2.0, ay=3.4, bI='DISCONTINUO', bS='CONTINUO', bL='DISCONTINUO', bR='CONTINUO', cforce='AUTO', col='MONOLÍTICO (caso I)', ho=12, scm=160),
    dict(key='EJEMPLO M-3 (volado 1.2 m)', tipo='VOLADO', dest='HABITACIÓN', ax=1.2, ay=3.0, bI='DISCONTINUO', bS='DISCONTINUO', bL='CONTINUO', bR='DISCONTINUO', cforce='AUTO', col='MONOLÍTICO (caso I)', ho=12, scm=160,
         pbor=100, Lb=3.0),
]


def sheet_macizas(wb):
    ws = wb.create_sheet('MACIZAS')
    write_table(ws, MC, COLS_M, GROUPS_M, NOTES_M, vals_macizas(), INP_M, EX_M,
                'MACIZAS — LOSA MACIZA, UN RENGLÓN POR TABLERO (FRANJA DE 1 m) · MÉTODO DE COEFICIENTES NTC-CONCRETO 2023, TABLA 3.4.2.1 (CASO I)',
                'Momentos por metro = coef·10⁻⁴·w_u·a1². Cada lecho lleva al menos As mín = 0.0018·b·h. La varilla sube de número si la separación baja de la mínima práctica; la separación no pasa de la máxima (DATOS sección 8). '
                'Sólo se captura lo AMARILLO; los renglones vacíos quedan en blanco. Memoria en DETALLE y dibujo en CROQUIS.',
                'ESTADO: «CUMPLE» = flexión, separación, ε_t ≥ 0.005, cortante y peralte bien. «REVISAR: calcular flecha» = h menor que el mínimo de la Tabla 6.7.3.1.2. '
                '«NO CUMPLE: cortante» = sube h (la losa no lleva estribos). «FUERA de 3.4.2.1 c)» = el M⁻ del vecino capturado difiere más de 50 %.',
                'nota')
    return ws
