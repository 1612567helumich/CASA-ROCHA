"""DETALLE (memoria paso a paso de un tablero), CROQUIS (dibujos que se redibujan solos) y CROQUIS_DATOS.

Mismo método que gen_c.py de Casa Rocha: gráficas de dispersión cuyas coordenadas son fórmulas vivas en CROQUIS_DATOS,
con el tablero escogido en DETALLE. Todo se dibuja en un cuadro normalizado con la misma escala en X y en Y.
"""
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.series import SeriesLabel, StrRef
from openpyxl.chart.label import DataLabelList, DataLabel
from openpyxl.chart.marker import Marker
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.title import Title
from openpyxl.chart.text import Text, RichText
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, RichTextProperties, RegularTextRun
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import cm_to_EMU
from openpyxl.styles import Alignment
from openpyxl.worksheet.pagebreak import Break
from losas import *            # noqa: F401,F403
from losas import NC, MC, BLKS, BLKNAME

CD = 'CROQUIS_DATOS'
Y0 = 8          # primera fila de coordenadas
LAB0 = 600      # fila donde viven los textos de leyendas y cotas
DEFAULT_SEL = 'N · EJEMPLO F80016 (de borde)'   # tablero que trae escogido el libro (render_croquis.py lo cambia)


def VV(k, fb='"—"'):
    """Valor de la columna k del tablero escogido (nervada o maciza)."""
    rn = 'INDEX(NERVADAS!$%s$%d:$%s$%d,{d_nr})' % (NC[k], R0, NC[k], R1) if k in NC else fb
    rm = 'INDEX(MACIZAS!$%s$%d:$%s$%d,{d_nr})' % (MC[k], R0, MC[k], R1) if k in MC else fb
    return 'IF({d_nr}=0,%s,IF({d_tl}="NERVADA",%s,%s))' % (fb, rn, rm)


def U(nerv, maz):
    return ('=IF({d_tl}="NERVADA","%s","%s")' % (nerv, maz)).format(**N)


# =====================================================================
def sheet_detalle(wb):
    ws = wb.create_sheet('DETALLE'); S = Narr(ws, N)
    S.title('DETALLE — MEMORIA DE CÁLCULO PASO A PASO DE UN TABLERO')
    S.subtitle('Escoge el tablero en la celda amarilla (N = losa nervada, M = losa maciza). Cada paso trae su fórmula y su referencia; los valores salen del renglón del tablero en NERVADAS o MACIZAS. '
               'CROQUIS dibuja este mismo tablero.', 30)
    S.gap()
    S.section('0. TABLERO')
    S.header()
    S.row('Tablero que se revisa y se dibuja', DEFAULT_SEL, '—', 'SELECTOR. La lista la arma el libro con los renglones capturados de NERVADAS (N) y MACIZAS (M).', 'TABLAS sección E', 'input', '@', 'd_sel')
    dv_list(ws, cel(N['d_sel']), '=LST_TAB')
    S.row('Tipo de losa', '=IF(LEFT({d_sel},1)="N","NERVADA",IF(LEFT({d_sel},1)="M","MACIZA","—"))', '—', '—', '—', 'key', '@', 'd_tl')
    S.row('Nombre', '=MID({d_sel},5,200)', '—', '—', '—', 'calc', '@', 'd_nom')
    S.row('Renglón en su hoja', '=IFERROR(IF({d_tl}="NERVADA",MATCH({d_nom},NERVADAS!$B$%d:$B$%d,0),MATCH({d_nom},MACIZAS!$B$%d:$B$%d,0)),0)' % (R0, R1, R0, R1), '—', 'Auxiliar. 0 = no se encontró.', '—', 'calc', '0', 'd_nr')
    S.row('Proyecto', '={proj}', '—', '—', 'DATOS', 'calc', '@')
    S.row('ESTADO DEL TABLERO', '=' + VV('estado'), '—', 'Resumen de todas las revisiones (sección 13).', 'NERVADAS / MACIZAS', 'key', '@', 'd_est')
    S.gap()
    S.section('1. DATOS DEL TABLERO')
    S.header()
    S.row('Tipo (perimetral o volado)', '=' + VV('tipo'), '—', 'PERIMETRAL = apoyado en sus cuatro bordes; VOLADO = empotrado en un borde.', NTCC + ' 3.4.2, p. 64', 'calc', '@')
    S.row('Destino', '=' + VV('dest'), '—', 'De la lista de DATOS sección 4.', NTCCA + ' Tabla 6.1.2.2, p. 11', 'calc', '@')
    S.row('a_x / a_y — claros a ejes', '=TEXT(' + VV('ax', '0') + ',"0.000")&" / "&TEXT(' + VV('ay', '0') + ',"0.000")', 'm', 'CAPTURA. A ejes, del lado seguro (la tabla admite claros libres en el caso I).', NTCC + ' Tabla 3.4.2.1, nota final, p. 65', 'calc', '@')
    S.row('Bordes (inferior · superior · izquierdo · derecho)', '=' + VV('bI') + '&" · "&' + VV('bS') + '&" · "&' + VV('bL') + '&" · "&' + VV('bR'), '—', 'CONTINUO = sigue losa del otro lado. En volado el CONTINUO es el borde empotrado.', '—', 'calc', '@')
    S.row('SCM — sobrecarga muerta', '=' + VV('scm', '0'), 'kg/m²', 'CAPTURA (acabados, rellenos, plafón, divisorios repartidos).', NTCCA + ' 5.1.1, p. 9', 'calc', '0')
    S.row('Carga en la orilla del volado / L_b de respaldo', '=TEXT(N(' + VV('pbor', '0') + '),"0.0")&" kg/m / "&IF(' + VV('tipo') + '="VOLADO",TEXT(IF(' + VV('Lb', '0') + '>0,' + VV('Lb', '0') + ',' + VV('a1', '0') + '),"0.00")&" m","—")', '—',
          'Sólo volados: carga lineal de servicio en la orilla (va con FC_CM) y claro del tablero de respaldo.', NTCCA + ' 3.4.1 a), p. 7', 'calc', '@')
    S.gap()
    S.section('2. GEOMETRÍA')
    S.header()
    S.row('a1 — claro corto (volado: largo del volado)', '=' + VV('a1', '0'), 'm', 'a1 = mín(a_x, a_y); en volado, el claro perpendicular al borde empotrado.', NTCC + ' 3.4.2, p. 64', 'key', '0.000')
    S.row('a2 — claro largo', '=' + VV('a2', '0'), 'm', 'a2 = máx(a_x, a_y).', '—', 'calc', '0.000')
    S.row('m = a1/a2 y dirección del claro corto', '=TEXT(' + VV('m', '0') + ',"0.000")&" · corto ‖ "&' + VV('dirc'), '—', 'El acero del claro corto va por fuera (más peralte); el del largo, encima.', '—', 'calc', '@')
    S.row('L_n — claros libres corto / largo', '=TEXT(' + VV('ln1', '0') + ',"0.0")&" / "&TEXT(' + VV('ln2', '0') + ',"0.0")', 'cm', 'L_n = a − b_p (volado: a − b_p/2), con b_p = ancho del apoyo (DATOS).', NTCC + ' 6.3.2.1, p. 117', 'calc', '@')
    S.row('Sección', '=IF({d_tl}="NERVADA","h = "&{h}&" cm · b_w = "&' + VV('bwr', '0') + '&" cm · a_c = "&' + VV('acr', '0') + '&" cm · s = "&{s}&" cm · t = "&{t}&" cm","Maciza, h = "&' + VV('hr', '0') + '&" cm (franja de 1 m)")', '—',
          'Nervada: h = h_c + t, s = a_c + b_w (con b_w propio el casetón se angosta a s − b_w). Maciza: h del renglón o el de DATOS.', NTCC + ' 6.7.8.1.2 y 6.7.8.1.3, p. 187', 'calc', '@')
    S.row('Patín b_f corto / largo (nervada)', '=IF({d_tl}="NERVADA",TEXT(' + VV('bfc', '0') + ',"0.0")&" / "&TEXT(' + VV('bfl', '0') + ',"0.0"),"— (b = 100 cm)")', 'cm', 'b_w + 2·el menor de 8t, a_c/2 y L_n/8, no más que s.', NTCC + ' Tabla 5.2.1.4.2, p. 89', 'calc', '@')
    S.row('¿Cumple 6.7.8 (losa encasetonada)?', '=IF({d_tl}="NERVADA",' + VV('g678') + ',"— (maciza)")', '—', 'Claro corto ≥ 6·a_c, a_c ≤ 80, b_w ≥ 10, h ≤ 3.5·b_w, t ≥ máx(a_c/13, 5). Si NO: losa apoyada en vigas: r de viga y V_cR sin ×1.1.',
          NTCC + ' 6.7.8.1.2 a 6.7.8.1.4, p. 187 · 6.7.8.1.9, p. 188', 'key', '@')
    S.row('Recubrimiento inferior / superior', '=TEXT(' + VV('rb', '0') + ',"0.0")&" / "&TEXT(' + VV('rt', '0') + ',"0.0")', 'cm', 'r interior; r de viga si no cumple 6.7.8; r_ext en caras a la intemperie (DATOS sección 5).', NTCC + ' Tabla 4.14.16.a, p. 85 · Tabla 4.3.1, p. 76', 'calc', '@')
    S.gap()
    S.section('3. CARGAS')
    S.header()
    S.row('PP — peso propio de diseño', '=' + VV('pp', '0'), 'kg/m²', 'Nervada: (h − a_c²·h_c/s²)·γc + casetón; maciza: h·γc. Más 20 + 20 de 5.1.2.1.', NTCCA + ' 5.1.2.1, p. 9', 'calc', '0')
    S.row('CM = PP + SCM', '=' + VV('cm', '0'), 'kg/m²', '—', NTCCA + ' 5.1.2.1, p. 9', 'calc', '0')
    S.row('W_m — carga viva máxima', '=' + VV('cv', '0'), 'kg/m²', 'Del destino o la capturada en el renglón.', NTCCA + ' Tabla 6.1.2.2, p. 11 · 6.1.2.1 a), p. 10', 'calc', '0')
    S.row('w_u — carga última', '=' + VV('wu', '0'), 'kg/m²', 'El mayor de 1.3·CM + 1.5·W_m y, si el destino es azotea, 1.1·(CM + W_g).', NTCCA + ' 3.4.1 a) y b), p. 7 · ec. 7.3.2.1, p. 17', 'key', '0')
    S.gap()
    S.section('4. CASO Y COEFICIENTES (TABLA 3.4.2.1, CASO I)')
    S.header()
    S.row('Bordes discontinuos (total / cortos)', '=' + VV('nd', '0') + '&" / "&' + VV('ndc', '0'), '—', '«Corto» = borde de largo a1.', '—', 'calc', '@')
    S.row('Caso de la tabla', '=' + VV('caso'), '—', 'Automático por bordes discontinuos, o el forzado en el renglón. Sólo caso I (colado monolítico).', NTCC + ' Tabla 3.4.2.1, p. 65', 'key', '@')
    for mm, lbl in (('NIC', 'M⁻ en bordes continuos, claro corto'), ('NDC', 'M⁻ en bordes discontinuos, claro corto'), ('PC', 'M⁺, claro corto'),
                    ('NIL', 'M⁻ en bordes continuos, claro largo'), ('NDL', 'M⁻ en bordes discontinuos, claro largo'), ('PL', 'M⁺, claro largo')):
        S.row('Coeficiente ' + lbl, '=' + VV('c' + mm, '0'), '×10⁻⁴', 'Interpolado linealmente en m.', 'TABLAS sección A', 'calc', '0')
    S.gap()
    S.section('5. MOMENTOS ÚLTIMOS')
    S.header()
    mu = U('t·m/nerv.', 't·m/m')
    how = 'M = coef·10⁻⁴·w_u·a1² por metro; en nervada × s (por nervadura).'
    S.row('M⁺ claro corto', '=' + VV('MPc', '0'), mu, how, NTCC + ' 3.4.2.1, p. 64', 'key', '0.000')
    S.row('M⁻ claro corto', '=' + VV('MNc', '0'), mu, 'El mayor de borde continuo, discontinuo y el del vecino capturado. Volado: w_u·L_n²/2 + 1.3·P·L_n (× s en nervada).', NTCC + ' 3.4.2.1, p. 64 · ' + NTCCA + ' 3.4.1 a), p. 7', 'key', '0.000')
    S.row('M⁺ claro largo', '=' + VV('MPl', '0'), mu, how, NTCC + ' 3.4.2.1, p. 64', 'key', '0.000')
    S.row('M⁻ claro largo', '=' + VV('MNl', '0'), mu, 'Igual que el corto.', NTCC + ' 3.4.2.1, p. 64', 'key', '0.000')
    S.row('3.4.2.1 c) — M⁻ del vecino', '=' + VV('c342'), '—', 'Si capturaste el M⁻ del vecino en el borde común: no debe diferir más de 50 % del menor; si difiere, distribuye (3.4.3.1) o confirma con un modelo.', NTCC + ' 3.4.2.1 c) y 3.4.3.1, p. 64', 'key', '@')
    sec = 6
    for blk in BLKS:
        neg = blk[0] == 'N'
        S.gap()
        S.section('%d. ACERO — LECHO %s' % (sec, BLKNAME[blk])); sec += 1
        S.header()
        S.row('d — peralte efectivo con la varilla que se usa', '=' + VV('d' + blk, '0'), 'cm',
              ('h − máx(r, (t − 2Ø)/2) − 2·Ø − d_b/2 en nervada (la malla va al centro de la capa y la varilla pasa debajo); h − r − d_b/2 en maciza; − 2 cm si d ≤ 20 cm sin la precaución de 3.9.2' if neg else 'h − r − d_b/2')
              + (' − d_b del lecho corto (va encima).' if blk[1] == 'l' else '.'),
              NTCC + ' Tabla 4.14.16.a, p. 85' + (' · 3.9.1 d), p. 74' if neg else ''), 'calc', '0.00')
        S.row('As requerida', '=' + VV('As' + blk, '0'), U('cm²/nerv.', 'cm²/m'),
              "q = 1 − √(1 − 2Mu/(F_R·b·d²·f''c)); As = q·f''c·b·d/fy; nervada: ≥ mín(p_min·b_w·d, 4/3·As); maciza: ≥ 0.0018·b·h. −1 = no alcanza o rebasa p_max (ε_t ≥ 0.005).",
              NTCC + ' ecs. 5.2.2.1.1.1, p. 89 · 6.3.5.1.3, p. 122 · 6.7.6.1.1, p. 175', 'calc', '0.00')
        S.row('Armado', '=' + VV('t' + blk), '—', 'Nervada: piezas por nervadura (la preferida o la siguiente que quepa). Maciza: varilla y separación (≤ máx., ≥ mín.).', NTCC + ' 6.7.6.1.3, p. 175 · 14.2.1, p. 347', 'key', '@')
        S.row('M_R — momento resistente', '=' + VV('MR' + blk, '0'), mu, 'F_R·As·fy·(d − a/2), a = As·fy/(f\'\'c·b).', NTCC + ' ec. 5.2.2.1.1.1.b, p. 89', 'calc', '0.000')
        S.row('Mu / M_R', '=' + VV('FU' + blk, '0'), '—', 'Debe ser ≤ 1 (9.99 = la sección no alcanza).', '—', 'key', '0.00')
        S.row('ε_t — deformación del acero', '=' + VV('e' + blk, '0'), '—', 'ε_t = 0.003·(d − c)/c, c = a/β1. Debe ser ≥ 0.005 (F_R = 0.9).', NTCC + ' Tabla 3.8.2.2, p. 70 · 6.7.3.3.1, p. 169', 'calc', '0.0000')
    S.gap()
    S.section('%d. CORTANTE' % sec); sec += 1
    S.header()
    S.row('V_u a d del paño', '=' + VV('Vu', '0'), U('t/nerv.', 't/m'), '(L_n/2 − d)(0.95 − 0.5m)·w_u, ×1.15 con bordes continuos y discontinuos (× s en nervada). Volado: w_u(L_n − d) + 1.3P.', NTCC + ' 3.4.1.1, p. 64 · 6.7.4.3.2, p. 172', 'key', '0.000')
    S.row('V_cR — cortante que toma el concreto', '=' + VV('VcR', '0'), U('t/nerv.', 't/m'), "F_R·2·λs·p^⅓·√f'c·b·d (b = b_w en nervada, × 1.1 si cumple 6.7.8; b = 100 en maciza).", NTCC + ' ec. 5.5.3.2.1, p. 95 · 6.7.8.1.5, p. 187 · 5.5.3.8.1, p. 97', 'key', '0.000')
    S.row('Estado (y hileras a macizar en nervada)', '=' + VV('cort'), '—', 'Nervada: lo que no alcanza se cubre quitando hileras de casetones junto al apoyo. Maciza: sube h.', NTCC + ' 5.5.3.8.1, p. 97', 'key', '@')
    S.gap()
    S.section('%d. ESPESOR MÍNIMO (PARA OMITIR EL CÁLCULO DE FLECHAS)' % sec); sec += 1
    S.header()
    S.row('α_fm', '=' + VV('afm', '0'), '—', 'Promedio de α_f de los cuatro bordes (99 = muros o trabes rígidas).', NTCC + ' Tabla 6.7.3.1.2 nota [1], p. 168', 'calc', '0.00')
    S.row('h mínimo', '=' + VV('hmin', '0'), 'cm', 'Tabla 6.7.3.1.2 (+10 % con borde discontinuo sin viga). Volado: L/8 en nervada, L/10 en maciza (criterio de esta hoja).', NTCC + ' Tabla 6.7.3.1.2 y 6.7.3.1.2.1, p. 168 · Tabla 6.3.3.1.1, p. 120', 'key', '0.0')
    S.row('h que se compara', '=IF({d_tl}="NERVADA",IF(' + VV('tipo') + '="VOLADO",{h},' + VV('heq', '0') + '),' + VV('hr', '0') + ')', 'cm', 'Nervada: h_eq (maciza de la misma inercia); volado nervado: h. Maciza: h.', 'Criterio de esta hoja · ' + NTCC + ' 6.7.3.1, p. 167–168', 'calc', '0.00')
    S.row('Estado', '=' + VV('peral'), '—', 'Si no cumple: calcular flechas (13.4) y revisar vibración (13.7).', NTCC + ' 6.7.3.2.1, p. 168', 'key', '@')
    S.gap()
    S.section('%d. BASTONES Y ANCLAJE' % sec); sec += 1
    S.header()
    S.row('Bastón ‖ claro corto: largo desde el paño', '=' + VV('Lbc', '0'), 'cm', 'Punto de inflexión + máx(d, 12d_b, L_n/16), ≥ L_d; si no hay punto de inflexión, de lado a lado. Volado: dentro del respaldo hasta M = 0 con 0.9·PP sin CV.',
          NTCC + ' 6.3.7.3.4, p. 127 · 6.3.7.3.10.4, p. 129 · Tabla 14.4.2.4, p. 353', 'key', '0')
    S.row('Bastón ‖ claro largo: largo desde el paño', '=' + VV('Lbl', '0'), 'cm', 'Mismo criterio.', NTCC + ' 6.3.7.3.10.4, p. 129', 'key', '0')
    S.row('¿Bastón también en bordes discontinuos? (corto / largo)', '=' + VV('bdc') + '&" / "&' + VV('bdl'), '—', 'Nervada: si el M⁻ de borde discontinuo no lo toma la varilla corrida. Maciza: siempre que la tabla da M⁻ en ese borde. Con gancho.', NTCC + ' 6.7.7.4.1.1 b), p. 179', 'calc', '@')
    S.row('Anclaje en bordes discontinuos', '=' + VV('anc'), '—', 'Dala de borde armada de DATOS: 1.25·L_dh de la varilla inferior (nervada; ≥ 15 cm dentro) y L_dh del bastón, menos r de viga; gancho de 180° de alto 8·d_b dentro del peralte de la dala.',
          NTCC + ' 6.7.8.1.7, p. 188 · 6.7.7.4.1.1, p. 179 · Tabla 14.3.1, p. 349 · 14.4.3.1, p. 355', 'key', '@')
    S.row('Respaldo de un volado con b_w propio: varillas del volado que caben en la nervadura normal / su M_R / su V_cR', '=IF(' + VV('tipo') + '<>"VOLADO","—",' + VV('rn1', '0') + '&" pzas · "&TEXT(' + VV('rM1', '0') + ',"0.000")&" t·m · "&TEXT(' + VV('rV1', '0') + ',"0.000")&" t")', '—',
          'Nervadura de b_w de DATOS con las varillas del volado que le caben y la malla al centro de la capa (F_R·As·fy·(d − a/2) con ε_t ≥ 0.005; ec. 5.5.3.2.1 sin ×1.1).', NTCC + ' ec. 5.2.2.1.1.1.b, p. 89 · ec. 5.5.3.2.1, p. 95', 'calc', '@')
    S.row('… dónde M y V del volado bajan a esos valores dentro del respaldo', '=IF(' + VV('tipo') + '<>"VOLADO","—","M: "&TEXT(' + VV('rxM', '0') + ',"0.00")&" m · V: "&TEXT(' + VV('rxV', '0') + ',"0.00")&" m desde el paño")', '—',
          'M(x) = M⁻(1 − x/L_b) − 0.9·PP·s·x(L_b − x)/2 (sin CV); V(x) = M⁻/L_b + w_u·s·(L_b/2 − x).', NTCCA + ' 3.4.1 c), p. 7 · 5.1.2.2, p. 9', 'calc', '@')
    S.row('Respaldo del volado', '=' + VV('rest'), '—', 'Las primeras hileras del tablero de respaldo siguen con las nervaduras anchas del volado (casetón s − b_w) hasta ese punto, más máx(d, 12d_b), y hasta donde terminan los bastones del volado si no le caben todas sus varillas. '
          'Junto al apoyo la nervadura ancha resiste M⁻/L_b + w_u·s·L_b/2 con el V_cR del volado. Pon estas hileras también en el renglón del tablero de respaldo.', 'Decisión D10 de Casa Rocha · ' + NTCC + ' 6.3.7.3.10.4, p. 129', 'key', '@')
    S.gap()
    S.section('%d. RESULTADO' % sec)
    S.header()
    for blk in BLKS:
        S.row('Armado — ' + BLKNAME[blk].lower(), '=' + VV('t' + blk), '—', '—', '—', 'key', '@')
    S.row('ESTADO', '=' + VV('estado'), '—', '—', '—', 'key', '@')
    r = S.row('NOTA', '=' + VV('nota'), '—', '—', '—', 'calc', '@', h=96)
    ws.cell(r, 3).alignment = Alignment(wrap_text=True, vertical='top', horizontal='left')
    ws.merge_cells('C%d:F%d' % (r, r))
    add_cf(ws, 'C1:C%d' % S.r)
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = 'A4'
    print_setup(ws, 'B1:F%d' % S.r, landscape=False)
    return ws


# =====================================================================
def sheet_croquis(wb):
    ws = wb.create_sheet('CROQUIS'); S = Narr(ws, N, {'A': 1.78, 'B': 46, 'C': 17, 'D': 9, 'E': 54, 'F': 44})
    S.title('CROQUIS — PLANTA, CORTE Y ALZADO DEL TABLERO ESCOGIDO (A ESCALA, SE REDIBUJAN SOLOS)')
    S.subtitle('Se dibuja el tablero escogido en DETALLE (celda amarilla de arriba). Arriba los dibujos de losa NERVADA y abajo los de losa MACIZA: sólo se llena el grupo del tipo escogido. '
               'Las coordenadas viven en CROQUIS_DATOS (al final del libro) y se calculan solas; los textos del armado salen de NERVADAS / MACIZAS.', 30)
    r = S.r
    ws.merge_cells('B%d:F%d' % (r, r))
    c = ws['B%d' % r]; c.value = '=HYPERLINK("#DETALLE!%s","▶ Cambiar de tablero (DETALLE, celda amarilla)")' % cel(N['d_sel']); c.font = font(10, True, color='0000CC')
    S.r += 1
    N['k_row_tit'] = S.r
    S.r += 1
    N['k_row_N'] = S.r            # encabezado del grupo NERVADA
    S.r += 39
    N['k_row_M'] = S.r            # encabezado del grupo MACIZA
    S.r += 39
    S.gap()
    S.section('DATOS DEL DIBUJO — TODO ENLAZADO (no se captura nada aquí)')
    S.header()
    P = []

    def kp(label, formula, nm, fmt='0.00', unit='', how=''):
        S.row(label, formula, unit or '—', how or 'Enlazado.', 'NERVADAS / MACIZAS / DATOS', 'calc', fmt, nm)
        P.append(nm)
    kp('¿Hay tablero válido?', '=IF(AND({d_nr}>0,' + VV('ok') + '="SÍ"),1,0)', 'k_ok', '0')
    kp('¿Es nervada?', '=IF(AND({k_ok}=1,{d_tl}="NERVADA"),1,0)', 'k_N', '0')
    kp('¿Es maciza?', '=IF(AND({k_ok}=1,{d_tl}="MACIZA"),1,0)', 'k_M', '0')
    kp('Tablero', '=IF({k_ok}=1,' + VV('key') + ',"(escoge un tablero con datos completos)")', 'k_key', '@')
    kp('Tipo', '=' + VV('tipo'), 'k_tipo', '@')
    kp('¿Volado?', '=IF(' + VV('tipo') + '="VOLADO",1,0)', 'k_vol', '0')
    kp('Claro corto ‖', '=IF({k_ok}=1,' + VV('dirc') + ',"X")', 'k_dir', '@')
    kp('Borde empotrado (1 izq., 2 der., 3 inf., 4 sup.)', '=IF({k_vol}=1,IF(' + VV('emp') + '="IZQUIERDO",1,IF(' + VV('emp') + '="DERECHO",2,IF(' + VV('emp') + '="INFERIOR",3,IF(' + VV('emp') + '="SUPERIOR",4,1)))),0)', 'k_es', '0')
    kp('¿Apoyo en u = 0? (volado)', '=IF(OR({k_es}=1,{k_es}=3),1,0)', 'k_sup0', '0')
    kp('a_x', '=IF({k_ok}=1,' + VV('ax', '0') + '*100,100)', 'k_ax', '0.0', 'cm')
    kp('a_y', '=IF({k_ok}=1,' + VV('ay', '0') + '*100,100)', 'k_ay', '0.0', 'cm')
    kp('A — claro ‖ corto (u)', '=IF({k_dir}="X",{k_ax},{k_ay})', 'k_A', '0.0', 'cm')
    kp('C — claro ‖ largo (v)', '=IF({k_dir}="X",{k_ay},{k_ax})', 'k_C', '0.0', 'cm')
    for s_, nm in (('bI', 'k_cI'), ('bS', 'k_cS'), ('bL', 'k_cL'), ('bR', 'k_cR')):
        kp('Borde %s continuo' % s_, '=IF(' + VV(s_) + '="CONTINUO",1,0)', nm, '0')
    kp('Continuidad en u = 0 / u = A / v = 0 / v = C', '=IF({k_dir}="X",{k_cL},{k_cI})', 'k_c0', '0')
    kp('…', '=IF({k_dir}="X",{k_cR},{k_cS})', 'k_c1', '0')
    kp('…', '=IF({k_dir}="X",{k_cI},{k_cL})', 'k_v0c', '0')
    kp('…', '=IF({k_dir}="X",{k_cS},{k_cR})', 'k_v1c', '0')
    kp('b_p', '={bp}', 'k_bp', '0.0', 'cm')
    kp('h', '=IF({k_M}=1,' + VV('hr', '0') + ',{h})', 'k_h', '0.0', 'cm')
    kp('r inferior / superior', '=N(' + VV('rb', '0') + ')', 'k_rb', '0.0', 'cm')
    kp('…', '=N(' + VV('rt', '0') + ')', 'k_rt', '0.0', 'cm')
    kp('b_w / a_c usados', '=IF({k_N}=1,' + VV('bwr', '0') + ',{bw})', 'k_bw', '0.0', 'cm')
    kp('…', '={s}-{k_bw}', 'k_ac', '0.0', 'cm')
    kp('s, h_c, t, Ø malla, separación libre', '={s}', 'k_s', '0.0', 'cm')
    kp('…', '={hc}', 'k_hc', '0.0', 'cm')
    kp('…', '={t}', 'k_tc', '0.0', 'cm')
    kp('…', '={fim}', 'k_fim', '0.000', 'cm')
    kp('r superior hasta la malla (nervada: malla al centro de la capa)', '=IF({k_N}=1,MAX({k_rt},({t}-2*{fim})/2),{k_rt})', 'k_rtm', '0.00', 'cm')
    kp('…', '={slib}', 'k_slib', '0.00', 'cm')
    for nm in ('ncx', 'ncy', 'mx', 'my', 'nh'):
        kp(nm, '=IF({k_N}=1,N(' + VV(nm, '0') + '),0)', 'k_' + nm, '0.0')
    for blk in BLKS:
        kp('Varilla ' + blk, '=IF({k_ok}=1,' + VV('v' + blk, '4') + ',4)', 'k_v' + blk, 'General')
        kp('d_b ' + blk, '=' + db('{k_v%s}' % blk), 'k_db' + blk, '0.00', 'cm')
        kp('Piezas (N) o separación (M) ' + blk, '=IF({k_N}=1,N(' + VV('n' + blk, '0') + '),IF({k_M}=1,N(' + VV('sp' + blk, '0') + '),0))', 'k_q' + blk, '0.0')
        kp('Armado ' + blk, '=' + VV('t' + blk), 'k_t' + blk, '@')
    kp('Bastón ‖ corto / largo desde el paño', '=N(' + VV('Lbc', '0') + ')', 'k_Lbc', '0', 'cm')
    kp('…', '=N(' + VV('Lbl', '0') + ')', 'k_Lbl', '0', 'cm')
    kp('Bastón en borde discontinuo corto / largo', '=IF(' + VV('bdc') + '="SÍ",1,0)', 'k_bdc', '0')
    kp('…', '=IF(' + VV('bdl') + '="SÍ",1,0)', 'k_bdl', '0')
    kp('Nervadura ancha del volado dentro del respaldo: hasta / hileras', '=IF({k_N}=1,N(' + VV('rxr', '0') + '),0)', 'k_rxr', '0', 'cm')
    kp('…', '=IF({k_N}=1,N(' + VV('rnh', '0') + '),0)', 'k_rnh', '0')
    kp('L_b de respaldo', '=IF({k_vol}=1,IF(N(' + VV('Lb', '0') + ')>0,' + VV('Lb', '0') + ',' + VV('a1', '0') + ')*100,0)', 'k_Lb', '0', 'cm')
    kp('d inferior ‖ corto', '=N(' + VV('dPc', '0') + ')', 'k_dPc', '0.00', 'cm')
    kp('Carga en la orilla', '=N(' + VV('pbor', '0') + ')', 'k_P', '0.0', 'kg/m')
    kp('Estado', '=' + VV('estado'), 'k_est', '@')
    # --- planta: marco y escala
    kp('Respaldo que se dibuja', '=IF({k_vol}=1,MIN({k_Lb},{k_bp}/2+{k_Lbc}+40),0)', 'k_BK', '0.0', 'cm')
    kp('x mín / x máx / y mín / y máx', '=IF({k_es}=1,-{k_BK},0)', 'k_xmin', '0.0', 'cm')
    kp('…', '={k_ax}+IF({k_es}=2,{k_BK},0)', 'k_xmax', '0.0', 'cm')
    kp('…', '=IF({k_es}=3,-{k_BK},0)', 'k_ymin', '0.0', 'cm')
    kp('…', '={k_ay}+IF({k_es}=4,{k_BK},0)', 'k_ymax', '0.0', 'cm')
    kp('Escala de la planta', '=76/MAX({k_xmax}-{k_xmin},{k_ymax}-{k_ymin},1)', 'k_sc', '0.0000', 'u/cm', 'Cuadro de 76 unidades, misma escala en X y Y.')
    kp('Origen X / Y', '=12+(76-({k_xmax}-{k_xmin})*{k_sc})/2-{k_xmin}*{k_sc}', 'k_ox', '0.00', 'u')
    kp('…', '=12+(76-({k_ymax}-{k_ymin})*{k_sc})/2-{k_ymin}*{k_sc}', 'k_oy', '0.00', 'u')
    # --- casetones con hileras macizadas
    kp('Casetones dibujados ‖ X', '=IF({k_dir}="X",MAX(0,{k_ncx}-{k_nh}*IF({k_vol}=1,1,2)),{k_ncx})', 'k_nXe', '0')
    kp('Inicio de casetones en X', '={k_bp}/2+{k_mx}+IF(AND({k_dir}="X",OR({k_vol}=0,{k_es}=1)),{k_nh}*{k_s},0)', 'k_xs', '0.0', 'cm')
    kp('Casetones dibujados ‖ Y', '=IF({k_dir}="Y",MAX(0,{k_ncy}-{k_nh}*IF({k_vol}=1,1,2)),{k_ncy})', 'k_nYe', '0')
    kp('Inicio de casetones en Y', '={k_bp}/2+{k_my}+IF(AND({k_dir}="Y",OR({k_vol}=0,{k_es}=3)),{k_nh}*{k_s},0)', 'k_ys', '0.0', 'cm')
    kp('Casetones / inicio en u y en v', '=IF({k_dir}="X",{k_nXe},{k_nYe})', 'k_nAe', '0')
    kp('…', '=IF({k_dir}="X",{k_xs},{k_ys})', 'k_us', '0.0', 'cm')
    kp('…', '=IF({k_dir}="X",{k_nYe},{k_nXe})', 'k_nCe', '0')
    kp('…', '=IF({k_dir}="X",{k_ys},{k_xs})', 'k_vs', '0.0', 'cm')
    kp('Nervadura representativa ‖ corto (posición en v)', '=IF({k_nCe}>=2,{k_vs}+{k_ac}+{k_bw}/2+INT(({k_nCe}-2)/2)*{k_s},{k_C}/2)', 'k_vrib', '0.0', 'cm')
    kp('Nervadura representativa ‖ largo (posición en u)', '=IF({k_nAe}>=2,{k_us}+{k_ac}+{k_bw}/2+INT(({k_nAe}-2)/2)*{k_s},{k_A}/2)', 'k_urib', '0.0', 'cm')
    # --- alzado
    kp('Respaldo en el alzado', '=IF({k_vol}=1,MIN({k_Lb},{k_bp}/2+{k_Lbc}+40),0)', 'k_BKa', '0.0', 'cm')
    kp('Alzado: x izq. / x der.', '=IF({k_vol}=1,-{k_BKa},-{k_bp}/2-30)', 'k_xL', '0.0', 'cm')
    kp('…', '=IF({k_vol}=1,{k_A}+12,{k_A}+{k_bp}/2+30)', 'k_xR', '0.0', 'cm')
    kp('Alzado: escala horizontal', '=92/MAX({k_xR}-{k_xL},1)', 'k_sa', '0.0000', 'u/cm')
    kp('Alzado: exageración vertical', '=MAX(1,ROUND(16/MAX({k_h}*{k_sa},0.01),0))', 'k_E', '0')
    kp('Corte nervada: escala', '=MIN(84/({k_s}+18),36/({k_h}+12))', 'k_cs', '0.0000', 'u/cm')
    kp('Corte maciza: escala', '=MIN(84/(60+18),36/({k_h}+12))', 'k_cm', '0.0000', 'u/cm')
    # encabezados de los grupos
    for rr, txt in ((N['k_row_tit'], '="Tablero: "&{k_key}&"   ·   "&{k_est}'),
                    (N['k_row_N'], '=IF({k_N}=1,"▼ LOSA NERVADA — ESTE ES EL TABLERO ESCOGIDO","LOSA NERVADA — (el tablero escogido no es nervado: estos dibujos quedan vacíos)")'),
                    (N['k_row_M'], '=IF({k_M}=1,"▼ LOSA MACIZA — ESTE ES EL TABLERO ESCOGIDO","LOSA MACIZA — (el tablero escogido no es macizo: estos dibujos quedan vacíos)")')):
        ws.merge_cells('B%d:F%d' % (rr, rr))
        x = ws['B%d' % rr]; x.value = txt.format(**N); x.font = font(11, True, color='FFFFFF'); x.fill = fill(SEC_BG)
        x.alignment = Alignment(vertical='center'); ws.row_dimensions[rr].height = 22
    ws.row_breaks.append(Break(id=N['k_row_M'] - 1))
    ws.row_breaks.append(Break(id=N['k_row_M'] + 38))
    add_cf(ws, 'C%d:C%d' % (N['k_row_M'] + 39, S.r))
    ws.sheet_view.zoomScale = 75
    print_setup(ws, 'B1:F%d' % (N['k_row_M'] + 38), landscape=True)
    ws.page_setup.fitToHeight = 0
    return ws


# =====================================================================
def build_croquis_datos(wb, wsK):
    ws = wb.create_sheet(CD); ws.sheet_view.showGridLines = False
    ws['A1'] = 'COORDENADAS DE LOS DIBUJOS DE LA HOJA CROQUIS (las genera el libro; no se tocan)'; ws['A1'].font = font(12, True, color=HDR_FG)
    ws['A2'] = ('Cada par de columnas es un trazo continuo; las filas que sobran repiten el último punto. Un trazo que no aplica se manda a −1000 (fuera del dibujo). '
                'Planta: cuadro de 0 a 100 con la misma escala en X y Y; u = dirección del claro corto, v = la del largo.')
    ws['A2'].font = font(9, i=True, color=GRAY)
    ws.column_dimensions['A'].width = 60
    for i in range(400):
        ws['A%d' % (Y0 + i)] = i + 1
    ws['A4'] = 'k'
    labels = {}; lab_r = [LAB0]; series = {}; col = [2]

    def label(name, formula):
        r = lab_r[0]; ws.cell(r, 1, formula.format(**N) if formula.startswith('=') else formula)
        labels[name] = '%s!$A$%d' % (CD, r); lab_r[0] += 1

    def dyn(name, xf, yf, npt, lab=None, hide='1'):
        cx, cy = CL(col[0]), CL(col[0] + 1); col[0] += 2
        ws['%s4' % cx] = name + ' x'; ws['%s4' % cy] = name + ' y'
        for i in range(npt):
            r = Y0 + i; k = '$A%d' % r
            ws['%s%d' % (cx, r)] = ('=IF(AND(%s),%s,-1000)' % (hide, xf(k))).format(**N)
            ws['%s%d' % (cy, r)] = ('=IF(AND(%s),%s,-1000)' % (hide, yf(k))).format(**N)
        series[name] = (cx, cy, Y0, Y0 + npt - 1)
        if lab:
            # la leyenda de un trazo que no se dibuja dice «—» (no los datos de otro tablero o de otro tipo de losa)
            if hide != '1' and lab.startswith('='):
                lab = '=IF(AND(%s),%s,"—")' % (hide, lab[1:])
            label(name, lab)

    ch_ = lambda k, n, vals: 'CHOOSE(MIN(MAX(%s,1),%d),%s)' % (k, n, ','.join(vals))
    PX = lambda x: '({k_ox}+(%s)*{k_sc})' % x
    PY = lambda y: '({k_oy}+(%s)*{k_sc})' % y

    def rect(x0, y0, x1, y1, tx=PX, ty=PY):
        return (lambda k: ch_(k, 5, [tx(x0), tx(x1), tx(x1), tx(x0), tx(x0)]),
                lambda k: ch_(k, 5, [ty(y0), ty(y0), ty(y1), ty(y1), ty(y0)]))

    def uv(u, v, offc='0'):
        """(u, v) local -> (X, Y) normalizados de la planta; offc = desplazamiento (unidades) en la dirección v."""
        return ('(%s+IF({k_dir}="X",0,%s))' % (PX('IF({k_dir}="X",%s,%s)' % (u, v)), offc),
                '(%s+IF({k_dir}="X",%s,0))' % (PY('IF({k_dir}="X",%s,%s)' % (v, u)), offc))

    def uv_l(u, v, offc='0'):
        """Igual, con el desplazamiento en la dirección u (barras ‖ largo)."""
        return ('(%s+IF({k_dir}="X",%s,0))' % (PX('IF({k_dir}="X",%s,%s)' % (u, v)), offc),
                '(%s+IF({k_dir}="X",0,%s))' % (PY('IF({k_dir}="X",%s,%s)' % (v, u)), offc))

    def serp(n, pos, a0, a1):
        """n trazos en v = pos(j) (j = 0..n−1), cada uno de u = a0 a a1, en serpentina. Devuelve (u(k), v(k))."""
        last = 'MAX((%s)-1,0)' % n
        j = lambda k: 'MIN(INT((%s-1)/2),%s)' % (k, last)
        w = lambda k: 'IF(INT((%s-1)/2)>%s,1,MOD(%s-1,2))' % (k, last, k)
        uf = lambda k: 'IF(MOD(%s,2)=0,IF(%s=0,%s,%s),IF(%s=0,%s,%s))' % (j(k), w(k), a0, a1, w(k), a1, a0)
        vf = lambda k: pos(j(k))
        return uf, vf

    def serp5(n, pos, a0, a1, U0):
        """n varillas en v = pos(j), de u = a0 a a1: cada una se dibuja de ida y vuelta y el trazo pasa a la siguiente por u = U0
        (el eje del borde, que tapa el contorno del tablero), así no aparecen uniones en «U» entre varillas."""
        last = 'MAX((%s)-1,0)' % n
        j = lambda k: 'MIN(INT((%s-1)/5),%s)' % (k, last)
        m = lambda k: 'IF(INT((%s-1)/5)>%s,0,MOD(%s-1,5))' % (k, last, k)
        uf = lambda k: 'CHOOSE(%s+1,%s,%s,%s,%s,%s)' % (m(k), U0, a0, a1, a0, U0)
        vf = lambda k: pos(j(k))
        return uf, vf

    def cota(p0, p1, along_x, at, lbl_name, lab, hide, tick=1.4):
        """Cota normalizada: de p0 a p1 sobre una línea horizontal (along_x) o vertical, en la coordenada at."""
        a = [p0, p0, p0, '(%s+%s)/2' % (p0, p1), p1, p1, p1]
        b = ['%s-%s' % (at, tick), '%s+%s' % (at, tick), at, at, at, '%s+%s' % (at, tick), '%s-%s' % (at, tick)]
        fx = (lambda k: ch_(k, 7, a)) if along_x else (lambda k: ch_(k, 7, b))
        fy = (lambda k: ch_(k, 7, b)) if along_x else (lambda k: ch_(k, 7, a))
        dyn(lbl_name, fx, fy, 7, lab, hide)

    # ================================================= PLANTA (dos grupos: N y M)
    for g, H in (('N', '{k_N}=1'), ('M', '{k_M}=1')):
        fx, fy = rect('0', '0', '{k_ax}', '{k_ay}')
        dyn(g + 'contorno', fx, fy, 5, '="Tablero "&{k_key}&": a_x = "&TEXT({k_ax}/100,"0.00")&" m × a_y = "&TEXT({k_ay}/100,"0.00")&" m (a ejes)"', H)
        x0 = 'IF({k_vol}=1,CHOOSE({k_es},-{k_bp}/2,{k_ax}-{k_bp}/2,0,0),{k_bp}/2)'
        x1 = 'IF({k_vol}=1,CHOOSE({k_es},{k_bp}/2,{k_ax}+{k_bp}/2,{k_ax},{k_ax}),{k_ax}-{k_bp}/2)'
        y0 = 'IF({k_vol}=1,CHOOSE({k_es},0,0,-{k_bp}/2,{k_ay}-{k_bp}/2),{k_bp}/2)'
        y1 = 'IF({k_vol}=1,CHOOSE({k_es},{k_ay},{k_ay},{k_bp}/2,{k_ay}+{k_bp}/2),{k_ay}-{k_bp}/2)'
        fx, fy = rect(x0, y0, x1, y1)
        dyn(g + 'apoyo', fx, fy, 5, '=IF({k_vol}=1,"Apoyo del volado (muro, dala o trabe de b_p = "&{bp}&" cm)","Paño de los apoyos (b_p = "&{bp}&" cm)")', H)
        # achurado del empotramiento
        ax_ = 'IF({k_sup0}=1,0,{k_A})'
        u_h = lambda k: '%s+IF(MOD(%s,2)=1,-{k_bp}/2,{k_bp}/2)' % (ax_, k)
        v_h = lambda k: '(MIN(%s,41)-1)/40*{k_C}' % k
        dyn(g + 'achurado', lambda k: uv(u_h(k), v_h(k))[0], lambda k: uv(u_h(k), v_h(k))[1], 41, '="Empotramiento"', H + ',{k_vol}=1')
        bx0 = 'CHOOSE(MAX({k_es},1),-{k_BK},{k_ax},0,0)'; bx1 = 'CHOOSE(MAX({k_es},1),0,{k_ax}+{k_BK},{k_ax},{k_ax})'
        by0 = 'CHOOSE(MAX({k_es},1),0,0,-{k_BK},{k_ay})'; by1 = 'CHOOSE(MAX({k_es},1),{k_ay},{k_ay},0,{k_ay}+{k_BK})'
        fx, fy = rect(bx0, by0, bx1, by1)
        dyn(g + 'respaldo', fx, fy, 5, '="Tablero de respaldo (L_b = "&TEXT({k_Lb}/100,"0.00")&" m; se dibuja una parte)"', H + ',{k_vol}=1')
        tipU = 'IF({k_sup0}=1,{k_A},0)'
        fpx = lambda k: uv(tipU, '%s*{k_C}' % ch_(k, 3, ['0', '0.5', '1']))[0]
        fpy = lambda k: uv(tipU, '%s*{k_C}' % ch_(k, 3, ['0', '0.5', '1']))[1]
        dyn(g + 'carga', fpx, fpy, 3, '="Carga en la orilla P = "&TEXT({k_P},"0.0")&" kg/m (servicio)"', H + ',{k_vol}=1,{k_P}>0')
        cota(PX('0'), PX('{k_ax}'), True, '(%s-5)' % PY('{k_ymin}'), g + 'cota_ax', '="a_x = "&TEXT({k_ax}/100,"0.00")&" m"', H)
        cota(PY('0'), PY('{k_ay}'), False, '(%s-5)' % PX('{k_xmin}'), g + 'cota_ay', '="a_y = "&TEXT({k_ay}/100,"0.00")&" m"', H)

    # ---- nervada: casetones, ejes de nervaduras, zona maciza y varillas representativas
    HN = '{k_N}=1'
    # respaldo de un volado con b_w propio: franja del tablero de respaldo donde siguen sus nervaduras anchas
    e_ = '{k_bp}/2'; f_ = '({k_bp}/2+{k_rxr})'
    zx0 = 'CHOOSE(MAX({k_es},1),-%s,{k_ax}+%s,0,0)' % (f_, e_); zx1 = 'CHOOSE(MAX({k_es},1),-%s,{k_ax}+%s,{k_ax},{k_ax})' % (e_, f_)
    zy0 = 'CHOOSE(MAX({k_es},1),0,0,-%s,{k_ay}+%s)' % (f_, e_); zy1 = 'CHOOSE(MAX({k_es},1),{k_ay},{k_ay},-%s,{k_ay}+%s)' % (e_, f_)
    fx, fy = rect(zx0, zy0, zx1, zy1)
    dyn('Nresp', fx, fy, 5, '="Respaldo: "&{k_rnh}&" hilera(s) con nervaduras de "&{k_bw}&" cm (hasta "&TEXT({k_rxr}/100,"0.00")&" m del paño)"', HN + ',{k_vol}=1,{k_rxr}>0')
    # aristas de casetones ‖ Y (verticales, en x) y ‖ X (horizontales, en y)
    for nm, nE, s0, nO, o0, isx in (('Ncas_v', '{k_nXe}', '{k_xs}', '{k_nYe}', '{k_ys}', True), ('Ncas_h', '{k_nYe}', '{k_ys}', '{k_nXe}', '{k_xs}', False)):
        pos = lambda j, s0=s0: '(%s+INT(%s/2)*{k_s}+MOD(%s,2)*{k_ac})' % (s0, j, j)
        a0 = o0; a1 = '(%s+%s*{k_s}-{k_bw})' % (o0, nO)
        uf, vf = serp('2*%s' % nE, pos, a0, a1)
        if isx:
            dyn(nm, lambda k, uf=uf, vf=vf: PX(vf(k)), lambda k, uf=uf: PY(uf(k)), 100,
                '="Casetones "&TEXT({k_ac},"0")&"×"&TEXT({k_ac},"0")&"×"&{hc}&" cm: "&{k_nXe}&" × "&{k_nYe}&" (nervaduras de "&{k_bw}&" cm entre ellos, a cada "&{s}&" cm)"', HN + ',%s>0,%s>0' % (nE, nO))
        else:
            dyn(nm, lambda k, uf=uf: PX(uf(k)), lambda k, uf=uf, vf=vf: PY(vf(k)), 100, None, HN + ',%s>0,%s>0' % (nE, nO))
    # ejes de nervaduras interiores
    for nm, nE, s0, isx in (('Neje_v', '{k_nXe}', '{k_xs}', True), ('Neje_h', '{k_nYe}', '{k_ys}', False)):
        pos = lambda j, s0=s0: '(%s+{k_ac}+{k_bw}/2+%s*{k_s})' % (s0, j)
        a0 = '{k_bp}/2'; a1 = ('({k_ay}-{k_bp}/2)' if isx else '({k_ax}-{k_bp}/2)')
        uf, vf = serp('%s-1' % nE, pos, a0, a1)
        if isx:
            dyn(nm, lambda k, uf=uf, vf=vf: PX(vf(k)), lambda k, uf=uf: PY(uf(k)), 60,
                '="Ejes de nervaduras (‖ X y ‖ Y)"', HN + ',%s>=2' % nE)
        else:
            dyn(nm, lambda k, uf=uf: PX(uf(k)), lambda k, uf=uf, vf=vf: PY(vf(k)), 60, None, HN + ',%s>=2' % nE)
    # zona maciza por cortante (junto a los apoyos del claro corto)
    uL = '{k_us}'; uR = '({k_us}+{k_nAe}*{k_s}-{k_bw})'; v0 = '{k_bp}/2'; v1 = '({k_C}-{k_bp}/2)'; p0 = '{k_bp}/2'; p1 = '({k_A}-{k_bp}/2)'
    showL = 'IF({k_vol}=1,{k_sup0},1)'; showR = 'IF({k_vol}=1,1-{k_sup0},1)'
    ul = [p0, uL, uL, p0, p0]; vl = [v0, v0, v1, v1, v0]
    ur = [uR, uR, p1, p1, uR]; vr = [v0, v1, v1, v0, v0]
    U10 = ['IF(%s=1,%s,%s)' % (showL, a, uR) for a in ul] + ['IF(%s=1,%s,%s)' % (showR, a, p0) for a in ur]
    V10 = ['IF(%s=1,%s,%s)' % (showL, a, v0) for a in vl] + ['IF(%s=1,%s,%s)' % (showR, a, v0) for a in vr]
    dyn('Nmacizo', lambda k: uv(ch_(k, 10, U10), ch_(k, 10, V10))[0], lambda k: uv(ch_(k, 10, U10), ch_(k, 10, V10))[1], 10,
        '="Zona maciza junto al apoyo: macizo de orilla + "&{k_nh}&" hilera(s) de casetones quitadas por cortante"', HN + ',{k_nh}>0,{k_nAe}>0')
    # varillas representativas en la nervadura central
    off = {'inf': '-0.9', 'sup': '0.9', 'bas': '2.1'}
    # ‖ corto (a lo largo de u, en v = k_vrib)
    supc0 = 'IF({k_vol}=1,IF({k_sup0}=1,-({k_bp}/2+{k_Lbc}),2),2)'; supc1 = 'IF({k_vol}=1,IF({k_sup0}=1,{k_A}-2,{k_A}+{k_bp}/2+{k_Lbc}),{k_A}-2)'
    for nm, a0, a1, o, lab, hide in (
            ('Nc_inf', '2', '{k_A}-2', off['inf'], '="‖ corto · inferior: "&{k_tPc}', HN),
            ('Nc_sup', supc0, supc1, off['sup'], '=IF({k_vol}=1,"‖ corto · superior del volado: "&{k_qNc}&" No. "&{k_vNc}&" (corrida y bastones: ver ARMADO en NERVADAS)","‖ corto · superior: "&{k_tNc})', HN),
            ('Nc_b0', 'IF({k_c0}=1,-12,2)', '{k_bp}/2+{k_Lbc}', off['bas'], '="‖ corto · bastones (borde continuo"&IF({k_bdc}=1," y discontinuo","")&"), L = "&TEXT({k_Lbc}/100,"0.00")&" m desde el paño"', HN + ',{k_vol}=0,{k_qNc}>=2,OR({k_c0}=1,{k_bdc}=1)'),
            ('Nc_b1', '{k_A}-{k_bp}/2-{k_Lbc}', 'IF({k_c1}=1,{k_A}+12,{k_A}-2)', off['bas'], None, HN + ',{k_vol}=0,{k_qNc}>=2,OR({k_c1}=1,{k_bdc}=1)')):
        dyn(nm, lambda k, a0=a0, a1=a1, o=o: uv(ch_(k, 2, [a0, a1]), '{k_vrib}', o)[0], lambda k, a0=a0, a1=a1, o=o: uv(ch_(k, 2, [a0, a1]), '{k_vrib}', o)[1], 2, lab, hide)
    for nm, a0, a1, o, lab, hide in (
            ('Nl_inf', '2', '{k_C}-2', off['inf'], '="‖ largo · inferior: "&{k_tPl}', HN),
            ('Nl_sup', '2', '{k_C}-2', off['sup'], '="‖ largo · superior: "&{k_tNl}', HN),
            ('Nl_b0', 'IF({k_v0c}=1,-12,2)', '{k_bp}/2+{k_Lbl}', off['bas'], '="‖ largo · bastones, L = "&TEXT({k_Lbl}/100,"0.00")&" m desde el paño"', HN + ',{k_vol}=0,{k_qNl}>=2,OR({k_v0c}=1,{k_bdl}=1)'),
            ('Nl_b1', '{k_C}-{k_bp}/2-{k_Lbl}', 'IF({k_v1c}=1,{k_C}+12,{k_C}-2)', off['bas'], None, HN + ',{k_vol}=0,{k_qNl}>=2,OR({k_v1c}=1,{k_bdl}=1)')):
        dyn(nm, lambda k, a0=a0, a1=a1, o=o: uv_l('{k_urib}', ch_(k, 2, [a0, a1]), o)[0], lambda k, a0=a0, a1=a1, o=o: uv_l('{k_urib}', ch_(k, 2, [a0, a1]), o)[1], 2, lab, hide)
    cota(PX('{k_xs}'), PX('{k_xs}+{k_s}'), True, '(%s+4)' % PY('{k_ymax}'), 'Ncota_s', '="s = "&{s}&" cm"', HN + ',{k_nXe}>=2')
    # ---- maciza: varillas a su separación real
    HM = '{k_M}=1'
    nb = lambda L, sp: 'MIN(80,MAX(1,INT(((%s)-{k_bp})/MAX(%s,1))+1))' % (L, sp)
    first = lambda L, sp: '({k_bp}/2+(((%s)-{k_bp})-(%s-1)*%s)/2)' % (L, nb(L, sp), sp)
    n1 = nb('{k_C}', '{k_qPc}'); f1 = first('{k_C}', '{k_qPc}')
    uf, vf = serp5(n1, lambda j: '(%s+%s*{k_qPc})' % (f1, j), '2', '{k_A}-2', '0')
    dyn('Mc_inf', lambda k, uf=uf, vf=vf: uv(uf(k), vf(k))[0], lambda k, uf=uf, vf=vf: uv(uf(k), vf(k))[1], 400, '="‖ corto · inferior: "&{k_tPc}', HM + ',{k_qPc}>0')
    n2 = nb('{k_A}', '{k_qPl}'); f2 = first('{k_A}', '{k_qPl}')
    uf, vf = serp5(n2, lambda j: '(%s+%s*{k_qPl})' % (f2, j), '2', '{k_C}-2', '0')
    dyn('Ml_inf', lambda k, uf=uf, vf=vf: uv(vf(k), uf(k))[0], lambda k, uf=uf, vf=vf: uv(vf(k), uf(k))[1], 400, '="‖ largo · inferior: "&{k_tPl}', HM + ',{k_qPl}>0')
    # bastones en los cuatro bordes (perimetral) y acero superior del volado
    nN = nb('{k_C}', '{k_qNc}'); fN = first('{k_C}', '{k_qNc}')
    for nm, a0, a1, U0, cond, lab in (('Mc_b0', 'IF({k_c0}=1,-12,2)', '{k_bp}/2+{k_Lbc}', '0', 'OR({k_c0}=1,{k_bdc}=1)', '="‖ corto · bastones superiores: "&{k_tNc}'),
                                      ('Mc_b1', '{k_A}-{k_bp}/2-{k_Lbc}', 'IF({k_c1}=1,{k_A}+12,{k_A}-2)', '{k_A}', 'OR({k_c1}=1,{k_bdc}=1)', None)):
        uf, vf = serp5(nN, lambda j: '(%s+%s*{k_qNc})' % (fN, j), a0, a1, U0)
        dyn(nm, lambda k, uf=uf, vf=vf: uv(uf(k), vf(k))[0], lambda k, uf=uf, vf=vf: uv(uf(k), vf(k))[1], 400, lab, HM + ',{k_vol}=0,{k_qNc}>0,' + cond)
    nL = nb('{k_A}', '{k_qNl}'); fL = first('{k_A}', '{k_qNl}')
    for nm, a0, a1, U0, cond, lab in (('Ml_b0', 'IF({k_v0c}=1,-12,2)', '{k_bp}/2+{k_Lbl}', '0', 'OR({k_v0c}=1,{k_bdl}=1)', '="‖ largo · bastones superiores: "&{k_tNl}'),
                                      ('Ml_b1', '{k_C}-{k_bp}/2-{k_Lbl}', 'IF({k_v1c}=1,{k_C}+12,{k_C}-2)', '{k_C}', 'OR({k_v1c}=1,{k_bdl}=1)', None)):
        uf, vf = serp5(nL, lambda j: '(%s+%s*{k_qNl})' % (fL, j), a0, a1, U0)
        dyn(nm, lambda k, uf=uf, vf=vf: uv(vf(k), uf(k))[0], lambda k, uf=uf, vf=vf: uv(vf(k), uf(k))[1], 400, lab, HM + ',{k_vol}=0,{k_qNl}>0,' + cond)
    uf, vf = serp5(nN, lambda j: '(%s+%s*{k_qNc})' % (fN, j), supc0, supc1, 'IF({k_sup0}=1,0,{k_A})')
    dyn('Mvol_sup', lambda k, uf=uf, vf=vf: uv(uf(k), vf(k))[0], lambda k, uf=uf, vf=vf: uv(uf(k), vf(k))[1], 400, '="Superior del volado: "&{k_tNc}', HM + ',{k_vol}=1,{k_qNc}>0')

    # ================================================= CORTE DE LA NERVADURA (cm reales, escala k_cs)
    CX = lambda x: '(50+(%s)*{k_cs})' % x
    CY = lambda y: '(12+(%s)*{k_cs})' % y
    hs = '{k_s}/2'
    poly = ['-' + hs, hs, hs, '{k_bw}/2', '{k_bw}/2', '-{k_bw}/2', '-{k_bw}/2', '-' + hs, '-' + hs]
    polyy = ['{k_h}', '{k_h}', '{k_hc}', '{k_hc}', '0', '0', '{k_hc}', '{k_hc}', '{k_h}']
    dyn('Kconc', lambda k: ch_(k, 9, [CX(p) for p in poly]), lambda k: ch_(k, 9, [CY(p) for p in polyy]), 9,
        '="Nervadura b_w = "&{k_bw}&" cm + capa t = "&{t}&" cm (h = "&{h}&" cm)"', HN)
    fx, fy = rect('-' + hs, '0', '-{k_bw}/2', '{k_hc}', CX, CY)
    dyn('Kcas1', fx, fy, 5, '="Casetón "&TEXT({k_ac},"0")&" × "&{hc}&" cm"', HN)
    fx, fy = rect('{k_bw}/2', '0', hs, '{k_hc}', CX, CY)
    dyn('Kcas2', fx, fy, 5, None, HN)
    ym = '({k_h}-{k_rtm}-{k_fim})'
    dyn('Kmalla', lambda k: ch_(k, 2, [CX('-%s+1' % hs), CX('%s-1' % hs)]), lambda k: CY(ym), 2, '="Malla "&{malla_sel}&" al centro de la capa ("&TEXT({k_rtm},"0.0")&" cm a la cara superior)"', HN)
    for top in (False, True):
        blk = 'Nc' if top else 'Pc'
        for i in (1, 2, 3):
            dbb = '{k_db%s}' % blk; n = '{k_q%s}' % blk
            xc = '((%d-(%s+1)/2)*(MAX({k_slib},%s)+%s))' % (i, n, dbb, dbb)
            yc = '({k_h}-{k_rtm}-2*{k_fim}-%s/2)' % dbb if top else '({k_rb}+%s/2)' % dbb
            nm = 'K%s%d' % (blk, i)
            lab = ('="Superior ‖ corto: "&{k_qNc}&" No. "&{k_vNc}&" junto al apoyo (armado completo en la planta)"' if top else '="Inferior ‖ corto: "&{k_qPc}&" No. "&{k_vPc}&" (r = "&{k_rb}&" cm)"') if i == 1 else None
            dyn(nm, lambda k, xc=xc, dbb=dbb: CX('%s+%s/2*COS(2*PI()*(MIN(%s,13)-1)/12)' % (xc, dbb, k)),
                lambda k, yc=yc, dbb=dbb: CY('%s+%s/2*SIN(2*PI()*(MIN(%s,13)-1)/12)' % (yc, dbb, k)), 13, lab, HN + ',%s>=%d' % (n, i))
    cota(CX('-{k_bw}/2'), CX('{k_bw}/2'), True, CY('-4'), 'Kc_bw', '="b_w = "&{k_bw}&" cm"', HN)
    cota(CX('-' + hs), CX(hs), True, CY('{k_h}+4'), 'Kc_s', '="s = "&{s}&" cm"', HN)
    cota(CY('0'), CY('{k_h}'), False, CX(hs + '+4'), 'Kc_h', '="h = "&{h}&" cm"', HN)
    cota(CY('{k_h}-{k_dPc}'), CY('{k_h}'), False, CX('-' + hs + '-4'), 'Kc_d', '="d = "&TEXT({k_dPc},"0.0")&" cm (inf. ‖ corto)"', HN)

    # ================================================= CORTE DE LA MACIZA (franja de 60 cm, escala k_cm)
    MX = lambda x: '(50+(%s)*{k_cm})' % x
    MY = lambda y: '(12+(%s)*{k_cm})' % y
    fx, fy = rect('-30', '0', '30', '{k_h}', MX, MY)
    dyn('Lconc', fx, fy, 5, '="Losa maciza h = "&{k_h}&" cm (franja de 60 cm)"', HM)
    yPc = '({k_rb}+{k_dbPc}/2)'; yNc = '({k_h}-{k_rt}-{k_dbNc}/2)'
    dyn('LPc', lambda k: ch_(k, 2, [MX('-29'), MX('29')]), lambda k: MY(yPc), 2, '="Inferior ‖ corto: "&{k_tPc}&" (r = "&{k_rb}&" cm)"', HM)
    nmk = lambda sp: 'MIN(12,INT(58/MAX(%s,1))+1)' % sp
    for nm, sp, yy, lab, cond in (('LPl', '{k_qPl}', '({k_rb}+{k_dbPc}+{k_dbPl}/2)', '="Inferior ‖ largo: No. "&{k_vPl}&" @ "&{k_qPl}&" cm"', HM),
                                  ('LNl', '{k_qNl}', '({k_h}-{k_rt}-{k_dbNc}-{k_dbNl}/2)', '="Superior ‖ largo: No. "&{k_vNl}&" @ "&{k_qNl}&" cm"', HM)):
        xi = lambda k, sp=sp: '(-(%s-1)*%s/2+(MIN(%s,%s)-1)*%s)' % (nmk(sp), sp, k, nmk(sp), sp)
        dyn(nm, lambda k, xi=xi: MX(xi(k)), lambda k, yy=yy: MY(yy), 12, lab, cond + ',%s>0' % sp)
    dyn('LNc', lambda k: ch_(k, 2, [MX('-29'), MX('29')]), lambda k: MY(yNc), 2, '="Superior ‖ corto: No. "&{k_vNc}&" @ "&{k_qNc}&" cm (r = "&{k_rt}&" cm)"', HM)
    cota(MY('0'), MY('{k_h}'), False, MX('33'), 'Lc_h', '="h = "&{k_h}&" cm"', HM)
    cota(MY('0'), MY('{k_rb}'), False, MX('-33'), 'Lc_rb', '="r = "&{k_rb}&" cm"', HM, 0.8)
    cota(MY('{k_h}-{k_rt}'), MY('{k_h}'), False, MX('-33'), 'Lc_rt', '="r = "&{k_rt}&" cm"', HM, 0.8)
    x1 = '(-(%s-1)*{k_qPl}/2)' % nmk('{k_qPl}')
    cota(MX(x1), MX(x1 + '+{k_qPl}'), True, MY('-4'), 'Lc_s', '="s = "&{k_qPl}&" cm"', HM + ',%s>=2' % nmk('{k_qPl}'))

    # ================================================= ALZADO ‖ CORTO (los dos grupos)
    AX = lambda x: '(4+((%s)-{k_xL})*{k_sa})' % x
    AY = lambda y: '(30+(%s)*{k_sa}*{k_E})' % y
    mir = lambda u: 'IF(AND({k_vol}=1,{k_sup0}=0),{k_A}-(%s),%s)' % (u, u)
    for g, H in (('N', HN), ('M', HM)):
        xa = 'IF({k_vol}=1,-{k_BKa},IF({k_c0}=1,-{k_bp}/2-30,-{k_bp}/2))'
        xb = 'IF({k_vol}=1,{k_A},IF({k_c1}=1,{k_A}+{k_bp}/2+30,{k_A}+{k_bp}/2))'
        fx, fy = rect(xa, '0', xb, '{k_h}', AX, AY)
        dyn(g + 'Alosa', fx, fy, 5, '=IF({k_vol}=1,"Volado y parte del tablero de respaldo","Losa ‖ corto (corte continuo donde el borde es continuo)")&" — escala vertical ×"&{k_E}', H)
        # apoyos (10 unidades hacia abajo)
        sL = [AX('-{k_bp}/2'), AX('{k_bp}/2'), AX('{k_bp}/2'), AX('-{k_bp}/2'), AX('-{k_bp}/2')]
        sR = [AX('{k_A}-{k_bp}/2'), AX('{k_A}+{k_bp}/2'), AX('{k_A}+{k_bp}/2'), AX('{k_A}-{k_bp}/2'), AX('{k_A}-{k_bp}/2')]
        sY = ['30', '30', '20', '20', '30']
        SX = sL + ['IF({k_vol}=1,%s,%s)' % (sL[4], a) for a in sR]
        SY = sY + ['IF({k_vol}=1,30,%s)' % a for a in sY]
        dyn(g + 'Aapoyo', lambda k, SX=SX: ch_(k, 10, SX), lambda k, SY=SY: ch_(k, 10, SY), 10, '="Apoyos (b_p = "&{bp}&" cm)"', H)
        # inferior corrida con gancho en bordes discontinuos
        dbP = '{k_dbPc}'; yb = '({k_rb}+%s/2)' % dbP
        hk0 = 'IF(AND({k_vol}=0,{k_c0}=0),MIN(12*%s,{k_h}-{k_rb}-{k_rt}),0)' % dbP
        hk1 = 'IF(AND({k_vol}=0,{k_c1}=0),MIN(12*%s,{k_h}-{k_rb}-{k_rt}),0)' % dbP
        xi0 = '-{k_bp}/2+3'; xi1 = 'IF({k_vol}=1,{k_A}-3,{k_A}+{k_bp}/2-3)'
        dyn(g + 'Ainf', lambda k: ch_(k, 4, [AX(xi0), AX(xi0), AX(xi1), AX(xi1)]), lambda k: ch_(k, 4, [AY('%s+%s' % (yb, hk0)), AY(yb), AY(yb), AY('%s+%s' % (yb, hk1))]), 4,
            '="Inferior ‖ corto: "&{k_tPc}', H)
        dbN = '{k_dbNc}'
        yt = ('({k_h}-{k_rtm}-2*{k_fim}-%s/2)' if g == 'N' else '({k_h}-{k_rt}-%s/2)') % dbN
        hkt = 'MIN(12*%s,{k_h}-{k_rb}-{k_rt})' % dbN
        # superior corrida (nervada) o principal del volado (nervada y maciza)
        xs0 = 'IF({k_vol}=1,-({k_bp}/2+{k_Lbc}),IF({k_c0}=1,-{k_bp}/2-30,-{k_bp}/2+3))'
        xs1 = 'IF({k_vol}=1,{k_A}-3,IF({k_c1}=1,{k_A}+{k_bp}/2+30,{k_A}+{k_bp}/2-3))'
        hs0 = 'IF(AND({k_vol}=0,{k_c0}=0),%s,0)' % hkt
        hs1 = 'IF(OR({k_vol}=1,{k_c1}=0),%s,0)' % hkt
        cond = H + (',1' if g == 'N' else ',{k_vol}=1')
        dyn(g + 'Asup', lambda k: ch_(k, 4, [AX(xs0), AX(xs0), AX(xs1), AX(xs1)]), lambda k: ch_(k, 4, [AY('%s-%s' % (yt, hs0)), AY(yt), AY(yt), AY('%s-%s' % (yt, hs1))]), 4,
            ('=IF({k_vol}=1,"Superior del volado: "&{k_qNc}&" No. "&{k_vNc}&" (armado completo en la planta)","Superior corrido: 1 No. "&{k_vNc}&" (6.7.6.1.3)")' if g == 'N'
             else '="Superior del volado: No. "&{k_vNc}&" @ "&{k_qNc}&" cm"'), cond)
        # bastones
        ybs = ('(%s-1.3*%s-0.4)' % (yt, dbN)) if g == 'N' else yt
        need = '{k_qNc}>=2' if g == 'N' else '{k_qNc}>0'
        for side in (0, 1):
            cside = '{k_c%d}' % side
            if side == 0:
                xA = 'IF(%s=1,-{k_bp}/2-30,-{k_bp}/2+3)' % cside; xB = '{k_bp}/2+{k_Lbc}'
                pts_x = [AX(xA), AX(xA), AX(xB)]
                pts_y = [AY('%s-IF(%s=1,0,%s)' % (ybs, cside, hkt)), AY(ybs), AY(ybs)]
            else:
                xA = '{k_A}-{k_bp}/2-{k_Lbc}'; xB = 'IF(%s=1,{k_A}+{k_bp}/2+30,{k_A}+{k_bp}/2-3)' % cside
                pts_x = [AX(xA), AX(xB), AX(xB)]
                pts_y = [AY(ybs), AY(ybs), AY('%s-IF(%s=1,0,%s)' % (ybs, cside, hkt))]
            lab = (('="Bastones ‖ corto: "&({k_qNc}-1)&" No. "&{k_vNc}&", L = "&TEXT({k_Lbc}/100,"0.00")&" m desde el paño"' if g == 'N'
                    else '="Bastones ‖ corto: No. "&{k_vNc}&" @ "&{k_qNc}&" cm, L = "&TEXT({k_Lbc}/100,"0.00")&" m desde el paño"') if side == 0 else None)
            dyn(g + 'Ab%d' % side, lambda k, px=pts_x: ch_(k, 3, px), lambda k, py=pts_y: ch_(k, 3, py), 3, lab,
                H + ',{k_vol}=0,%s,OR(%s=1,{k_bdc}=1)' % (need, cside))
        if g == 'N':
            # casetones detrás de la nervadura y límite de la zona maciza
            cu0 = '{k_us}'; cu1 = '({k_us}+{k_nAe}*{k_s}-{k_bw})'
            ca = mir(cu0); cb = mir(cu1)
            fx, fy = rect('MIN(%s,%s)' % (ca, cb), '0', 'MAX(%s,%s)' % (ca, cb), '{k_hc}', AX, AY)
            dyn('NAcas', fx, fy, 5, '="Casetones (detrás de la nervadura)"', H + ',{k_nAe}>0')
            dyn('NAmac', lambda k: ch_(k, 4, [AX('MIN(%s,%s)' % (ca, cb)), AX('MIN(%s,%s)' % (ca, cb)), AX('MAX(%s,%s)' % (ca, cb)), AX('MAX(%s,%s)' % (ca, cb))]),
                lambda k: ch_(k, 4, [AY('0'), AY('{k_h}'), AY('{k_h}'), AY('0')]), 4,
                '="Fin de la zona maciza (macizo de orilla"&IF({k_nh}>0," + "&{k_nh}&" hilera(s) por cortante","")&")"', H + ',{k_nAe}>0')
        cota(AX('0'), AX('{k_A}'), True, '15', g + 'Ac_a', '=IF({k_vol}=1,"Volado: "&TEXT({k_A}/100,"0.00")&" m del eje del apoyo a la orilla","a1 = "&TEXT({k_A}/100,"0.00")&" m (a ejes)")', H)
        LBx0 = 'IF({k_vol}=1,-{k_bp}/2,{k_bp}/2)'; LBx1 = 'IF({k_vol}=1,-{k_bp}/2-{k_Lbc},{k_bp}/2+{k_Lbc})'
        cota(AX(LBx0), AX(LBx1), True, '(%s+5)' % AY('{k_h}'), g + 'Ac_L', '="L = "&TEXT({k_Lbc}/100,"0.00")&" m desde el paño"',
             H + ',OR({k_vol}=1,AND(%s,OR({k_c0}=1,{k_bdc}=1)))' % need)
        tip = AX('{k_A}'); ty = AY('{k_h}')
        dyn(g + 'AP', lambda k: ch_(k, 5, [tip, tip, '%s-1.2' % tip, tip, '%s+1.2' % tip]), lambda k: ch_(k, 5, ['%s+13' % ty, '%s+1' % ty, '%s+4' % ty, '%s+1' % ty, '%s+4' % ty]), 5,
            '="P = "&TEXT({k_P},"0.0")&" kg/m"', H + ',{k_vol}=1,{k_P}>0')
    ws['A%d' % (LAB0 - 1)] = 'TEXTOS DE LEYENDAS Y COTAS'; ws['A%d' % (LAB0 - 1)].font = font(10, True, color=HDR_FG)

    # ================================================= gráficas
    small = lambda sz: RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=sz)), endParaRPr=CharacterProperties(sz=sz))])

    def mk_title(txt, sz=900):
        cp = CharacterProperties(sz=sz, b=True, solidFill='1F3864')
        return Title(tx=Text(rich=RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), r=[RegularTextRun(rPr=cp, t=txt)])])), overlay=False)

    def mk_chart(title, xmin, xmax, ymin, ymax, w_cm, h_cm, hfrac, ytop=0.10):
        # título fijo (el tablero, su armado y su estado van en el encabezado de la hoja y en las leyendas)
        ch = ScatterChart(); ch.style = 13; ch.scatterStyle = 'lineMarker'
        ch.title = mk_title(title) if title else None
        ch.x_axis.scaling.min = xmin; ch.x_axis.scaling.max = xmax; ch.y_axis.scaling.min = ymin; ch.y_axis.scaling.max = ymax
        ch.x_axis.delete = True; ch.y_axis.delete = True
        ch.x_axis.majorGridlines = None; ch.y_axis.majorGridlines = None
        ch.width = w_cm; ch.height = h_cm
        ch.legend.position = 'b'
        ch.legend.txPr = small(700)
        ch.display_blanks = 'gap'
        wfrac = (xmax - xmin) / (ymax - ymin) * hfrac * h_cm / w_cm
        if wfrac > 0.96:
            hfrac = hfrac * 0.96 / wfrac; wfrac = 0.96
        ch.plot_area.layout = Layout(manualLayout=ManualLayout(layoutTarget='inner', xMode='edge', yMode='edge', x=(1 - wfrac) / 2, y=ytop, w=wfrac, h=hfrac))
        return ch

    def add(ch, name, color='1F4E79', w=1.5, dash='solid', marker=None, msize=6, lbl_idx=None, noleg=False, pos='t', lsz=750, bold=False):
        cx, cy, r0, r1 = series[name]
        xr = Reference(ws, min_col=CI(cx), min_row=r0, max_row=r1)
        yr = Reference(ws, min_col=CI(cy), min_row=r0, max_row=r1)
        s = Series(yr, xr)
        if name in labels:
            s.tx = SeriesLabel(strRef=StrRef(labels[name]))
        s.smooth = False
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = int(w * 12700)
        s.graphicalProperties.line.prstDash = dash
        if marker:
            s.marker = Marker(symbol=marker, size=msize)
            s.marker.graphicalProperties = GraphicalProperties(solidFill=color)
            s.marker.graphicalProperties.line.solidFill = color
            s.graphicalProperties.line = LineProperties(noFill=True)
        else:
            s.marker = Marker(symbol='none')
        if lbl_idx is not None:
            tp = RichText(bodyPr=RichTextProperties(wrap='none'), p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=lsz, b=bold, solidFill=(color if bold else None))), endParaRPr=CharacterProperties(sz=lsz, b=bold))])
            s.dLbls = DataLabelList(dLbl=[DataLabel(idx=lbl_idx, showSerName=True, showVal=False, showCatName=False, showLegendKey=False, showPercent=False, dLblPos=pos, txPr=tp)])
            s.dLbls.showVal = False; s.dLbls.showSerName = False; s.dLbls.showLegendKey = False; s.dLbls.showCatName = False
        if (noleg or name not in labels or lbl_idx is not None) and ch.legend is not None:
            ch.legend.legendEntry.append(LegendEntry(idx=len(ch.series), delete=True))
        ch.series.append(s)
        return s

    def place(ch, row, right=False, dy_cm=0.0):
        if right:
            px = sum(int(wsK.column_dimensions[k].width * 7 + 5) for k in 'BCD')
            off = max(0.0, 15.5 - px * 2.54 / 96)
            anchor = OneCellAnchor(_from=AnchorMarker(col=4, colOff=cm_to_EMU(off), row=row - 1, rowOff=cm_to_EMU(dy_cm)),
                                   ext=XDRPositiveSize2D(cm_to_EMU(ch.width), cm_to_EMU(ch.height)))
        else:
            anchor = OneCellAnchor(_from=AnchorMarker(col=1, colOff=0, row=row - 1, rowOff=cm_to_EMU(dy_cm)),
                                   ext=XDRPositiveSize2D(cm_to_EMU(ch.width), cm_to_EMU(ch.height)))
        wsK.add_chart(ch, anchor)

    T = lambda nm: 'CROQUIS!' + N[nm].split('!')[1]
    for g, rowg in (('N', N['k_row_N'] + 1), ('M', N['k_row_M'] + 1)):
        ch = mk_chart('PLANTA DEL TABLERO — ' + ('CASETONES, NERVADURAS Y ARMADO POR NERVADURA' if g == 'N' else 'ARMADO POR METRO'), 0, 100, 0, 100, 15.2, 19.6, 0.66, 0.08)
        add(ch, g + 'respaldo', '7F7F7F', 1.0, 'dash')
        if g == 'N':
            add(ch, 'Nresp', 'ED7D31', 1.5, 'dash')
        add(ch, g + 'contorno', '1F3864', 2.2)
        add(ch, g + 'apoyo', 'C00000', 1.0, 'dash')
        add(ch, g + 'achurado', '595959', 0.75, noleg=True)
        if g == 'N':
            add(ch, 'Ncas_v', 'A6A6A6', 0.75); add(ch, 'Ncas_h', 'A6A6A6', 0.75, noleg=True)
            add(ch, 'Neje_v', 'BDD7EE', 0.5, 'sysDot'); add(ch, 'Neje_h', 'BDD7EE', 0.5, 'sysDot', noleg=True)
            add(ch, 'Nmacizo', 'ED7D31', 1.25, 'dash')
            add(ch, 'Nc_inf', '0070C0', 2.5); add(ch, 'Nc_sup', '7F4F00', 1.75)
            add(ch, 'Nc_b0', 'C00000', 2.5); add(ch, 'Nc_b1', 'C00000', 2.5, noleg=True)
            add(ch, 'Nl_inf', '00B050', 2.0); add(ch, 'Nl_sup', '375623', 1.5)
            add(ch, 'Nl_b0', 'ED7D31', 2.0); add(ch, 'Nl_b1', 'ED7D31', 2.0, noleg=True)
            add(ch, 'Ncota_s', '7030A0', 0.75, lbl_idx=3)
        else:
            add(ch, 'Mc_inf', '0070C0', 0.75); add(ch, 'Ml_inf', '00B050', 0.75)
            add(ch, 'Mc_b0', 'C00000', 1.25); add(ch, 'Mc_b1', 'C00000', 1.25, noleg=True)
            add(ch, 'Ml_b0', 'ED7D31', 1.25); add(ch, 'Ml_b1', 'ED7D31', 1.25, noleg=True)
            add(ch, 'Mvol_sup', 'C00000', 1.25)
        add(ch, g + 'contorno', '1F3864', 2.2, noleg=True)
        add(ch, g + 'carga', '7030A0', 3.5)
        add(ch, g + 'cota_ax', '7030A0', 0.75, lbl_idx=3)
        add(ch, g + 'cota_ay', '7030A0', 0.75, lbl_idx=3)
        place(ch, rowg)
        if g == 'N':
            ch = mk_chart('CORTE DE LA NERVADURA ‖ CORTO (cm)', 0, 100, 0, 48, 13.8, 9.6, 0.58, 0.12)
            add(ch, 'Kconc', '1F3864', 2.2)
            add(ch, 'Kcas1', 'A6A6A6', 1.0, 'dash'); add(ch, 'Kcas2', 'A6A6A6', 1.0, 'dash', noleg=True)
            add(ch, 'Kmalla', '00B050', 1.0, 'sysDash')
            for i in (1, 2, 3):
                add(ch, 'KPc%d' % i, '0070C0', 1.5, noleg=i > 1)
            for i in (1, 2, 3):
                add(ch, 'KNc%d' % i, 'C00000', 1.5, noleg=i > 1)
            for nm in ('Kc_bw', 'Kc_s', 'Kc_h', 'Kc_d'):
                add(ch, nm, '7F7F7F', 0.75, lbl_idx=3)
            place(ch, rowg, right=True)
        else:
            ch = mk_chart('CORTE DE LA LOSA JUNTO AL APOYO (cm)', 0, 100, 0, 48, 13.8, 9.6, 0.58, 0.12)
            add(ch, 'Lconc', '1F3864', 2.2)
            add(ch, 'LPc', '0070C0', 2.0); add(ch, 'LNc', 'C00000', 2.0)
            add(ch, 'LPl', '00B050', marker='circle', msize=5); add(ch, 'LNl', 'ED7D31', marker='circle', msize=5)
            for nm in ('Lc_h', 'Lc_rb', 'Lc_rt', 'Lc_s'):
                add(ch, nm, '7F7F7F', 0.75, lbl_idx=3)
            place(ch, rowg, right=True)
        ch = mk_chart('ALZADO ‖ CLARO CORTO' + (' DE LA NERVADURA' if g == 'N' else ' (FRANJA DE 1 m)'), 0, 100, 0, 64, 13.8, 9.8, 0.58, 0.12)
        add(ch, g + 'Alosa', '1F3864', 2.0)
        add(ch, g + 'Aapoyo', '7F7F7F', 2.0)
        if g == 'N':
            add(ch, 'NAcas', 'A6A6A6', 0.75, 'dash'); add(ch, 'NAmac', 'ED7D31', 1.25, 'dash')
        add(ch, g + 'Ainf', '0070C0', 2.25)
        if g == 'N':
            add(ch, g + 'Asup', '7F4F00', 1.75)
        else:
            add(ch, g + 'Asup', 'C00000', 1.75)
        add(ch, g + 'Ab0', 'C00000', 2.25); add(ch, g + 'Ab1', 'C00000', 2.25, noleg=True)
        add(ch, g + 'Ac_a', '7030A0', 0.75, lbl_idx=3)
        add(ch, g + 'Ac_L', '7030A0', 0.75, lbl_idx=3)
        add(ch, g + 'AP', '7030A0', 2.0, lbl_idx=0)
        place(ch, rowg, right=True, dy_cm=9.8)
    ws.sheet_state = 'visible'
    return ws
