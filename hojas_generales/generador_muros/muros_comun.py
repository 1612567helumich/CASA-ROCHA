"""Nombres compartidos, TABLAS, DATOS, INSTRUCCIONES y REFERENCIAS de la hoja de muros de mampostería confinada (NTC 2023)."""
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as CL
from xlstyle import *            # noqa: F401,F403
from muros_spec import *         # noqa: F401,F403

N = {}                           # nombre -> referencia absoluta (DATOS!$C$12 ...)
TB = {}                          # rangos de TABLAS

BARS = [  # designación, d_b (cm), área (cm²), kg/m, uso
    (2, 0.635, 0.32, 0.249, 'Alambrón No. 2 (1/4"): estribos de castillos y dalas; también sirve en juntas (6.4 mm ≤ 0.75·h_j sólo con juntas ≥ 9 mm)'),
    (2.5, 0.794, 0.49, 0.388, 'No. 2.5 (5/16"): estribos'),
    (3, 0.953, 0.71, 0.560, 'No. 3 (3/8"): castillos y dalas'),
    (4, 1.270, 1.27, 0.994, 'No. 4 (1/2")'),
    (5, 1.588, 1.99, 1.552, 'No. 5 (5/8")'),
    (6, 1.905, 2.87, 2.235, 'No. 6 (3/4")'),
    (8, 2.540, 5.07, 3.973, 'No. 8 (1")'),
    ('5/32 in', 0.397, 0.124, 0.097, 'Alambre corrugado laminado en frío de 5/32" (4.0 mm), para juntas'),
    ('3/16 in', 0.476, 0.178, 0.140, 'Alambre corrugado laminado en frío de 3/16" (4.8 mm), para juntas'),
]
MALLAS = [  # designación, Ø alambre (mm), A_s horizontal (cm²/m), kg/m²
    ('6x6-10/10', 3.43, 0.61, 0.95), ('6x6-8/8', 4.11, 0.87, 1.37), ('6x6-6/6', 4.88, 1.23, 1.93), ('6x6-4/4', 5.72, 1.69, 2.65)]
FM_CONC = [(60, 25, 20), (75, 40, 35), (100, 50, 45), (150, 75, 60), (200, 100, 90)]
FM_ARC = [(60, 20, 20), (75, 30, 30), (100, 40, 40), (150, 60, 60), (200, 80, 70), (300, 120, 90)]
PIEZAS = [  # pieza, material, ¿hueca?, f'm Tipo II, v'm Tipo II
    ('BLOQUE DE CONCRETO HUECO', 'CONCRETO', 'SÍ', 15, 2), ('BLOQUE DE CONCRETO MACIZO', 'CONCRETO', 'NO', 15, 2),
    ('TABICÓN (TABIQUE DE CONCRETO)', 'CONCRETO', 'NO', 15, 2), ('TABIQUE DE ARCILLA MACIZO ARTESANAL', 'ARCILLA', 'NO', 15, 2),
    ('TABIQUE DE ARCILLA EXTRUIDO HUECO', 'ARCILLA', 'SÍ', 20, 2)]
SISMO = [  # estructuración (Tabla 4.3.3 de NTC-Sismo), Q, γSV, γOI
    ('CONFINADA DE PIEZAS MACIZAS CON REFUERZO HORIZONTAL', 2.0, 0.010, 0.0015), ('CONFINADA DE PIEZAS MACIZAS', 2.0, 0.005, 0.0015),
    ('CONFINADA DE PIEZAS HUECAS CON REFUERZO HORIZONTAL', 2.0, 0.008, 0.0015), ('CONFINADA DE PIEZAS HUECAS', 1.5, 0.004, 0.0015)]
LISTAS = {
    'LST_SINO': ['SÍ', 'NO'], 'LST_GRUPO': ['A', 'B'],
    'LST_USO': ['HABITACIÓN UNIFAMILIAR', 'HABITACIÓN PLURIFAMILIAR', 'SERVICIOS', 'INDUSTRIA', 'INFRAESTRUCTURA', 'AGRÍCOLA, PECUARIO O FORESTAL', 'OTRO'],
    'LST_MORT': ['I', 'II', 'III'],
    'LST_FMSRC': ['TABLA (2.8.1.4)', 'ENSAYES DE PILAS (2.8.1.3)', 'TIPO II (TABLA 12.2.3.1.2)'],
    'LST_VMSRC': ['ENSAYES DE MURETES (2.8.2.4)', 'TIPO II (TABLA 12.2.3.2.2)', 'PROVISIONAL (CAPTURA)'],
    'LST_FSRC': ['CAPTURA EN MUROS', 'ETABS (PIER_FORCES)'],
    'LST_DIR': ['X', 'Y'], 'LST_POS': ['EXTREMO', 'INTERIOR'], 'LST_VANO': ['PUERTA', 'VENTANA'],
    'LST_RH': ['NINGUNO', 'VARILLA EN JUNTAS', 'MALLA', 'ESCALERILLA'], 'LST_CLASE': [1, 2],
}


# --------------------------------------------------------------- ayudantes
def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def cel(ref):
    """'DATOS!$C$5' -> 'C5'."""
    return ref.split('!')[1].replace('$', '')


def db(v):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (TB['DB'], v, TB['BAR'])


def ab(v):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (TB['AB'], v, TB['BAR'])


def wb_(v):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (TB['WB'], v, TB['BAR'])


def mas(v):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (TB['MAS'], v, TB['MAL'])


def mkg(v):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (TB['MKG'], v, TB['MAL'])


def mdb(v):
    return 'INDEX(%s,MATCH(%s,%s,0))' % (TB['MDB'], v, TB['MAL'])


def interp(x, X, Y):
    """Interpolación lineal en una tabla ascendente X→Y (se queda en los extremos)."""
    i = 'MATCH(%s,%s,1)' % (x, X)
    n = 'ROWS(%s)' % X
    return ('IF(%s<=INDEX(%s,1),INDEX(%s,1),IF(%s>=INDEX(%s,%s),INDEX(%s,%s),INDEX(%s,%s)+(INDEX(%s,%s+1)-INDEX(%s,%s))*(%s-INDEX(%s,%s))/(INDEX(%s,%s+1)-INDEX(%s,%s))))'
            % (x, X, Y, x, X, n, Y, n, Y, i, Y, i, Y, i, x, X, i, X, i, X, i))


def print_setup(ws, area, landscape=True):
    ws.print_area = area
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.center.text = '&A — página &P de &N'
    ws.oddFooter.center.size = 8


# =====================================================================
def sheet_tablas(wb):
    ws = wb.create_sheet('TABLAS')
    S = Narr(ws, N, {'A': 1.78, 'B': 30, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 60}, last='G')
    S.title('TABLAS — VARILLAS, MALLAS, RESISTENCIAS DE LA NTC Y LISTAS')
    S.subtitle('Todas las hojas leen de aquí con INDEX/MATCH. Los valores son los de la NTC-Mampostería 2023 y la NTC-Sismo 2023 (CDMX); no se tocan.', 18)

    def table(title, note, heads, rows, names, fmts):
        S.section(title)
        if note:
            S.secnote(note, 30)
        r = S.r
        for k, h in enumerate(heads):
            hdr_cell(ws, r, 2 + k, h)
        ws.row_dimensions[r].height = 30
        r += 1
        r0 = r
        for row in rows:
            for k, v in enumerate(row):
                cell(ws, r, 2 + k, v, 'data', fmts[k] if k < len(fmts) else 'General', al='left' if isinstance(v, str) and len(v) > 12 else 'center', wrap=True)
            ws.row_dimensions[r].height = 16 if all(not isinstance(v, str) or len(v) < 60 for v in row) else 28
            r += 1
        for k, nm in enumerate(names):
            if nm:
                TB[nm] = "TABLAS!$%s$%d:$%s$%d" % (CL(2 + k), r0, CL(2 + k), r - 1)
        S.r = r
        S.gap()

    table('A. VARILLAS Y ALAMBRES', 'Varilla corrugada NMX-C-407-ONNCCE y alambre corrugado laminado en frío. Las varillas se escogen por su número (2, 2.5, 3, 4 …); los alambres por su nombre.',
          ['Designación', 'd_b (cm)', 'Área (cm²)', 'kg/m', '', 'Uso'], [(b[0], b[1], b[2], b[3], '', b[4]) for b in BARS], ['BAR', 'DB', 'AB', 'WB'], ['General', '0.000', '0.000', '0.000', '@', '@'])
    table('B. MALLAS ELECTROSOLDADAS (refuerzo horizontal con mortero, 6.5.4)', 'A_s horizontal por metro de altura de una cara (alambres @ 15.24 cm). Malla en las dos caras = 2 «caras» en MUROS.',
          ['Designación', 'Ø alambre (mm)', 'A_s horizontal (cm²/m)', 'kg/m²'], MALLAS, ['MAL', 'MDB', 'MAS', 'MKG'], ['@', '0.00', '0.00', '0.00'])
    table("C. f'm DE PIEZAS DE CONCRETO — TABLA 2.8.1.4.1 (NTC-M p. 24)", "f'm sobre área bruta en kg/cm² según f'p de la pieza; valores intermedios se interpolan. f'p de 60 y 75 sólo para piezas huecas.",
          ["f'p (kg/cm²)", 'Mortero I', 'Mortero II'], FM_CONC, ['FPC', 'FMCI', 'FMCII'], ['0', '0', '0'])
    table("D. f'm DE PIEZAS DE ARCILLA — TABLA 2.8.1.4.2 (NTC-M p. 25)", None, ["f'p (kg/cm²)", 'Mortero I', 'Mortero II'], FM_ARC, ['FPA', 'FMAI', 'FMAII'], ['0', '0', '0'])
    table("E. PIEZAS — MATERIAL Y VALORES PARA ESTRUCTURAS TIPO II (TABLAS 12.2.3.1.2 Y 12.2.3.2.2, NTC-M p. 105)", "Los valores de Tipo II sólo valen si la estructura es Tipo II (1.6.1) y se cumplen a) a c) de 2.8.1.4.3.",
          ['Pieza', 'Material', '¿Hueca?', "f'm Tipo II (kg/cm²)", "v'm Tipo II (kg/cm²)"], PIEZAS, ['PZ', 'PZMAT', 'PZHUE', 'PZFM2', 'PZVM2'], ['@', '@', '@', '0.0', '0.0'])
    table('F. FACTOR DE COMPORTAMIENTO Y DISTORSIONES LÍMITE — TABLA 4.3.3 DE LA NTC-SISMO (p. 60)', 'Para edificios de hasta 6 niveles. «Con refuerzo horizontal» sólo si TODOS los muros estructurales lo llevan [2]. Piezas huecas confinadas: no se permiten en el Grupo A [3].',
          ['Estructuración', 'Q', 'γ_SV', 'γ_OI'], SISMO, ['SIS', 'SISQ', 'SISSV', 'SISOI'], ['@', '0.0', '0.000', '0.0000'])
    S.section('G. LISTAS DE LOS SELECTORES')
    r = S.r
    for k, (nm, vals) in enumerate(LISTAS.items()):
        c = 2 + (k % 6); rr = r + (k // 6) * 9
        hdr_cell(ws, rr, c, nm)
        for i, v in enumerate(vals):
            cell(ws, rr + 1 + i, c, v, 'data', '@' if isinstance(v, str) else 'General', al='left', sz=8)
        TB[nm] = "TABLAS!$%s$%d:$%s$%d" % (CL(c), rr + 1, CL(c), rr + len(vals))
    S.r = r + ((len(LISTAS) + 5) // 6) * 9
    for nm, ref in TB.items():
        add_name(wb, nm if nm.startswith('LST_') else 'T_' + nm, ref)
    TB['LST_MUROS'] = mr('key')
    add_name(wb, 'LST_MUROS', mr('key'))
    ws.sheet_view.zoomScale = 90
    print_setup(ws, 'B1:G%d' % S.r, landscape=False)
    return ws


# =====================================================================
def sheet_datos(wb):
    ws = wb.create_sheet('DATOS'); S = Narr(ws, N, {'A': 1.78, 'B': 44, 'C': 26, 'D': 8, 'E': 40, 'F': 34})
    S.title('DATOS GENERALES — MUROS DE MAMPOSTERÍA CONFINADA, NTC 2023 (CDMX)')
    S.subtitle('Sólo se tocan las celdas amarillas. Valen para todos los muros; lo de cada muro va en MUROS.', 16)
    S.gap()
    S.section('1. PROYECTO Y CLASIFICACIÓN (NTC-Mampostería 1.6)')
    S.header()
    S.row('Proyecto', 'Casa Rocha', '—', '', '—', 'input', '@', 'proj')
    S.row('Grupo de la construcción (Reglamento, art. 139)', 'B', '—', 'A o B', 'NTC-Criterios 3.4.1 a), p. 7', 'input', '@', 'grupo')
    dv_list(ws, cel(N['grupo']), '=LST_GRUPO')
    S.row('Superficie construida total', 380, 'm²', 'Tipo II: hasta 250 m²', 'NTC-M 1.6.1 a), p. 15', 'input', '0', 'sup')
    S.row('Niveles (incluye estacionamiento)', 2, '—', 'Tipo II: hasta 2', 'NTC-M 1.6.1 b), p. 15', 'input', '0', 'niv')
    S.row('¿Hay muros estructurales desplantados sobre vigas o sobre la losa?', 'SÍ', '—', 'Si hay, es Tipo I', 'NTC-M 1.6.1 c), p. 15', 'input', '@', 'mviga')
    dv_list(ws, cel(N['mviga']), '=LST_SINO')
    S.row('¿Hay muros diafragma?', 'NO', '—', '', 'NTC-M 1.6.1 d), p. 15', 'input', '@', 'mdiaf')
    dv_list(ws, cel(N['mdiaf']), '=LST_SINO')
    S.row('Uso', 'HABITACIÓN UNIFAMILIAR', '—', '', 'NTC-M 1.6.1 e), p. 15', 'input', '@', 'uso')
    dv_list(ws, cel(N['uso']), '=LST_USO')
    S.row('Viviendas en el predio (si es plurifamiliar)', 1, '—', 'Plurifamiliar: máximo 10', 'NTC-M 1.6.1 f), p. 15', 'input', '0', 'nviv')
    S.row('TIPO DE ESTRUCTURA', '=IF(AND({grupo}="B",{sup}<=250,{niv}<=2,{mviga}="NO",{mdiaf}="NO",{uso}<>"OTRO",OR({uso}<>"HABITACIÓN PLURIFAMILIAR",{nviv}<=10)),"TIPO II","TIPO I")',
          '—', 'Tipo II sólo si cumple todo 1.6.1', 'NTC-M 1.6.1 y 1.6.2, p. 15–16', 'key', '@', 'tipoE')
    S.row('Por qué', '=IF({tipoE}="TIPO II","Cumple a) a f) de 1.6.1",""&IF({grupo}="A","Grupo A; ","")&IF({sup}>250,"más de 250 m²; ","")&IF({niv}>2,"más de dos niveles; ","")&IF({mviga}="SÍ","muros sobre vigas o losa; ","")&IF({mdiaf}="SÍ","muros diafragma; ","")&IF({uso}="OTRO","uso no listado; ","")&IF(AND({uso}="HABITACIÓN PLURIFAMILIAR",{nviv}>10),"más de diez viviendas; ",""))',
          '—', '', 'NTC-M 1.6.1, p. 15', 'text', '@', 'tipoR')
    S.gap()
    S.section('2. MAMPOSTERÍA — PIEZA, MORTERO, f\'m Y v\'m (NTC-M 2.1, 2.8)')
    S.header()
    S.row('Pieza', 'BLOQUE DE CONCRETO HUECO', '—', '', 'TABLAS sección E', 'input', '@', 'pieza')
    dv_list(ws, cel(N['pieza']), '=' + TB['PZ'])
    S.row('Material de la pieza', '=INDEX(%s,MATCH({pieza},%s,0))' % (TB['PZMAT'], TB['PZ']), '—', '', 'TABLAS sección E', 'calc', '@', 'mat')
    S.row('¿Pieza hueca?', '=INDEX(%s,MATCH({pieza},%s,0))' % (TB['PZHUE'], TB['PZ']), '—', '', 'NTC-M Tabla 2.1.1.2, p. 17', 'calc', '@', 'hueca')
    S.row("f'p — resistencia de diseño a compresión de la pieza", 100, 'kg/cm²', 'De ensayes de las piezas', 'NTC-M 2.1.2, p. 18', 'input', '0', 'fp')
    S.row('Mortero', 'I', '—', '', 'NTC-M 2.5', 'input', '@', 'mort')
    dv_list(ws, cel(N['mort']), '=LST_MORT')
    S.row('f_an — área neta / área bruta de la pieza', 0.5, '—', 'Ficha del fabricante (hueca ≈ 0.5)', 'NTC-M 6.5.3.3.2, p. 69', 'input', '0.00', 'fan')
    S.row('h_j — espesor de la junta horizontal', 10, 'mm', '10 a 12 mm', 'NTC-M 4.2.2.2.1, p. 47', 'input', '0', 'hj')
    S.row('Altura de una hilada (pieza + junta)', 20, 'cm', 'Bloque 15×20×40 → 20 cm', 'NTC-M 6.5.3.2.1, p. 69', 'input', '0.0', 'hh')
    fmt_tab = ('IF({mort}="III",-1,IF({mat}="CONCRETO",IF({mort}="I",%s,%s),IF({mort}="I",%s,%s)))'
               % (interp('{fp}', TB['FPC'], TB['FMCI']), interp('{fp}', TB['FPC'], TB['FMCII']), interp('{fp}', TB['FPA'], TB['FMAI']), interp('{fp}', TB['FPA'], TB['FMAII'])))
    S.row("f'm de la tabla", '=' + fmt_tab, 'kg/cm²', "Tabla con f'p y el mortero", 'NTC-M 2.8.1.4, p. 24–25', 'calc', '0.0', 'fm_tab')
    S.row("f'm de ensayes de pilas: f_m medio (corregido por esbeltez)", None, 'kg/cm²', 'Si hay ensayes (9 pilas)', 'NTC-M 2.8.1.3, p. 23', 'input', '0.0', 'fm_bar')
    S.row('c_m — coeficiente de variación de las pilas', 0.15, '—', '≥ 0.15', 'NTC-M 2.8.1.3.7, p. 24', 'input', '0.00', 'cm_')
    S.row("f'm de ensayes = f_m / (1 + 2.5·c_m)", '=IF(ISNUMBER({fm_bar}),{fm_bar}/(1+2.5*MAX({cm_},0.15)),0)', 'kg/cm²', '', 'NTC-M ec. 2.8.1.3.7, p. 24', 'calc', '0.0', 'fm_ens')
    S.row("f'm Tipo II (sin ensayes)", '=INDEX(%s,MATCH({pieza},%s,0))' % (TB['PZFM2'], TB['PZ']), 'kg/cm²', 'Sólo Tipo II', 'NTC-M Tabla 12.2.3.1.2, p. 105', 'calc', '0.0', 'fm_t2')
    S.row("¿De dónde sale f'm?", 'TABLA (2.8.1.4)', '—', '', 'NTC-M 2.8.1, p. 22', 'input', '@', 'fm_src')
    dv_list(ws, cel(N['fm_src']), '=LST_FMSRC')
    S.row("f'm — RESISTENCIA DE DISEÑO A COMPRESIÓN QUE SE USA", '=IF({fm_src}="ENSAYES DE PILAS (2.8.1.3)",{fm_ens},IF({fm_src}="TIPO II (TABLA 12.2.3.1.2)",{fm_t2},MAX({fm_tab},0)))', 'kg/cm²', '', 'NTC-M 2.8.1', 'key', '0.0', 'fm')
    S.row("¿El f'm es válido?", '=IF({fm}<=0,"NO VÁLIDO: sin f\'m (mortero III o faltan ensayes)",IF({fm_src}="TIPO II (TABLA 12.2.3.1.2)",IF({tipoE}="TIPO II","VÁLIDO (Tipo II)","NO VÁLIDO: la tabla de Tipo II no aplica a estructura Tipo I"),'
          'IF({fm_src}="ENSAYES DE PILAS (2.8.1.3)","VÁLIDO (ensayes)",IF(OR({fp}<60,AND({mat}="CONCRETO",{hueca}="NO",{fp}<100)),"NO VÁLIDO: f\'p fuera de la tabla (2.8.1.4.1)",IF({fp}>IF({mat}="CONCRETO",200,300),"PENDIENTE: f\'p arriba de la tabla, se tomó el último renglón","PENDIENTE: confirmar f\'p con ensayes de piezas (2.8.1.4.3 d)")))))',
          '—', '', 'NTC-M 2.8.1.4.3, p. 25', 'text', '@', 'fm_st')
    S.row("v'm de ensayes de muretes: v_m medio", None, 'kg/cm²', 'Si hay ensayes (9 muretes)', 'NTC-M 2.8.2.4, p. 26', 'input', '0.00', 'vm_bar')
    S.row('c_v — coeficiente de variación de los muretes', 0.2, '—', '≥ 0.20', 'NTC-M 2.8.2.4.6, p. 26', 'input', '0.00', 'cv_')
    S.row("v'm de ensayes = v_m / (1 + 2.5·c_v)", '=IF(ISNUMBER({vm_bar}),{vm_bar}/(1+2.5*MAX({cv_},0.2)),0)', 'kg/cm²', '', 'NTC-M ec. 2.8.2.4.6, p. 26', 'calc', '0.00', 'vm_ens')
    S.row("v'm Tipo II (sin ensayes)", '=INDEX(%s,MATCH({pieza},%s,0))' % (TB['PZVM2'], TB['PZ']), 'kg/cm²', 'Sólo Tipo II', 'NTC-M Tabla 12.2.3.2.2, p. 105', 'calc', '0.00', 'vm_t2')
    S.row("v'm provisional (mientras llegan los ensayes)", 2, 'kg/cm²', 'Mientras llegan los ensayes', 'Criterio provisional — no viene de norma', 'input', '0.00', 'vm_prov')
    S.row("¿De dónde sale v'm?", 'PROVISIONAL (CAPTURA)', '—', '', 'NTC-M 2.8.2.1, p. 25', 'input', '@', 'vm_src')
    dv_list(ws, cel(N['vm_src']), '=LST_VMSRC')
    S.row("v'm — RESISTENCIA DE DISEÑO A COMPRESIÓN DIAGONAL QUE SE USA", '=IF({vm_src}="ENSAYES DE MURETES (2.8.2.4)",{vm_ens},IF({vm_src}="TIPO II (TABLA 12.2.3.2.2)",{vm_t2},N({vm_prov})))', 'kg/cm²', '', 'NTC-M 2.8.2', 'key', '0.00', 'vm')
    S.row("¿El v'm es válido?", '=IF({vm}<=0,"NO VÁLIDO: sin v\'m",IF({vm_src}="ENSAYES DE MURETES (2.8.2.4)","VÁLIDO (ensayes)",IF({vm_src}="TIPO II (TABLA 12.2.3.2.2)",IF({tipoE}="TIPO II","VÁLIDO (Tipo II)","NO VÁLIDO: estructura Tipo I, el v\'m sale de 9 muretes (2.8.2.1)"),"PENDIENTE: v\'m provisional; estructura "&{tipoE}&", hacen falta los ensayes de muretes (2.8.2.4)")))',
          '—', 'Tipo I: sólo vale de ensayes', 'NTC-M 2.8.2.1 y 2.8.2.4.5, p. 25–26', 'text', '@', 'vm_st')
    S.row('E_m — módulo de elasticidad (corta duración)', '=IF({mat}="CONCRETO",800,600)*{fm}', 'kg/cm²', "800·f'm (concreto) · 600·f'm (arcilla)", 'NTC-M 2.8.5.4, p. 27', 'calc', '0', 'Em')
    S.gap()
    S.section('3. CONCRETO Y ACERO')
    S.header()
    S.row("f'c — concreto de castillos y dalas", 200, 'kg/cm²', '≥ 200', 'NTC-M 6.2.1.3.1, p. 61', 'input', '0', 'fc')
    S.row('f_y — varillas longitudinales', 4200, 'kg/cm²', '', 'NMX-C-407-ONNCCE', 'input', '0', 'fy')
    S.row('f_y — estribos de castillos y dalas', 2530, 'kg/cm²', 'Alambrón No. 2: 2 530', 'NTC-M 6.2.2.3, p. 62', 'input', '0', 'fye')
    S.row('f_yh — refuerzo horizontal (juntas o malla)', 6000, 'kg/cm²', 'La hoja usa 6 000 como tope en juntas y 5 000 en malla', 'NTC-M 6.5.3.1.3 y 6.5.4.1.3, p. 69 y 71', 'input', '0', 'fyh')
    S.row("f'c — vigas sobre las que desplantan muros", 250, 'kg/cm²', '', 'NTC-M 4.1.18.2.2, p. 44', 'input', '0', 'fcv')
    S.row('Clase del concreto de esas vigas', 1, '—', '', 'NTC-Concreto 2.1.3', 'input', '0', 'clv')
    dv_list(ws, cel(N['clv']), '=LST_CLASE')
    S.row('E_f — módulo de elasticidad de la viga', '=IF({clv}=1,14000,8000)*SQRT({fcv})', 'kg/cm²', "14 000·√f'c (clase 1) · 8 000·√f'c (clase 2)", 'NTC-Concreto 2.1.3', 'calc', '0', 'Ecv')
    S.gap()
    S.section('4. MURO, CASTILLO Y DALA TÍPICOS (los usan todos los muros; MUROS deja cambiar H, t y los castillos extremos)')
    S.header()
    S.row('H — altura libre típica', 2.75, 'm', 'Altura libre entre losas', 'NTC-M 6.1.3, p. 60', 'input', '0.00', 'Hdef')
    S.row('t — espesor típico de la mampostería', 15, 'cm', '≥ 12 cm en Tipo I', 'NTC-M 6.1.3, p. 60', 'input', '0.0', 'tdef')
    S.row('h_c — castillo, lado paralelo al muro', 15, 'cm', '≥ 15 cm', 'NTC-M 6.2.1.1 b), p. 60', 'input', '0.0', 'hc')
    S.row('b_c — castillo, lado perpendicular al muro', 15, 'cm', '≥ 12 cm y ≥ t', 'NTC-M 6.2.1.1 b), p. 60', 'input', '0.0', 'bc')
    S.row('Varillas longitudinales por castillo', 4, 'pzas', '≥ 4', 'NTC-M 6.2.2.2, p. 62', 'input', '0', 'nc')
    S.row('No. de varilla del castillo', 3, 'No.', '', 'TABLAS', 'input', 'General', 'vc')
    dv_list(ws, cel(N['vc']), '=' + TB['BAR'])
    S.row('No. del estribo de castillos y dalas', 2, 'No.', '', 'NTC-M 6.2.2.3, p. 62', 'input', 'General', 've')
    dv_list(ws, cel(N['ve']), '=' + TB['BAR'])
    S.row('s — separación de estribos en el castillo', 20, 'cm', '≤ 1.5·t y ≤ 20 cm', 'NTC-M 6.2.2.3, p. 62', 'input', '0.0', 'se')
    S.row('s₀ — separación de estribos en los extremos (H₀)', 7.5, 'cm', '≤ h_c/2, en H₀ = máx(H/6, 40 cm)', 'NTC-M 6.2.3.2, p. 62', 'input', '0.0', 'se0')
    S.row('Recubrimiento libre de las varillas de castillos y dalas', 2.0, 'cm', '≥ 2 cm', 'NTC-M 4.2.4.1.1, p. 49', 'input', '0.0', 'rc')
    S.row('h_d — peralte de la dala', 20, 'cm', '≥ 15 cm', 'NTC-M 6.2.1.2 b), p. 60', 'input', '0.0', 'hd')
    S.row('b_d — ancho de la dala', 15, 'cm', '≥ 12 cm', 'NTC-M 6.2.1.2 b), p. 60', 'input', '0.0', 'bd')
    S.row('Varillas longitudinales de la dala', 4, 'pzas', '≥ 4', 'NTC-M 6.2.2.2, p. 62', 'input', '0', 'nd')
    S.row('No. de varilla de la dala', 3, 'No.', '', 'TABLAS', 'input', 'General', 'vd')
    dv_list(ws, cel(N['vd']), '=' + TB['BAR'])
    S.row('s — separación de estribos en la dala', 20, 'cm', '≤ 1.5·t y ≤ 20 cm', 'NTC-M 6.2.2.3, p. 62', 'input', '0.0', 'sed')
    S.row('L mín. — tramo más corto que se cuenta como muro', '=ROUND(2*{hc}/100,2)', 'm', 'Tramo más corto que cuenta como muro (dos castillos)',
          'Criterio de esta hoja — no viene de norma', 'input', '0.00', 'Lmin')
    S.row('Varilla o alambre del refuerzo horizontal (si MUROS no dice otro)', '5/32 in', '—', '', 'NTC-M 4.2.2.2.1, p. 47', 'input', '@', 'rhvd')
    dv_list(ws, cel(N['rhvd']), '=' + TB['BAR'])
    S.row('Malla (si MUROS dice MALLA y no dice cuál)', '6x6-10/10', '—', '', 'TABLAS', 'input', '@', 'mallad')
    dv_list(ws, cel(N['mallad']), '=' + TB['MAL'])
    S.gap()
    S.section('5. REVISIÓN DEL CASTILLO Y LA DALA TÍPICOS (6.2)')
    S.header()
    S.row('A_s del castillo', '={nc}*' + ab('{vc}'), 'cm²', '', 'TABLAS', 'calc', '0.00', 'Asc')
    S.row("A_s mínima = 0.2·f'c/f_y·b_c·h_c", '=0.2*{fc}/{fy}*{bc}*{hc}', 'cm²', "0.2·f'c/f_y·b_c·h_c", 'NTC-M 6.2.2.2, p. 62', 'calc', '0.00', 'Ascmin')
    S.row('A_sc de estribo (2 ramas) ≥ 1 000·s/(f_y·h_c)', '=2*' + ab('{ve}'), 'cm²', ('="Mínimo: "&TEXT(1000*{se}/({fye}*{hc}),"0.00")&" cm² (castillo) · "&TEXT(1000*{sed}/({fye}*{hd}),"0.00")&" cm² (dala)"').format(**N), 'NTC-M ec. 6.2.3.3, p. 62', 'calc', '0.00', 'Ase2')
    chk = ('IF(AND({hc}>=15,{bc}>=12,{bc}>={tdef},{nc}>=4,{Asc}>=0.99*{Ascmin},%s<={hc}/6*MIN(1,{bc}/{hc}),{Ase2}>=1000*{se}/({fye}*{hc})-0.0001,{se}<=MIN(1.5*{tdef},20),{se0}<={hc}/2,{fc}>=200,{rc}>=2,{rc}-%s>=1),"CUMPLE",'
           '"NO CUMPLE: "&IF(OR({hc}<15,{bc}<12,{bc}<{tdef}),"dimensiones (6.2.1.1 b) ","")&IF(OR({nc}<4,{Asc}<0.99*{Ascmin}),"A_s (6.2.2.2) ","")&IF(%s>MIN({bc},{hc})/6,"varilla > 1/6 del lado (4.2.2.1.2) ","")'
           '&IF({Ase2}<1000*{se}/({fye}*{hc})-0.0001,"estribo (6.2.2.3) ","")&IF({se}>MIN(1.5*{tdef},20),"s > mín(1.5t, 20) ","")&IF({se0}>{hc}/2,"s₀ > h_c/2 (6.2.3.2) ","")&IF({fc}<200,"f\'c < 200 ","")&IF(OR({rc}<2,{rc}-%s<1),"recubrimiento (4.2.4.1.1)",""))') % (db('{vc}'), db('{ve}'), db('{vc}'), db('{ve}'))
    S.row('CASTILLO TÍPICO', '=' + chk, '—', '', 'NTC-M 6.2.1 a 6.2.3 y 4.2, p. 47–62', 'key', '@', 'kchk')
    chkd = ('IF(AND({hd}>=15,{bd}>=12,{bd}>={tdef},{nd}>=4,{nd}*%s>=0.99*0.2*{fc}/{fy}*{bd}*{hd},%s<=MIN({bd},{hd})/6,{Ase2}>=1000*{sed}/({fye}*{hd})-0.0001,{sed}<=MIN(1.5*{tdef},20)),"CUMPLE",'
            '"NO CUMPLE: "&IF(OR({hd}<15,{bd}<12,{bd}<{tdef}),"dimensiones (6.2.1.2 b) ","")&IF(OR({nd}<4,{nd}*%s<0.99*0.2*{fc}/{fy}*{bd}*{hd}),"A_s (6.2.2.2) ","")&IF(%s>MIN({bd},{hd})/6,"varilla > 1/6 del lado ","")'
            '&IF({Ase2}<1000*{sed}/({fye}*{hd})-0.0001,"estribo (6.2.2.3) ","")&IF({sed}>MIN(1.5*{tdef},20),"s > mín(1.5t, 20)",""))') % (ab('{vd}'), db('{vd}'), ab('{vd}'), db('{vd}'))
    S.row('DALA TÍPICA', '=' + chkd, '—', '', 'NTC-M 6.2.1.2 y 6.2.2, p. 60–62', 'key', '@', 'dchk')
    S.gap()
    S.section('6. FACTORES DE RESISTENCIA Y DE CARGA')
    S.header()
    S.row('F_R — cortante', 0.7, '—', '', 'NTC-M 4.1.6.4, p. 40', 'calc', '0.00', 'FRv')
    S.row('F_R — compresión (P_R) y f_CM de 4.1.18', 0.6, '—', '', 'NTC-M 4.1.6.2, p. 39', 'calc', '0.00', 'FRc')
    S.row('F_R — flexocompresión (método optativo)', '0.8 / 0.6', '—', '0.8 si P_u ≤ P_R/3; si no, 0.6', 'NTC-M 4.1.6.3.1 b), p. 39', 'calc', '@', 'FRf')
    S.row('FC — combinaciones con sismo', 1.1, '—', 'Para sacar P sin factor', 'NTC-Criterios 3.4.1 b), p. 7', 'input', '0.00', 'FCs')
    S.row('FC — cargas permanentes en gravedad', '=IF({grupo}="A",1.5,1.3)', '—', '', 'NTC-Criterios 3.4.1 a), p. 7', 'calc', '0.00', 'FCcm')
    S.row('α — del refuerzo horizontal', 0.045, 'cm²/kg', '', 'NTC-M 6.5.3.4.1, p. 70', 'calc', '0.000', 'alfa')
    S.gap()
    S.section('7. SISMO — FACTOR DE COMPORTAMIENTO Y DISTORSIONES (NTC-Sismo 1.7 y Tabla 4.3.3)')
    S.header()
    S.row('¿Todos los muros llevan refuerzo horizontal?', '=IF(COUNTIF(%s,"SÍ")=0,"NO",IF(COUNTIFS(%s,"SÍ",%s,"VARILLA EN JUNTAS")+COUNTIFS(%s,"SÍ",%s,"MALLA")=COUNTIF(%s,"SÍ"),"SÍ","NO"))'
          % (mr('ok'), mr('ok'), mr('rht'), mr('ok'), mr('rht'), mr('ok')), '—', '', 'NTC-Sismo Tabla 4.3.3, nota [2], p. 60', 'calc', '@', 'rhall')
    S.row('Estructuración (fila de la Tabla 4.3.3)', '="CONFINADA DE PIEZAS "&IF({hueca}="SÍ","HUECAS","MACIZAS")&IF({rhall}="SÍ"," CON REFUERZO HORIZONTAL","")', '—', '', 'NTC-Sismo Tabla 4.3.3, p. 60', 'calc', '@', 'sis')
    S.row('Q — factor de comportamiento sísmico', '=INDEX(%s,MATCH({sis},%s,0))' % (TB['SISQ'], TB['SIS']), '—', '', 'NTC-Sismo Tabla 4.3.3, p. 60', 'key', '0.0', 'Q')
    S.row('γ_SV — distorsión límite de seguridad de vida', '=INDEX(%s,MATCH({sis},%s,0))' % (TB['SISSV'], TB['SIS']), '—', '', 'NTC-Sismo Tabla 4.3.3, p. 60', 'key', '0.000', 'gSV')
    S.row('γ_OI — distorsión límite de ocupación inmediata', '=INDEX(%s,MATCH({sis},%s,0))' % (TB['SISOI'], TB['SIS']), '—', 'Sólo Grupo A', 'NTC-Sismo Tabla 4.3.3 y 1.7, p. 29 y 60', 'calc', '0.0000', 'gOI')
    S.row("R' — sobrerresistencia del periodo fundamental", 1.68, '—', "Q·R' = 2.52 en Casa Rocha", 'NTC-Sismo 3.3', 'input', '0.00', 'Rp')
    S.row("Factor de la distorsión de seguridad de vida = Q·R'·FC", '={Q}*{Rp}*{FCs}', '—', "Q·R'·FC", 'NTC-Sismo 1.7, p. 29', 'calc', '0.000', 'kSV')
    S.row('¿Los elementos no estructurales están separados de la estructura?', 'NO', '—', 'NO → 0.002 · SÍ → 0.004', 'NTC-Sismo 1.7, p. 29', 'input', '@', 'sep')
    dv_list(ws, cel(N['sep']), '=LST_SINO')
    S.row('γ_LD — distorsión límite de daño no estructural', '=IF({sep}="SÍ",0.004,0.002)', '—', '', 'NTC-Sismo 1.7, p. 29', 'key', '0.000', 'gLD')
    S.row('¿Se permite la estructuración en este grupo?', '=IF(AND({grupo}="A",{hueca}="SÍ"),"NO CUMPLE: piezas huecas confinadas no se permiten en el Grupo A (nota [3])","CUMPLE")', '—', '', 'NTC-Sismo Tabla 4.3.3, nota [3], p. 60', 'text', '@', 'sisok')
    S.gap()
    S.section('8. FUERZAS — DE DÓNDE SALEN')
    S.header()
    S.row('Fuerzas de los muros', 'CAPTURA EN MUROS', '—', 'ETABS: pega la tabla Pier Forces en PIER_FORCES',
          'Hoja PIER_FORCES', 'input', '@', 'fsrc')
    dv_list(ws, cel(N['fsrc']), '=LST_FSRC')
    S.secnote('Combinaciones que se buscan en PIER_FORCES.', 16)
    r = S.r
    hdr_cell(ws, r, 2, 'Combinaciones con SISMO'); hdr_cell(ws, r, 3, 'Combinaciones de GRAVEDAD')
    ws.merge_cells('C%d:D%d' % (r, r))
    sis = ['Comb4.%d' % i for i in range(1, 9)] + [None] * 4
    grav = ['Comb2.2'] + [None] * 11
    for i in range(12):
        cell(ws, r + 1 + i, 2, sis[i], 'input', '@', al='left')
        cell(ws, r + 1 + i, 3, grav[i], 'input', '@', al='left')
        ws.merge_cells('C%d:D%d' % (r + 1 + i, r + 1 + i))
    N['CMB_S'] = "DATOS!$B$%d:$B$%d" % (r + 1, r + 12); N['CMB_G'] = "DATOS!$C$%d:$C$%d" % (r + 1, r + 12)
    add_name(wb, 'CMB_S', N['CMB_S']); add_name(wb, 'CMB_G', N['CMB_G'])
    S.r = r + 13
    S.gap()
    S.section('9. DISTORSIONES — CÓMO RECONOCER CADA CASO EN LA TABLA DE ETABS')
    S.header()
    S.row('Texto de los casos de seguridad de vida', 'SV', '—', '', 'Hoja DISTORSIONES', 'input', '@', 'tSV')
    S.row('Texto de los casos de limitación de daño', 'LD', '—', '', 'Hoja DISTORSIONES', 'input', '@', 'tLD')
    S.row('Texto de los casos de ocupación inmediata (Grupo A)', 'OI', '—', '', 'Hoja DISTORSIONES', 'input', '@', 'tOI')
    add_cf(ws, 'C1:C%d' % S.r)
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = 'A4'
    print_setup(ws, 'B1:F%d' % S.r, landscape=False)
    return ws


# =====================================================================
def sheet_instrucciones(wb):
    ws = wb.active; ws.title = 'INSTRUCCIONES'
    S = Narr(ws, N)
    S.title('REVISIÓN DE MUROS DE MAMPOSTERÍA CONFINADA — NTC 2023 (CDMX)')
    S.gap()
    S.section('CÓMO SE USA')
    for t_ in ['1. DATOS: pieza, mortero, f\'m, v\'m, concreto, castillo y dala típicos.',
               '2. MUROS: un renglón por muro (longitud, vanos de izquierda a derecha, viga de apoyo si la hay y fuerzas).',
               '3. Fuerzas: captúralas en MUROS, o pega la tabla «Pier Forces» de ETABS en PIER_FORCES y en DATOS cambia «Fuerzas de los muros» a ETABS.',
               '4. DETALLE: escoge un muro y sale su revisión paso a paso. CROQUIS lo dibuja.',
               '5. RESUMEN: qué cumple, qué no y qué le falta a cada muro. ENTREPISO y DISTORSIONES: revisión global.']:
        S.bullet(t_)
    S.gap()
    S.section('COLORES')
    S.bullet('Amarillo = se captura · Verde = CUMPLE · Ámbar = PENDIENTE (por ejemplo v\'m provisional) · Rojo = NO CUMPLE.')
    S.bullet('En MUROS las columnas de cálculo intermedio están agrupadas: ábrelas con el botón [+] de arriba si quieres ver las cuentas.')
    S.gap()
    S.section('NO REVISA')
    S.bullet('Flexión fuera del plano, muros diafragma, mampostería reforzada interiormente y el diseño de la viga de apoyo (va en la hoja de losas).')
    S.gap()
    S.section('EJEMPLO: CASA ROCHA')
    S.bullet('37 muros del modelo R3.2 con los vanos del plano arquitectónico y las fuerzas de tu hoja de revisión. Alturas de vanos típicas: confírmalas.')
    ws.sheet_view.zoomScale = 100
    print_setup(ws, 'B1:F%d' % S.r, landscape=False)
    return ws


# =====================================================================
REFS = [
    ('NTC-Mampostería 2023', [
        ('1.6.1 y 1.6.2', 'Clasificación Tipo I / Tipo II', 15), ('2.1.2', "Resistencia de las piezas f'p (ensayes)", 18), ('2.8.1.3', "f'm de ensayes de pilas, ec. 2.8.1.3.7", 23),
        ('2.8.1.4 (tablas 2.8.1.4.1 y 2.8.1.4.2)', "f'm a partir de f'p y el mortero", 24), ('2.8.1.4.3', 'Condiciones para usar las tablas de f\'m', 25), ('2.8.2.1 y 2.8.2.4', "v'm: 9 muretes, ec. 2.8.2.4.6", 25),
        ('2.8.5.4', 'Módulo de elasticidad E_m', 27), ('3.1.2.3.1', 'Factor F_E por excentricidad y esbeltez', 32), ('4.1.6', 'Factores de resistencia', 39),
        ('4.1.10 a 4.1.12', 'Cortante resistente de entrepiso, ec. 4.1.12.1', 40), ('4.1.18', 'Muros sobre vigas: f_CM, flecha y peralte mínimo', 44), ('4.2.2.1.2', 'Diámetro de varillas en castillos y dalas', 47),
        ('4.2.2.2.1', 'Diámetro del refuerzo horizontal', 47), ('4.2.4.1.1', 'Recubrimientos en castillos y dalas', 49), ('6.1.3', 'Espesor mínimo y H/t ≤ 25', 60),
        ('6.2.1.1 y 6.2.1.2', 'Castillos y dalas: dónde van y dimensiones', 60), ('6.2.1.3.1', "f'c de castillos y dalas", 61), ('6.2.2.2 y 6.2.2.3', 'A_s mínima y estribos', 62),
        ('6.2.3.2', 'Estribos a h_c/2 en H₀', 62), ('6.2.5', 'Muros con aberturas', 64), ('6.4.1.1', 'Compresión P_R', 65), ('6.4.2.2', 'Flexocompresión, método optativo', 66),
        ('6.5.1.3', 'P en tensión: el refuerzo toma todo el cortante', 67), ('6.5.2.1 a 6.5.2.4', 'V_mR y factor f', 67), ('6.5.3', 'Refuerzo horizontal en juntas', 69),
        ('6.5.3.4.1', 'V_sR y eficiencia η', 70), ('6.5.4', 'Malla de alambre soldado', 71), ('12.2.3', "f'm y v'm de estructuras Tipo II", 105)]),
    ('NTC-Sismo 2023', [('1.7', 'Revisión de desplazamientos laterales (seguridad de vida, ocupación inmediata, daño)', 29), ('4.3.3 (Tabla 4.3.3)', 'Q, γ_SV y γ_OI de mampostería', 60)]),
    ('NTC-Criterios y Acciones 2023', [('3.4.1', 'Factores de carga', 7)]),
]


def sheet_referencias(wb):
    ws = wb.create_sheet('REFERENCIAS')
    S = Narr(ws, N, {'A': 1.78, 'B': 30, 'C': 34, 'D': 70, 'E': 12}, last='E')
    S.title('REFERENCIAS — QUÉ NORMA, QUÉ INCISO Y QUÉ PÁGINA CHECAR')
    S.subtitle('Páginas impresas de las Normas Técnicas Complementarias 2023 de la CDMX (Gaceta Oficial). Se cita el inciso, que no cambia entre tirajes.', 18)
    for norma, items in REFS:
        S.section(norma)
        r = S.r
        for k, h in enumerate(('Norma', 'Inciso', 'Qué se usa', 'Página impresa')):
            hdr_cell(ws, r, 2 + k, h)
        r += 1
        for inc, que, pg in items:
            cell(ws, r, 2, norma, 'ref', '@', al='left'); cell(ws, r, 3, inc, 'calc', '@', al='left'); cell(ws, r, 4, que, 'calc', '@', al='left', wrap=True); cell(ws, r, 5, pg, 'calc', '0')
            r += 1
        S.r = r
        S.gap()
    print_setup(ws, 'B1:E%d' % S.r, landscape=False)
    return ws
