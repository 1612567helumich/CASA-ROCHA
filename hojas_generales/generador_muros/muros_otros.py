"""ENTREPISO (4.1.10 / 4.1.12), DISTORSIONES (NTC-Sismo 1.7) y RESUMEN."""
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter as CL
from muros_comun import *        # noqa: F401,F403

NST = 4          # entrepisos


def sheet_entrepiso(wb, stories):
    ws = wb.create_sheet('ENTREPISO')
    S = Narr(ws, N, {'A': 1.78, 'B': 14, 'C': 10, 'D': 12, 'E': 12, 'F': 12, 'G': 12, 'H': 12, 'I': 12, 'J': 12, 'K': 12, 'L': 30}, last='L')
    S.title('ENTREPISO — CORTANTE DE CADA ENTREPISO (4.1.10)')
    S.subtitle('V_R = F_R·(0.5·v\'m + 0.3·σ)·ΣA_T + ΣV_sR ≥ 0.8·V_u en cada dirección (ec. 4.1.12.1). ΣA_T sale de MUROS.', 16)
    S.gap()
    S.section('1. DATOS DE CADA ENTREPISO (de ETABS, V_u con factor de carga)')
    r = S.r
    heads = ['Story de ETABS', 'Nombre', 'W — peso arriba (t)', 'V_u X (t)', 'V_u Y (t)']
    for k, h in enumerate(heads):
        hdr_cell(ws, r, 2 + k, h)
    ws.row_dimensions[r].height = 30
    rows0 = r + 1
    for i in range(NST):
        st = stories[i] if i < len(stories) else {}
        for k, key in enumerate(('st', 'nom', 'W', 'Vx', 'Vy')):
            cell(ws, rows0 + i, 2 + k, st.get(key), 'input', '@' if key in ('st', 'nom') else '0.0')
    S.r = rows0 + NST
    S.gap()
    S.row('¿Despreciar los muros de menos de 1 m?', 'NO', '—', 'Lo recomienda el comentario de 4.1.12', 'NTC-M 4.1.12, comentario, p. 41', 'input', '@', 'e_m1')
    dv_list(ws, cel(N['e_m1']), '=LST_SINO')
    S.gap()
    S.section('2. REVISIÓN')
    r = S.r
    heads = ['Story', 'Dirección', 'ΣA_T dir. (m²)', 'ΣA_T total (m²)', 'σ (kg/cm²)', 'σ usado ≤ 3.33·v\'m', 'ΣV_sR (t)', 'V_R (t)', 'V_u (t)', '0.8·V_u / V_R', 'ESTADO']
    for k, h in enumerate(heads):
        hdr_cell(ws, r, 2 + k, h)
    ws.row_dimensions[r].height = 30
    lim = 'IF({e_m1}="SÍ",">=1",">=0")'.format(**N)
    for i in range(NST):
        for d_, dd in enumerate(('X', 'Y')):
            rr = r + 1 + 2 * i + d_
            stc = '$B$%d' % (rows0 + i)
            cell(ws, rr, 2, '=IF(%s="","",%s)' % (stc, stc), 'calc', '@')
            cell(ws, rr, 3, dd if True else '', 'calc', '@')
            at = 'SUMIFS(%s,%s,%s,%s,"%s",%s,"SÍ",%s,%s)/10000' % (mr('AT'), mr('st'), stc, mr('dir'), dd, mr('ok'), mr('L'), lim)
            att = 'SUMIFS(%s,%s,%s,%s,"SÍ",%s,%s)/10000' % (mr('AT'), mr('st'), stc, mr('ok'), mr('L'), lim)
            cell(ws, rr, 4, '=IF(%s="","",%s)' % (stc, at), 'calc', '0.000')
            cell(ws, rr, 5, '=IF(%s="","",%s)' % (stc, att), 'calc', '0.000')
            W = '$D$%d' % (rows0 + i)
            cell(ws, rr, 6, '=IF(OR(%s="",N(E%d)=0),"",N(%s)/E%d/10)' % (stc, rr, W, rr), 'calc', '0.00')
            cell(ws, rr, 7, ('=IF(F%d="","",MIN(F%d,3.33*{vm}))' % (rr, rr)).format(**N), 'calc', '0.00')
            cell(ws, rr, 8, '=IF(%s="","",SUMIFS(%s,%s,%s,%s,"%s",%s,"SÍ",%s,%s))' % (stc, mr('VsR'), mr('st'), stc, mr('dir'), dd, mr('ok'), mr('L'), lim), 'calc', '0.00')
            cell(ws, rr, 9, ('=IF(G%d="","",{FRv}*(0.5*{vm}+0.3*G%d)*D%d*10000/1000+H%d)' % (rr, rr, rr, rr)).format(**N), 'key', '0.0')
            cell(ws, rr, 10, '=IF(%s="","",N($%s$%d))' % (stc, 'E' if dd == 'X' else 'F', rows0 + i), 'calc', '0.0')
            cell(ws, rr, 11, '=IF(OR(I%d="",N(I%d)=0),"",0.8*J%d/I%d)' % (rr, rr, rr, rr), 'key', '0.00')
            est = ('=IF(%s="","",IF(N(E%d)=0,"PENDIENTE: ningún muro de MUROS es de este story",IF(K%d="",IF(N(J%d)>0,"NO CUMPLE: no hay muros en esta dirección",""),'
                   'IF(K%d<=1,IF(LEFT({vm_st},9)="PENDIENTE","CUMPLE — PENDIENTE: v\'m provisional","CUMPLE"),"NO CUMPLE: V_R < 0.8·V_u (4.1.10.1)"))))' % (stc, rr, rr, rr, rr)).format(**N)
            cell(ws, rr, 12, est, 'calc', '@', al='left')
    N['e_r0'] = r + 1; N['e_r1'] = r + 2 * NST
    add_cf(ws, 'L%d:L%d' % (r + 1, r + 2 * NST))
    S.r = r + 2 * NST + 1
    add_cf(ws, 'C1:C%d' % S.r)
    print_setup(ws, 'B1:L%d' % S.r, landscape=True)
    return ws


def sheet_distorsiones(wb, drifts):
    ws = wb.create_sheet('DISTORSIONES')
    S = Narr(ws, N, {'A': 1.78, 'B': 10, 'C': 30, 'D': 12, 'E': 8, 'F': 9, 'G': 11, 'H': 8, 'I': 8, 'J': 8, 'K': 8, 'L': 2, 'M': 7, 'N': 9, 'O': 11, 'P': 9, 'Q': 8, 'R': 18}, last='R')
    S.title('DISTORSIONES DE ENTREPISO — NTC-SISMO 1.7 Y TABLA 4.3.3')
    S.subtitle('Pega la tabla «Story Drifts» de ETABS desde el renglón %d. Seguridad de vida: × Q·R\'·FC ≤ γ_SV. Daño: ≤ γ_LD.' % DR0, 16)
    r = S.r + 1
    for k, h in enumerate(('Revisión', 'Dirección', 'Distorsión revisada máx.', 'Límite', 'Relación', 'ESTADO')):
        hdr_cell(ws, r, 2 + k * 1 + (1 if k > 0 else 0) * 0, h)
    ws.row_dimensions[r].height = 30
    i = 0
    for cls, lim in (('SV', '{gSV}'), ('LD', '{gLD}'), ('OI', '{gOI}')):
        for dd in ('X', 'Y'):
            rr = r + 1 + i; i += 1
            cell(ws, rr, 2, cls, 'calc', '@'); cell(ws, rr, 3, dd, 'calc', '@')
            mx = 'SUMPRODUCT(MAX(($M$%d:$M$%d="%s")*($F$%d:$F$%d="%s")*$O$%d:$O$%d))' % (DR0, DR1, cls, DR0, DR1, dd, DR0, DR1)
            cell(ws, rr, 4, '=' + mx, 'key', '0.00000')
            cell(ws, rr, 5, ('=' + lim).format(**N), 'calc', '0.0000')
            cell(ws, rr, 6, '=IF(E%d>0,D%d/E%d,0)' % (rr, rr, rr), 'key', '0.000')
            cnt = 'COUNTIFS($M$%d:$M$%d,"%s",$F$%d:$F$%d,"%s")' % (DR0, DR1, cls, DR0, DR1, dd)
            cell(ws, rr, 7, ('=IF(%s=0,IF(AND("%s"="OI",{grupo}<>"A"),"— (sólo Grupo A)","PENDIENTE: no hay renglones de este caso"),IF(AND("%s"<>"LD",N({Rp})<=0),"PENDIENTE: falta R\' en DATOS",IF(F%d<=1,"CUMPLE","NO CUMPLE")))' % (cnt, cls, cls, rr)).format(**N), 'calc', '@', al='left')
    add_cf(ws, 'G%d:G%d' % (r + 1, r + 6))
    N['dist_r0'] = r + 1
    rh = DR0 - 2
    ws.merge_cells('B%d:K%d' % (rh - 1, rh - 1)); ws['B%d' % (rh - 1)] = 'TABLA «STORY DRIFTS» DE ETABS — CAPTURA (pegar valores)'; ws['B%d' % (rh - 1)].font = font(9, True, color='FFFFFF'); ws['B%d' % (rh - 1)].fill = fill('7F6000')
    ws.merge_cells('M%d:R%d' % (rh - 1, rh - 1)); ws['M%d' % (rh - 1)] = 'REVISIÓN DE CADA RENGLÓN'; ws['M%d' % (rh - 1)].font = font(9, True, color='FFFFFF'); ws['M%d' % (rh - 1)].fill = fill(SEC_BG)
    for k, h in enumerate(('Story', 'Output Case', 'Case Type', 'Step Type', 'Direction', 'Drift', 'Label', 'X', 'Y', 'Z')):
        hdr_cell(ws, rh, 2 + k, h)
    for k, h in enumerate(('Caso', 'Factor', 'Distorsión revisada', 'Límite', 'Relación', 'Estado')):
        hdr_cell(ws, rh, 13 + k, h)
    ws.row_dimensions[rh].height = 30
    for rr in range(DR0, DR1 + 1):
        ex = drifts[rr - DR0] if rr - DR0 < len(drifts) else None
        for k in range(10):
            v = ex[k] if ex and k < len(ex) else None
            x = cell(ws, rr, 2 + k, v, 'input', '0.000000' if k == 5 else ('0.00' if k in (7, 8, 9) else '@' if k in (0, 1, 2, 3, 4) else 'General'), sz=8)
        cls = ('=IF($C{r}="","",IF(ISNUMBER(SEARCH({tSV},$C{r})),"SV",IF(ISNUMBER(SEARCH({tLD},$C{r})),"LD",IF(ISNUMBER(SEARCH({tOI},$C{r})),"OI",""))))').replace('{r}', str(rr)).format(**N)
        ws.cell(rr, 13, cls).font = font(8)
        ws.cell(rr, 14, ('=IF($M{r}="SV",{kSV},IF($M{r}="LD",1,IF($M{r}="OI",{Rp}*{FCs},0)))').replace('{r}', str(rr)).format(**N)).font = font(8)
        ws.cell(rr, 15, '=IF($M%d="",0,N($G%d)*$N%d)' % (rr, rr, rr)).font = font(8)
        ws.cell(rr, 16, ('=IF($M{r}="SV",{gSV},IF($M{r}="LD",{gLD},IF($M{r}="OI",{gOI},"")))').replace('{r}', str(rr)).format(**N)).font = font(8)
        ws.cell(rr, 17, '=IF(OR($M%d="",N($P%d)=0),"",$O%d/$P%d)' % (rr, rr, rr, rr)).font = font(8)
        ws.cell(rr, 18, ('=IF($M{r}="","",IF(AND($M{r}<>"LD",N({Rp})<=0),"PENDIENTE: falta R\'",IF($Q{r}<=1,"CUMPLE","NO CUMPLE")))').replace('{r}', str(rr)).format(**N)).font = font(8)
        for c in (15, 17):
            ws.cell(rr, c).number_format = '0.00000;-0.00000;' if c == 15 else '0.000'
    add_cf(ws, 'R%d:R%d' % (DR0, DR1))
    ws.freeze_panes = 'B%d' % DR0
    print_setup(ws, 'B1:R%d' % (DR0 + 40), landscape=True)
    return ws


def sheet_resumen(wb):
    ws = wb.create_sheet('RESUMEN')
    S = Narr(ws, N, {'A': 1.78, 'B': 14, 'C': 9, 'D': 5, 'E': 7, 'F': 8, 'G': 8, 'H': 8, 'I': 8, 'J': 8, 'K': 22, 'L': 40, 'M': 36}, last='M')
    S.title('RESUMEN — QUÉ MUROS CUMPLEN, CUÁLES NO Y POR QUÉ')
    S.subtitle('Sale solo de MUROS, ENTREPISO y DISTORSIONES.', 16)
    S.section('1. DATOS QUE MANDAN')
    S.header(('Concepto', 'Valor', '', '', '', ''))
    for rr in range(S.r - 1, S.r):
        ws.merge_cells('C%d:F%d' % (rr, rr))
    for txt, f in (('Proyecto', '={proj}&" · "&{tipoE}&" · Grupo "&{grupo}'), ("f'm", '=TEXT({fm},"0.0")&" kg/cm² — "&{fm_st}'), ("v'm", '=TEXT({vm},"0.00")&" kg/cm² — "&{vm_st}'),
                   ('Fuerzas', '={fsrc}'), ('Sistema (Tabla 4.3.3)', '={sis}&": Q = "&{Q}&", γ_SV = "&{gSV}')):
        r = S.r
        cell(ws, r, 2, txt, 'calc', '@', al='left', bold=True)
        ws.merge_cells('C%d:M%d' % (r, r))
        cell(ws, r, 3, f.format(**N), 'text', '@', al='left')
        add_cf(ws, 'C%d:C%d' % (r, r))
        S.r += 1
    S.gap()
    S.section('2. CONTEO')
    est = mr('est'); ok = mr('ok')
    items = [('Muros capturados', '=COUNTIF(%s,"SÍ")' % ok), ('CUMPLE', '=COUNTIF(%s,"CUMPLE")' % est), ('CUMPLE — PENDIENTE', '=COUNTIF(%s,"CUMPLE — PENDIENTE*")' % est),
             ('NO CUMPLE', '=COUNTIF(%s,"NO CUMPLE*")' % est)]
    for t_, f in items:
        r = S.r
        cell(ws, r, 2, t_, 'calc', '@', al='left', bold=True); ws.merge_cells('B%d:D%d' % (r, r))
        cell(ws, r, 5, f, 'key', '0')
        S.r += 1
    for t_, k in (('Peor V_u/V_R', 'rV'), ('Peor M_u/M_R', 'rM'), ('Peor P_u/P_R', 'rPR')):
        r = S.r
        cell(ws, r, 2, t_, 'calc', '@', al='left', bold=True); ws.merge_cells('B%d:D%d' % (r, r))
        cell(ws, r, 5, '=MAX(%s)' % mr(k), 'key', '[>=99]"∞";0.00')
        ws.merge_cells('F%d:H%d' % (r, r))
        cell(ws, r, 6, '=IFERROR(INDEX(%s,MATCH(E%d,%s,0)),"")' % (mr('key'), r, mr(k)), 'calc', '@', al='left')
        S.r += 1
    S.gap()
    S.section('3. MUROS')
    r = S.r
    heads = ['Muro', 'Story', 'Dir.', 'L (m)', 'V_u (t)', 'V_R (t)', 'V_u/V_R', 'M_u/M_R', 'P_u/P_R', 'Muro sobre viga', 'ESTADO', 'Qué refuerzo horizontal le falta']
    for k, h in enumerate(heads):
        hdr_cell(ws, r, 2 + k, h)
    ws.row_dimensions[r].height = 30
    ks = ['key', 'st', 'dir', 'L', 'Vu', 'VR', 'rV', 'rM', 'rPR', 'm418', 'est', 'rhsug']
    fm_ = ['@', '@', '@', '0.00', '0.00', '0.00', '[>=99]"V_R = 0";0.00', '[>=99]"M_R = 0";0.00', '0.00', '@', '@', '@']
    for i in range(NW):
        rr = r + 1 + i; mrow = R0 + i
        for k, (key, f_) in enumerate(zip(ks, fm_)):
            src = "MUROS!$%s$%d" % (COL[key], mrow)
            c = cell(ws, rr, 2 + k, '=IF(MUROS!$%s$%d<>"SÍ","",%s)' % (COL['ok'], mrow, src), 'key' if key in ('rV', 'rM', 'rPR') else 'calc', f_, al='left' if f_ == '@' else 'center', sz=8, wrap=key in ('m418', 'est', 'rhsug'))
        ws.row_dimensions[rr].height = 24
    from openpyxl.formatting.rule import CellIsRule
    for c_ in 'HIJ':
        ws.conditional_formatting.add('%s%d:%s%d' % (c_, r + 1, c_, r + NW), CellIsRule(operator='greaterThan', formula=['1'], font=font(8, True, color='FFFFFF'), fill=fill('C00000')))
    for c_ in 'KLM':
        add_cf(ws, '%s%d:%s%d' % (c_, r + 1, c_, r + NW))
    S.r = r + NW + 1
    S.gap()
    S.section('4. ENTREPISO Y DISTORSIONES')
    for k in range(2 * 4):
        rr = S.r
        er = N['e_r0'] + k
        cell(ws, rr, 2, '=IF(ENTREPISO!$B$%d="","",ENTREPISO!$B$%d&" "&ENTREPISO!$C$%d)' % (er, er, er), 'calc', '@', al='left')
        ws.merge_cells('C%d:D%d' % (rr, rr))
        cell(ws, rr, 3, '=IF(ENTREPISO!$K$%d="","","0.8·V_u/V_R = "&TEXT(ENTREPISO!$K$%d,"0.00"))' % (er, er), 'calc', '@', al='left')
        ws.merge_cells('E%d:M%d' % (rr, rr))
        cell(ws, rr, 5, '=ENTREPISO!$L$%d' % er, 'calc', '@', al='left')
        add_cf(ws, 'E%d:E%d' % (rr, rr))
        S.r += 1
    for k in range(6):
        rr = S.r
        dr = N['dist_r0'] + k
        cell(ws, rr, 2, '="Distorsión "&DISTORSIONES!$B$%d&" "&DISTORSIONES!$C$%d' % (dr, dr), 'calc', '@', al='left')
        ws.merge_cells('C%d:D%d' % (rr, rr))
        cell(ws, rr, 3, '=IF(DISTORSIONES!$F$%d=0,"","relación "&TEXT(DISTORSIONES!$F$%d,"0.00"))' % (dr, dr), 'calc', '@', al='left')
        ws.merge_cells('E%d:M%d' % (rr, rr))
        cell(ws, rr, 5, '=DISTORSIONES!$G$%d' % dr, 'calc', '@', al='left')
        add_cf(ws, 'E%d:E%d' % (rr, rr))
        S.r += 1
    S.gap()
    S.section('5. CANTIDADES (todos los muros capturados)')
    for t_, k, u in (('Mampostería neta', 'Am', 'm²'), ('Concreto de castillos', 'Vcc', 'm³'), ('Acero de castillos', 'Kgc', 'kg'), ('Refuerzo horizontal', 'Kgh', 'kg')):
        rr = S.r
        cell(ws, rr, 2, t_, 'calc', '@', al='left', bold=True); ws.merge_cells('B%d:D%d' % (rr, rr))
        cell(ws, rr, 5, '=SUM(%s)' % mr(k), 'key', '0.00'); cell(ws, rr, 6, u, 'calc', '@')
        S.r += 1
    ws.sheet_view.zoomScale = 90
    print_setup(ws, 'B1:M%d' % S.r, landscape=True)
    return ws
