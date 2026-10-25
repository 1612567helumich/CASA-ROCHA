"""DETALLE: memoria paso a paso del muro escogido (sale del renglón de MUROS)."""
from openpyxl.styles import Alignment
from muros_comun import *        # noqa: F401,F403

DEFAULT_SEL = 'M-2'


def V(k):
    return 'INDEX(%s,{d_r})' % mr(k)


def sheet_detalle(wb):
    ws = wb.create_sheet('DETALLE'); S = Narr(ws, N, {'A': 1.78, 'B': 36, 'C': 46, 'D': 7, 'E': 36, 'F': 28})
    S.title('DETALLE — MEMORIA DE CÁLCULO PASO A PASO DEL MURO ESCOGIDO')
    S.subtitle('Escoge el muro en la celda amarilla. CROQUIS dibuja este mismo muro.', 16)
    S.gap()
    S.section('0. MURO')
    S.header()
    S.row('Muro que se revisa y se dibuja', DEFAULT_SEL, '—', 'SELECTOR. La lista son los muros capturados en MUROS.', 'MUROS', 'input', '@', 'd_sel')
    dv_list(ws, cel(N['d_sel']), '=LST_MUROS')
    S.row('Renglón en MUROS', '=IFERROR(MATCH({d_sel},%s,0),0)' % mr('key'), '—', 'Auxiliar. 0 = no está en MUROS.', '—', 'calc', '0', 'd_r')
    S.row('¿Datos completos?', '=IF({d_r}=0,"NO",' + V('ok') + ')', '—', 'Story, dirección y L capturados.', 'MUROS', 'calc', '@', 'd_ok')
    S.row('ESTADO DEL MURO', '=IF({d_ok}<>"SÍ","— (escoge un muro con datos completos)",' + V('est') + ')', '—', 'Junta todas las revisiones (sección 10).', 'MUROS', 'key', '@', 'd_est')
    S.gap()

    def sec(t):
        S.section(t); S.header()

    sec('1. DATOS DEL MURO')
    S.row('Story de ETABS / dirección / posición', '=' + V('st') + '&" · "&' + V('dir') + '&" · "&' + V('pos'), '—', 'Capturados en MUROS.', 'MUROS', 'text', '@')
    S.row('L — longitud total', '=' + V('L'), 'm', 'De extremo a extremo, con los castillos.', 'NTC-M 6.5.2.1, p. 67', 'calc', '0.00', 'd_L')
    S.row('H — altura libre', '=' + V('H'), 'm', 'Entre losas. Vacío en MUROS = la de DATOS.', 'NTC-M 6.1.3, p. 60', 'calc', '0.00', 'd_H')
    S.row('t — espesor de la mampostería', '=' + V('t'), 'cm', 'Vacío en MUROS = el de DATOS.', 'NTC-M 6.1.3, p. 60', 'calc', '0.0', 'd_t')
    van = '&'.join('IF(N(%s)>0,%s&" en x = "&TEXT(%s,"0.00")&": "&TEXT(%s,"0.00")&" × "&TEXT(%s,"0.00")&" m"&IF(N(%s)>0," sobre antepecho de "&TEXT(%s,"0.00"),"")&"; ","")'
                   % (V('v%dw' % j), V('v%dt' % j), V('v%dx' % j), V('v%dw' % j), V('v%dh' % j), V('v%da' % j), V('v%da' % j)) for j in range(1, NV + 1))
    S.row('Vanos (ancho × alto)', '=IF(%s=0,"Sin vanos",%s)' % (V('nv'), van), '—', 'Con castillos a los lados', 'NTC-M 6.2.5.1, p. 64', 'text', '@')
    S.row('Vanos bien capturados', '=' + V('vok'), '—', 'Orden, que quepan en L y en H, tipo PUERTA o VENTANA.', 'MUROS', 'text', '@')
    S.row('Refuerzo horizontal', '=IF(OR(%s="",%s="NINGUNO"),"NINGUNO",%s&": "&IF(%s="MALLA",IF(ISNUMBER(%s),%s,1)&" cara(s) de malla "&IF(%s<>"",%s,{mallad}),IF(ISNUMBER(%s),%s,2)&" × "&IF(%s<>"",%s,{rhvd})&" cada "&IF(ISNUMBER(%s),%s,2)&" hiladas (s_h = "&TEXT(%s,"0")&" cm)"))'
          % (V('rht'), V('rht'), V('rht'), V('rht'), V('rhn'), V('rhn'), V('rhv'), V('rhv'), V('rhn'), V('rhn'), V('rhv'), V('rhv'), V('rhs'), V('rhs'), V('sh')), '—', 'Capturado en MUROS.', 'NTC-M 6.5.3 / 6.5.4', 'text', '@')
    S.row('Viga de apoyo', '=IF(%s="SÍ",IF(%s<>"",%s&" ","")&%s&" × "&%s&" cm, L_v = "&TEXT(%s,"0.00")&" m, c = "&TEXT(N(%s),"0.00")&" m","No desplanta sobre viga")'
          % (V('sv'), V('vnom'), V('vnom'), V('bv'), V('hv'), V('Lv'), V('cv')), '—', 'Capturada en MUROS.', 'NTC-M 4.1.18, p. 44', 'text', '@')
    S.gap()
    sec('2. MATERIALES (de DATOS)')
    S.row('Tipo de estructura', '={tipoE}&" — "&{tipoR}', '—', 'Con Tipo I, f\'m y v\'m salen de ensayes.', 'NTC-M 1.6.1, p. 15', 'text', '@')
    S.row("f'm", '={fm}', 'kg/cm²', '={fm_src}&" · "&{fm_st}', 'NTC-M 2.8.1', 'calc', '0.0')
    S.row("v'm", '={vm}', 'kg/cm²', '={vm_src}&" · "&{vm_st}', 'NTC-M 2.8.2', 'key', '0.00')
    for rr in (S.r - 2, S.r - 1):
        ws.cell(rr, 5).value = ws.cell(rr, 5).value.format(**N)
    S.row("f'c / f_y / f_yh", '={fc}&" / "&{fy}&" / "&{fyh}', 'kg/cm²', 'Castillos y dalas / varillas / refuerzo horizontal.', 'DATOS', 'text', '@')
    S.gap()
    sec('3. FUERZAS DE DISEÑO')
    S.row('Fuerzas de', '=' + V('fsr'), '—', 'Según DATOS', 'DATOS sección 8', 'text', '@')
    S.row('V_u — cortante de diseño', '=' + V('Vu'), 't', '=IF(' + V('fsr') + '="ETABS","Máximo |V2| de las combinaciones con sismo: "&' + V('eCb') + ',"Capturado.")', 'NTC-M 6.3.1, p. 65', 'key', '0.00', 'd_Vu')
    ws.cell(S.r - 1, 5).value = ws.cell(S.r - 1, 5).value.format(**N)
    S.row('M_u — momento en el plano', '=' + V('Mu'), 't·m', 'Máximo |M3| de las combinaciones con sismo (o capturado).', 'NTC-M 6.3.1, p. 65', 'calc', '0.00')
    S.row('P — carga axial mínima SIN factor', '=' + V('P'), 't', 'P_u mínima / FC', 'NTC-M 6.5.2.3, p. 68', 'key', '0.00', 'd_P')
    S.row('P_u mínima / máxima (factorizadas)', '=TEXT(' + V('Pmn') + ',"0.00")&" / "&TEXT(' + V('Pmx') + ',"0.00")', 't', 'La mínima y la máxima entran en M_R; la máxima en P_R.', 'NTC-M 6.4', 'text', '@')
    S.row('P_u de gravedad', '=' + V('Pg'), 't', 'Combinación de carga vertical; entra en 4.1.18.', 'NTC-M 4.1.18.2.2, p. 44', 'calc', '0.00')
    S.row('V_u por carga vertical', '=TEXT(' + V('Vg') + ',"0.00")&" t ("&TEXT(' + V('pg') + ',"0")&" % del V_u)"', '—', '> 50 %: muro colgado de una viga flexible', 'Aviso de esta hoja', 'text', '@')
    S.gap()
    sec('4. SECCIÓN CRÍTICA — TRAMOS (6.5.2)')
    S.row('Tramos que cuentan / ΣL / A_T', '=' + V('ncnt') + '&" de "&(' + V('nv') + '+1)&" · ΣL = "&TEXT(' + V('Lc') + ',"0.00")&" m · A_T = "&TEXT(' + V('AT') + ',"#,##0")&" cm²"', '—',
          '=" Un tramo menor que L mín. = "&TEXT({Lmin},"0.00")&" m es un machón: no cuenta."', 'NTC-M 6.5.2.1, p. 67', 'text', '@')
    ws.cell(S.r - 1, 5).value = ws.cell(S.r - 1, 5).value.format(**N)
    for i in range(1, NT + 1):
        T = V('T%d' % i)
        txt = ('=IF(%s<=0,"—",IF(%s<{Lmin}-0.0001,TEXT(%s,"0.00")&" m: machón, no cuenta",TEXT(%s,"0.00")&" m · f = "&TEXT(%s,"0.00")&" · V_mR = "&TEXT(%s,"0.00")&" t"&IF(%s>0," · V_sR = "&TEXT(%s,"0.00")&" t","")))'
               % (T, T, T, T, V('f%d' % i), V('V%d' % i), V('S%d' % i), V('S%d' % i)))
        S.row('Tramo %d' % i, txt, '—', 'F_R(0.5·v\'m·A_T + 0.3·P)·f' if i == 1 else '', 'NTC-M 6.5.2.1 y 6.5.2.4, p. 67–68' if i == 1 else '', 'text', '@')
    S.gap()
    sec('5. REQUISITOS GENERALES (6.1 y 6.2)')
    S.row('Espesor y esbeltez', '="t = "&TEXT({d_t},"0")&" cm · H/t = "&TEXT(' + V('Ht') + ',"0.0")&" → "&' + V('gok'), '—', 't ≥ 12 cm en Tipo I y H/t ≤ 25.', 'NTC-M 6.1.3, p. 60', 'text', '@')
    S.row('Castillos', '=' + V('nck') + '&" castillos ("&' + V('nci') + '&" intermedios) · separación máx. = mín(1.5H, 4 m) = "&TEXT(' + V('smax') + ',"0.00")&" m"', '—', 'En los extremos de cada tramo y a ≤ 1.5H y ≤ 4 m', 'NTC-M 6.2.1.1 a) y 6.2.5, p. 60–64', 'text', '@')
    S.row('Castillo típico', '={kchk}', '—', '={hc}&"×"&{bc}&" cm, "&{nc}&" No. "&{vc}&", E No. "&{ve}&" @ "&{se}&" cm (@ "&{se0}&" cm en H₀ = "&TEXT(' + V('H0') + ',"0")&" cm)"', 'NTC-M 6.2.2 y 6.2.3, p. 62', 'text', '@')
    S.row('Dala típica', '={dchk}', '—', '={bd}&"×"&{hd}&" cm, "&{nd}&" No. "&{vd}&", E No. "&{ve}&" @ "&{sed}&" cm"', 'NTC-M 6.2.1.2, p. 60', 'text', '@')
    for rr in (S.r - 2, S.r - 1):
        ws.cell(rr, 5).value = ws.cell(rr, 5).value.format(**N)
    S.row('Dalas intermedias', '=IF({d_H}>3,"FALTA una dala intermedia: H > 3 m (6.2.1.2 a)","CUMPLE: H ≤ 3 m, basta la dala superior")', '—', 'Dala en todo extremo horizontal y a no más de 3 m.', 'NTC-M 6.2.1.2 a), p. 60', 'text', '@')
    S.gap()
    sec('6. COMPRESIÓN (6.4.1)')
    S.row('F_E — excentricidad y esbeltez', '=' + V('FE'), '—', '0.6 extremo · 0.7 interior', 'NTC-M 3.1.2.3.1, p. 32', 'calc', '0.000')
    S.row('ΣA_s de los castillos', '=' + V('AsT'), 'cm²', 'Todos los castillos de la sección, con los extremos de MUROS.', 'NTC-M 6.4.1.1, p. 65', 'calc', '0.00')
    S.row('P_R = F_R·F_E·(f\'m·A_T + ΣA_s·f_y)', '=' + V('PR'), 't', 'F_R = 0.6; A_T de todos los tramos.', 'NTC-M ec. 6.4.1.1, p. 65', 'calc', '0.0')
    S.row('P_u / P_R', '=' + V('rPR'), '≤ 1', 'P_u máxima.', 'NTC-M 6.4.1.1, p. 65', 'key', '0.00')
    S.gap()
    sec('7. FLEXOCOMPRESIÓN EN EL PLANO — MÉTODO OPTATIVO (6.4.2.2)')
    S.row('A_s de un castillo extremo', '=' + V('Ase'), 'cm²', 'Del castillo típico o de lo capturado en MUROS.', 'NTC-M 6.4.2.2.1, p. 66', 'calc', '0.00')
    S.row("d' / d", '=TEXT(' + V('dp') + ',"0")&" / "&TEXT(' + V('dd') + ',"0")', 'cm', "d' entre centroides de los castillos extremos; d al acero de tensión.", 'NTC-M fig. 6.4.2.2.1, p. 66', 'text', '@')
    S.row("M₀ = A_s·f_y·d'", '=' + V('M0'), 't·m', 'Resistencia a flexión pura.', 'NTC-M 6.4.2.2.1, p. 66', 'calc', '0.00')
    S.row('M_R con P_u mínima / máxima', '=TEXT(' + V('MRa') + ',"0.00")&" / "&TEXT(' + V('MRb') + ',"0.00")', 't·m', 'F_R·M₀ + 0.3·P_u·d (si P_u ≤ P_R/3)', 'NTC-M ecs. 6.4.2.2.1.a y b, p. 66', 'text', '@')
    S.row('M_u / M_R', '=' + V('rM'), '≤ 1', 'Con el menor de los dos M_R.', 'NTC-M 6.4.2.2, p. 66', 'key', '0.00')
    S.gap()
    sec('8. CORTANTE (6.5)')
    S.row('V_mR — mampostería', '=' + V('VmR'), 't', 'Suma de los tramos (sección 4). P en tensión → 0 (6.5.2.3).', 'NTC-M 6.5.2, p. 67', 'calc', '0.00')
    S.row('(p_h·f_yh)e · η_s · k₁', '=TEXT(' + V('phfe') + ',"0.00")&" kg/cm² · "&TEXT(' + V('etas') + ',"0.000")&" · "&TEXT(' + V('k1') + ',"0.000")', '—', '(p_h·f_yh)e ≤ 0.1·f\'m·f_an', 'NTC-M 6.5.3.4.1, p. 70', 'text', '@')
    S.row('V_sR — refuerzo horizontal', '=' + V('VsR'), 't', 'Suma de los tramos: F_R·η·(p_h·f_yh)e·A_T.', 'NTC-M ec. 6.5.3.4.1.a, p. 70', 'calc', '0.00')
    S.row('Refuerzo horizontal (requisitos)', '=' + V('rhok'), '—', 'Separación, mínimo, máximo y diámetro', 'NTC-M 6.5.3.1 a 6.5.3.3, p. 69', 'text', '@')
    S.row('V_R = V_mR + V_sR', '=' + V('VR'), 't', '—', 'NTC-M ec. 6.5.1.1, p. 66', 'calc', '0.00')
    S.row('V_u / V_R', '=' + V('rV'), '≤ 1', '—', 'NTC-M 6.5, p. 66', 'key', '0.00')
    S.row('Qué refuerzo horizontal le falta', '=' + V('rhsug'), '—', 'El p_h·f_yh que hace V_R = V_u con las mismas ecuaciones de 6.5.3.4.', 'NTC-M 6.5.3.4.1, p. 70', 'text', '@')
    S.gap()
    sec('9. MURO SOBRE VIGA (4.1.18)')
    S.row('¿Desplanta sobre viga?', '=IF(' + V('sv') + '="SÍ","SÍ — "&IF(' + V('vnom') + '<>"",' + V('vnom') + '&" ","")&' + V('bv') + '&" × "&' + V('hv') + '&" cm, L_v = "&TEXT(' + V('Lv') + ',"0.00")&" m","NO")', '—', 'Viga que no es de cimentación.', 'NTC-M 4.1.18.1, p. 44', 'text', '@')
    S.row('H/L', '={d_H}/{d_L}', '—', '< 0.6: falta análisis muro–viga', 'NTC-M 4.1.18.2.1, p. 44', 'calc', '0.00')
    S.row('a = L/L_v · W_p · f_I', '=TEXT(' + V('av') + ',"0.000")&" · "&TEXT(' + V('Wp') + ',"0.0")&" kg/cm · "&TEXT(' + V('fI') + ',"0.000")', '—', 'W_p = P_u/A_T·t·(a⁴ − 4a³ + 8a); f_I = W_p/(W_p + 5w_u).', 'NTC-M ecs. 4.1.18.2.2.g y h, p. 45', 'text', '@')
    S.row('K · k_c · F_CE', '=TEXT(' + V('K') + ',"0.00")&" · "&TEXT(' + V('kc') + ',"0.00")&" · "&TEXT(' + V('FCE') + ',"0.00")', '—', 'K = [E_m·t·L_v³/(E_f·I·f_I)]^¼; F_CE = 1.6K − 3.9 o (3.2K − 7.8)·k_c.', 'NTC-M ecs. 4.1.18.2.2.c a f, p. 44–45', 'text', '@')
    S.row('f_CM = F_CE·P_u/A_T  vs  F_R·f\'m', '=TEXT(' + V('fCM') + ',"0.0")&" vs "&TEXT({FRc}*{fm},"0.0")', 'kg/cm²', 'Esfuerzo en los extremos del muro.', 'NTC-M ec. 4.1.18.2.2.a, p. 44', 'text', '@')
    S.row('L_v/h_v  vs  14 y vs límite por flecha', '=TEXT(' + V('Lhr') + ',"0.0")&" vs 14 · "&TEXT(' + V('Lhd') + ',"0.0")', '—', 'Peralte mínimo (h_v ≥ 25 cm) y flecha con cargas permanentes.', 'NTC-M 4.1.18.3 y 4.1.18.4, p. 46', 'text', '@')
    S.row('MURO SOBRE VIGA', '=' + V('m418'), '—', '—', 'NTC-M 4.1.18', 'key', '@')
    S.gap()
    sec('10. RESULTADO')
    S.row('Lo que no cumple', '=' + V('fal'), '—', '—', 'MUROS', 'text', '@')
    S.row('Pendiente o aviso', '=' + V('pen'), '—', '—', 'MUROS', 'text', '@')
    S.row('ESTADO', '={d_est}', '—', '—', 'MUROS', 'key', '@')
    S.row('Cantidades', '="Mampostería "&TEXT(' + V('Am') + ',"0.00")&" m² · concreto de castillos "&TEXT(' + V('Vcc') + ',"0.000")&" m³ · acero de castillos "&TEXT(' + V('Kgc') + ',"0.0")&" kg · refuerzo horizontal "&TEXT(' + V('Kgh') + ',"0.0")&" kg"', '—', 'Estimación para presupuesto.', 'MUROS', 'text', '@')
    for row in ws.iter_rows(min_row=1, max_row=S.r, min_col=3, max_col=5):
        for c in row:
            if isinstance(c.value, str) and '{' in c.value and c.value.startswith('='):
                c.value = c.value.format(**N)
    add_cf(ws, 'C1:C%d' % S.r)
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = 'A4'
    print_setup(ws, 'B1:F%d' % S.r, landscape=False)
    return ws
