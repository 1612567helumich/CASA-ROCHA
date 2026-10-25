"""MUROS (un renglón por muro, todas las revisiones) y PIER_FORCES (tabla de ETABS que se pega)."""
import re
from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string as CI
from muros_comun import *        # noqa: F401,F403

PF = {}          # rangos de PIER_FORCES


def F(expr, r):
    s = expr.format(**N) if '{' in expr else expr
    s = s.replace('#n#', str(r - R0 + 1))
    return re.sub(r'\[(\w+)\]', lambda m: '%s%d' % (COL[m.group(1)], r), s)


def G(expr):
    """Renglón vacío o incompleto: celda en blanco."""
    return '=IF([ok]<>"SÍ","",%s)' % expr


def mask(cls):
    return '(%s=#n#)*(%s=%d)' % (PF['IDX'], PF['CLS'], cls)


def vals_muros():
    v = {}
    v['kst'] = ('=IF([key]="","",[st]&"|"&[key])', 'calc', '@')
    v['eN'] = ('=IF([key]="","",COUNTIFS(%s,#n#,%s,1))' % (PF['IDX'], PF['CLS']), 'calc', '0')
    v['eNg'] = ('=IF([key]="","",COUNTIFS(%s,#n#,%s,2))' % (PF['IDX'], PF['CLS']), 'calc', '0')
    v['eVu'] = ('=IF(N([eN])=0,0,SUMPRODUCT(MAX(%s*%s)))' % (mask(1), PF['AV']), 'calc', '0.00')
    v['eMu'] = ('=IF(N([eN])=0,0,SUMPRODUCT(MAX(%s*%s)))' % (mask(1), PF['AM']), 'calc', '0.00')
    v['ePmn'] = ('=IF(N([eN])=0,0,-SUMPRODUCT(MAX(%s*(-%s)+(1-%s)*(-1E9))))' % (mask(1), PF['PC'], mask(1)), 'calc', '0.00')
    v['ePmx'] = ('=IF(N([eN])=0,0,SUMPRODUCT(MAX(%s*%s+(1-%s)*(-1E9))))' % (mask(1), PF['PC'], mask(1)), 'calc', '0.00')
    v['ePg'] = ('=IF(N([eNg])=0,[ePmx],SUMPRODUCT(MAX(%s*%s+(1-%s)*(-1E9))))' % (mask(2), PF['PC'], mask(2)), 'calc', '0.00')
    v['eVg'] = ('=IF(N([eNg])=0,0,SUMPRODUCT(MAX(%s*%s)))' % (mask(2), PF['AV']), 'calc', '0.00')
    v['eRw'] = ('=IF(N([eN])=0,0,SUMPRODUCT(MAX(%s*(ABS(%s-[eVu])<0.000001)*%s)))' % (mask(1), PF['AV'], PF['ROWN']), 'calc', '0')
    v['eCb'] = ('=IF(N([eRw])=0,"",INDEX(%s,[eRw])&" "&INDEX(%s,[eRw])&" / "&INDEX(%s,[eRw]))' % (PF['CASE'], PF['STEP'], PF['LOC']), 'calc', '@')
    v['ok'] = ('=IF([key]="","",IF(AND([st]<>"",OR([dir]="X",[dir]="Y"),ISNUMBER([L]),N([L])>0),"SÍ","NO"))', 'calc', '@')
    v['fsr'] = (G('IF({fsrc}="ETABS (PIER_FORCES)",IF(N([eN])>0,"ETABS","NO ESTÁ EN PIER_FORCES"),"CAPTURA")'), 'calc', '@')
    v['Vu'] = (G('IF([fsr]="ETABS",[eVu],IF([fsr]="CAPTURA",N([cVu]),0))'), 'key', '0.00')
    v['Mu'] = (G('IF([fsr]="ETABS",[eMu],IF([fsr]="CAPTURA",N([cMu]),0))'), 'key', '0.00')
    v['P'] = (G('IF([fsr]="ETABS",[ePmn]/{FCs},IF([fsr]="CAPTURA",N([cP]),0))'), 'key', '0.00')
    v['Pmn'] = (G('IF([fsr]="ETABS",[ePmn],IF([fsr]="CAPTURA",IF(ISNUMBER([cPmn]),[cPmn],N([cP])*{FCs}),0))'), 'calc', '0.00')
    v['Pmx'] = (G('IF([fsr]="ETABS",MAX([ePmx],[ePg]),IF([fsr]="CAPTURA",N([cPmx]),0))'), 'calc', '0.00')
    v['Pg'] = (G('IF([fsr]="ETABS",[ePg],IF([fsr]="CAPTURA",IF(ISNUMBER([cPg]),[cPg],N([cPmx])),0))'), 'calc', '0.00')
    v['Vg'] = (G('IF([fsr]="ETABS",[eVg],IF([fsr]="CAPTURA",N([cVg]),0))'), 'calc', '0.00')
    v['pg'] = (G('IF([Vu]>0,100*[Vg]/[Vu],0)'), 'calc', '0')
    v['H'] = (G('IF(ISNUMBER([Hi]),[Hi],{Hdef})'), 'calc', '0.00')
    v['t'] = (G('IF(ISNUMBER([ti]),[ti],{tdef})'), 'calc', '0.0')
    v['nv'] = (G('IF(N([v1w])>0,1+IF(N([v2w])>0,1+IF(N([v3w])>0,1+IF(N([v4w])>0,1,0),0),0),0)'), 'calc', '0')
    conds = []
    for j in range(1, NV + 1):
        e_prev = '0' if j == 1 else '([v%dx]+[v%dw])' % (j - 1, j - 1)
        conds.append('OR([nv]<%d,AND(N([v%dx])>=%s-0.001,[v%dx]+[v%dw]<=[L]+0.001,N([v%da])>=0,N([v%da])+N([v%dh])<=[H]+0.001,N([v%dh])>0,OR([v%dt]="PUERTA",[v%dt]="VENTANA")))'
                     % (j, j, e_prev, j, j, j, j, j, j, j, j))
    loose = '+'.join('(N([v%dw])>0)' % j for j in range(1, NV + 1))
    v['vok'] = (G('IF(AND(%s,%s=[nv]),"CUMPLE","REVISAR VANOS: tipo, orden de izquierda a derecha, que quepan en L y en H")' % (','.join(conds), loose)), 'calc', '@')
    v['T1'] = (G('MAX(0,IF([nv]>=1,N([v1x]),[L]))'), 'calc', '0.00')
    for j in range(2, NT + 1):
        nxt = 'N([v%dx])-(N([v%dx])+N([v%dw]))' % (j, j - 1, j - 1)
        last = '[L]-(N([v%dx])+N([v%dw]))' % (j - 1, j - 1)
        if j <= NV:
            v['T%d' % j] = (G('MAX(0,IF([nv]>=%d,%s,IF([nv]=%d,%s,0)))' % (j, nxt, j - 1, last)), 'calc', '0.00')
        else:
            v['T%d' % j] = (G('MAX(0,IF([nv]=%d,%s,0))' % (j - 1, last)), 'calc', '0.00')
    TT = '[T1]:[T%d]' % NT
    v['Lc'] = (G('SUMPRODUCT((%s>={Lmin}-0.0001)*%s)' % (TT, TT)), 'key', '0.00')
    v['ncnt'] = (G('SUMPRODUCT(--(%s>={Lmin}-0.0001))' % TT), 'calc', '0')
    v['Ltot'] = (G('SUM(%s)' % TT), 'calc', '0.00')
    v['AT'] = (G('[t]*[Lc]*100'), 'calc', '0')
    v['ATt'] = (G('[t]*[Ltot]*100'), 'calc', '0')
    v['Ht'] = (G('[H]*100/[t]'), 'calc', '0.0')
    v['gok'] = (G('IF(AND(OR({tipoE}="TIPO II",[t]>=12),[Ht]<=25),"CUMPLE","NO CUMPLE: "&IF(AND({tipoE}="TIPO I",[t]<12),"t < 12 cm ","")&IF([Ht]>25,"H/t > 25",""))'), 'calc', '@')
    # castillos: en los extremos de cada tramo (uno solo en un tramo más corto que 2·h_c) más los intermedios por separación
    v['smax'] = (G('MIN(1.5*[H],4)'), 'calc', '0.00')
    v['nci'] = (G('SUMPRODUCT((%s>=2*{hc}/100)*INT((%s-{hc}/100)/[smax]-0.000001)*((%s-{hc}/100)>0))' % (TT, TT, TT)), 'calc', '0')
    v['nck'] = (G('SUMPRODUCT((%s>0.0001)*1+(%s>=2*{hc}/100)*1)+[nci]' % (TT, TT)), 'calc', '0')
    v['Ase'] = (G('IF(ISNUMBER([nce]),[nce],{nc})*' + ab('IF([vce]<>"",[vce],{vc})')), 'calc', '0.00')
    v['AsT'] = (G('[nck]*{Asc}+IF([nck]>=2,2,[nck])*([Ase]-{Asc})'), 'calc', '0.00')
    v['H0'] = (G('MAX([H]*100/6,40)'), 'calc', '0.0')
    # compresión
    v['FE'] = (G('IF([Ht]<=20,IF([pos]="INTERIOR",0.7,0.6),MIN(IF([pos]="INTERIOR",0.7,0.6),(1-2/24)*(1-(IF([pos]="INTERIOR",0.8,1)*[Ht]/30)^2)))'), 'calc', '0.000')
    v['PR'] = (G('{FRc}*[FE]*({fm}*[ATt]+[AsT]*{fy})/1000'), 'calc', '0.0')
    v['rPR'] = (G('IF([PR]>0,[Pmx]/[PR],99)'), 'key', '0.00')
    # flexocompresión, método optativo (muro completo entre castillos extremos)
    v['dp'] = (G('MAX(0,[L]*100-{hc})'), 'calc', '0.0')
    v['dd'] = (G('MAX(0,[L]*100-{hc}/2)'), 'calc', '0.0')
    v['M0'] = (G('[Ase]*{fy}*[dp]/100000'), 'calc', '0.00')

    def MRf(P):
        return ('IF(%s<0,0.8*[M0]*MAX(0,1+%s/(0.8*[AsT]*{fy}/1000)),IF(%s<=[PR]/3,0.8*[M0]+0.3*%s*[dd]/100,(1.5*0.6*[M0]+0.15*[PR]*[dd]/100)*MAX(0,1-%s/[PR])))' % (P, P, P, P, P))
    v['MRa'] = (G(MRf('[Pmn]')), 'calc', '0.00')
    v['MRb'] = (G(MRf('[Pmx]')), 'calc', '0.00')
    v['MR'] = (G('MAX(0,MIN([MRa],[MRb]))'), 'key', '0.00')
    v['rM'] = (G('IF([MR]>0,[Mu]/[MR],IF([Mu]>0,99,0))'), 'key', '0.00')
    # cortante de la mampostería por tramo
    for i in range(1, NT + 1):
        v['f%d' % i] = (G('IF([T{0}]>0,MAX(1,MIN(1.5,1.5-0.625*([H]/[T{0}]-0.2))),0)'.format(i)), 'calc', '0.00')
        v['V%d' % i] = (G('IF(OR([T{0}]<{{Lmin}}-0.0001,[P]<0,[Lc]<=0),0,MIN({{FRv}}*(0.5*{{vm}}*[t]*[T{0}]*100+0.3*[P]*[T{0}]/[Lc]*1000)*[f{0}],1.5*{{FRv}}*{{vm}}*[t]*[T{0}]*100*[f{0}])/1000)'.format(i)), 'calc', '0.00')
    v['VmR'] = (G('SUM([V1]:[V%d])' % NT), 'key', '0.00')
    # refuerzo horizontal
    rhv = 'IF([rhv]<>"",[rhv],{rhvd})'
    v['Ash'] = (G('IF([rht]="VARILLA EN JUNTAS",IF(ISNUMBER([rhn]),[rhn],2)*%s,IF([rht]="MALLA",IF(ISNUMBER([rhn]),[rhn],1)*%s/100,IF([rht]="ESCALERILLA",IF(ISNUMBER([rhn]),[rhn],2)*%s,0)))'
                  % (ab(rhv), mas('IF([rhv]<>"",[rhv],{mallad})'), ab(rhv))), 'calc', '0.000')
    v['sh'] = (G('IF(OR([rht]="VARILLA EN JUNTAS",[rht]="ESCALERILLA"),IF(ISNUMBER([rhs]),[rhs],2)*{hh},IF([rht]="MALLA",1,0))'), 'calc', '0.0')
    v['phf'] = (G('IF([sh]>0,[Ash]/([sh]*[t])*MIN({fyh},IF([rht]="MALLA",5000,6000)),0)'), 'calc', '0.00')
    v['phfe'] = (G('IF([rht]="ESCALERILLA",0,MIN([phf],0.1*{fm}*{fan}))'), 'calc', '0.00')
    v['etas'] = (G('IF([rht]="MALLA",0.5,IF({fm}<=60,0.55,IF({fm}>=90,0.75,0.55+0.2*({fm}-60)/30)))'), 'calc', '0.000')
    v['k1'] = (G('1-{alfa}*[phfe]'), 'calc', '0.000')
    for i in range(1, NT + 1):
        v['k0%d' % i] = (G('IF([T{0}]>0,IF([H]/[T{0}]<=1,1.3,IF([H]/[T{0}]>=1.5,1,1.3-0.6*([H]/[T{0}]-1))),0)'.format(i)), 'calc', '0.00')
        eta = 'MAX(0,IF([P]>=0,[V{0}]/({{FRv}}*[phfe]*[t]*[T{0}]*100/1000)*([k0{0}]*[k1]-1)+[etas],[k1]*[etas]))'.format(i)
        v['S%d' % i] = (G('IF(OR([phfe]<=0,[T{0}]<{{Lmin}}-0.0001),0,{{FRv}}*{1}*[phfe]*[t]*[T{0}]*100/1000)'.format(i, eta)), 'calc', '0.00')
    v['VsR'] = (G('SUM([S1]:[S%d])' % NT), 'key', '0.00')
    dbh = db(rhv)
    v['rhok'] = (G('IF(OR([rht]="",[rht]="NINGUNO"),IF({vm}>=10,"FALTA refuerzo horizontal: v\'m ≥ 10 kg/cm² (6.5.2.2)",IF(AND([P]<0,[Vu]>0),"FALTA: P en tensión, el refuerzo debe tomar todo el cortante (6.5.1.3)","—")),'
                   'IF([rht]="ESCALERILLA","NO CUENTA: la escalerilla no se permite para cortante por sismo (6.5.3.1.2)",'
                   'IF([rht]="MALLA",IF(AND([phf]>=2.5,[phf]<=MAX(0.15*{fm}*{fan},2.5),{fyh}<=5000),"CUMPLE","NO CUMPLE: "&IF([phf]<2.5,"p_h·f_yh < 2.5 ","")&IF([phf]>MAX(0.15*{fm}*{fan},2.5),"p_h·f_yh > 0.15·f\'m·f_an ","")&IF({fyh}>5000,"f_yh > 5000 (6.5.4.1.3)","")),'
                   'IF(AND([sh]<=MIN(6*{hh},45),[phf]>=3,[phf]<=MIN(MAX(0.15*{fm}*{fan},3),0.05*{hj}/10*MIN({fyh},6000)/[sh]),%s>=0.35,%s<=0.75*{hj}/10,{fyh}<=6000),"CUMPLE",'
                   '"NO CUMPLE: "&IF([sh]>MIN(6*{hh},45),"s_h > 6 hiladas o 45 cm (6.5.3.2.1) ","")&IF([phf]<3,"p_h·f_yh < 3 (6.5.3.3.1) ","")&IF([phf]>MIN(MAX(0.15*{fm}*{fan},3),0.05*{hj}/10*MIN({fyh},6000)/[sh]),"p_h·f_yh > máximo (6.5.3.3.2) ","")'
                   '&IF(OR(%s<0.35,%s>0.75*{hj}/10),"diámetro fuera de 3.5 mm a 0.75·h_j (4.2.2.2.1) ","")&IF({fyh}>6000,"f_yh > 6000 (6.5.3.1.3)","")))))' % (dbh, dbh, dbh, dbh)), 'calc', '@')
    v['VR'] = (G('[VmR]+[VsR]'), 'key', '0.00')
    v['rV'] = (G('IF([VR]>0,[Vu]/[VR],IF([Vu]>0,99,0))'), 'key', '0.00')
    # refuerzo horizontal que hace falta: V_mR·k₀ + x·(F_R·η_s·A_T − α·Σk₀·V_mR) ≥ V_u (P ≥ 0), o F_R·η_s·x(1 − αx)·A_T ≥ V_u (P en tensión)
    SK = 'SUMPRODUCT([k01]:[k0%d],[V1]:[V%d])' % (NT, NT)
    es = 'IF([rht]="MALLA",0.5,IF({fm}<=60,0.55,IF({fm}>=90,0.75,0.55+0.2*({fm}-60)/30)))'
    den = '({FRv}*%s*[AT]/1000-{alfa}*%s)' % (es, SK)
    disc = '(1-4*{alfa}*[Vu]*1000/({FRv}*%s*MAX([AT],1)))' % es
    x = 'IF([P]>=0,IF(%s<=0,-1,MAX(3,([Vu]-%s)/%s)),IF(%s<0,-1,MAX(3,(1-SQRT(%s))/(2*{alfa}))))' % (den, SK, den, disc, disc)
    xr = 'IF([P]>=0,IF(%s<=0,1E9,([Vu]-%s)/%s),IF(%s<0,1E9,(1-SQRT(%s))/(2*{alfa})))' % (den, SK, den, disc, disc)
    # x efectivo que hace falta; si cabe en 0.1·f'm·f_an, lo físico es el mayor entre ese x y el mínimo (3, o 2.5 con malla)
    v['phreq'] = (G('IF([AT]<=0,-1,IF([Vu]<=[VmR],0,IF(%s>0.1*{fm}*{fan}+0.000001,-1,MAX(IF([rht]="MALLA",2.5,3),%s))))' % (xr, xr)), 'calc', '0.0')
    fye = 'MIN({fyh},6000)'
    opts = [(2, '3/16 in'), (2, '5/32 in'), (1, '3/16 in'), (1, '5/32 in')]
    pmax = lambda k: 'MIN(MAX(0.15*{fm}*{fan},3),0.05*{hj}/10*%s/(%s*{hh}))' % (fye, k)
    txt = '"ninguna combinación de 1 o 2 alambres por junta queda entre el mínimo y el máximo: revisa con malla (6.5.4)"'
    for n_, b_ in reversed(opts):
        A_ = '(%d*%s)' % (n_, ab('"%s"' % b_))
        k_ = 'MIN(6,FLOOR(45/{hh},1),FLOOR(%s*%s/([t]*{hh}*MAX([phreq],0.001)),1))' % (A_, fye)
        ph_ = '(%s*%s/(MAX(%s,1)*{hh}*[t]))' % (A_, fye, k_)
        ok_ = 'AND(%s>=1,%s<=%s)' % (k_, ph_, pmax(k_))
        lbl = '"%d alambre%s de %s por junta cada "&%s&" hilada(s) (p_h·f_yh = "&TEXT(%s,"0.00")&")"' % (n_, 's' if n_ > 1 else '', b_.replace(' in', '""'), k_, ph_)
        txt = 'IF(%s,%s,%s)' % (ok_, lbl, txt)
    v['rhsug'] = (G('IF([phreq]=0,"NO REQUIERE: V_mR ≥ V_u",IF([phreq]<0,"NO ALCANZA con refuerzo horizontal (haría falta más de 0.1·f\'m·f_an efectivo): sube v\'m con ensayes, el espesor o la longitud del muro",'
                    '"p_h·f_yh ≥ "&TEXT([phreq],"0.0")&" kg/cm² → "&%s))' % txt), 'calc', '@')
    # muro sobre viga
    a_ = '[av]'
    poly = '(%s^4-4*%s^3+8*%s)' % (a_, a_, a_)
    v['Iv'] = (G('IF([sv]="SÍ",N([bv])*N([hv])^3/12,0)'), 'calc', '0')
    v['av'] = (G('IF(AND([sv]="SÍ",N([Lv])>0),MIN(1,[L]/[Lv]),0)'), 'calc', '0.000')
    v['Wp'] = (G('IF([sv]="SÍ",[Pg]*1000/([t]*[L]*100)*[t]*%s,0)' % poly), 'calc', '0.0')
    v['fI'] = (G('IF(AND([sv]="SÍ",[Wp]>0),[Wp]/([Wp]+5*N([wv])*10),0)'), 'calc', '0.000')
    v['K'] = (G('IF(AND([sv]="SÍ",[fI]>0,[Iv]>0),({Em}*[t]*([Lv]*100)^3/({Ecv}*[Iv]*[fI]))^0.25,0)'), 'calc', '0.00')
    rc_ = 'MIN(1,N([cv])/(([Lv]-[L])/2))'
    v['kc'] = (G('IF([sv]<>"SÍ",0,IF(N([Lv])<=[L],1,IF(%s<=0.5,0.7+0.6*%s,1.5-%s)))' % (rc_, rc_, rc_)), 'calc', '0.00')
    f9 = '((3.2*[K]-7.8)*[kc])'; f1 = '(1.6*[K]-3.9)'
    v['FCE'] = (G('IF([sv]<>"SÍ",0,MAX(1,IF([av]>=1,%s,IF([av]<0.9,%s,%s+(%s-%s)*([av]-0.9)/0.1))))' % (f1, f9, f9, f1, f9)), 'calc', '0.00')
    v['fCM'] = (G('IF([sv]="SÍ",[FCE]*[Pg]*1000/([t]*[L]*100),0)'), 'calc', '0.0')
    v['Lhr'] = (G('IF(AND([sv]="SÍ",N([hv])>0),[Lv]*100/[hv],0)'), 'calc', '0.0')
    PD = 'IF(ISNUMBER([PD]),[PD],[Pg]/{FCcm})'
    WpD = '(%s*1000/([t]*[L]*100)*[t]*%s)' % (PD, poly)
    v['Lhd'] = (G('IF([sv]="SÍ",(1/15*{Ecv}*N([bv])/MAX(%s+5*N([wv])/{FCcm}*10,0.001))^(1/3),0)' % WpD), 'calc', '0.0')
    v['m418'] = (G('IF([sv]<>"SÍ","—",IF(OR(N([bv])<=0,N([hv])<=0,N([Lv])<=0),"REVISAR: faltan b, h o L_v de la viga",IF(AND([fCM]<={FRc}*{fm},[Lhr]<=14,N([hv])>=25,[Lhr]<=[Lhd]),'
                   'IF([H]/[L]<0.6,"PENDIENTE: H/L < 0.6, falta el análisis de interacción muro–viga (4.1.18.2.1); f_CM es estimado","CUMPLE"),'
                   '"NO CUMPLE: "&IF([fCM]>{FRc}*{fm},"f_CM > F_R·f\'m ","")&IF([Lhr]>14,"L_v/h_v > 14 ","")&IF(N([hv])<25,"h_v < 25 cm ","")&IF([Lhr]>[Lhd],"flecha (4.1.18.3)",""))))'), 'calc', '@')
    # cantidades
    vanos = '+'.join('N([v%dw])*N([v%dh])' % (j, j) for j in range(1, NV + 1))
    v['Am'] = (G('MAX(0,[L]*[H]-(%s)-[nck]*{hc}/100*[H])' % vanos), 'calc', '0.00')
    v['Vcc'] = (G('[nck]*{hc}*{bc}/10000*[H]'), 'calc', '0.000')
    est_n = '(2*ROUNDUP([H0]/{se0},0)+ROUNDUP(MAX(0,[H]*100-2*[H0])/{se},0))'
    est_L = '((2*({bc}+{hc}-4*({rc}-%s))+15)/100)' % db('{ve}')
    v['Kgc'] = (G('[nck]*({nc}*([H]+0.6)*%s+%s*%s*%s)' % (wb_('{vc}'), est_n, est_L, wb_('{ve}'))), 'calc', '0.0')
    v['Kgh'] = (G('IF(OR([rht]="VARILLA EN JUNTAS",[rht]="ESCALERILLA"),INT([H]*100/MAX([sh],1))*IF(ISNUMBER([rhn]),[rhn],2)*[Lc]*%s,IF([rht]="MALLA",IF(ISNUMBER([rhn]),[rhn],1)*[Lc]*[H]*%s,0))'
                  % (wb_(rhv), mkg('IF([rhv]<>"",[rhv],{mallad})'))), 'calc', '0.0')
    # resultado
    v['fal'] = (G('""&IF([fsr]="NO ESTÁ EN PIER_FORCES","sin fuerzas en PIER_FORCES; ","")&IF([vok]<>"CUMPLE","vanos; ","")&IF(LEFT([gok],2)="NO","espesor o esbeltez; ","")&IF([rPR]>1,"compresión; ","")'
                  '&IF([rM]>1,"flexión; ","")&IF([Lc]<=0,"sin tramos de muro (todos son machones menores que L mín.: diséñalos como columnas o baja L mín.); ",IF([rV]>1,"cortante; ",""))&IF(OR(LEFT([rhok],2)="NO",LEFT([rhok],5)="FALTA"),"refuerzo horizontal; ","")&IF(OR(LEFT([m418],2)="NO",LEFT([m418],7)="REVISAR"),"muro sobre viga; ","")'
                  '&IF(LEFT({kchk},2)="NO","castillo típico; ","")&IF(LEFT({dchk},2)="NO","dala típica; ","")&IF(AND([t]<>{tdef},{bc}<[t]),"castillo más angosto que el muro; ","")'), 'calc', '@')
    v['pen'] = (G('""&IF(LEFT([m418],9)="PENDIENTE","análisis muro–viga (H/L < 0.6); ","")&IF([pg]>50,"V_u casi todo por carga vertical; ","")'
                  '&IF(OR(LEFT({vm_st},9)="PENDIENTE",LEFT({fm_st},9)="PENDIENTE"),"v\'m provisional; ","")&IF(OR(LEFT({vm_st},2)="NO",LEFT({fm_st},2)="NO"),"f\'m o v\'m no válido; ","")'), 'calc', '@')
    v['est'] = (G('IF([fal]<>"","NO CUMPLE: "&LEFT([fal],LEN([fal])-2),IF([pen]<>"","CUMPLE — PENDIENTE: "&LEFT([pen],LEN([pen])-2),"CUMPLE"))'), 'key', '@')
    return v


NOTES = {
    'key': 'Muro y Story iguales que en ETABS. H y t vacíos = los de DATOS.',
    'v1t': 'De izquierda a derecha. x = del extremo izquierdo al vano. Puerta: antepecho 0.',
    'nce': 'Opcional. Refuerzo horizontal: varilla en juntas o malla.',
    'sv': 'Sólo si el muro desplanta sobre una viga (4.1.18).',
    'cVu': 'Sólo si DATOS dice CAPTURA. P = la mínima sin factor.',
}
# columnas de resultado que se ven; las demás de cálculo quedan agrupadas (botón [+])
VISIBLE = {'Vu', 'Mu', 'P', 'Pmx', 'Lc', 'PR', 'rPR', 'MR', 'rM', 'VmR', 'VsR', 'rhok', 'VR', 'rV', 'rhsug', 'm418', 'est', 'nota'}


def sheet_muros(wb, examples):
    ws = wb.create_sheet('MUROS')
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3.2
    last = COL['nota']
    title = 'MUROS — UN RENGLÓN POR MURO (NTC-MAMPOSTERÍA 2023)'
    sub = 'Amarillo = captura. Las cuentas intermedias están agrupadas: ábrelas con [+] arriba de las columnas. DETALLE explica un muro paso a paso y CROQUIS lo dibuja.'
    leg = ('="Proyecto: "&{proj}&" · "&{tipoE}&" · f\'m = "&TEXT({fm},"0")&" kg/cm² · v\'m = "&TEXT({vm},"0.00")&" kg/cm² ("&{vm_src}&") · fuerzas: "&{fsrc}&" · castillo "&{hc}&"×"&{bc}&" con "&{nc}&" No. "&{vc}').format(**N)
    for rr, txt, f_, h in ((1, title, font(13, True, color='FFFFFF'), 26), (2, sub, font(9, i=True, color=GRAY), 22), (3, leg, font(9, True, color=WARN_FG), 22)):
        ws.merge_cells('B%d:%s%d' % (rr, last, rr))
        c = ws['B%d' % rr]; c.value = txt; c.font = f_
        c.alignment = Alignment(wrap_text=True, vertical='center')
        if rr == 1:
            c.fill = fill(TITLE_BG)
        if rr == 3:
            c.fill = fill(WARN_BG)
        ws.row_dimensions[rr].height = h
    ws.row_dimensions[4].height = 6
    for a, b, t in GROUPS:
        ws.merge_cells('%s6:%s6' % (COL[a], COL[b]))
        x = ws['%s6' % COL[a]]; x.value = t; x.font = font(9, True, color='FFFFFF'); x.fill = fill('7F6000' if 'CAPTURA' in t or 'CAPTURADAS' in t else SEC_BG)
        x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.merge_cells('%s5:%s5' % (COL[a], COL[b]))
        y = ws['%s5' % COL[a]]; y.value = NOTES.get(a, ''); y.font = font(7.5, i=True, color=GRAY); y.fill = fill(NOTE_BG)
        y.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 18
    for k, (t, w, u) in COLS.items():
        hdr_cell(ws, 7, CI(COL[k]), t, w)
        unit_cell(ws, 8, CI(COL[k]), u)
    ws.row_dimensions[7].height = 66
    ws.merge_cells('B9:%s9' % last)
    ws['B9'].value = 'Ejemplo: Casa Rocha. Para otro proyecto escribe encima de las columnas amarillas.'
    ws['B9'].font = font(8, i=True, color=GRAY)
    ws.row_dimensions[10].height = 4
    hdr_cell(ws, 7, 1, '#'); unit_cell(ws, 8, 1, '')
    vals = vals_muros()
    txt_inputs = {'key', 'st', 'dir', 'pos', 'rht', 'rhv', 'sv', 'vnom', 'nota', 'vce'} | {'v%dt' % j for j in range(1, NV + 1)}
    for r in range(R0, R1 + 1):
        x = ws.cell(r, 1, '=IF(B%d="","",ROW()-%d)' % (r, R0 - 1)); x.font = font(7, color='A6A6A6'); x.alignment = Alignment(horizontal='center')
        ex = examples[r - R0] if r - R0 < len(examples) else {}
        for k in INPUTS:
            fmt = '@' if k in txt_inputs - {'vce', 'rhv'} else ('0.00' if k in ('L', 'Hi', 'Lv', 'cv', 'wv', 'cVu', 'cMu', 'cP', 'cPmx', 'cPmn', 'cPg', 'cVg', 'PD') or k[:1] == 'v' and k[-1] in 'xwah' else 'General')
            cell(ws, r, CI(COL[k]), ex.get(k), 'input', fmt, al='left' if k in ('key', 'nota') else 'center', sz=9, wrap=(k == 'nota'))
        for k in COLS:
            if k in INPUTS:
                continue
            f, kind, fmt = vals[k]
            c = cell(ws, r, CI(COL[k]), F(f, r), 'key' if kind == 'key' else 'calc', fmt, al='left' if fmt == '@' else 'center', sz=9)
            if k in ('est', 'fal', 'pen', 'rhsug', 'm418', 'rhok', 'vok', 'gok', 'eCb'):
                c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')
        ws.row_dimensions[r].height = 30
    for k, lst in (('dir', 'LST_DIR'), ('pos', 'LST_POS'), ('rht', 'LST_RH'), ('sv', 'LST_SINO')):
        dv_list(ws, '%s%d:%s%d' % (COL[k], R0, COL[k], R1), '=' + lst)
    for j in range(1, NV + 1):
        dv_list(ws, '%s%d:%s%d' % (COL['v%dt' % j], R0, COL['v%dt' % j], R1), '=LST_VANO')
    dv_list(ws, '%s%d:%s%d' % (COL['vce'], R0, COL['vce'], R1), '=' + TB['BAR'])
    # agrupar (y esconder) las columnas de cálculo intermedio
    run = []
    keys = [k for k, _ in [(k_, 0) for k_ in COLS]]
    for k in keys + [None]:
        hide = k is not None and k not in INPUTS and k not in VISIBLE
        if hide:
            run.append(k)
        elif run:
            ws.column_dimensions.group(COL[run[0]], COL[run[-1]], outline_level=1, hidden=True)
            run = []
    ws.sheet_properties.outlinePr.summaryRight = False
    ws.freeze_panes = 'C%d' % R0
    for k in ('est', 'fal', 'pen', 'rhok', 'rhsug', 'm418', 'vok', 'gok', 'fsr', 'ok'):
        add_cf(ws, '%s%d:%s%d' % (COL[k], R0, COL[k], R1))
    from openpyxl.formatting.rule import CellIsRule
    for k in ('rV', 'rM', 'rPR'):
        rng = '%s%d:%s%d' % (COL[k], R0, COL[k], R1)
        ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThan', formula=['1'], font=font(9, True, color='FFFFFF'), fill=fill('C00000')))
    ws.sheet_view.zoomScale = 80
    print_setup(ws, 'A1:%s%d' % (last, R1), landscape=True)
    ws.print_title_rows = '5:8'
    ws.page_setup.paperSize = ws.PAPERSIZE_TABLOID
    return ws


def define_pf_names(wb):
    for k, c in (('IDX', 'O'), ('CLS', 'P'), ('PC', 'Q'), ('AV', 'R'), ('AM', 'S'), ('ROWN', 'T'), ('CASE', 'D'), ('STEP', 'F'), ('LOC', 'G')):
        PF[k] = 'PIER_FORCES!$%s$%d:$%s$%d' % (c, PF0, c, PF1)


def sheet_pier_forces(wb):
    ws = wb.create_sheet('PIER_FORCES')
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3
    for c, w in zip('BCDEFGHIJKLMNOPQRST', (10, 13, 16, 12, 9, 9, 9, 9, 9, 9, 9, 9, 2, 8, 6, 9, 9, 9, 7)):
        ws.column_dimensions[c].width = w
    ws.merge_cells('B1:T1'); ws['B1'] = 'PIER_FORCES — TABLA «PIER FORCES» DE ETABS (SE PEGA AQUÍ)'; ws['B1'].font = font(13, True, color='FFFFFF'); ws['B1'].fill = fill(TITLE_BG)
    ws.row_dimensions[1].height = 26
    ws.merge_cells('B2:T3')
    ws['B2'] = ('En ETABS: Display → Show Tables → Analysis Results → Wall Results → Pier Forces, unidades tonf y m, con las combinaciones de diseño. Copia la tabla SIN encabezados y pégala como valores desde B%d. '
                'Las columnas deben quedar como en el renglón 8 (si tu tabla trae «Step Number» u otra columna extra, bórrala antes de pegar). Luego, en DATOS sección 8, cambia «Fuerzas de los muros» a ETABS. '
                'Las columnas grises de la derecha se calculan solas: renglón del muro en MUROS (por Story|Pier), tipo de combinación (1 = sismo, 2 = gravedad), P a compresión positiva, |V2| y |M3|. Caben %d renglones.' % (PF0, NPF))
    ws['B2'].font = font(9, i=True, color=GRAY); ws['B2'].alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[2].height = 30; ws.row_dimensions[3].height = 30
    ws.merge_cells('B5:M5'); ws['B5'] = 'TABLA DE ETABS — CAPTURA (pegar valores)'; ws['B5'].font = font(9, True, color='FFFFFF'); ws['B5'].fill = fill('7F6000')
    ws.merge_cells('O5:T5'); ws['O5'] = 'AUXILIARES (no se tocan)'; ws['O5'].font = font(9, True, color='FFFFFF'); ws['O5'].fill = fill(SEC_BG)
    heads = ['Story', 'Pier', 'Output Case', 'Case Type', 'Step Type', 'Location', 'P', 'V2', 'V3', 'T', 'M2', 'M3']
    units = ['', '', '', '', '', '', 'tonf', 'tonf', 'tonf', 'tonf·m', 'tonf·m', 'tonf·m']
    for i, (h, u) in enumerate(zip(heads, units)):
        hdr_cell(ws, 7, 2 + i, h); unit_cell(ws, 8, 2 + i, u)
    for i, (h, u) in enumerate(zip(['Renglón en MUROS', 'Tipo (1 sismo, 2 grav.)', 'P (compresión +)', '|V2|', '|M3|', '#'], ['', '', 'tonf', 'tonf', 'tonf·m', ''])):
        hdr_cell(ws, 7, 15 + i, h); unit_cell(ws, 8, 15 + i, u)
    ws.row_dimensions[7].height = 30
    kst = mr('kst')
    for r in range(PF0, PF1 + 1):
        for c in range(2, 14):
            x = ws.cell(r, c); x.fill = fill(INP_BG); x.font = font(8, color=INP_FG)
        ws.cell(r, 15, '=IF($C%d="","",IFERROR(MATCH($B%d&"|"&$C%d,%s,0),0))' % (r, r, r, kst)).font = font(8, color=GRAY)
        ws.cell(r, 16, '=IF($D%d="","",IF(ISNUMBER(MATCH($D%d,CMB_S,0)),1,IF(ISNUMBER(MATCH($D%d,CMB_G,0)),2,0)))' % (r, r, r)).font = font(8, color=GRAY)
        ws.cell(r, 17, '=IF(ISNUMBER($H%d),-$H%d,0)' % (r, r)).font = font(8, color=GRAY)
        ws.cell(r, 18, '=IF(ISNUMBER($I%d),ABS($I%d),0)' % (r, r)).font = font(8, color=GRAY)
        ws.cell(r, 19, '=IF(ISNUMBER($M%d),ABS($M%d),0)' % (r, r)).font = font(8, color=GRAY)
        ws.cell(r, 20, r - PF0 + 1).font = font(7, color='A6A6A6')
    ws.freeze_panes = 'B%d' % PF0
    ws.sheet_view.zoomScale = 85
    print_setup(ws, 'B1:T%d' % (PF0 + 60), landscape=True)
    return ws
