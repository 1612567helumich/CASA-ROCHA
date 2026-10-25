"""Arma la hoja de revisión de muros de mampostería confinada (NTC 2023 CDMX) en dos versiones con las mismas fórmulas:

    hojas_generales/Muros_Mamposteria_NTC2023.xlsx   hoja general, con tres muros de ejemplo (para cualquier proyecto)
    muros_casa_rocha/Muros_CasaRocha_NTC2023.xlsx    los 37 muros de Casa Rocha

Uso:
    python3 gen_muros.py                          # arma las dos, las recalcula con LibreOffice y deja los valores en caché
    python3 gen_muros.py general | casa_rocha     # sólo una
    python3 gen_muros.py ... --sin-recalculo      # sólo arma (Excel calcula al abrir)

Ver README.md.
"""
import glob
import json
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, HERE)
import muros_comun as C          # noqa: E402
import muros_tabla as T          # noqa: E402
import muros_detalle as D        # noqa: E402
import muros_otros as O          # noqa: E402
import muros_croquis as K        # noqa: E402

EJ = os.path.join(HERE, 'ejemplo_casa_rocha.json')

# Story drifts de la hoja Revision_muros_NTC2023.xlsx (ETABS R3.2): Story, Output Case, Case Type, Step Type, Direction, Drift, Label, X, Y, Z
DRIFTS_CR = [
    ('N+6.00', 'SXLD', 'LinRespSpec', 'Max', 'X', 0.000018, 223, 4.33, 1.92, None), ('N+6.00', 'SYLD', 'LinRespSpec', 'Max', 'X', 0.000001, 238, 0, 9.95, None),
    ('N+6.00', 'SYLD', 'LinRespSpec', 'Max', 'Y', 0.000005, 77, 6.42, 6.95, None), ('N+6.00', 'SXSV', 'LinRespSpec', 'Max', 'X', 0.000042, 223, 4.33, 1.92, None),
    ('N+6.00', 'SYSV', 'LinRespSpec', 'Max', 'X', 0.000003, 238, 0, 9.95, None), ('N+6.00', 'SYSV', 'LinRespSpec', 'Max', 'Y', 0.000011, 77, 6.42, 6.95, None),
    ('N+3.00', 'SXLD', 'LinRespSpec', 'Max', 'X', 0.000101, 92, 8.85, 0, None), ('N+3.00', 'SYLD', 'LinRespSpec', 'Max', 'X', 0.000001, 9, 0, 0, None),
    ('N+3.00', 'SYLD', 'LinRespSpec', 'Max', 'Y', 0.000005, 300, 8.43, 6.95, None), ('N+3.00', 'SXSV', 'LinRespSpec', 'Max', 'X', 0.000241, 92, 8.85, 0, None),
    ('N+3.00', 'SYSV', 'LinRespSpec', 'Max', 'X', 0.000004, 9, 0, 0, None), ('N+3.00', 'SYSV', 'LinRespSpec', 'Max', 'Y', 0.000013, 300, 8.43, 6.95, None)]
STORIES_CR = [dict(st='N+3.00', nom='PB', W=253, Vx=22.1, Vy=17.3), dict(st='N+6.00', nom='PA', W=139, Vx=12.1, Vy=10.0)]


def muros_casa_rocha():
    out = []
    for w in json.load(open(EJ, encoding='utf-8')):
        e = dict(key=w['key'], st=w['story'], dir=w['dir'], pos='EXTREMO', L=w['L'], rht='NINGUNO',
                 cVu=w['Vu'], cMu=w['Mu'], cP=w['P'], cPmx=w['Pu'], cPg=w['Pu'], cVg=w['Vg'])
        for j, o in enumerate(w['openings'][:C.NV], start=1):
            e.update({'v%dt' % j: o['tipo'], 'v%dx' % j: o['x'], 'v%dw' % j: o['w'], 'v%da' % j: o['ant'], 'v%dh' % j: o['alto']})
        nota = []
        if w['openings']:
            nota.append('Vanos del plano arquitectónico; alturas típicas por confirmar.')
        if w.get('user_tramos'):
            nota.append('Tramos de tu hoja anterior: %s m.' % w['user_tramos'])
        b = w.get('beam')
        if b:
            e.update(sv='SÍ', vnom=b['name'], bv=b['b'], hv=b['h'], Lv=b['Lv'], wv=b['wu'], cv=b['c'])
            nota.append('Viga %s %s×%s propuesta en PRINCIPALES (el modelo R3.2 tiene dala 15×20).' % (b['name'], b['b'], b['h']))
        else:
            e['sv'] = 'NO'
        e['nota'] = ' '.join(nota)
        out.append(e)
    return out


# tres muros de ejemplo para la hoja general: sin vanos, con puerta y ventana, y sobre una viga
MUROS_GENERAL = [
    dict(key='M-1', st='N1', dir='X', pos='EXTREMO', L=4.00, rht='NINGUNO', sv='NO', cVu=6.0, cMu=9.0, cP=7.0, cPmx=18.0,
         nota='Ejemplo: muro sin vanos.'),
    dict(key='M-2', st='N1', dir='Y', pos='EXTREMO', L=5.00, v1t='PUERTA', v1x=0.40, v1w=0.90, v1a=0, v1h=2.10, v2t='VENTANA', v2x=2.30, v2w=1.50, v2a=0.90, v2h=1.20,
         rht='NINGUNO', sv='NO', cVu=5.5, cMu=12.0, cP=9.0, cPmx=22.0, nota='Ejemplo: muro con puerta y ventana; sin refuerzo horizontal no le alcanza el cortante.'),
    dict(key='M-3', st='N2', dir='X', pos='EXTREMO', L=3.60, rht='NINGUNO', sv='SÍ', vnom='T-1', bv=20, hv=45, Lv=4.00, wv=1.2, cv=0, cVu=3.0, cMu=6.0, cP=4.0, cPmx=9.0,
         nota='Ejemplo: muro sobre la viga T-1 (4.1.18).'),
]

CASOS = {
    'general': dict(out=os.path.join(REPO, 'hojas_generales', 'Muros_Mamposteria_NTC2023.xlsx'), sel='M-2', muros=lambda: MUROS_GENERAL,
                    stories=[dict(st='N1', nom='Ejemplo', W=45, Vx=8, Vy=6)], drifts=[], DEF=dict()),
    'casa_rocha': dict(out=os.path.join(REPO, 'muros_casa_rocha', 'Muros_CasaRocha_NTC2023.xlsx'), sel='CAD_PA_21', muros=muros_casa_rocha, stories=STORIES_CR, drifts=DRIFTS_CR,
                       DEF=dict(proj='Casa Rocha', sup=380, niv=2, mviga='SÍ', nviv=1, Hdef=2.75, Rp=1.68, Rp_txt="Q·R' = 2.52 en Casa Rocha",
                                cmb_s=['Comb4.%d' % i for i in range(1, 9)], cmb_g=['Comb2.2'], ej_tit='CASA ROCHA',
                                ej_txt='37 muros del modelo R3.2 con los vanos del plano arquitectónico y las fuerzas de tu hoja de revisión. Alturas de vanos típicas: confírmalas.',
                                ej_fila='Casa Rocha: 37 muros del modelo R3.2. Lo amarillo se puede cambiar.')),
}
DEF0 = dict(C.DEF)


def build(path, sel=None, only=None, caso='general'):
    cfg = CASOS[caso]
    C.DEF.clear(); C.DEF.update(DEF0); C.DEF.update(cfg['DEF'])
    D.DEFAULT_SEL = sel or cfg['sel']
    C.N.clear(); C.TB.clear()
    wb = Workbook()
    C.sheet_instrucciones(wb)
    C.sheet_tablas(wb)
    C.sheet_datos(wb)
    T.define_pf_names(wb)
    T.sheet_muros(wb, cfg['muros']())
    T.sheet_pier_forces(wb)
    D.sheet_detalle(wb)
    wsK = K.sheet_croquis(wb)
    O.sheet_entrepiso(wb, cfg['stories'])
    O.sheet_distorsiones(wb, cfg['drifts'])
    O.sheet_resumen(wb)
    C.sheet_referencias(wb)
    K.build_croquis_datos(wb, wsK)
    order = ['INSTRUCCIONES', 'DATOS', 'MUROS', 'DETALLE', 'CROQUIS', 'RESUMEN', 'ENTREPISO', 'DISTORSIONES', 'PIER_FORCES', 'REFERENCIAS', 'TABLAS', 'CROQUIS_DATOS']
    wb._sheets = [wb[n] for n in order if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in order]
    wb.active = 0
    if only:
        for w in wb.worksheets:
            if w.title != only:
                w.sheet_state = 'hidden'
        wb.active = wb.sheetnames.index(only)
    wb.calculation.fullCalcOnLoad = True
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


def recalc_and_inject(path):
    rc = sorted(glob.glob('/root/.claude/skills/synced/*/xlsx/scripts/recalc.py'))
    if not rc:
        print('No encontré recalc.py: el libro queda sin valores en caché (Excel los calcula al abrir).')
        return None
    tmp = tempfile.mkdtemp()
    calc = os.path.join(tmp, 'rc.xlsx')
    shutil.copy(path, calc)
    res = subprocess.run([sys.executable, rc[0], calc, '600'], capture_output=True, text=True)
    print(res.stdout[-3000:])
    if res.returncode != 0:
        print(res.stderr[-3000:])
    final = os.path.join(tmp, 'final.xlsx')
    subprocess.run([sys.executable, os.path.join(HERE, 'inject_values.py'), path, calc, final], check=True)
    shutil.copy(final, path)
    return res.stdout


if __name__ == '__main__':
    casos = [a for a in sys.argv[1:] if a in CASOS] or list(CASOS)
    for caso in casos:
        out = build(CASOS[caso]['out'], caso=caso)
        print('armado:', out)
        if '--sin-recalculo' not in sys.argv:
            recalc_and_inject(out)
            print('listo:', out)
